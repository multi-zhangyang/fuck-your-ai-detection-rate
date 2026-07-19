from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import sys
import shutil
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Sequence

from docx import Document  # type: ignore[import]

from fyadr_records import (
    _collect_round_file_paths,
    build_artifact_summary,
    delete_document,
    delete_rounds,
    list_referenced_history_artifact_paths,
    list_records,
    normalize_doc_id,
    preview_delete_document,
)
from chunking import (
    DEFAULT_CHUNK_LIMIT,
    Chunk,
    ChunkManifest,
    ParagraphManifest,
    build_manifest,
    load_manifest,
    restore_text_from_chunks,
    split_text_to_paragraphs,
)
from academic_readability import (
    ACADEMIC_READABILITY_DELTA_SCHEMA,
    ACADEMIC_READABILITY_DELTA_VERSION,
    assess_academic_readability_delta,
)
from source_relative_style_delta import (
    assess_source_relative_document_delta,
    assess_source_relative_style_delta,
    source_relative_document_delta_passed,
    source_pattern_profile_valid,
    source_relative_style_delta_passed,
)
from fyadr_round_service import (
    CANDIDATE_SELECTION_SCHEMA,
    CANDIDATE_SELECTION_VERSION,
    DOCUMENT_PATTERN_ACCUMULATION_BLOCKED,
    STYLE_CARD_VERSION,
    build_global_style_profile_from_texts,
    build_language_guard,
    build_local_style_card,
    build_paragraph_guard,
    build_prompt_input,
    get_max_rounds,
    get_round_checkpoint_path,
    get_round_compare_path,
    get_chunk_metric,
    load_prompt,
    normalize_chunk_output,
    normalize_path,
    normalize_prompt_profile,
    normalize_prompt_sequence,
    protect_structure_tokens,
    resolve_round_dimension,
    restore_structure_tokens,
    validate_chunk_output,
    validate_immutable_text_anchors,
    validate_structure_placeholders,
    _build_chunk_quality,
    _apply_candidate_selection_quality,
    _build_candidate_selection_event,
    _build_candidate_selection_retry_note,
    _build_checkpoint_signature,
    _build_validation_repair_steps,
    _candidate_requires_conditional_retry,
    _evaluate_rewrite_candidate,
    _extract_required_terms,
    _is_checkpoint_compatible,
    _select_rewrite_candidate,
    _score_rewrite_output,
    _classify_failed_attempt_diagnostic,
    _normalize_failed_attempt_evidence,
    _public_candidate_selection_event,
    _public_validation_events,
    _serialize_failed_output,
    _sha256_json,
    _sha256_text,
)
from prompt_library import (
    DEFAULT_PROMPT_PROFILE,
    LEGACY_PROMPT_PROFILE,
    ROUND_PERTURBATION_DIMENSIONS,
    get_prompt_id_for_round,
    get_rate_audit_dimension_definition,
    is_prompt_sequence_customizable,
    prompt_sequence_match_rank,
    resolve_prompt_path,
)
from docx_bodymap import (
    docx_body_map_from_payload,
    get_body_map_unit_model_format_anchors,
    load_docx_body_map,
    save_docx_body_map,
    update_docx_body_map_texts,
    validate_docx_body_map,
)
from docx_audit import (
    audit_docx_export,
    audit_docx_format_lock,
    audit_docx_ooxml_integrity,
    get_docx_audit_report_path,
    get_docx_format_lock_report_path,
    get_docx_ooxml_audit_report_path,
)
from docx_export_guard import (
    run_docx_pre_export_guard,
    summarize_docx_export_guard_failure,
)
from docx_pipeline import (
    _load_docx_snapshot,
    _resolve_target_paragraph,
    _restore_rewritten_text_with_source_whitespace,
    _split_text_into_blocks,
    build_docx_scope_diagnostics,
    ensure_docx_processing_assets,
    get_docx_scope_diagnostics_path,
    get_docx_snapshot_path,
    rebuild_docx_from_snapshot,
    rebuild_docx_from_body_map_units,
    verify_docx_snapshot_derivation,
    write_docx_text,
)
from factual_guards import build_factual_relation_guard, build_factual_scope_repair_guard
from docx_protection_map import build_docx_protection_map
from document_edit_contract import (
    assert_document_edit_contract_ready,
    build_document_edit_contract,
    get_document_edit_contract_path,
    get_export_edit_contract_path,
)
from llm_client import LLMRequestError, list_llm_models, llm_completion, test_llm_connection
from app_config import normalize_streaming
from path_utils import is_path_under, truncate_utf8_filename_component
from runtime_error_safety import (
    safe_public_error_message,
    safe_retry_progress,
    sanitize_persisted_error,
)
from round_helper import (
    build_round_artifact_stem,
    build_round_context,
    ensure_round_input_text,
    get_document_round_state,
    round_record_has_usable_artifacts,
)
from rate_audit import build_rate_audit_report


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_PREVIEW_MAX_CHARS = 12000
TARGETED_RERUN_SELECTION_ATTEMPTS = 2
TARGETED_RERUN_VALIDATION_ATTEMPTS = 3
# A compare artifact only needs profiles referenced by its current candidate
# selections.  Content-addressing plus bounded garbage collection prevents
# repeated targeted reruns from growing the sidecar without limit.
MAX_SOURCE_PATTERN_PROFILE_REGISTRY_ENTRIES = 256
MAX_SOURCE_PATTERN_PROFILE_BYTES = 8 * 1024 * 1024
MAX_SOURCE_PATTERN_PROFILE_REGISTRY_BYTES = 32 * 1024 * 1024
# Dimension-convergence closure: a chunk flagged dimension_direction_not_effective
# is rerun with a dimension-targeted instruction, then re-scored on THAT dimension.
# If it still hasn't moved toward lower-AI, retry with a strengthened note, up to
# this many dimension-convergence attempts (on top of the validation attempts).
DIMENSION_RERUN_CONVERGE_ATTEMPTS = 2
# Provider rate-limit / cooldown state now lives in provider_guard. Keep module
# references here so existing callers (and provider_guard_regression.py via
# app_service.<symbol>) keep working against the shared singleton.
from provider_guard import (  # noqa: E402
    _RATE_LIMIT_STATE,
    _RATE_LIMIT_LOCK,
    _PROVIDER_GUARD_STATE,
    _PROVIDER_GUARD_LOCK,
    _coerce_int_config,
    _coerce_rewrite_timeout_seconds,
    _coerce_rate_limit,
    _coerce_rate_window_minutes,
    _coerce_rate_max_requests,
    _provider_guard_key,
    _provider_display_name,
    _build_provider_rate_limiter,
    _wait_for_provider_cooldown,
    _register_provider_success,
    _register_provider_failure,
)
DEFAULT_REQUEST_TIMEOUT_SECONDS = 600
MIN_REWRITE_REQUEST_TIMEOUT_SECONDS = 600
DEFAULT_MAX_RETRIES = 3
MIN_REWRITE_TRANSIENT_RETRIES = 6
REWRITE_RETRY_BACKOFF_SECONDS = 5.0
DEFAULT_REWRITE_CONCURRENCY = 2
MAX_REWRITE_CONCURRENCY = 16
API_READ_ALLOWED_ROOTS = tuple((ROOT_DIR / name).resolve() for name in ("finish", "origin", "prompts", "references"))
API_OUTPUT_ALLOWED_ROOTS = ((ROOT_DIR / "finish").resolve(),)
API_EXPORT_ALLOWED_ROOTS = ((ROOT_DIR / "finish").resolve(),)
RATE_AUDIT_STRATEGY_BINDING_VERSION = 2
ROUND_ARTIFACT_SNAPSHOT_VERSION = 1
ROUND_ARTIFACT_SNAPSHOT_LOCK_TIMEOUT_SECONDS = 2.0
ROUND_ARTIFACT_SNAPSHOT_NO_REVIEW_SENTINEL = b"fyadr:no-review-sidecar:v1"


class StaleRateAuditStrategyPlanError(ValueError):
    """Raised before model construction when an execution plan is no longer current."""

    def __init__(self, mismatch_codes: list[str], message: str = "RateAudit strategy plan is stale.") -> None:
        normalized_codes: list[str] = []
        for value in mismatch_codes:
            code = str(value or "").strip()
            if code and code not in normalized_codes:
                normalized_codes.append(code)
        super().__init__(message)
        self.mismatch_codes = normalized_codes or ["strategy_binding_mismatch"]


class ReviewRevisionRequiredError(ValueError):
    """Raised when an external review write omits its compare precondition."""


class StaleReviewDecisionsError(ValueError):
    """Raised when a review snapshot was prepared against an older candidate."""

    def __init__(self, current_compare_revision: str) -> None:
        super().__init__("Review decisions were prepared against a stale compare revision.")
        self.current_compare_revision = str(current_compare_revision or "")


class InconsistentReviewStateError(ValueError):
    """Raised when compare.reviewUpdatedAt no longer links to its sidecar."""

    def __init__(self, current_compare_revision: str) -> None:
        super().__init__("Review decision sidecar is not linked to the current compare review revision.")
        self.current_compare_revision = str(current_compare_revision or "")


class RoundArtifactSnapshotError(ValueError):
    """Raised when a coherent round snapshot cannot be captured safely."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        retryable: bool = False,
        details: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = str(code or "round_snapshot_invalid").strip() or "round_snapshot_invalid"
        self.retryable = bool(retryable)
        self.details = dict(details or {})


class DocumentReleaseGateError(ValueError):
    """Raised when review materialization would publish uncertified text."""

    code = "document_release_gate_failed"

    def __init__(
        self,
        chunk_id: str,
        mode: str,
        issue_codes: Sequence[str],
        *,
        evidence: dict[str, Any] | None = None,
    ) -> None:
        normalized_codes: list[str] = []
        for value in issue_codes:
            code = str(value or "").strip()
            if code and code not in normalized_codes:
                normalized_codes.append(code)
        self.chunk_id = str(chunk_id or "").strip()
        self.mode = str(mode or "unknown").strip() or "unknown"
        self.issue_codes = normalized_codes or ["release_evidence_invalid"]
        self.details: dict[str, Any] = {
            "chunkId": self.chunk_id,
            "mode": self.mode,
            "issueCodes": list(self.issue_codes),
        }
        for key, value in dict(evidence or {}).items():
            if key in {
                "materializedTextSha256",
                "outputTextSha256",
                "selectedTextSha256",
            } and isinstance(value, str) and value:
                self.details[key] = value
        super().__init__(
            "Document release gate rejected uncertified materialization "
            f"for chunk {self.chunk_id or '<unknown>'}."
        )


class RoundInputRevisionRequiredError(ValueError):
    """Raised when a downstream run is not bound to its parent revision."""


class StaleRoundInputError(ValueError):
    """Raised when a downstream run no longer matches its parent snapshot."""

    def __init__(
        self,
        current_compare_revision: str,
        mismatch_codes: list[str],
        message: str = "Downstream round input is stale.",
    ) -> None:
        super().__init__(message)
        self.current_compare_revision = str(current_compare_revision or "")
        self.mismatch_codes = [
            str(code or "").strip()
            for code in mismatch_codes
            if str(code or "").strip()
        ] or ["parent_compare_revision_mismatch"]


PARENT_INPUT_BINDING_FIELDS: tuple[tuple[str, str], ...] = (
    ("compareRevision", "parent_compare_revision_mismatch"),
    ("reviewRevision", "parent_review_revision_mismatch"),
    ("contentRevision", "parent_content_revision_mismatch"),
    ("artifactSnapshotDigest", "parent_artifact_snapshot_digest_mismatch"),
    ("effectiveTextSha256", "parent_effective_text_sha256_mismatch"),
)


def _normalize_expected_parent_input_binding(
    expected_previous_compare_revision: str | None,
    expected_parent_input_binding: dict[str, Any] | None,
) -> dict[str, str]:
    if expected_parent_input_binding is None:
        raw_binding: dict[str, Any] = {}
    elif isinstance(expected_parent_input_binding, dict):
        raw_binding = expected_parent_input_binding
    else:
        raise ValueError("expected parent input binding must be an object.")

    normalized: dict[str, str] = {}
    for field, _mismatch_code in PARENT_INPUT_BINDING_FIELDS:
        raw_value = raw_binding.get(field)
        if raw_value is None:
            continue
        if not isinstance(raw_value, str):
            raise ValueError(f"expected parent input binding field {field} must be a string.")
        value = raw_value.strip()
        if value:
            normalized[field] = value

    legacy_compare_revision = str(expected_previous_compare_revision or "").strip()
    bound_compare_revision = normalized.get("compareRevision", "")
    if legacy_compare_revision and bound_compare_revision and legacy_compare_revision != bound_compare_revision:
        raise ValueError("expectedPreviousCompareRevision conflicts with the parent input binding.")
    if legacy_compare_revision:
        normalized["compareRevision"] = legacy_compare_revision
    return normalized


def _round_snapshot_parent_input_binding(snapshot: dict[str, Any] | None) -> dict[str, str]:
    if not isinstance(snapshot, dict):
        return {}
    return {
        field: str(snapshot.get(field, "") or "").strip()
        for field, _mismatch_code in PARENT_INPUT_BINDING_FIELDS
    }


_OUTPUT_RERUN_LOCKS_GUARD = threading.Lock()
_OUTPUT_RERUN_LOCKS: dict[str, threading.RLock] = {}


def get_output_rerun_lock(output_path: str | Path) -> threading.RLock:
    """Return the process-local mutation lock shared by every rerun path."""

    normalized = str(normalize_path(Path(output_path)).resolve())
    with _OUTPUT_RERUN_LOCKS_GUARD:
        lock = _OUTPUT_RERUN_LOCKS.get(normalized)
        if lock is None:
            lock = threading.RLock()
            _OUTPUT_RERUN_LOCKS[normalized] = lock
        return lock


class ExportRoundError(ValueError):
    def __init__(self, message: str, export_failure: dict[str, Any]) -> None:
        super().__init__(message)
        self.export_failure = export_failure


def _resolve_api_path(path_value: str | Path, *, allowed_roots: tuple[Path, ...], label: str) -> Path:
    normalized_path = normalize_path(Path(path_value)).resolve()
    if not any(is_path_under(normalized_path, root) for root in allowed_roots):
        allowed = ", ".join(str(root) for root in allowed_roots)
        raise ValueError(f"{label} must stay under allowed workspace directories: {allowed}")
    return normalized_path


def _truncate_issue_text(value: Any, limit: int = 180) -> str:
    text = " ".join(str(value or "").split())
    return text[:limit] + ("…" if len(text) > limit else "")


def _format_issue_location(issue: dict[str, Any]) -> str:
    parts: list[str] = []
    for key, label in (
        ("paragraphIndex", "段落"),
        ("unitIndex", "单元"),
        ("chunkId", "块"),
        ("tableIndex", "表"),
        ("rowIndex", "行"),
        ("cellIndex", "单元格"),
        ("role", "角色"),
    ):
        value = issue.get(key)
        if value is not None and value != "":
            parts.append(f"{label} {value}")
    return " · ".join(parts)


def _normalize_issue_sample(issue: Any) -> dict[str, str] | None:
    if not isinstance(issue, dict):
        text = _truncate_issue_text(issue)
        return {"message": text} if text else None
    message = _truncate_issue_text(issue.get("message") or issue.get("type") or issue.get("code") or "检查项")
    sample = _truncate_issue_text(
        issue.get("sample")
        or issue.get("rewrittenSample")
        or issue.get("originalSample")
        or issue.get("actual")
        or issue.get("expected")
    )
    normalized = {
        "code": _truncate_issue_text(issue.get("code") or issue.get("type") or issue.get("severity"), 80),
        "severity": _truncate_issue_text(issue.get("severity"), 40),
        "message": message,
        "location": _format_issue_location(issue),
        "sample": sample,
    }
    return {key: value for key, value in normalized.items() if value}


def _read_json_report(path: Any) -> dict[str, Any] | None:
    try:
        if not path:
            return None
        report_path = Path(str(path))
        if not report_path.exists():
            return None
        payload = json.loads(report_path.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else None
    except Exception:
        return None


def _collect_issue_samples(*reports: dict[str, Any] | None, keys: tuple[str, ...] = ("issues",), limit: int = 3) -> list[dict[str, str]]:
    samples: list[dict[str, str]] = []
    for report in reports:
        if not isinstance(report, dict):
            continue
        for key in keys:
            raw_issues = report.get(key)
            if not isinstance(raw_issues, list):
                continue
            for issue in raw_issues:
                sample = _normalize_issue_sample(issue)
                if sample:
                    samples.append(sample)
                if len(samples) >= limit:
                    return samples
    return samples


def _build_export_failure(
    *,
    stage: str,
    label: str,
    message: str,
    report: dict[str, Any] | None = None,
    report_path: Any = None,
    issue_count: int | None = None,
    warning_count: int | None = None,
    samples: list[dict[str, str]] | None = None,
    sample_keys: tuple[str, ...] = ("issues",),
) -> dict[str, Any]:
    normalized_report = report if isinstance(report, dict) else None
    normalized_report_path = (
        str(report_path or "").strip()
        or str((normalized_report or {}).get("reportPath", "") or "").strip()
        or str((normalized_report or {}).get("path", "") or "").strip()
    )
    normalized_issue_count = issue_count
    if normalized_issue_count is None and normalized_report is not None:
        normalized_issue_count = int(normalized_report.get("blockingIssueCount", normalized_report.get("issueCount", 0)) or 0)
    normalized_warning_count = warning_count
    if normalized_warning_count is None and normalized_report is not None:
        normalized_warning_count = int(normalized_report.get("warningCount", 0) or 0)
    normalized_samples = samples if samples is not None else _collect_issue_samples(normalized_report, keys=sample_keys, limit=5)
    return {
        "stage": stage,
        "label": label,
        "message": _truncate_issue_text(message, 400),
        "reportPath": normalized_report_path,
        "issueCount": int(normalized_issue_count or 0),
        "warningCount": int(normalized_warning_count or 0),
        "samples": normalized_samples[:5],
    }


def _coerce_history_artifact_stats(value: Any) -> dict[str, Any] | None:
    if not isinstance(value, dict):
        return None
    empty_stats = _empty_artifact_stats()
    if not any(key in value for key in empty_stats):
        return None
    stats = dict(empty_stats)
    for key in stats:
        item = value.get(key)
        if isinstance(item, (int, float)):
            stats[key] = int(item)
    return stats


def _map_history_round(item: dict[str, Any]) -> dict[str, Any]:
    prompt_sequence = item.get("prompt_sequence")
    quality_summary = _load_history_quality_summary(item)
    run_audit = _build_history_run_audit(item, quality_summary)
    artifact_stats = _coerce_history_artifact_stats(item.get("artifactStats"))
    if artifact_stats is None:
        artifact_stats = _coerce_history_artifact_stats(item.get("artifact_stats"))
    return {
        "round": int(item.get("round", 0)),
        "prompt": str(item.get("prompt", "")),
        "promptProfile": str(item.get("prompt_profile", LEGACY_PROMPT_PROFILE) or LEGACY_PROMPT_PROFILE).strip().lower() or LEGACY_PROMPT_PROFILE,
        "promptSequence": prompt_sequence if isinstance(prompt_sequence, list) else [],
        "inputPath": str(item.get("input_path", "")),
        "outputPath": str(item.get("output_path", "")),
        "manifestPath": str(item.get("manifest_path", "")),
        "comparePath": str(item.get("compare_path", "")),
        "qualityPath": str(item.get("quality_path", "")),
        "scoreTotal": item.get("score_total"),
        "chunkLimit": item.get("chunk_limit"),
        "inputSegmentCount": item.get("input_segment_count"),
        "outputSegmentCount": item.get("output_segment_count"),
        "bodyMapPath": str(item.get("body_map_path", "")),
        "validationPath": str(item.get("validation_path", "")),
        "timestamp": str(item.get("timestamp", "")),
        "artifactStats": artifact_stats if artifact_stats is not None else build_artifact_summary([item]),
        "qualitySummary": quality_summary or None,
        "runAudit": run_audit or None,
    }


def _history_path_to_absolute(value: object) -> Path | None:
    raw_path = str(value or "").strip()
    if not raw_path:
        return None
    candidate = Path(raw_path)
    if not candidate.is_absolute():
        candidate = ROOT_DIR / candidate
    return candidate.resolve()


def _load_history_quality_summary(item: dict[str, Any]) -> dict[str, Any]:
    quality_path = _history_path_to_absolute(item.get("quality_path"))
    if quality_path is None or not quality_path.exists() or not quality_path.is_file():
        return {}
    try:
        payload = json.loads(quality_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return payload if isinstance(payload, dict) else {}


def _build_history_run_audit(item: dict[str, Any], quality_summary: dict[str, Any]) -> dict[str, Any]:
    stored_audit = item.get("run_audit")
    audit = dict(stored_audit) if isinstance(stored_audit, dict) else {}
    for legacy_key in ("offlineMode", "rewriteCandidateMode", "candidateMaxPerChunk", "twoCandidateChunkCount"):
        audit.pop(legacy_key, None)
    if quality_summary:
        split_summary = quality_summary.get("paragraphSplitSummary")
        audit.setdefault("estimatedApiCalls", quality_summary.get("estimatedApiCalls"))
        audit.setdefault("validationRetryCount", quality_summary.get("validationRetryCount"))
        audit.setdefault("sourceFallbackCount", quality_summary.get("sourceFallbackCount"))
        audit.setdefault("validationEventCount", quality_summary.get("validationEventCount"))
        audit.setdefault("machineLikeRiskCount", quality_summary.get("machineLikeRiskCount"))
        audit.setdefault("protectedTokenCount", quality_summary.get("protectedTokenCount"))
        if isinstance(split_summary, dict):
            audit.setdefault("paragraphCount", split_summary.get("paragraphCount"))
            audit.setdefault("chunkCount", split_summary.get("chunkCount"))
            audit.setdefault("splitParagraphCount", split_summary.get("splitParagraphCount"))
    if item.get("input_segment_count") is not None:
        audit.setdefault("chunkCount", item.get("input_segment_count"))
    audit.setdefault("promptProfile", str(item.get("prompt_profile", LEGACY_PROMPT_PROFILE) or LEGACY_PROMPT_PROFILE))
    prompt_sequence = item.get("prompt_sequence")
    if isinstance(prompt_sequence, list):
        audit.setdefault("promptSequence", prompt_sequence)
    return {key: value for key, value in audit.items() if value not in ("", None, [])}


def _record_entry_to_history(
    doc_id: str,
    entry: dict[str, Any],
    *,
    all_records: dict[str, Any] | None = None,
) -> dict[str, Any]:
    rounds = entry.get("rounds") if isinstance(entry.get("rounds"), list) else []
    usable_rounds = [
        item
        for item in rounds
        if isinstance(item, dict)
        and round_record_has_usable_artifacts(item, expected_doc_id=doc_id, all_records=all_records)
    ]
    history_rounds = [_map_history_round(item) for item in usable_rounds]
    history_rounds.sort(key=lambda item: item["round"], reverse=True)
    completed_rounds = sorted({item["round"] for item in history_rounds})
    latest_round = max(
        history_rounds,
        key=lambda item: (item.get("timestamp", ""), item["round"]),
        default=None,
    )
    origin_path = str(entry.get("origin_path", doc_id))
    source_kind = str(entry.get("sourceKind", entry.get("source_kind", Path(origin_path).suffix.lower() or ".txt")))
    artifact_stats = _coerce_history_artifact_stats(entry.get("artifactStats"))
    if artifact_stats is None:
        artifact_stats = _coerce_history_artifact_stats(entry.get("artifact_stats"))

    return {
        "docId": doc_id,
        "sourcePath": origin_path,
        "originPath": origin_path,
        "sourceKind": source_kind,
        "completedRounds": completed_rounds,
        "latestOutputPath": latest_round.get("outputPath", "") if latest_round else "",
        "lastTimestamp": latest_round.get("timestamp", "") if latest_round else "",
        "artifactStats": artifact_stats if artifact_stats is not None else build_artifact_summary(usable_rounds),
        "rounds": history_rounds,
    }


def _indexed_history_document_to_record_entry(document: dict[str, Any]) -> dict[str, Any]:
    rounds = document.get("rounds") if isinstance(document.get("rounds"), list) else []
    origin_path = str(document.get("originPath", document.get("sourcePath", document.get("docId", ""))))
    return {
        "origin_path": origin_path,
        "sourceKind": str(document.get("sourceKind", Path(origin_path).suffix.lower() or ".txt")),
        "artifactStats": document.get("artifactStats"),
        "rounds": [item for item in rounds if isinstance(item, dict)],
    }


def _list_indexed_history_documents() -> list[dict[str, Any]] | None:
    try:
        from fyadr_history_db import list_history_documents_from_index

        documents = list_history_documents_from_index()
    except Exception:
        return None
    return documents if isinstance(documents, list) else None


def _get_indexed_history_document(doc_id: str) -> dict[str, Any] | None:
    try:
        from fyadr_history_db import get_history_document_from_index

        document = get_history_document_from_index(doc_id)
    except Exception:
        return None
    return document if isinstance(document, dict) else None


def _find_record_context_for_output(output_path: Path) -> tuple[dict[str, Any], dict[str, Any]] | None:
    normalized_output_path = normalize_path(output_path)
    records = list_records()
    matches: list[tuple[dict[str, Any], dict[str, Any]]] = []
    for doc_id, entry in records.items():
        if not isinstance(entry, dict):
            continue
        rounds = entry.get("rounds")
        if not isinstance(rounds, list):
            continue
        for round_item in rounds:
            if not isinstance(round_item, dict):
                continue
            round_output_path = round_item.get("output_path")
            if not isinstance(round_output_path, str) or not round_output_path.strip():
                continue
            if normalize_path(Path(round_output_path)) != normalized_output_path:
                continue
            if round_record_has_usable_artifacts(
                round_item,
                expected_doc_id=str(doc_id),
                all_records=records,
            ):
                matches.append((entry, round_item))
    if len(matches) > 1:
        raise ValueError("The selected legacy artifact is referenced by multiple documents and has ambiguous provenance.")
    return matches[0] if matches else None


PLAIN_EXPORT_CERTIFICATION = "plain_uncertified"


def _find_declared_round_output_provenance(output_path: Path) -> list[tuple[str, int]]:
    """Find history declarations even when their round artifacts are unusable.

    ``_find_record_context_for_output`` intentionally returns only usable
    rounds.  Export downgrade decisions need the opposite authority model: an
    exact ``rounds[].output_path`` declaration remains provenance after its
    compare/evidence files are lost or damaged.
    """

    normalized_output_path = normalize_path(output_path).resolve()
    matches: list[tuple[str, int]] = []
    for doc_id, entry in list_records().items():
        if not isinstance(entry, dict):
            continue
        rounds = entry.get("rounds")
        if not isinstance(rounds, list):
            continue
        for round_index, round_item in enumerate(rounds):
            if not isinstance(round_item, dict):
                continue
            raw_output_path = round_item.get("output_path")
            if not isinstance(raw_output_path, str) or not raw_output_path.strip():
                continue
            declared_output_path = normalize_path(Path(raw_output_path)).resolve()
            if declared_output_path == normalized_output_path:
                matches.append((str(doc_id), round_index))
    return matches


def _canonical_fyadr_sidecar_paths(output_path: Path) -> tuple[Path, ...]:
    """Return canonical sibling artifacts whose mere presence is provenance."""

    normalized_output_path = normalize_path(output_path).resolve()
    stem = normalized_output_path.stem
    compare_path = normalized_output_path.with_name(f"{stem}_compare.json")
    paths = (
        compare_path,
        normalized_output_path.with_name(f"{stem}_manifest.json"),
        normalized_output_path.with_name(f"{stem}_quality.json"),
        normalized_output_path.with_name(f"{stem}_checkpoint.json"),
        normalized_output_path.with_name(f"{stem}_body_map.json"),
        normalized_output_path.with_name(f"{stem}_bodymap.json"),
        normalized_output_path.with_name(f"{stem}_validation.json"),
        compare_path.with_name(f"{compare_path.stem}_review_decisions.json"),
    )
    return tuple(dict.fromkeys(paths))


def _has_fyadr_round_or_docx_provenance(output_path: Path) -> bool:
    """Fail closed when history or any canonical round/Word trace exists."""

    if _find_declared_round_output_provenance(output_path):
        return True
    return any(path.exists() or path.is_symlink() for path in _canonical_fyadr_sidecar_paths(output_path))


def _find_origin_docx_for_output(output_path: Path) -> tuple[Path, Path] | None:
    record_context = _find_record_context_for_output(output_path)
    if record_context is None:
        return None
    entry, _ = record_context
    origin_path = entry.get("origin_path")
    if not isinstance(origin_path, str) or not origin_path.strip():
        return None
    resolved_origin_path = normalize_path(Path(origin_path))
    if resolved_origin_path.suffix.lower() != ".docx":
        return None
    return resolved_origin_path, get_docx_snapshot_path(resolved_origin_path)


def _find_declared_origin_docx_for_output(output_path: Path) -> Path | None:
    """Find a unique DOCX provenance declaration without trusting artifact usability."""

    normalized_output_path = normalize_path(output_path)
    records = list_records()
    matches: list[Path] = []
    for entry in records.values():
        if not isinstance(entry, dict):
            continue
        raw_origin = str(entry.get("origin_path", "") or "").strip()
        if not raw_origin:
            continue
        origin_path = normalize_path(Path(raw_origin))
        if origin_path.suffix.lower() != ".docx":
            continue
        rounds = entry.get("rounds")
        if not isinstance(rounds, list):
            continue
        for round_item in rounds:
            if not isinstance(round_item, dict):
                continue
            raw_round_output = str(round_item.get("output_path", "") or "").strip()
            if raw_round_output and normalize_path(Path(raw_round_output)) == normalized_output_path:
                matches.append(origin_path)
                break
    unique_matches = list(dict.fromkeys(matches))
    if len(unique_matches) > 1:
        raise ValueError("The selected output has ambiguous DOCX provenance declarations.")
    return unique_matches[0] if unique_matches else None


def _find_compare_path_for_output(output_path: Path) -> Path:
    normalized_output_path = normalize_path(output_path)
    default_compare_path = get_round_compare_path(normalized_output_path)
    if default_compare_path.exists():
        return default_compare_path
    record_context = _find_record_context_for_output(normalized_output_path)
    if record_context is not None:
        _, round_item = record_context
        compare_path = round_item.get("compare_path")
        if isinstance(compare_path, str) and compare_path.strip():
            return normalize_path(Path(compare_path))
    return default_compare_path


def _find_manifest_path_for_output(output_path: Path, *, include_records: bool = True) -> Path | None:
    normalized_output_path = normalize_path(output_path)
    default_manifest_path = normalized_output_path.with_name(f"{normalized_output_path.stem}_manifest.json")
    if default_manifest_path.exists():
        return default_manifest_path
    if not include_records:
        return None
    record_context = _find_record_context_for_output(normalized_output_path)
    if record_context is None:
        return None
    _, round_item = record_context
    manifest_path = round_item.get("manifest_path")
    if not isinstance(manifest_path, str) or not manifest_path.strip():
        return None
    return normalize_path(Path(manifest_path))


def _find_body_map_path_for_output(output_path: Path, *, include_records: bool = True) -> Path | None:
    normalized_output_path = normalize_path(output_path)
    default_body_map_path = normalized_output_path.with_name(f"{normalized_output_path.stem}_body_map.json")
    if default_body_map_path.exists():
        return default_body_map_path
    if not include_records:
        return None
    record_context = _find_record_context_for_output(normalized_output_path)
    if record_context is None:
        return None
    _, round_item = record_context
    body_map_path = round_item.get("body_map_path")
    if not isinstance(body_map_path, str) or not body_map_path.strip():
        return None
    return normalize_path(Path(body_map_path))


def _find_validation_path_for_output(output_path: Path) -> Path | None:
    normalized_output_path = normalize_path(output_path)
    default_validation_path = normalized_output_path.with_name(f"{normalized_output_path.stem}_validation.json")
    if default_validation_path.exists():
        return default_validation_path
    record_context = _find_record_context_for_output(normalized_output_path)
    if record_context is None:
        return None
    _, round_item = record_context
    validation_path = round_item.get("validation_path")
    if not isinstance(validation_path, str) or not validation_path.strip():
        return None
    return normalize_path(Path(validation_path))


def _load_body_map_for_output(output_path: Path):
    body_map_path = _find_body_map_path_for_output(output_path)
    if body_map_path is None:
        return None
    return load_docx_body_map(body_map_path)


def _load_compare_payload_for_output(output_path: Path) -> dict[str, Any] | None:
    compare_path = _find_compare_path_for_output(output_path)
    if not compare_path.exists():
        return None
    try:
        raw_payload = json.loads(compare_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(raw_payload, dict):
        return None
    return _normalize_compare_failed_attempts(copy.deepcopy(raw_payload))


def _normalize_failed_attempts(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    attempts: list[dict[str, Any]] = []
    for item in value[-4:]:
        normalized = _normalize_failed_attempt_evidence(item)
        if normalized is not None:
            attempts.append(dict(normalized))
    return attempts


_TEXT_FREE_FAILURE_MESSAGES = {
    "structure": "候选未通过结构与格式保护校验；失败正文和原始错误已隐藏。",
    "factual": "候选未通过事实关系保护校验；失败正文和原始错误已隐藏。",
    "readability": "候选出现学术可读性回退；失败正文和原始错误已隐藏。",
    "style": "候选未通过写作结构启发式校验；失败正文和原始错误已隐藏。",
    "provider": "模型服务调用失败；上游响应、思考内容和原始错误已隐藏。",
    "local_validation": "候选未通过本地安全校验；失败正文和原始错误已隐藏。",
}


def _build_text_free_app_failure(
    *,
    chunk_id: str,
    error: object,
    failed_attempts: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    attempts = _normalize_failed_attempts(failed_attempts)
    latest = attempts[-1] if attempts else {}
    guard_category, issue_codes = _classify_failed_attempt_diagnostic(
        error,
        guard_category=latest.get("guardCategory"),
        issue_codes=latest.get("issueCodes"),
    )
    result: dict[str, Any] = {
        "chunkId": str(chunk_id or "")[:128],
        "error": _TEXT_FREE_FAILURE_MESSAGES.get(
            guard_category,
            _TEXT_FREE_FAILURE_MESSAGES["local_validation"],
        ),
        "guardCategory": guard_category,
        "issueCodes": issue_codes,
        "errorStored": False,
        "reasoningSuppressed": True,
        "providerContentStored": False,
    }
    if attempts:
        result["failedAttempts"] = attempts
    return result


def _migrate_legacy_compare_diagnostic(
    chunk: dict[str, Any],
    *,
    field: str,
    prefix: str,
) -> None:
    raw_error = chunk.pop(field, None)
    if raw_error in (None, ""):
        return
    guard_category, issue_codes = _classify_failed_attempt_diagnostic(raw_error)
    chunk[f"{prefix}GuardCategory"] = guard_category
    chunk[f"{prefix}IssueCodes"] = issue_codes
    chunk[f"{prefix}ErrorStored"] = False


def _normalize_compare_failed_attempts(payload: dict[str, Any]) -> dict[str, Any]:
    raw_validation_events = payload.get("validationEvents")
    if isinstance(raw_validation_events, list):
        payload["validationEvents"] = _public_validation_events(raw_validation_events)
    chunks = payload.get("chunks")
    if not isinstance(chunks, list):
        return payload
    for chunk in chunks:
        if not isinstance(chunk, dict):
            continue
        selection = chunk.get("candidateSelection")
        if isinstance(selection, dict):
            sanitized_selection = _public_candidate_selection_event(selection)
            if isinstance(sanitized_selection, dict):
                chunk["candidateSelection"] = sanitized_selection
        failed_attempts = _normalize_failed_attempts(chunk.get("failedAttempts"))
        if not failed_attempts:
            failed_attempts = _normalize_failed_attempts(chunk.get("rejectedCandidates"))
        chunk.pop("rejectedCandidates", None)
        if failed_attempts:
            chunk["failedAttempts"] = failed_attempts
        else:
            chunk.pop("failedAttempts", None)
        _migrate_legacy_compare_diagnostic(chunk, field="fallbackError", prefix="fallback")
        _migrate_legacy_compare_diagnostic(chunk, field="rerunFallbackError", prefix="rerunFallback")
        # These fields are prompt material (including user feedback and
        # source-derived style-card text), not release evidence.  Old compare
        # files are scrubbed as soon as they cross a supported read boundary.
        removed_prompt_material = False
        for prompt_field in (
            "rerunPromptNote",
            "rerunUserFeedback",
            "rerunIssueCards",
            "rerunStyleCard",
            "rerunGlobalStyleProfile",
        ):
            if prompt_field in chunk:
                chunk.pop(prompt_field, None)
                removed_prompt_material = True
        if removed_prompt_material:
            chunk["rerunPromptStored"] = False
    return payload


def _compare_payload_is_usable(output_path: Path, payload: dict[str, Any]) -> bool:
    try:
        if not output_path.exists() or not output_path.is_file() or output_path.stat().st_size <= 0:
            return False
    except OSError:
        return False
    chunks = payload.get("chunks")
    if not isinstance(chunks, list) or not chunks:
        return False
    expected_chunk_count = _coerce_optional_nonnegative_int(payload.get("chunkCount"))
    expected_paragraph_count = _coerce_optional_nonnegative_int(payload.get("paragraphCount"))
    if expected_chunk_count is None or expected_chunk_count <= 0 or expected_chunk_count != len(chunks):
        return False
    if expected_paragraph_count is not None and expected_paragraph_count <= 0:
        return False
    manifest = _load_manifest_for_compare(output_path, payload)
    if manifest is not None:
        if manifest.chunk_count != expected_chunk_count:
            return False
        if expected_paragraph_count is not None and manifest.paragraph_count != expected_paragraph_count:
            return False
    return True


def _load_manifest_for_compare(output_path: Path, compare_payload: dict[str, Any]):
    manifest_path: Path | None = None
    raw_manifest_path = compare_payload.get("manifestPath")
    if isinstance(raw_manifest_path, str) and raw_manifest_path.strip():
        manifest_path = normalize_path(Path(raw_manifest_path))
    if manifest_path is None:
        manifest_path = _find_manifest_path_for_output(output_path, include_records=False)
    if manifest_path is None or not manifest_path.exists():
        return None
    try:
        return load_manifest(manifest_path)
    except Exception:
        return None


def _normalize_contract_text(value: str) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n")


def _assert_docx_targeted_rerun_contract(
    output_path: Path,
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
) -> dict[str, Any] | None:
    """Fail before any model call if a DOCX rerun lost frozen-body provenance."""

    body_map = _load_body_map_for_output(output_path)
    origin_bundle = _find_origin_docx_for_output(output_path)
    if body_map is None:
        if origin_bundle is not None:
            raise ValueError("DOCX 定点重跑已阻断：冻结 body map 缺失，无法证明目标块只属于正文。")
        return None

    source_path = normalize_path(Path(str(body_map.source_path)))
    if source_path.suffix.lower() != ".docx":
        raise ValueError("DOCX targeted rerun contract rejected a non-DOCX body-map source.")
    if origin_bundle is not None and normalize_path(origin_bundle[0]) != source_path:
        raise ValueError("DOCX 定点重跑已阻断：body map 与历史记录中的源 Word 不一致。")

    current_contract = build_document_edit_contract(
        source_path,
        body_map=body_map,
        candidate_texts=body_map.current_texts(),
        stage="targeted_rerun_preflight",
    )
    assert_document_edit_contract_ready(current_contract, label="DOCX 定点重跑")

    raw_input_path = str(compare_payload.get("inputPath", "") or "").strip()
    if not raw_input_path:
        raise ValueError("DOCX 定点重跑已阻断：compare 缺少冻结轮次输入路径。")
    input_path = _resolve_api_path(
        raw_input_path,
        allowed_roots=API_OUTPUT_ALLOWED_ROOTS,
        label="Round input path",
    )
    if not input_path.exists() or not input_path.is_file():
        raise ValueError("DOCX 定点重跑已阻断：冻结轮次输入文件不存在。")

    input_body_map = _load_body_map_for_output(input_path)
    round_number = int(compare_payload.get("round", 0) or 0)
    if input_body_map is not None:
        if normalize_path(Path(str(input_body_map.source_path))) != source_path:
            raise ValueError("DOCX 定点重跑已阻断：上一轮正文映射属于另一份源文档。")
        input_contract = build_document_edit_contract(
            source_path,
            body_map=input_body_map,
            candidate_texts=input_body_map.current_texts(),
            stage="targeted_rerun_input_preflight",
        )
        assert_document_edit_contract_ready(input_contract, label="DOCX 定点重跑输入")
        expected_input_text = "\n\n".join(input_body_map.current_texts())
        actual_input_text = input_path.read_text(encoding="utf-8")
        if _normalize_contract_text(actual_input_text) != _normalize_contract_text(expected_input_text):
            raise ValueError("DOCX 定点重跑已阻断：上一轮正文文件与冻结 body map 不一致。")
    else:
        if round_number != 1:
            raise ValueError("DOCX 定点重跑已阻断：上一轮正文映射缺失，无法证明模型输入只包含正文。")
        input_contract = build_document_edit_contract(
            source_path,
            extracted_text_path=input_path,
            stage="targeted_rerun_input_preflight",
        )
        assert_document_edit_contract_ready(input_contract, label="DOCX 定点重跑输入")

    manifest = _load_manifest_for_compare(output_path, compare_payload)
    if manifest is None:
        raise ValueError("DOCX 定点重跑已阻断：chunk manifest 缺失或无法读取。")
    if manifest.paragraph_count != len(body_map.units):
        raise ValueError("DOCX 定点重跑已阻断：manifest 自然段数与冻结正文单元数不一致。")

    identity_chunks = {chunk.chunk_id: chunk.text for chunk in manifest.chunks}
    restored_input = restore_text_from_chunks(manifest, identity_chunks)
    actual_input = input_path.read_text(encoding="utf-8")
    if _normalize_contract_text(restored_input) != _normalize_contract_text(actual_input):
        raise ValueError("DOCX 定点重跑已阻断：chunk manifest 无法还原冻结轮次输入。")

    chunk_id = str(target_chunk.get("chunkId", "") or "").strip()
    manifest_chunk = next((chunk for chunk in manifest.chunks if chunk.chunk_id == chunk_id), None)
    if manifest_chunk is None:
        raise ValueError("DOCX 定点重跑已阻断：目标块不属于冻结 chunk manifest。")
    try:
        paragraph_index = int(target_chunk.get("paragraphIndex"))
        chunk_index = int(target_chunk.get("chunkIndex"))
    except (TypeError, ValueError) as exc:
        raise ValueError("DOCX 定点重跑已阻断：目标块位置无效。") from exc
    if paragraph_index != manifest_chunk.paragraph_index or chunk_index != manifest_chunk.chunk_index:
        raise ValueError("DOCX 定点重跑已阻断：目标块位置与冻结 chunk manifest 不一致。")
    if str(target_chunk.get("inputText", "")) != manifest_chunk.text:
        raise ValueError("DOCX 定点重跑已阻断：目标块输入与冻结 chunk manifest 不一致。")
    return current_contract


def _resolve_docx_target_model_format_anchors(
    output_path: Path,
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
    effective_input_text: str,
) -> list[str]:
    """Bind one targeted model call to its frozen Word styled-text anchors.

    A body-map unit represents one editable Word paragraph, while a manifest
    may split that paragraph into several chunks.  First prove that every
    styled anchor belongs to exactly one frozen chunk in the target paragraph;
    then pass only anchors that occur exactly once in this call's fresh,
    review-materialized text.  A formerly target-local anchor may not disappear
    or become ambiguous between review materialization and model execution.
    """

    body_map = _load_body_map_for_output(output_path)
    if body_map is None:
        return []
    try:
        paragraph_index = int(target_chunk.get("paragraphIndex"))
    except (TypeError, ValueError) as exc:
        raise ValueError("DOCX 定点重跑已阻断：目标段落位置无效，无法绑定 Word 格式锚点。") from exc
    if paragraph_index < 0 or paragraph_index >= len(body_map.units):
        raise ValueError("DOCX 定点重跑已阻断：目标段落超出冻结 body map，无法绑定 Word 格式锚点。")

    anchors = get_body_map_unit_model_format_anchors(body_map.units[paragraph_index])
    if not anchors:
        return []
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        raise ValueError("DOCX 定点重跑已阻断：compare 缺少冻结 chunks，无法绑定 Word 格式锚点。")
    paragraph_chunks = [
        item
        for item in raw_chunks
        if isinstance(item, dict)
        and item.get("paragraphIndex") == paragraph_index
        and isinstance(item.get("inputText"), str)
    ]
    if not paragraph_chunks:
        raise ValueError("DOCX 定点重跑已阻断：目标段落没有冻结 chunk，无法绑定 Word 格式锚点。")

    target_chunk_id = str(target_chunk.get("chunkId", "") or "").strip()
    selected: list[str] = []
    for anchor in anchors:
        frozen_owners: list[str] = []
        frozen_occurrence_count = 0
        for item in paragraph_chunks:
            occurrence_count = str(item.get("inputText", "") or "").count(anchor)
            frozen_occurrence_count += occurrence_count
            if occurrence_count == 1:
                frozen_owners.append(str(item.get("chunkId", "") or "").strip())
        if frozen_occurrence_count != 1 or len(frozen_owners) != 1 or not frozen_owners[0]:
            raise ValueError(
                "DOCX 定点重跑已阻断：Word 格式锚点在冻结 chunk 中缺失、跨块或存在歧义。"
            )

        effective_occurrence_count = effective_input_text.count(anchor)
        if effective_occurrence_count > 1:
            raise ValueError("DOCX 定点重跑已阻断：Word 格式锚点在当前模型输入中存在歧义。")
        if effective_occurrence_count == 1:
            selected.append(anchor)
        elif frozen_owners[0] == target_chunk_id:
            raise ValueError("DOCX 定点重跑已阻断：当前模型输入缺少目标 chunk 的 Word 格式锚点。")
    return selected


def _mask_format_anchors_in_prompt_context(
    text: str,
    protected_chunk: Any,
    anchors: Sequence[str],
) -> str:
    """Replace raw styled text in non-input prompt context with its placeholder."""

    masked = str(text or "")
    for anchor in anchors:
        placeholder = next(
            (
                token
                for token, original in protected_chunk.tokens.items()
                if str(original) == anchor or anchor in str(original)
            ),
            "",
        )
        if not placeholder:
            raise ValueError("DOCX 定点重跑已阻断：Word 格式锚点没有对应的模型占位符。")
        # Do not repeat the live placeholder in context sections: seeing the
        # same token in both "previous output" and "input" can encourage a
        # model to emit it twice.  The context only needs a non-content marker;
        # the one authoritative @@FYADR_*@@ token remains in [INPUT TEXT].
        masked = masked.replace(anchor, "[FROZEN_WORD_FORMAT_TEXT]")
    return masked


def _validate_docx_custom_review_format_anchors(
    output_path: Path,
    compare_payload: dict[str, Any],
    decisions: dict[str, Any],
) -> None:
    """Reject a custom review edit that detaches frozen Word styled text."""

    if _load_body_map_for_output(output_path) is None:
        return
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        raise ValueError("DOCX 审阅保存已阻断：compare 缺少冻结 chunks。")
    chunks = {
        str(item.get("chunkId", "") or ""): item
        for item in raw_chunks
        if isinstance(item, dict) and str(item.get("chunkId", "") or "")
    }
    for chunk_id, decision in decisions.items():
        if not isinstance(decision, dict) or str(decision.get("mode", "") or "") != "custom":
            continue
        target_chunk = chunks.get(str(chunk_id))
        if target_chunk is None:
            continue
        custom_text = str(decision.get("text", "") or "")
        anchors = _resolve_docx_target_model_format_anchors(
            output_path,
            compare_payload,
            target_chunk,
            custom_text,
        )
        validate_immutable_text_anchors(
            str(target_chunk.get("inputText", "") or ""),
            custom_text,
            anchors,
            str(chunk_id),
        )


def _normalize_review_decision_value(decision: Any) -> str | dict[str, Any]:
    if isinstance(decision, dict):
        mode = str(decision.get("mode", "")).strip().lower()
        text = decision.get("text")
        if mode == "custom" and isinstance(text, str) and text.strip():
            source = str(decision.get("source", "")).strip()
            if source == "rejected_candidate":
                source = "failed_output"
            return {
                "mode": "custom",
                "text": text,
                "source": source,
                "confirmed": bool(decision.get("confirmed")),
                "attempt": decision.get("attempt"),
                "error": str(decision.get("error", "")).strip(),
            }
    decision_text = str(decision)
    if decision_text in {"source", "source_confirmed"}:
        return decision_text
    if decision_text in {"rewrite", "rewrite_confirmed"}:
        return decision_text
    return "rewrite"


def _default_export_decision_for_chunk(chunk: dict[str, Any]) -> str:
    """Prefer source for unresolved high-risk chunks instead of silently exporting rewrite."""
    if chunk.get("rateAuditStrategyReviewRequired") is True:
        return "source"
    selection = chunk.get("candidateSelection")
    if isinstance(selection, dict) and selection.get("publishedRewrite") is False:
        # "baseline" means the already accepted review-materialized text. It
        # can be the frozen source or an earlier confirmed/custom rewrite.
        return (
            "source"
            if str(chunk.get("outputText", "") or "") == str(chunk.get("inputText", "") or "")
            else "rewrite"
        )
    quality = chunk.get("quality") if isinstance(chunk.get("quality"), dict) else {}
    flags = {str(flag) for flag in (quality.get("flags") or []) if str(flag).strip()}
    if chunk.get("fallbackMode") == "source" or "source_fallback" in flags or "targeted_rerun_fallback" in flags:
        return "source"
    if chunk.get("rerunStatus") == "fallback" or chunk.get("rerunFallbackMode"):
        return "source"
    failed_attempts = chunk.get("failedAttempts")
    if isinstance(failed_attempts, list) and failed_attempts and bool(quality.get("needsReview")):
        return "source"
    hard_flags = {
        "academic_register_drift",
        "citation_missing",
        "machine_style_drift",
        "machine_like_expression",
        "repeated_content",
        "template_phrase_density",
        "abstract_padding_density",
    }
    if flags & hard_flags:
        return "source"
    return "rewrite"


def _select_review_text(input_text: Any, output_text: Any, decision: Any, *, chunk: dict[str, Any] | None = None) -> str:
    source_text = input_text if isinstance(input_text, str) else ""
    rewrite_text = output_text if isinstance(output_text, str) else ""
    if isinstance(chunk, dict) and chunk.get("rateAuditStrategyReviewRequired") is True:
        normalized_strategy_decision = _normalize_review_decision_value(decision) if decision is not None else "source"
        if normalized_strategy_decision == "rewrite_confirmed":
            return rewrite_text
        if (
            isinstance(normalized_strategy_decision, dict)
            and normalized_strategy_decision.get("mode") == "custom"
            and normalized_strategy_decision.get("confirmed") is True
        ):
            return str(normalized_strategy_decision.get("text", "") or "")
        return source_text
    if decision is None and chunk is not None:
        decision = _default_export_decision_for_chunk(chunk)
    normalized_decision = _normalize_review_decision_value(decision)
    if isinstance(normalized_decision, dict):
        if normalized_decision.get("source") == "failed_output" and normalized_decision.get("confirmed") is not True:
            return source_text
        return str(normalized_decision.get("text", ""))
    return source_text if normalized_decision in {"source", "source_confirmed"} else rewrite_text


def _release_gate_normalize_text(input_text: str, text: str) -> str:
    """Use the same deterministic text form that review materialization emits."""

    return normalize_chunk_output(input_text, text)


def _release_gate_exact_int(value: Any) -> int | None:
    if isinstance(value, bool) or not isinstance(value, int):
        return None
    return value


def _canonical_json_bytes(value: object) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _source_pattern_profile_sha_from_evidence(value: object) -> str:
    if not isinstance(value, dict) or value.get("contextScope") != "document":
        return ""
    binding = value.get("binding")
    if not isinstance(binding, dict):
        return ""
    profile_sha256 = str(binding.get("sourceProfileSha256", "") or "").strip().lower()
    return profile_sha256 if re.fullmatch(r"[0-9a-f]{64}", profile_sha256) else ""


def _candidate_selection_profile_references(
    selection: object,
) -> tuple[set[str], list[str]]:
    """Collect document-profile references from one current v2 selection."""

    if not isinstance(selection, dict) or (
        selection.get("schema") != CANDIDATE_SELECTION_SCHEMA
        or selection.get("schemaVersion") != CANDIDATE_SELECTION_VERSION
    ):
        return set(), []
    evidence_values: list[object] = [selection.get("resultSourceRelativeStyleDelta")]
    raw_candidates = selection.get("candidates")
    if isinstance(raw_candidates, list):
        evidence_values.extend(
            candidate.get("sourceRelativeStyleDelta")
            for candidate in raw_candidates
            if isinstance(candidate, dict)
        )
    references: set[str] = set()
    issues: list[str] = []
    for evidence in evidence_values:
        if not isinstance(evidence, dict):
            issues.append("source_relative_style_evidence_missing")
            continue
        context_scope = evidence.get("contextScope")
        if context_scope == "document":
            profile_sha256 = _source_pattern_profile_sha_from_evidence(evidence)
            if not profile_sha256:
                issues.append("source_pattern_profile_binding_invalid")
            else:
                references.add(profile_sha256)
        elif context_scope != "local":
            issues.append("source_relative_style_context_invalid")
    return references, list(dict.fromkeys(issues))


def _compare_source_pattern_profile_references(
    compare_payload: dict[str, Any],
) -> tuple[set[str], list[str]]:
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        return set(), ["release_compare_chunks_invalid"]
    references: set[str] = set()
    issues: list[str] = []
    for raw_chunk in raw_chunks:
        if not isinstance(raw_chunk, dict):
            issues.append("release_chunk_invalid")
            continue
        chunk_references, chunk_issues = _candidate_selection_profile_references(
            raw_chunk.get("candidateSelection")
        )
        references.update(chunk_references)
        issues.extend(chunk_issues)
    return references, list(dict.fromkeys(issues))


def _normalize_source_pattern_profile_registry(
    value: object,
) -> tuple[dict[str, dict[str, object]], list[str]]:
    """Validate a bounded content-addressed registry without truncation."""

    if value is None:
        return {}, []
    if not isinstance(value, dict):
        return {}, ["source_pattern_profile_registry_invalid"]
    if len(value) > MAX_SOURCE_PATTERN_PROFILE_REGISTRY_ENTRIES:
        return {}, ["source_pattern_profile_registry_entry_limit_exceeded"]
    normalized: dict[str, dict[str, object]] = {}
    total_bytes = 0
    for raw_key, raw_profile in value.items():
        profile_sha256 = str(raw_key or "").strip().lower()
        if (
            not isinstance(raw_key, str)
            or re.fullmatch(r"[0-9a-f]{64}", profile_sha256) is None
            or not source_pattern_profile_valid(raw_profile)
            or not isinstance(raw_profile, dict)
            or str(raw_profile.get("profileSha256", "") or "") != profile_sha256
        ):
            return {}, ["source_pattern_profile_registry_entry_invalid"]
        profile_size = len(_canonical_json_bytes(raw_profile))
        if profile_size > MAX_SOURCE_PATTERN_PROFILE_BYTES:
            return {}, ["source_pattern_profile_size_limit_exceeded"]
        total_bytes += profile_size
        if total_bytes > MAX_SOURCE_PATTERN_PROFILE_REGISTRY_BYTES:
            return {}, ["source_pattern_profile_registry_size_limit_exceeded"]
        normalized[profile_sha256] = dict(raw_profile)
    return normalized, []


def _refresh_source_pattern_profile_registry(
    compare_payload: dict[str, Any],
    *,
    global_style_profile: dict[str, object],
) -> None:
    """Register the current context and garbage-collect unreferenced profiles."""

    document_profile = (
        global_style_profile.get("documentPatternBaseline")
        if isinstance(global_style_profile, dict)
        else None
    )
    if not source_pattern_profile_valid(document_profile) or not isinstance(document_profile, dict):
        raise ValueError("Source pattern profile is missing or invalid.")
    profile_sha256 = str(document_profile.get("profileSha256", "") or "")
    profile_size = len(_canonical_json_bytes(document_profile))
    if profile_size > MAX_SOURCE_PATTERN_PROFILE_BYTES:
        raise ValueError("Source pattern profile exceeds the per-profile size limit.")

    registry, registry_issues = _normalize_source_pattern_profile_registry(
        compare_payload.get("sourcePatternProfiles")
    )
    if registry_issues:
        raise ValueError("Source pattern profile registry is invalid: " + ", ".join(registry_issues))
    registry[profile_sha256] = dict(document_profile)

    references, reference_issues = _compare_source_pattern_profile_references(compare_payload)
    if reference_issues:
        raise ValueError("Candidate profile binding is invalid: " + ", ".join(reference_issues))
    if len(references) > MAX_SOURCE_PATTERN_PROFILE_REGISTRY_ENTRIES:
        raise ValueError("Source pattern profile registry entry limit exceeded.")
    missing = sorted(reference for reference in references if reference not in registry)
    if missing:
        raise ValueError("Candidate selection references an unavailable source pattern profile.")

    retained = {reference: registry[reference] for reference in sorted(references)}
    retained_size = sum(len(_canonical_json_bytes(profile)) for profile in retained.values())
    if retained_size > MAX_SOURCE_PATTERN_PROFILE_REGISTRY_BYTES:
        raise ValueError("Source pattern profile registry size limit exceeded.")
    if retained:
        compare_payload["sourcePatternProfiles"] = retained
    else:
        compare_payload.pop("sourcePatternProfiles", None)


def _assess_compare_candidate_document_delta(
    compare_payload: dict[str, Any],
    *,
    result_texts_by_chunk: dict[str, str] | None = None,
    output_overrides: dict[str, str] | None = None,
) -> dict[str, object]:
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        raise ValueError("Compare data is invalid: chunks missing.")
    baseline_texts: list[str] = []
    result_texts: list[str] = []
    seen_chunk_ids: set[str] = set()
    overrides = output_overrides or {}
    for raw_chunk in raw_chunks:
        if not isinstance(raw_chunk, dict):
            raise ValueError("Compare data contains an invalid chunk.")
        chunk_id = str(raw_chunk.get("chunkId", "") or "").strip()
        input_text = raw_chunk.get("inputText")
        output_text = overrides.get(
            chunk_id,
            (
                result_texts_by_chunk.get(chunk_id)
                if isinstance(result_texts_by_chunk, dict)
                else raw_chunk.get("outputText")
            ),
        )
        if (
            not chunk_id
            or chunk_id in seen_chunk_ids
            or not isinstance(input_text, str)
            or not input_text.strip()
            or not isinstance(output_text, str)
            or not output_text.strip()
        ):
            raise ValueError("Compare document-delta text binding is invalid.")
        seen_chunk_ids.add(chunk_id)
        baseline_texts.append(input_text.strip())
        result_texts.append(output_text.strip())
    if set(overrides) - seen_chunk_ids:
        raise ValueError("Document-delta override references an unknown chunk.")
    if isinstance(result_texts_by_chunk, dict) and set(result_texts_by_chunk) != seen_chunk_ids:
        raise ValueError("Document-delta effective text map is incomplete.")
    return assess_source_relative_document_delta(baseline_texts, result_texts)


def _arbitrate_app_candidate_document_delta(
    *,
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
    selection_result: dict[str, Any],
    candidate_baseline_text: str,
    effective_document_texts: dict[str, str],
    round_number: int,
) -> dict[str, object]:
    """Apply the same cumulative document gate to targeted selector output."""

    chunk_id = str(target_chunk.get("chunkId", "") or "").strip()
    selection_event = selection_result.get("selectionEvent")
    if not chunk_id or not isinstance(selection_event, dict):
        raise ValueError("Targeted candidate selection has no document binding.")
    proposed_text = str(
        selection_result.get("text", candidate_baseline_text) or candidate_baseline_text
    ).strip()
    document_delta = _assess_compare_candidate_document_delta(
        compare_payload,
        result_texts_by_chunk=effective_document_texts,
        output_overrides={chunk_id: proposed_text},
    )
    if (
        selection_event.get("publishedRewrite") is True
        and not source_relative_document_delta_passed(document_delta)
    ):
        raw_candidates = selection_result.get("candidates")
        candidates = [
            candidate
            for candidate in raw_candidates
            if isinstance(candidate, dict)
        ] if isinstance(raw_candidates, list) else []
        baseline_candidates = [
            candidate
            for candidate in candidates
            if str(candidate.get("candidateId", "") or "") == "baseline"
            and str(candidate.get("origin", "") or "") == "baseline"
        ]
        if len(baseline_candidates) != 1:
            raise ValueError("Targeted document arbitration lost its authoritative baseline.")
        reason_codes = [
            *list(selection_event.get("reasonCodes") or []),
            DOCUMENT_PATTERN_ACCUMULATION_BLOCKED,
            *list(document_delta.get("blockingIssueCodes") or []),
        ]
        replacement = _build_candidate_selection_event(
            chunk_id=chunk_id,
            round_number=round_number,
            candidates=candidates,
            selected=baseline_candidates[0],
            reason_codes=list(
                dict.fromkeys(str(code) for code in reason_codes if str(code or "").strip())
            ),
            conditional_retry_count=int(
                selection_event.get("conditionalRetryCount", 0) or 0
            ),
        )
        replacement["documentArbitration"] = {
            "decision": "baseline_preserved",
            "reasonCode": DOCUMENT_PATTERN_ACCUMULATION_BLOCKED,
            "rejectedDocumentDelta": document_delta,
        }
        selection_result["selectionEvent"] = replacement
        selection_result["text"] = candidate_baseline_text.strip()
        _replace_chunk_candidate_selection_event(
            compare_payload,
            target_chunk,
            replacement,
        )
        _append_compare_validation_event(
            compare_payload,
            {
                "event": "document-pattern-delta-arbitration",
                "round": round_number,
                "chunkId": chunk_id,
                "decision": "baseline_preserved",
                "reasonCodes": list(document_delta.get("blockingIssueCodes") or []),
                "documentDelta": document_delta,
            },
        )
        proposed_text = candidate_baseline_text.strip()
        document_delta = _assess_compare_candidate_document_delta(
            compare_payload,
            result_texts_by_chunk=effective_document_texts,
            output_overrides={chunk_id: proposed_text},
        )
    if not source_relative_document_delta_passed(document_delta):
        raise ValueError("Targeted document pattern-delta arbitration did not converge.")
    raw_candidate_document_delta = _assess_compare_candidate_document_delta(
        compare_payload,
        output_overrides={chunk_id: proposed_text},
    )
    compare_payload["sourceRelativeDocumentDelta"] = raw_candidate_document_delta
    return document_delta


def _release_gate_current_text_checks(
    input_text: str,
    materialized_text: str,
    chunk_id: str,
) -> list[str]:
    issues: list[str] = []
    try:
        validate_chunk_output(input_text, materialized_text, chunk_id)
    except Exception:
        issues.append("hard_validation_failed")
    try:
        readability = assess_academic_readability_delta(input_text, materialized_text)
    except Exception:
        issues.append("academic_readability_assessment_failed")
    else:
        if readability.get("ok") is not True:
            issues.append("academic_readability_delta_failed")
            for value in readability.get("issueCodes", []):
                code = str(value or "").strip()
                if code and code not in issues:
                    issues.append(code)
    return issues


def _assess_document_release_chunk(
    chunk: dict[str, Any],
    decision: Any,
    *,
    source_pattern_profile: dict[str, object] | None = None,
    source_pattern_profiles: dict[str, dict[str, object]] | None = None,
    source_pattern_profile_registry_issues: Sequence[str] = (),
    document_context_available: bool = False,
) -> dict[str, Any]:
    """Return text-free release evidence for one review-materialized chunk."""

    chunk_id = str(chunk.get("chunkId", "") or "").strip()
    raw_input = chunk.get("inputText")
    raw_output = chunk.get("outputText")
    if not chunk_id or not isinstance(raw_input, str) or not isinstance(raw_output, str):
        return {
            "ok": False,
            "chunkId": chunk_id,
            "mode": "invalid_chunk",
            "issueCodes": ["release_chunk_invalid"],
        }

    try:
        input_text = _release_gate_normalize_text(raw_input, raw_input)
        output_text = _release_gate_normalize_text(raw_input, raw_output)
    except Exception:
        return {
            "ok": False,
            "chunkId": chunk_id,
            "mode": "invalid_chunk",
            "issueCodes": ["release_text_normalization_failed"],
        }

    normalized_decision = (
        _normalize_review_decision_value(decision)
        if decision is not None
        else _default_export_decision_for_chunk(chunk)
    )
    strategy_review_pending = chunk.get("rateAuditStrategyReviewRequired") is True

    if isinstance(normalized_decision, dict):
        custom_source = str(normalized_decision.get("source", "") or "").strip().lower()
        if custom_source in {"failed_output", "rejected_candidate"}:
            return {
                "ok": False,
                "chunkId": chunk_id,
                "mode": "manual_custom",
                "issueCodes": ["failed_output_custom_forbidden"],
                "outputTextSha256": _sha256_text(output_text),
            }
        custom_is_materialized = bool(
            not strategy_review_pending
            or normalized_decision.get("confirmed") is True
        )
        if not custom_is_materialized:
            return {
                "ok": True,
                "chunkId": chunk_id,
                "mode": "source",
                "issueCodes": [],
                "materializedTextSha256": _sha256_text(input_text),
                "outputTextSha256": _sha256_text(output_text),
            }

        raw_custom_text = str(normalized_decision.get("text", "") or "")
        try:
            custom_text = _release_gate_normalize_text(raw_input, raw_custom_text)
        except Exception:
            custom_text = ""
        issues: list[str] = []
        if not custom_text:
            issues.append("custom_text_invalid")
        custom_sha256 = _sha256_text(custom_text)
        if custom_text and custom_text in {input_text, output_text}:
            issues.append("custom_text_not_independent")

        raw_failed_attempts = chunk.get("failedAttempts")
        if not isinstance(raw_failed_attempts, list):
            raw_failed_attempts = chunk.get("rejectedCandidates")
        failed_attempts = raw_failed_attempts[-4:] if isinstance(raw_failed_attempts, list) else []
        for failed_attempt in failed_attempts:
            if not isinstance(failed_attempt, dict):
                continue
            failed_text_raw = str(failed_attempt.get("outputText", "") or "")
            try:
                failed_text = _release_gate_normalize_text(raw_input, failed_text_raw)
            except Exception:
                failed_text = failed_text_raw.strip()
            declared_hashes = {
                str(failed_attempt.get(key, "") or "").strip()
                for key in ("outputTextSha256", "textSha256", "sha256")
            }
            declared_hashes.discard("")
            if custom_text and (
                custom_text == failed_text
                or custom_sha256 == _sha256_text(failed_text)
                or custom_sha256 in declared_hashes
            ):
                issues.append("custom_matches_failed_attempt")
                break

        selection = chunk.get("candidateSelection")
        if isinstance(selection, dict):
            candidates = selection.get("candidates")
            if isinstance(candidates, list):
                model_hashes = {
                    str(candidate.get("textSha256", "") or "").strip()
                    for candidate in candidates
                    if isinstance(candidate, dict)
                    and str(candidate.get("origin", "") or "") == "model"
                    and str(candidate.get("textSha256", "") or "").strip()
                }
                if custom_sha256 in model_hashes:
                    issues.append("custom_matches_model_candidate")

        if custom_text:
            issues.extend(_release_gate_current_text_checks(raw_input, custom_text, chunk_id))
        return {
            "ok": not issues,
            "chunkId": chunk_id,
            "mode": "manual_custom",
            "issueCodes": list(dict.fromkeys(issues)),
            "materializedTextSha256": custom_sha256,
            "outputTextSha256": _sha256_text(output_text),
        }

    materializes_source = bool(
        normalized_decision in {"source", "source_confirmed"}
        or (strategy_review_pending and normalized_decision != "rewrite_confirmed")
    )
    if materializes_source:
        # Source selection is always a safe escape hatch from stale or forged
        # candidate metadata, so users can repair an already-invalid sidecar.
        return {
            "ok": True,
            "chunkId": chunk_id,
            "mode": "source",
            "issueCodes": [],
            "materializedTextSha256": _sha256_text(input_text),
            "outputTextSha256": _sha256_text(output_text),
        }

    selection = chunk.get("candidateSelection")
    if not isinstance(selection, dict):
        if output_text == input_text:
            frozen_issues = _release_gate_current_text_checks(raw_input, output_text, chunk_id)
            return {
                "ok": not frozen_issues,
                "chunkId": chunk_id,
                "mode": "frozen_identity",
                "issueCodes": list(dict.fromkeys(frozen_issues)),
                "materializedTextSha256": _sha256_text(output_text),
                "outputTextSha256": _sha256_text(output_text),
            }
        return {
            "ok": False,
            "chunkId": chunk_id,
            "mode": "model_output",
            "issueCodes": ["candidate_selection_missing"],
            "materializedTextSha256": _sha256_text(output_text),
            "outputTextSha256": _sha256_text(output_text),
        }

    published_rewrite = selection.get("publishedRewrite")
    mode = "published_rewrite" if published_rewrite is True else "preserved_baseline"
    issues = []
    if (
        selection.get("event") != "candidate-selection"
        or selection.get("schema") != CANDIDATE_SELECTION_SCHEMA
        or selection.get("schemaVersion") != CANDIDATE_SELECTION_VERSION
        or str(selection.get("chunkId", "") or "") != chunk_id
    ):
        issues.append("candidate_selection_schema_invalid")

    raw_candidate_baseline = chunk.get("candidateBaselineText")
    candidate_baseline_text = input_text
    if not isinstance(raw_candidate_baseline, str) or not raw_candidate_baseline.strip():
        issues.append("candidate_baseline_text_missing_or_invalid")
    else:
        # The selector hashes the exact review-materialized baseline after
        # outer trimming.  Do not re-normalize its internal line structure
        # against frozen inputText: that would silently turn B back into an
        # A-shaped value for multiline/custom-review targeted reruns.
        if raw_candidate_baseline != raw_candidate_baseline.strip():
            issues.append("candidate_baseline_text_not_canonical")
        else:
            candidate_baseline_text = raw_candidate_baseline
    candidate_baseline_sha256 = _sha256_text(candidate_baseline_text)

    decision_name = str(selection.get("decision", "") or "")
    selected_origin = str(selection.get("selectedOrigin", "") or "")
    selected_candidate_id = str(selection.get("selectedCandidateId", "") or "")
    if published_rewrite is True:
        if (
            decision_name != "generated_selected"
            or selected_origin != "model"
            or selection.get("runFailed") is not False
        ):
            issues.append("candidate_selection_publication_invalid")
    elif published_rewrite is False:
        if (
            decision_name != "preserved_baseline"
            or selected_origin != "baseline"
            or selected_candidate_id != "baseline"
            or selection.get("runFailed") is not False
        ):
            issues.append("candidate_selection_preservation_invalid")
    else:
        issues.append("candidate_selection_publication_invalid")

    candidates = selection.get("candidates")
    baseline_candidates = (
        [
            candidate
            for candidate in candidates
            if isinstance(candidate, dict)
            and str(candidate.get("candidateId", "") or "") == "baseline"
        ]
        if isinstance(candidates, list)
        else []
    )
    baseline_candidate = baseline_candidates[0] if len(baseline_candidates) == 1 else None
    if baseline_candidate is None:
        issues.append("baseline_candidate_missing_or_duplicated")
    else:
        baseline_source_relative = baseline_candidate.get("sourceRelativeStyleDelta")
        baseline_source_binding = (
            baseline_source_relative.get("binding")
            if isinstance(baseline_source_relative, dict)
            else None
        )
        if (
            str(baseline_candidate.get("origin", "") or "") != "baseline"
            or baseline_candidate.get("changedFromBaseline") is not False
            or baseline_candidate.get("hardValid") is not True
            or baseline_candidate.get("safetyEligible") is not True
            or str(baseline_candidate.get("textSha256", "") or "")
            != candidate_baseline_sha256
            or not source_relative_style_delta_passed(baseline_source_relative)
            or baseline_candidate.get("sourceRelativeStyleGuardPassed") is not True
            or not isinstance(baseline_source_binding, dict)
            or str(baseline_source_binding.get("baselineTextSha256", "") or "")
            != candidate_baseline_sha256
            or str(baseline_source_binding.get("candidateTextSha256", "") or "")
            != candidate_baseline_sha256
        ):
            issues.append("baseline_candidate_evidence_invalid")
    selected_candidates = (
        [
            candidate
            for candidate in candidates
            if isinstance(candidate, dict)
            and str(candidate.get("candidateId", "") or "") == selected_candidate_id
        ]
        if isinstance(candidates, list)
        else []
    )
    selected_candidate = selected_candidates[0] if len(selected_candidates) == 1 else None
    if selected_candidate is None:
        issues.append("selected_candidate_missing_or_duplicated")
    else:
        readability_evidence = selected_candidate.get("academicReadabilityDelta")
        readability_issue_codes = selected_candidate.get("readabilityIssueCodes")
        if (
            str(selected_candidate.get("origin", "") or "") != selected_origin
            or selected_candidate.get("hardValid") is not True
            or selected_candidate.get("safetyEligible") is not True
            or selected_candidate.get("factualGuardPassed") is not True
            or selected_candidate.get("readabilityGuardPassed") is not True
            or not isinstance(readability_evidence, dict)
            or readability_evidence.get("schema") != ACADEMIC_READABILITY_DELTA_SCHEMA
            or readability_evidence.get("schemaVersion") != ACADEMIC_READABILITY_DELTA_VERSION
            or readability_evidence.get("ok") is not True
            or not isinstance(readability_issue_codes, list)
            or bool(readability_issue_codes)
            or not isinstance(readability_evidence.get("issueCodes"), list)
            or readability_evidence.get("issueCodes") != readability_issue_codes
        ):
            issues.append("selected_candidate_safety_invalid")
        expected_changed = published_rewrite is True
        if selected_candidate.get("changedFromBaseline") is not expected_changed:
            issues.append("selected_candidate_origin_invalid")
        source_relative = selected_candidate.get("sourceRelativeStyleDelta")
        source_relative_binding = (
            source_relative.get("binding") if isinstance(source_relative, dict) else None
        )
        if (
            not source_relative_style_delta_passed(source_relative)
            or selected_candidate.get("sourceRelativeStyleGuardPassed") is not True
            or not isinstance(source_relative_binding, dict)
            or str(source_relative_binding.get("baselineTextSha256", "") or "")
            != candidate_baseline_sha256
            or str(source_relative_binding.get("candidateTextSha256", "") or "")
            != str(selected_candidate.get("textSha256", "") or "")
        ):
            issues.append("selected_candidate_source_relative_style_invalid")

    selected_sha256 = str(selection.get("selectedTextSha256", "") or "").strip()
    result_sha256 = str(selection.get("resultTextSha256", "") or "").strip()
    published_sha256 = str(selection.get("publishedTextSha256", "") or "").strip()
    output_sha256 = _sha256_text(output_text)
    if (
        not selected_sha256
        or selected_candidate is None
        or str(selected_candidate.get("textSha256", "") or "") != selected_sha256
    ):
        issues.append("selected_candidate_hash_mismatch")
    if result_sha256 != output_sha256:
        issues.append("result_output_hash_mismatch")
    result_source_relative = selection.get("resultSourceRelativeStyleDelta")
    result_source_relative_binding = (
        result_source_relative.get("binding")
        if isinstance(result_source_relative, dict)
        else None
    )
    if (
        not source_relative_style_delta_passed(result_source_relative)
        or not isinstance(result_source_relative_binding, dict)
        or str(result_source_relative_binding.get("baselineTextSha256", "") or "")
        != candidate_baseline_sha256
        or str(result_source_relative_binding.get("candidateTextSha256", "") or "") != output_sha256
    ):
        issues.append("result_source_relative_style_invalid")

    # Direct unit callers may still provide one profile.  Production document
    # release always supplies the bounded content-addressed registry and must
    # resolve the exact profile SHA bound by the persisted result evidence.
    available_profiles = source_pattern_profiles
    if available_profiles is None and source_pattern_profile is not None:
        legacy_profile_sha256 = str(source_pattern_profile.get("profileSha256", "") or "")
        available_profiles = (
            {legacy_profile_sha256: source_pattern_profile}
            if legacy_profile_sha256
            else {}
        )
    result_context_scope = (
        str(result_source_relative.get("contextScope", "") or "")
        if isinstance(result_source_relative, dict)
        else ""
    )
    result_profile_sha256 = _source_pattern_profile_sha_from_evidence(result_source_relative)
    effective_source_pattern_profile: dict[str, object] | None = None
    if source_pattern_profile_registry_issues:
        issues.extend(str(code) for code in source_pattern_profile_registry_issues)
    if result_context_scope == "document":
        if not result_profile_sha256:
            issues.append("result_source_pattern_profile_binding_invalid")
        elif not isinstance(available_profiles, dict):
            issues.append("source_pattern_profile_registry_missing")
        else:
            registered_profile = available_profiles.get(result_profile_sha256)
            if (
                not isinstance(registered_profile, dict)
                or not source_pattern_profile_valid(registered_profile)
                or str(registered_profile.get("profileSha256", "") or "")
                != result_profile_sha256
            ):
                issues.append("result_source_pattern_profile_missing_or_invalid")
            else:
                effective_source_pattern_profile = registered_profile
    elif result_context_scope == "local":
        if document_context_available:
            issues.append("local_source_pattern_context_forbidden")
    else:
        issues.append("result_source_relative_style_context_invalid")

    selected_profile_sha256 = _source_pattern_profile_sha_from_evidence(
        selected_candidate.get("sourceRelativeStyleDelta")
        if isinstance(selected_candidate, dict)
        else None
    )
    baseline_profile_sha256 = _source_pattern_profile_sha_from_evidence(
        baseline_candidate.get("sourceRelativeStyleDelta")
        if isinstance(baseline_candidate, dict)
        else None
    )
    if result_context_scope == "document" and (
        selected_profile_sha256 != result_profile_sha256
        or baseline_profile_sha256 != result_profile_sha256
    ):
        issues.append("candidate_source_pattern_profile_binding_mismatch")

    fresh_result_source_relative = assess_source_relative_style_delta(
        candidate_baseline_text,
        output_text,
        source_pattern_profile=effective_source_pattern_profile,
    )
    if (
        not source_relative_style_delta_passed(fresh_result_source_relative)
        or not isinstance(result_source_relative, dict)
        or json.dumps(fresh_result_source_relative, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        != json.dumps(result_source_relative, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    ):
        issues.append("result_source_relative_style_fresh_assessment_mismatch")

    postprocess_applied = selection.get("postprocessApplied")
    if not isinstance(postprocess_applied, bool):
        issues.append("postprocess_evidence_invalid")
    if published_rewrite is True:
        if published_sha256 != output_sha256:
            issues.append("published_output_hash_mismatch")
        if postprocess_applied is False and selected_sha256 != output_sha256:
            issues.append("selected_output_hash_mismatch")
    else:
        if postprocess_applied is not False:
            issues.append("postprocess_evidence_invalid")
        if selected_sha256 != output_sha256:
            issues.append("selected_output_hash_mismatch")
        if published_sha256:
            issues.append("unexpected_published_hash")

    selected_char_count = _release_gate_exact_int(selection.get("selectedCharCount"))
    result_char_count = _release_gate_exact_int(selection.get("resultCharCount"))
    candidate_char_count = (
        _release_gate_exact_int(selected_candidate.get("charCount"))
        if selected_candidate is not None
        else None
    )
    if (
        selected_char_count is None
        or candidate_char_count is None
        or selected_char_count != candidate_char_count
    ):
        issues.append("selected_char_count_mismatch")
    if result_char_count != len(output_text):
        issues.append("result_char_count_mismatch")
    if published_rewrite is True:
        if _release_gate_exact_int(selection.get("publishedCharCount")) != len(output_text):
            issues.append("published_char_count_mismatch")
    elif selection.get("publishedCharCount") not in (None, ""):
        issues.append("unexpected_published_char_count")

    issues.extend(
        _release_gate_current_text_checks(
            candidate_baseline_text,
            output_text,
            chunk_id,
        )
    )
    return {
        "ok": not issues,
        "chunkId": chunk_id,
        "mode": mode,
        "issueCodes": list(dict.fromkeys(issues)),
        "materializedTextSha256": output_sha256,
        "outputTextSha256": output_sha256,
        "selectedTextSha256": selected_sha256,
    }


def _assert_document_release_chunk(
    chunk: dict[str, Any],
    decision: Any,
    *,
    source_pattern_profile: dict[str, object] | None = None,
    source_pattern_profiles: dict[str, dict[str, object]] | None = None,
    source_pattern_profile_registry_issues: Sequence[str] = (),
    document_context_available: bool = False,
) -> dict[str, Any]:
    assessment = _assess_document_release_chunk(
        chunk,
        decision,
        source_pattern_profile=source_pattern_profile,
        source_pattern_profiles=source_pattern_profiles,
        source_pattern_profile_registry_issues=source_pattern_profile_registry_issues,
        document_context_available=document_context_available,
    )
    if assessment.get("ok") is not True:
        raise DocumentReleaseGateError(
            str(assessment.get("chunkId", "") or ""),
            str(assessment.get("mode", "unknown") or "unknown"),
            [str(code) for code in assessment.get("issueCodes", [])],
            evidence=assessment,
        )
    return assessment


def _assert_document_release_payload(
    compare_payload: dict[str, Any],
    decisions: dict[str, Any],
) -> list[dict[str, Any]]:
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        raise DocumentReleaseGateError(
            "",
            "invalid_compare",
            ["release_compare_chunks_invalid"],
        )
    source_pattern_profiles, registry_issues = _normalize_source_pattern_profile_registry(
        compare_payload.get("sourcePatternProfiles")
    )
    assessments: list[dict[str, Any]] = []
    baseline_texts: list[str] = []
    candidate_texts: list[str] = []
    materialized_texts: list[str] = []
    for raw_chunk in raw_chunks:
        if not isinstance(raw_chunk, dict):
            raise DocumentReleaseGateError(
                "",
                "invalid_chunk",
                ["release_chunk_invalid"],
            )
        chunk_id = str(raw_chunk.get("chunkId", "") or "").strip()
        raw_input = raw_chunk.get("inputText")
        raw_output = raw_chunk.get("outputText")
        if not isinstance(raw_input, str) or not isinstance(raw_output, str):
            raise DocumentReleaseGateError(
                chunk_id,
                "invalid_chunk",
                ["release_chunk_invalid"],
            )
        try:
            baseline_text = _release_gate_normalize_text(raw_input, raw_input)
            candidate_text = _release_gate_normalize_text(raw_input, raw_output)
            materialized_text = _release_gate_normalize_text(
                raw_input,
                _select_review_text(
                    raw_input,
                    raw_output,
                    decisions.get(chunk_id),
                    chunk=raw_chunk,
                ),
            )
        except Exception as exc:
            raise DocumentReleaseGateError(
                chunk_id,
                "invalid_chunk",
                ["release_text_normalization_failed"],
            ) from exc
        baseline_texts.append(baseline_text)
        candidate_texts.append(candidate_text)
        materialized_texts.append(materialized_text)
        assessments.append(
            _assert_document_release_chunk(
                raw_chunk,
                decisions.get(chunk_id),
                source_pattern_profiles=source_pattern_profiles,
                source_pattern_profile_registry_issues=registry_issues,
                document_context_available=True,
            )
        )
    fresh_candidate_document_delta = assess_source_relative_document_delta(
        baseline_texts,
        candidate_texts,
    )
    fresh_materialized_document_delta = assess_source_relative_document_delta(
        baseline_texts,
        materialized_texts,
    )
    all_candidate_outputs_materialized = materialized_texts == candidate_texts
    has_v2_candidate_selection = any(
        isinstance(raw_chunk, dict)
        and isinstance(raw_chunk.get("candidateSelection"), dict)
        and raw_chunk["candidateSelection"].get("schema") == CANDIDATE_SELECTION_SCHEMA
        and raw_chunk["candidateSelection"].get("schemaVersion")
        == CANDIDATE_SELECTION_VERSION
        for raw_chunk in raw_chunks
    )
    persisted_document_delta = compare_payload.get("sourceRelativeDocumentDelta")
    document_issues: list[str] = []
    if has_v2_candidate_selection and all_candidate_outputs_materialized:
        if not source_relative_document_delta_passed(persisted_document_delta):
            document_issues.append("source_relative_document_delta_missing_or_invalid")
        elif _canonical_json_bytes(persisted_document_delta) != _canonical_json_bytes(
            fresh_candidate_document_delta
        ):
            document_issues.append("source_relative_document_delta_fresh_assessment_mismatch")
    if not source_relative_document_delta_passed(fresh_materialized_document_delta):
        document_issues.append("materialized_source_relative_document_delta_failed")
        document_issues.extend(
            str(code)
            for code in fresh_materialized_document_delta.get("blockingIssueCodes", [])
            if str(code or "").strip()
        )
    if document_issues:
        evidence = {
            "ok": False,
            "chunkId": "",
            "mode": "document",
            "issueCodes": list(dict.fromkeys(document_issues)),
            "candidateDocumentDelta": fresh_candidate_document_delta,
            "materializedDocumentDelta": fresh_materialized_document_delta,
        }
        raise DocumentReleaseGateError(
            "",
            "document",
            evidence["issueCodes"],
            evidence=evidence,
        )
    assessments.append(
        {
            "ok": True,
            "chunkId": "",
            "mode": "document",
            "issueCodes": [],
            "materializedDocumentDelta": fresh_materialized_document_delta,
        }
    )
    return assessments


def _restore_paragraphs_from_compare_manifest(
    output_path: Path,
    compare_payload: dict[str, Any],
    decisions: dict[str, Any] | None = None,
    *,
    captured_manifest: ChunkManifest | None = None,
    allow_manifest_disk_read: bool = True,
) -> list[str] | None:
    manifest = captured_manifest
    if manifest is None and allow_manifest_disk_read:
        manifest = _load_manifest_for_compare(output_path, compare_payload)
    if manifest is None:
        return None
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        return None
    chunk_results: dict[str, str] = {}
    seen_chunk_ids: set[str] = set()
    for raw_chunk in raw_chunks:
        if not isinstance(raw_chunk, dict):
            continue
        chunk_id = str(raw_chunk.get("chunkId", "")).strip()
        if not chunk_id:
            continue
        if chunk_id in seen_chunk_ids:
            return None
        seen_chunk_ids.add(chunk_id)
        input_text = raw_chunk.get("inputText", "")
        output_text = raw_chunk.get("outputText", "")
        selected_text = _select_review_text(input_text, output_text, (decisions or {}).get(chunk_id), chunk=raw_chunk)
        chunk_results[chunk_id] = normalize_chunk_output(
            input_text if isinstance(input_text, str) else "",
            selected_text,
        )
    manifest_chunk_ids = [chunk.chunk_id for chunk in manifest.chunks]
    manifest_chunk_id_set = set(manifest_chunk_ids)
    if (
        not manifest_chunk_ids
        or any(chunk_id not in chunk_results for chunk_id in manifest_chunk_ids)
        or any(chunk_id not in manifest_chunk_id_set for chunk_id in chunk_results)
    ):
        return None
    restored_text = restore_text_from_chunks(manifest, chunk_results)
    paragraphs = split_text_to_paragraphs(restored_text)
    if len(paragraphs) != manifest.paragraph_count:
        return None
    return paragraphs or None


def _build_paragraphs_from_compare_payload(
    output_path: Path,
    compare_payload: dict[str, Any],
    decisions: dict[str, Any] | None = None,
    *,
    captured_manifest: ChunkManifest | None = None,
    allow_manifest_disk_read: bool = True,
) -> list[str] | None:
    restored_paragraphs = _restore_paragraphs_from_compare_manifest(
        output_path,
        compare_payload,
        decisions,
        captured_manifest=captured_manifest,
        allow_manifest_disk_read=allow_manifest_disk_read,
    )
    if restored_paragraphs is not None:
        return restored_paragraphs

    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        return None

    expected_paragraph_count = _coerce_optional_nonnegative_int(compare_payload.get("paragraphCount"))
    expected_chunk_count = _coerce_optional_nonnegative_int(compare_payload.get("chunkCount"))
    if expected_chunk_count is not None and expected_chunk_count != len(raw_chunks):
        return None

    paragraphs_by_index: dict[int, list[tuple[int, str]]] = {}
    seen_chunk_ids: set[str] = set()
    seen_positions: set[tuple[int, int]] = set()
    for raw_chunk in raw_chunks:
        if not isinstance(raw_chunk, dict):
            return None
        chunk_id = str(raw_chunk.get("chunkId", "")).strip()
        if not chunk_id or chunk_id in seen_chunk_ids:
            return None
        seen_chunk_ids.add(chunk_id)
        try:
            paragraph_index = int(raw_chunk.get("paragraphIndex"))
            chunk_index = int(raw_chunk.get("chunkIndex"))
        except (TypeError, ValueError):
            return None
        if paragraph_index < 0 or chunk_index < 0:
            return None
        if expected_paragraph_count is not None and paragraph_index >= expected_paragraph_count:
            return None
        position = (paragraph_index, chunk_index)
        if position in seen_positions:
            return None
        seen_positions.add(position)
        input_text = raw_chunk.get("inputText", "")
        output_text = raw_chunk.get("outputText", "")
        selected_text = _select_review_text(input_text, output_text, (decisions or {}).get(chunk_id), chunk=raw_chunk)
        normalized_text = normalize_chunk_output(
            input_text if isinstance(input_text, str) else "",
            selected_text,
        )
        paragraphs_by_index.setdefault(paragraph_index, []).append((chunk_index, normalized_text))

    if not paragraphs_by_index:
        return None

    if not _compare_paragraph_chunks_are_complete(paragraphs_by_index, expected_paragraph_count):
        return None

    if expected_paragraph_count is not None:
        return [
            _join_compare_chunk_texts(paragraphs_by_index[paragraph_index])
            for paragraph_index in range(expected_paragraph_count)
        ]

    return [
        _join_compare_chunk_texts(paragraphs_by_index[paragraph_index])
        for paragraph_index in sorted(paragraphs_by_index)
    ]


def _coerce_optional_nonnegative_int(value: Any) -> int | None:
    try:
        normalized = int(value)
    except (TypeError, ValueError):
        return None
    return normalized if normalized >= 0 else None


def _compare_paragraph_chunks_are_complete(
    paragraphs_by_index: dict[int, list[tuple[int, str]]],
    expected_paragraph_count: int | None,
) -> bool:
    if expected_paragraph_count is not None and set(paragraphs_by_index) != set(range(expected_paragraph_count)):
        return False
    for chunks in paragraphs_by_index.values():
        indexes = sorted(chunk_index for chunk_index, _ in chunks)
        if indexes != list(range(len(indexes))):
            return False
    return True


def _join_compare_chunk_texts(chunks: list[tuple[int, str]]) -> str:
    ordered = [
        text.strip()
        for _, text in sorted(chunks, key=lambda item: item[0])
        if text and text.strip()
    ]
    if not ordered:
        return ""
    text = ordered[0]
    for part in ordered[1:]:
        text = _join_compare_text_fragments(text, part)
    return text.strip()


def _join_compare_text_fragments(left: str, right: str) -> str:
    left = left.strip()
    right = right.strip()
    if not left:
        return right
    if not right:
        return left
    if _needs_compare_join_space(left[-1], right[0]):
        return f"{left} {right}"
    return f"{left}{right}"


def _needs_compare_join_space(left: str, right: str) -> bool:
    if not left or not right:
        return False
    if left.isspace() or right.isspace():
        return False
    if not (left.isascii() and right.isascii()):
        return False
    if left.isalnum() and right.isalnum():
        return True
    if left in ".!?;:," and right.isalnum():
        return True
    if left.isalnum() and right in "([{":
        return True
    return False


def _sync_compare_payload_to_text_artifacts(
    output_path: Path,
    compare_payload: dict[str, Any],
    decisions: dict[str, Any] | None = None,
) -> bool:
    paragraphs = _build_paragraphs_from_compare_payload(output_path, compare_payload, decisions)
    if paragraphs is None:
        return False
    try:
        output_path.write_text("\n\n".join(paragraphs), encoding="utf-8")
    except OSError:
        return False

    body_map_path = _find_body_map_path_for_output(output_path, include_records=False)
    body_map = load_docx_body_map(body_map_path) if body_map_path is not None else None
    if body_map is not None and body_map_path is not None:
        try:
            reviewed_paragraphs = _ensure_reviewed_paragraph_count_for_body_map(
                output_path,
                compare_payload,
                paragraphs,
                len(body_map.units),
                decisions or {},
            )
            save_docx_body_map(
                update_docx_body_map_texts(body_map, reviewed_paragraphs, round_number=body_map.round_number),
                body_map_path,
            )
        except Exception:
            return False
    return True


def _ensure_reviewed_paragraph_count_for_body_map(
    output_path: Path,
    compare_payload: dict[str, Any],
    reviewed_paragraphs: list[str],
    expected_count: int,
    decisions: dict[str, Any] | None = None,
) -> list[str]:
    if len(reviewed_paragraphs) == expected_count:
        return reviewed_paragraphs
    restored_paragraphs = _restore_paragraphs_from_compare_manifest(output_path, compare_payload, decisions)
    if restored_paragraphs is not None and len(restored_paragraphs) == expected_count:
        return restored_paragraphs
    raise ValueError(
        "DOCX body map paragraph count mismatch after manifest restore. "
        f"Expected {expected_count}, got {len(reviewed_paragraphs)}."
    )


def _read_output_paragraphs_for_export(output_path: Path) -> tuple[list[str], str]:
    compare_payload = _load_compare_payload_for_output(output_path)
    if compare_payload is not None:
        paragraphs = _build_paragraphs_from_compare_payload(output_path, compare_payload)
        if paragraphs is not None:
            return paragraphs, "compare"

    text = output_path.read_text(encoding="utf-8")
    return _split_text_into_blocks(text), "text"


def _build_reviewed_paragraphs_from_compare(output_path: Path, decisions: dict[str, Any]) -> list[str]:
    compare_payload = _load_compare_payload_for_output(output_path)
    if compare_payload is None:
        raise ValueError("Review export requires compare data for the selected round.")
    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list):
        raise ValueError("Review export compare data is invalid: chunks missing.")
    paragraphs = _build_paragraphs_from_compare_payload(output_path, compare_payload, decisions)
    if paragraphs is None:
        raise ValueError("Review export did not find any selectable chunks.")
    return paragraphs


def _round_model_key(prompt_profile: str, round_number: int) -> str:
    return f"{normalize_prompt_profile(prompt_profile)}:{int(round_number)}"


def _record_prompt_sequence_match_rank(item: dict[str, Any], prompt_profile: str, prompt_sequence: list[str]) -> int:
    if not is_prompt_sequence_customizable(prompt_profile):
        return 0
    return prompt_sequence_match_rank(item.get("prompt_sequence"), prompt_sequence, int(item.get("round", 0) or 0))


def _record_matches_prompt(
    item: dict[str, Any],
    prompt_profile: str,
    prompt_sequence: list[str],
) -> bool:
    normalized_profile = normalize_prompt_profile(prompt_profile)
    if str(item.get("prompt_profile", LEGACY_PROMPT_PROFILE) or LEGACY_PROMPT_PROFILE).strip().lower() != normalized_profile:
        return False
    if not is_prompt_sequence_customizable(normalized_profile):
        return True
    return _record_prompt_sequence_match_rank(item, normalized_profile, prompt_sequence) >= 0


def _source_to_doc_id(source_path: str) -> tuple[Path, str]:
    normalized_source = normalize_path(Path(source_path))
    try:
        relative_doc_id = normalized_source.relative_to(ROOT_DIR)
        doc_id = normalize_doc_id(str(relative_doc_id).replace("\\", "/"))
    except ValueError:
        doc_id = normalize_doc_id(str(normalized_source))
    return normalized_source, doc_id


def _route_key(prompt_profile: str, prompt_sequence: list[str]) -> tuple[str, tuple[str, ...]]:
    normalized_profile = normalize_prompt_profile(prompt_profile)
    if not is_prompt_sequence_customizable(normalized_profile):
        return normalized_profile, ()
    return normalized_profile, tuple(prompt_sequence)


def find_conflicting_history_route(source_path: str, model_config: dict[str, Any]) -> dict[str, Any] | None:
    """Return an existing completed route when a fresh route would hide history."""
    _, doc_id = _source_to_doc_id(source_path)
    requested_profile = normalize_prompt_profile(model_config.get("promptProfile", DEFAULT_PROMPT_PROFILE))
    requested_sequence = normalize_prompt_sequence(requested_profile, model_config.get("promptSequence"))
    requested_key = _route_key(requested_profile, requested_sequence)
    records = list_records()
    entry = records.get(doc_id, {}) if isinstance(records, dict) else {}
    rounds = entry.get("rounds", []) if isinstance(entry, dict) else []
    normalized_rounds = [
        item
        for item in rounds
        if isinstance(item, dict)
        and isinstance(item.get("round"), int)
        and round_record_has_usable_artifacts(item, expected_doc_id=doc_id, all_records=records)
    ]
    if not normalized_rounds:
        return None
    if any(_record_matches_prompt(item, requested_profile, requested_sequence) for item in normalized_rounds):
        return None

    route_groups: dict[tuple[str, tuple[str, ...]], dict[str, Any]] = {}
    for item in normalized_rounds:
        item_profile = normalize_prompt_profile(str(item.get("prompt_profile", LEGACY_PROMPT_PROFILE) or LEGACY_PROMPT_PROFILE))
        item_sequence = normalize_prompt_sequence(item_profile, item.get("prompt_sequence"))
        item_key = _route_key(item_profile, item_sequence)
        if item_key == requested_key:
            continue
        group = route_groups.setdefault(
            item_key,
            {
                "promptProfile": item_profile,
                "promptSequence": item_sequence,
                "completedRounds": [],
                "latestOutputPath": "",
            },
        )
        group["completedRounds"].append(int(item["round"]))
        if item.get("output_path"):
            group["latestOutputPath"] = str(item.get("output_path") or "")

    if not route_groups:
        return None
    preferred = max(
        route_groups.values(),
        key=lambda item: (max(item["completedRounds"] or [0]), len(item["completedRounds"])),
    )
    preferred["completedRounds"] = sorted(set(preferred["completedRounds"]))
    preferred["docId"] = doc_id
    preferred["requestedPromptProfile"] = requested_profile
    preferred["requestedPromptSequence"] = requested_sequence
    preferred["message"] = (
        "该文档已有另一条提示词路线的完成记录。为避免误开新任务，请先切回历史记录对应路线再继续。"
    )
    return preferred


def _find_round_provider(model_config: dict[str, Any], override: dict[str, Any]) -> tuple[dict[str, Any] | None, bool]:
    provider_id = str(override.get("providerId", "") or "").strip()
    providers = model_config.get("modelProviders")
    if not provider_id or not isinstance(providers, list):
        return None, False
    for provider in providers:
        if not isinstance(provider, dict):
            continue
        if str(provider.get("id", "") or "").strip() == provider_id:
            return provider, True
    return None, True


def _copy_provider_route_fields(resolved: dict[str, Any], provider: dict[str, Any]) -> None:
    field_map = {
        "id": "providerId",
        "name": "providerName",
        "baseUrl": "baseUrl",
        "apiKey": "apiKey",
        "apiType": "apiType",
        "streaming": "streaming",
        "temperature": "temperature",
        "requestTimeoutSeconds": "requestTimeoutSeconds",
        "maxRetries": "maxRetries",
        "rateLimitWindowMinutes": "rateLimitWindowMinutes",
        "rateLimitMaxRequests": "rateLimitMaxRequests",
        "rateLimitPerMinute": "rateLimitPerMinute",
        "rateLimitPerFiveMinutes": "rateLimitPerFiveMinutes",
    }
    for source_key, target_key in field_map.items():
        if source_key in provider:
            resolved[target_key] = provider.get(source_key)
    default_model = str(provider.get("defaultModel", "") or "").strip()
    if not default_model:
        models = provider.get("models")
        if isinstance(models, list):
            default_model = next((str(item).strip() for item in models if str(item).strip()), "")
    if default_model:
        resolved["model"] = default_model


def _resolve_round_model_config(model_config: dict[str, Any], prompt_profile: str, round_number: int) -> dict[str, Any]:
    base_config = dict(model_config)
    round_models = model_config.get("roundModels")
    if not isinstance(round_models, dict):
        return base_config
    override = round_models.get(_round_model_key(prompt_profile, round_number))
    if not isinstance(override, dict) or not bool(override.get("enabled", False)):
        return base_config
    resolved = dict(base_config)
    provider, provider_lookup_attempted = _find_round_provider(model_config, override)
    if provider is not None:
        if provider.get("enabled") is False:
            raise ValueError(
                f"Round {round_number} model provider is disabled: "
                f"{str(provider.get('name') or provider.get('id') or '').strip() or 'unknown provider'}"
            )
        _copy_provider_route_fields(resolved, provider)
        resolved["routeSource"] = "provider"
    elif provider_lookup_attempted and not (
        str(override.get("baseUrl", "") or "").strip()
        and str(override.get("apiKey", "") or "").strip()
        and str(override.get("model", "") or "").strip()
    ):
        raise ValueError(f"Round {round_number} model provider no longer exists; reselect the model route.")
    else:
        resolved["routeSource"] = "round_snapshot"
    override_identity_keys = ("providerId", "model") if provider is not None else ("providerId", "providerName", "model")
    for key in override_identity_keys:
        override_value = str(override.get(key, "") or "").strip()
        if override_value:
            resolved[key] = override.get(key)
    if provider is None:
        for key in (
            "baseUrl",
            "apiKey",
            "apiType",
            "streaming",
            "temperature",
            "requestTimeoutSeconds",
            "maxRetries",
            "rateLimitWindowMinutes",
            "rateLimitMaxRequests",
            "rateLimitPerMinute",
            "rateLimitPerFiveMinutes",
        ):
            if key in override:
                resolved[key] = override.get(key)
    return resolved


def emit_progress_event(event: dict[str, Any]) -> None:
    payload = {"event": "round-progress", "payload": event}
    print(json.dumps(payload, ensure_ascii=False), flush=True)


def emit_result_payload(payload: dict[str, Any]) -> None:
    print(json.dumps({"event": "result", "payload": payload}, ensure_ascii=False), flush=True)


def emit_error_payload(message: str) -> None:
    print(json.dumps({"event": "error", "payload": {"message": message}}, ensure_ascii=False), flush=True)


def import_document(source_path: str) -> dict[str, Any]:
    normalized_source = normalize_path(Path(source_path))
    try:
        relative_doc_id = normalized_source.relative_to(ROOT_DIR)
        doc_id = normalize_doc_id(str(relative_doc_id).replace("\\", "/"))
    except ValueError:
        doc_id = normalize_doc_id(str(normalized_source))

    round_state = get_document_round_state(doc_id)
    input_text_path, extracted_from_docx = ensure_round_input_text(normalized_source)
    output_text_path = ""
    manifest_path = ""

    if round_state.next_round is not None:
        context = build_round_context(normalized_source, round_number=round_state.next_round)
        output_text_path = str(context.output_text_path)
        manifest_path = str(context.manifest_path)

    return {
        "docId": doc_id,
        "sourcePath": str(normalized_source),
        "sourceKind": normalized_source.suffix.lower() or ".txt",
        "completedRounds": round_state.completed_rounds,
        "nextRound": round_state.next_round,
        "plannedRounds": len(round_state.prompt_sequence),
        "maxRounds": get_max_rounds(round_state.prompt_profile, round_state.prompt_sequence),
        "hasNextRound": round_state.next_round is not None,
        "isComplete": round_state.is_complete,
        "inputTextPath": str(input_text_path),
        "outputTextPath": output_text_path,
        "manifestPath": manifest_path,
        "extractedFromDocx": extracted_from_docx,
    }


def get_document_status(
    source_path: str,
    prompt_profile: str = DEFAULT_PROMPT_PROFILE,
    prompt_sequence: object | None = None,
) -> dict[str, Any]:
    normalized_source = normalize_path(Path(source_path))
    try:
        relative_doc_id = normalized_source.relative_to(ROOT_DIR)
        doc_id = normalize_doc_id(str(relative_doc_id).replace("\\", "/"))
    except ValueError:
        doc_id = normalize_doc_id(str(normalized_source))

    normalized_prompt_profile = normalize_prompt_profile(prompt_profile)
    normalized_prompt_sequence = normalize_prompt_sequence(normalized_prompt_profile, prompt_sequence)
    round_state = get_document_round_state(
        doc_id,
        prompt_profile=normalized_prompt_profile,
        prompt_sequence=normalized_prompt_sequence,
    )
    records = list_records()
    rounds: list[dict[str, Any]] = []
    indexed_document = _get_indexed_history_document(doc_id)
    if indexed_document is not None:
        entry = _indexed_history_document_to_record_entry(indexed_document)
        rounds = entry.get("rounds") if isinstance(entry.get("rounds"), list) else []
    else:
        fallback_entry = records.get(doc_id, {}) if isinstance(records, dict) else {}
        rounds = fallback_entry.get("rounds", []) if isinstance(fallback_entry, dict) else []
    normalized_prompt_profile = round_state.prompt_profile
    normalized_prompt_sequence = round_state.prompt_sequence
    max_rounds = get_max_rounds(normalized_prompt_profile, normalized_prompt_sequence)
    completed_rounds = [
        item.get("round")
        for item in rounds
        if isinstance(item, dict)
        and isinstance(item.get("round"), int)
        and _record_matches_prompt(item, normalized_prompt_profile, normalized_prompt_sequence)
        and 1 <= int(item.get("round")) <= max_rounds
        and round_record_has_usable_artifacts(item, expected_doc_id=doc_id, all_records=records)
    ]
    completed_rounds.sort()
    latest_output_path = ""
    current_input_path, extracted_from_docx = ensure_round_input_text(normalized_source)
    current_output_path = ""
    manifest_path = ""

    if round_state.next_round is not None:
        context = build_round_context(
            normalized_source,
            round_number=round_state.next_round,
            prompt_profile=normalized_prompt_profile,
            prompt_sequence=normalized_prompt_sequence,
        )
        current_input_path = context.input_text_path
        current_output_path = str(context.output_text_path)
        manifest_path = str(context.manifest_path)

    if rounds:
        latest_round = max(
            (
                item
                for item in rounds
                if isinstance(item, dict)
                and isinstance(item.get("round"), int)
                and _record_matches_prompt(item, normalized_prompt_profile, normalized_prompt_sequence)
                and 1 <= int(item.get("round")) <= max_rounds
                and round_record_has_usable_artifacts(item, expected_doc_id=doc_id, all_records=records)
            ),
            key=lambda item: item["round"],
            default=None,
        )
        if latest_round:
            latest_output_path = str(normalize_path(Path(str(latest_round.get("output_path", ""))))) if latest_round.get("output_path") else ""
    return {
        "docId": doc_id,
        "promptProfile": normalized_prompt_profile,
        "promptSequence": normalized_prompt_sequence,
        "sourcePath": str(normalized_source),
        "sourceKind": normalized_source.suffix.lower() or ".txt",
        "completedRounds": completed_rounds,
        "nextRound": round_state.next_round,
        "plannedRounds": len(normalized_prompt_sequence),
        "maxRounds": max_rounds,
        "hasNextRound": round_state.next_round is not None,
        "isComplete": round_state.is_complete,
        "currentInputPath": str(current_input_path),
        "currentOutputPath": current_output_path,
        "manifestPath": manifest_path,
        "latestOutputPath": latest_output_path,
        "extractedFromDocx": extracted_from_docx,
    }


def get_document_history(source_path: str) -> dict[str, Any]:
    normalized_source = normalize_path(Path(source_path))
    try:
        relative_doc_id = normalized_source.relative_to(ROOT_DIR)
        doc_id = normalize_doc_id(str(relative_doc_id).replace("\\", "/"))
    except ValueError:
        doc_id = normalize_doc_id(str(normalized_source))

    records = list_records()
    indexed_document = _get_indexed_history_document(doc_id)
    if indexed_document is not None:
        history = _record_entry_to_history(
            doc_id,
            _indexed_history_document_to_record_entry(indexed_document),
            all_records=records,
        )
        return {
            "docId": doc_id,
            "sourcePath": str(normalized_source),
            "artifactStats": history.get("artifactStats", _empty_artifact_stats()),
            "rounds": history.get("rounds", []),
        }

    entry = records.get(doc_id, {}) if isinstance(records, dict) else {}
    rounds = entry.get("rounds", []) if isinstance(entry, dict) else []

    usable_rounds = [
        item
        for item in rounds
        if isinstance(item, dict)
        and round_record_has_usable_artifacts(item, expected_doc_id=doc_id, all_records=records)
    ]
    history_rounds = [_map_history_round(item) for item in usable_rounds]

    history_rounds.sort(key=lambda item: item["round"], reverse=True)

    return {
        "docId": doc_id,
        "sourcePath": str(normalized_source),
        "artifactStats": build_artifact_summary(usable_rounds),
        "rounds": history_rounds,
    }


def get_document_protection_map(source_path: str) -> dict[str, Any]:
    normalized_source = normalize_path(Path(source_path))
    return build_docx_protection_map(normalized_source)


def get_document_scope_diagnostics(source_path: str) -> dict[str, Any]:
    normalized_source = normalize_path(Path(source_path))
    source_kind = normalized_source.suffix.lower() or ".txt"
    if source_kind != ".docx":
        return {
            "available": False,
            "ok": True,
            "sourcePath": str(normalized_source),
            "sourceKind": source_kind,
            "message": "当前文档不是 DOCX，无法生成正文边界诊断。",
            "path": "",
            "totalTextUnitCount": 0,
            "editableUnitCount": 0,
            "protectedUnitCount": 0,
            "scope": {},
            "reasonCounts": {},
            "issues": [],
            "units": [],
            "issueCount": 0,
            "errorCount": 0,
            "warningCount": 0,
        }

    _, snapshot_path, snapshot = ensure_docx_processing_assets(normalized_source)
    diagnostics_path = get_docx_scope_diagnostics_path(normalized_source)
    payload = _read_json_report(diagnostics_path)
    if payload is None:
        payload = build_docx_scope_diagnostics(snapshot, snapshot_path=snapshot_path)
        diagnostics_path.parent.mkdir(parents=True, exist_ok=True)
        diagnostics_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    payload["available"] = True
    payload["sourceKind"] = ".docx"
    payload["path"] = str(diagnostics_path)
    payload["message"] = "已生成 DOCX 正文边界诊断。"
    return payload


def _read_rate_audit_text(path: Path, *, label: str) -> str:
    normalized_path = normalize_path(path)
    if not normalized_path.exists() or not normalized_path.is_file():
        raise ValueError(f"{label} does not exist: {normalized_path}")
    try:
        return normalized_path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError(f"{label} is not valid UTF-8 text.") from exc


def _same_rate_audit_route(left: dict[str, Any], right: dict[str, Any]) -> bool:
    left_profile = str(left.get("promptProfile", "") or "").strip()
    right_profile = str(right.get("promptProfile", "") or "").strip()
    left_sequence = [str(item) for item in left.get("promptSequence", [])] if isinstance(left.get("promptSequence"), list) else []
    right_sequence = [str(item) for item in right.get("promptSequence", [])] if isinstance(right.get("promptSequence"), list) else []
    return left_profile == right_profile and left_sequence == right_sequence


def _materialize_rate_audit_output(
    output_path: Path,
    *,
    compare_payload: dict[str, Any] | None = None,
    artifact_snapshot: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Resolve the exact text that a current export would use.

    Saved review choices, confirmed source retention, custom text, and the
    safe-source default for unresolved high-risk chunks all affect export.  A
    rate audit must use that same materialized state for both document-level
    scoring and chunk hotspots instead of diagnosing a discarded model draft.
    """

    normalized_output = normalize_path(output_path)
    if artifact_snapshot is not None:
        if not isinstance(artifact_snapshot, dict):
            raise ValueError("Round artifact snapshot must be an object.")
        snapshot_output = str(artifact_snapshot.get("outputPath", "") or "").strip()
        if not snapshot_output or normalize_path(Path(snapshot_output)).resolve() != normalized_output.resolve():
            raise ValueError("Round artifact snapshot belongs to another output.")
        internal = artifact_snapshot.get("_internal")
        if not isinstance(internal, dict):
            raise ValueError("Internal round artifact snapshot payload is required.")
        resolved_compare = internal.get("comparePayload")
        decisions = (artifact_snapshot.get("review") or {}).get("decisions")
        paragraphs = internal.get("effectiveParagraphs")
        effective_chunks = internal.get("effectiveChunks")
        effective_text = internal.get("effectiveText")
        if (
            not isinstance(resolved_compare, dict)
            or not isinstance(decisions, dict)
            or not isinstance(paragraphs, list)
            or not all(isinstance(item, str) for item in paragraphs)
            or not isinstance(effective_chunks, list)
            or not all(isinstance(item, dict) for item in effective_chunks)
            or not isinstance(effective_text, str)
            or effective_text != "\n\n".join(paragraphs)
        ):
            raise ValueError("Round artifact snapshot materialization is incomplete.")
        if hashlib.sha256(effective_text.encode("utf-8")).hexdigest() != str(
            artifact_snapshot.get("effectiveTextSha256", "") or ""
        ):
            raise ValueError("Round artifact snapshot effective-text hash is inconsistent.")
        return {
            "text": effective_text,
            "paragraphs": list(paragraphs),
            "chunks": copy.deepcopy(effective_chunks),
            "comparePayload": copy.deepcopy(resolved_compare),
            "decisions": copy.deepcopy(decisions),
            "decisionsPath": str(internal.get("reviewPath", "") or ""),
            "source": "review_materialized_compare",
            # Private, in-process provenance. Callers must not attach this
            # object to an API response; it contains captured artifact bytes.
            "artifactSnapshot": artifact_snapshot,
        }

    resolved_compare = compare_payload if isinstance(compare_payload, dict) else _load_compare_payload_for_output(normalized_output)
    decisions: dict[str, Any] = {}
    decisions_path = ""
    if resolved_compare is not None:
        review_payload = load_review_decisions(str(normalized_output))
        raw_decisions = review_payload.get("decisions")
        if isinstance(raw_decisions, dict):
            decisions = raw_decisions
        decisions_path = str(review_payload.get("path", "") or "")
        # Direct targeted-rerun/RateAudit callers can supply compare data
        # without first capturing a round snapshot.  Apply the same release
        # gate before that review-materialized text can become a model baseline.
        _assert_document_release_payload(resolved_compare, decisions)

    paragraphs = (
        _build_paragraphs_from_compare_payload(normalized_output, resolved_compare, decisions)
        if resolved_compare is not None
        else None
    )
    if paragraphs is None:
        return {
            "text": _read_rate_audit_text(normalized_output, label="Round output"),
            "paragraphs": None,
            "chunks": None,
            "comparePayload": resolved_compare,
            "decisions": decisions,
            "decisionsPath": decisions_path,
            "source": "output_file",
        }

    effective_chunks: list[dict[str, Any]] = []
    raw_chunks = resolved_compare.get("chunks") if isinstance(resolved_compare, dict) else None
    if isinstance(raw_chunks, list):
        for index, item in enumerate(raw_chunks):
            if not isinstance(item, dict):
                continue
            chunk_id = str(item.get("chunkId", f"chunk-{index + 1}") or f"chunk-{index + 1}")
            input_text = str(item.get("inputText", "") or "")
            selected_text = _select_review_text(
                input_text,
                item.get("outputText", ""),
                decisions.get(chunk_id),
                chunk=item,
            )
            effective_text = normalize_chunk_output(input_text, selected_text)
            if not effective_text.strip():
                continue
            effective_chunks.append(
                {
                    "chunkId": chunk_id,
                    "paragraphIndex": int(item.get("paragraphIndex", index) or 0),
                    "chunkIndex": int(item.get("chunkIndex", 0) or 0),
                    "text": effective_text,
                }
            )

    return {
        "text": "\n\n".join(paragraphs),
        "paragraphs": paragraphs,
        "chunks": effective_chunks,
        "comparePayload": resolved_compare,
        "decisions": decisions,
        "decisionsPath": decisions_path,
        "source": "review_materialized_compare",
    }


def _sha256_file_bytes(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _sha256_canonical_json(payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _build_rate_audit_strategy_binding(
    report: dict[str, Any],
    *,
    source_path: Path,
    output_path: Path | None,
    materialized: dict[str, Any] | None,
) -> dict[str, Any]:
    """Bind an executable targeted plan to every mutable input it depends on.

    ``planDigest`` deliberately includes hashes that are not repeated in the
    POST request (compare/effective output/prompt).  A client therefore cannot
    replay an otherwise identical looking plan after a review decision, prompt
    file, compare sidecar, or output artifact changed.
    """

    strategy = report.get("strategyPlan") if isinstance(report.get("strategyPlan"), dict) else {}
    contract = report.get("contentContract") if isinstance(report.get("contentContract"), dict) else {}
    compare_payload = (
        materialized.get("comparePayload")
        if isinstance(materialized, dict) and isinstance(materialized.get("comparePayload"), dict)
        else None
    )
    artifact_snapshot = (
        materialized.get("artifactSnapshot")
        if isinstance(materialized, dict) and isinstance(materialized.get("artifactSnapshot"), dict)
        else None
    )
    snapshot_internal = (
        artifact_snapshot.get("_internal")
        if isinstance(artifact_snapshot, dict) and isinstance(artifact_snapshot.get("_internal"), dict)
        else None
    )
    target_chunk_ids = [
        str(value).strip()
        for value in strategy.get("targetChunkIds", [])
        if str(value).strip()
    ] if isinstance(strategy.get("targetChunkIds"), list) else []
    dimension_id = str(strategy.get("dimensionId", "") or "").strip()
    recommended_prompt_id = str(strategy.get("recommendedPromptId", "") or "").strip()
    dimension_definition = get_rate_audit_dimension_definition(dimension_id)

    blocked_codes: list[str] = []
    normalized_output = normalize_path(output_path).resolve() if output_path is not None else None
    compare_path = _find_compare_path_for_output(normalized_output) if normalized_output is not None else None
    compare_sha256 = ""
    output_sha256 = ""
    prompt_sha256 = ""
    manifest_sha256 = ""
    body_map_sha256 = ""
    if normalized_output is None:
        blocked_codes.append("output_missing")
    elif artifact_snapshot is not None:
        if str(artifact_snapshot.get("outputPath", "") or "") != str(normalized_output):
            blocked_codes.append("artifact_snapshot_output_mismatch")
        output_sha256 = str(artifact_snapshot.get("outputSha256", "") or "")
        compare_sha256 = str(artifact_snapshot.get("compareSha256", "") or "")
        if not output_sha256:
            blocked_codes.append("output_unreadable")
        if not compare_sha256:
            blocked_codes.append("compare_unreadable")
    else:
        try:
            output_sha256 = _sha256_file_bytes(normalized_output)
        except OSError:
            blocked_codes.append("output_unreadable")
        if compare_path is None or not compare_path.exists() or not compare_path.is_file():
            blocked_codes.append("compare_missing")
        else:
            try:
                compare_sha256 = _sha256_file_bytes(compare_path)
            except OSError:
                blocked_codes.append("compare_unreadable")

    if compare_payload is None:
        blocked_codes.append("compare_invalid")
    elif (
        artifact_snapshot is None
        and normalized_output is not None
        and not _compare_payload_is_usable(normalized_output, compare_payload)
    ):
        blocked_codes.append("compare_not_usable")
    if source_path.suffix.lower() == ".docx" and normalized_output is not None and artifact_snapshot is not None:
        body_map_sha256 = str(artifact_snapshot.get("bodyMapSha256", "") or "")
        manifest_sha256 = str(artifact_snapshot.get("manifestSha256", "") or "")
        if not isinstance(snapshot_internal, dict) or not isinstance(snapshot_internal.get("bodyMapPayload"), dict):
            blocked_codes.append("docx_body_map_missing")
        elif not body_map_sha256:
            blocked_codes.append("docx_body_map_unreadable")
        if not isinstance(snapshot_internal, dict) or not isinstance(snapshot_internal.get("manifestPayload"), dict):
            blocked_codes.append("docx_manifest_missing")
        elif not manifest_sha256:
            blocked_codes.append("docx_manifest_unreadable")
    elif source_path.suffix.lower() == ".docx" and normalized_output is not None:
        body_map_path = _find_body_map_path_for_output(normalized_output)
        if (
            _load_body_map_for_output(normalized_output) is None
            or body_map_path is None
            or not is_path_under(body_map_path, ROOT_DIR / "finish")
        ):
            blocked_codes.append("docx_body_map_missing")
        elif body_map_path.exists() and body_map_path.is_file():
            try:
                body_map_sha256 = _sha256_file_bytes(body_map_path)
            except OSError:
                blocked_codes.append("docx_body_map_unreadable")
        raw_manifest_path = str((compare_payload or {}).get("manifestPath", "") or "").strip()
        manifest_path = normalize_path(Path(raw_manifest_path)) if raw_manifest_path else _find_manifest_path_for_output(normalized_output)
        if manifest_path is None or not is_path_under(manifest_path, ROOT_DIR / "finish"):
            blocked_codes.append("docx_manifest_path_invalid")
        elif _load_manifest_for_compare(normalized_output, compare_payload or {}) is None:
            blocked_codes.append("docx_manifest_missing")
        elif manifest_path.exists() and manifest_path.is_file():
            try:
                manifest_sha256 = _sha256_file_bytes(manifest_path)
            except OSError:
                blocked_codes.append("docx_manifest_unreadable")

    if recommended_prompt_id:
        try:
            prompt_path = resolve_prompt_path(recommended_prompt_id)
            prompt_sha256 = _sha256_file_bytes(prompt_path)
        except (OSError, ValueError):
            blocked_codes.append("prompt_unavailable")
    else:
        blocked_codes.append("prompt_missing")

    compare_revision = str(
        (artifact_snapshot or {}).get("compareRevision", "")
        or (compare_payload or {}).get("updatedAt", "")
        or ""
    ).strip()
    if not compare_revision and compare_sha256:
        compare_revision = f"sha256:{compare_sha256}"

    source_sha256 = str(contract.get("sourceSha256", "") or "").strip()
    try:
        current_source_sha256 = _sha256_file_bytes(source_path)
    except OSError:
        current_source_sha256 = ""
        blocked_codes.append("source_unreadable")
    if not source_sha256:
        source_sha256 = current_source_sha256
    elif current_source_sha256 and source_sha256 != current_source_sha256:
        blocked_codes.append("source_contract_stale")

    scope_digest = str(contract.get("scopeDigest", "") or "").strip()
    format_digest = str(contract.get("formatDigest", "") or "").strip()
    effective_text = str((materialized or {}).get("text", "") or "")
    effective_text_sha256 = hashlib.sha256(effective_text.encode("utf-8")).hexdigest() if effective_text else ""
    if artifact_snapshot is not None and effective_text_sha256 != str(
        artifact_snapshot.get("effectiveTextSha256", "") or ""
    ):
        blocked_codes.append("artifact_snapshot_effective_text_mismatch")

    compare_chunk_ids: list[str] = []
    if isinstance(compare_payload, dict) and isinstance(compare_payload.get("chunks"), list):
        compare_chunk_ids = [
            str(item.get("chunkId", "") or "").strip()
            for item in compare_payload["chunks"]
            if isinstance(item, dict) and str(item.get("chunkId", "") or "").strip()
        ]
    effective_chunks = (
        materialized.get("chunks")
        if isinstance(materialized, dict) and isinstance(materialized.get("chunks"), list)
        else []
    )
    effective_chunk_ids = [
        str(item.get("chunkId", "") or "").strip()
        for item in effective_chunks
        if isinstance(item, dict) and str(item.get("chunkId", "") or "").strip()
    ]
    if not target_chunk_ids:
        blocked_codes.append("targets_missing")
    if len(target_chunk_ids) != len(set(target_chunk_ids)):
        blocked_codes.append("targets_duplicate")
    if any(chunk_id not in compare_chunk_ids or chunk_id not in effective_chunk_ids for chunk_id in target_chunk_ids):
        blocked_codes.append("targets_not_materialized")
    if len(compare_chunk_ids) != len(set(compare_chunk_ids)):
        blocked_codes.append("compare_chunk_ids_duplicate")
    if len(effective_chunk_ids) != len(set(effective_chunk_ids)):
        blocked_codes.append("effective_chunk_ids_duplicate")
    effective_chunk_sha256 = {
        str(item.get("chunkId", "") or "").strip(): hashlib.sha256(
            str(item.get("text", "") or "").encode("utf-8")
        ).hexdigest()
        for item in effective_chunks
        if isinstance(item, dict) and str(item.get("chunkId", "") or "").strip()
    }
    hotspot_dimensions = {
        str(item.get("chunkId", "") or "").strip(): {
            str(value).strip()
            for value in item.get("dimensionIds", [])
            if str(value).strip()
        }
        for item in report.get("hotspots", [])
        if isinstance(item, dict)
        and str(item.get("chunkId", "") or "").strip()
        and isinstance(item.get("dimensionIds"), list)
    }
    if any(dimension_id not in hotspot_dimensions.get(chunk_id, set()) for chunk_id in target_chunk_ids):
        blocked_codes.append("targets_not_dimension_hotspots")
    saved_decisions = (
        materialized.get("decisions")
        if isinstance(materialized, dict) and isinstance(materialized.get("decisions"), dict)
        else {}
    )
    review_locked_targets: list[str] = []
    for chunk_id in target_chunk_ids:
        if chunk_id not in saved_decisions:
            continue
        # Any saved choice is part of the review contract. In particular,
        # source_confirmed/rewrite_confirmed and confirmed custom text must
        # never be overwritten, while even an unconfirmed source/rewrite choice
        # would otherwise make candidate materialization ambiguous.
        review_locked_targets.append(chunk_id)
    if review_locked_targets:
        blocked_codes.append("review_locked_target")
    review_pending_targets: list[str] = []
    if isinstance(compare_payload, dict) and isinstance(compare_payload.get("chunks"), list):
        for item in compare_payload["chunks"]:
            if not isinstance(item, dict) or item.get("rateAuditStrategyReviewRequired") is not True:
                continue
            chunk_id = str(item.get("chunkId", "") or "").strip()
            decision = _normalize_review_decision_value(saved_decisions.get(chunk_id)) if chunk_id in saved_decisions else None
            decision_confirmed = bool(
                decision in {"source_confirmed", "rewrite_confirmed"}
                if isinstance(decision, str)
                else isinstance(decision, dict) and decision.get("confirmed") is True
            )
            if chunk_id and not decision_confirmed:
                review_pending_targets.append(chunk_id)
    if review_pending_targets:
        blocked_codes.append("review_pending_target")
    strategy_attempts = (
        compare_payload.get("rateAuditStrategyAttempts")
        if isinstance(compare_payload, dict) and isinstance(compare_payload.get("rateAuditStrategyAttempts"), dict)
        else {}
    )
    attempt_limited_targets: list[str] = []
    maximum_attempts = int(dimension_definition.get("maxAttempts", 0) or 0)
    for chunk_id in target_chunk_ids:
        entry = strategy_attempts.get(f"{dimension_id}:{chunk_id}")
        if not isinstance(entry, dict):
            continue
        if (
            str(entry.get("dimensionId", "") or "") == dimension_id
            and str(entry.get("recommendedPromptId", "") or "") == recommended_prompt_id
            and str(entry.get("effectiveTextSha256", "") or "") == effective_chunk_sha256.get(chunk_id, "")
            and int(entry.get("attemptCount", 0) or 0) >= maximum_attempts > 0
        ):
            attempt_limited_targets.append(chunk_id)
    if attempt_limited_targets:
        blocked_codes.append("strategy_attempt_limit")

    if str(strategy.get("decision", "") or "") != "targeted_rerun":
        blocked_codes.append("decision_not_targeted_rerun")
    if not bool(strategy.get("canExecute")):
        blocked_codes.append("strategy_not_executable")
    if not bool(contract.get("ready")):
        blocked_codes.append("content_contract_not_ready")
    if not bool(contract.get("scopeReady")):
        blocked_codes.append("scope_contract_not_ready")
    if not bool(contract.get("formatLockReady")):
        blocked_codes.append("format_contract_not_ready")
    if not bool(dimension_definition.get("canExecute")):
        blocked_codes.append("dimension_not_executable")
    if str(dimension_definition.get("targetScope", "") or "") != "chunk":
        blocked_codes.append("dimension_scope_not_chunk")
    if str(dimension_definition.get("repairPromptId", "") or "") != recommended_prompt_id:
        blocked_codes.append("dimension_prompt_mismatch")
    if str(strategy.get("repairPromptId", "") or "") != recommended_prompt_id:
        blocked_codes.append("strategy_prompt_mismatch")
    if str(strategy.get("evaluatorDimensionId", "") or "") != str(dimension_definition.get("evaluatorDimensionId", "") or ""):
        blocked_codes.append("strategy_evaluator_mismatch")
    if int(dimension_definition.get("maxAttempts", 0) or 0) <= 0:
        blocked_codes.append("dimension_attempts_unavailable")

    digest_payload = {
        "version": RATE_AUDIT_STRATEGY_BINDING_VERSION,
        "sourcePath": str(source_path.resolve()),
        "outputPath": str(normalized_output or ""),
        "compareRevision": compare_revision,
        "sourceSha256": source_sha256,
        "scopeDigest": scope_digest,
        "formatDigest": format_digest,
        "dimensionId": dimension_id,
        "recommendedPromptId": recommended_prompt_id,
        "targetChunkIds": target_chunk_ids,
        "effectiveTextSha256": effective_text_sha256,
        "outputSha256": output_sha256,
        "compareSha256": compare_sha256,
        "promptSha256": prompt_sha256,
        "manifestSha256": manifest_sha256,
        "bodyMapSha256": body_map_sha256,
        "reviewRevision": str((artifact_snapshot or {}).get("reviewRevision", "") or ""),
        "contentRevision": str((artifact_snapshot or {}).get("contentRevision", "") or ""),
        "artifactSnapshotDigest": str((artifact_snapshot or {}).get("artifactSnapshotDigest", "") or ""),
        "dimensionRegistryVersion": strategy.get("dimensionRegistryVersion"),
        "evaluatorDimensionId": str(dimension_definition.get("evaluatorDimensionId", "") or ""),
        "primaryMetric": str(dimension_definition.get("primaryMetric", "") or ""),
        "secondaryMetric": str(dimension_definition.get("secondaryMetric", "") or ""),
        "maxAttempts": int(dimension_definition.get("maxAttempts", 0) or 0),
        "plateauPolicy": str(dimension_definition.get("plateauPolicy", "") or ""),
    }
    ready = not blocked_codes
    blocked_reason = ""
    if not ready:
        for priority_code in ("review_pending_target", "review_locked_target", "strategy_attempt_limit"):
            if priority_code in blocked_codes:
                blocked_reason = priority_code
                break
        if not blocked_reason:
            blocked_reason = blocked_codes[0]
    return {
        **digest_payload,
        "planDigest": _sha256_canonical_json(digest_payload),
        "ready": ready,
        "blockedReason": blocked_reason,
    }


def _materialize_rate_audit_plateau_state(
    report: dict[str, Any],
    binding: dict[str, Any],
) -> None:
    """Promote an exhausted execution ledger into the public strategy state.

    The binding layer is still the authority that proves the attempt ledger is
    for this exact dimension, prompt, target, and effective-text generation.
    Once it reports ``strategy_attempt_limit``, leaving the public plan as an
    executable ``targeted_rerun`` is contradictory and encourages repeated
    clicks.  Materialize a hard stop while retaining the blocked binding as
    stale-request evidence.
    """

    if str(binding.get("blockedReason", "") or "") != "strategy_attempt_limit":
        return
    plan = report.get("strategyPlan")
    if not isinstance(plan, dict) or str(plan.get("decision", "") or "") != "targeted_rerun":
        return
    dimension_id = str(plan.get("dimensionId", "") or "")
    target_chunk_ids = [
        str(value).strip()
        for value in binding.get("targetChunkIds", [])
        if str(value).strip()
    ] if isinstance(binding.get("targetChunkIds"), list) else []
    maximum_attempts = int(binding.get("maxAttempts", plan.get("maxAttempts", 0)) or 0)
    reason = (
        f"“{str(plan.get('dimensionLabel', dimension_id) or dimension_id)}”维度已对当前文本代际达到"
        f"{maximum_attempts or '规定'}次定向尝试上限，仍未通过同维度收敛评估。"
        "系统已硬停止该自动策略并保留此前已接受文本，后续须人工复核；文本变化后才可形成新的绑定计划。"
    )
    plan.update(
        {
            "decision": "manual_review",
            "label": "达到尝试上限，转人工复核",
            "recommendedPromptId": "",
            "promptSelectionSource": "none",
            "canExecute": False,
            "selectedExecutableDimensionId": "",
            "executableQueue": [],
            "executableQueueCount": 0,
            "manualReviewRequired": True,
            "manualReviewStillRequired": True,
            "hardStop": True,
            "plateauReached": True,
            "plateauReason": "strategy_attempt_limit",
            "plateauDimensionId": dimension_id,
            "plateauTargetChunkIds": target_chunk_ids,
            "plateauTargetChunkCount": len(target_chunk_ids),
            "plateauAttemptLimit": maximum_attempts,
            "reason": reason,
        }
    )
    readiness = report.get("readiness")
    if isinstance(readiness, dict):
        readiness.update(
            {
                "status": "attention",
                "strategyDecisionReady": True,
            }
        )
    report["plateau"] = {
        "reached": True,
        "reason": "strategy_attempt_limit",
        "hardStop": True,
        "dimensionId": dimension_id,
        "targetChunkIds": target_chunk_ids,
        "targetChunkCount": len(target_chunk_ids),
        "attemptLimit": maximum_attempts,
        "preservedPreviousText": True,
        "manualReviewRequired": True,
    }


def validate_rate_audit_strategy_request(request_payload: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """Recompute and exactly compare an execution request with the live plan."""

    if not isinstance(request_payload, dict):
        raise ValueError("RateAudit strategy request must be an object.")
    if not isinstance(request_payload.get("sourcePath"), str):
        raise ValueError("sourcePath must be a string.")
    if not isinstance(request_payload.get("outputPath"), str):
        raise ValueError("outputPath must be a string.")
    source_value = request_payload["sourcePath"].strip()
    output_value = request_payload["outputPath"].strip()
    if not source_value:
        raise ValueError("sourcePath is required.")
    if not output_value:
        raise ValueError("outputPath is required.")
    normalized_source = _resolve_api_path(source_value, allowed_roots=((ROOT_DIR / "origin").resolve(),), label="Source file")
    normalized_output = _resolve_api_path(output_value, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    if not normalized_source.exists() or not normalized_source.is_file():
        raise ValueError("Source file does not exist.")
    if not normalized_output.exists() or not normalized_output.is_file():
        raise ValueError("Output file does not exist.")

    requested_targets = request_payload.get("targetChunkIds")
    if not isinstance(requested_targets, list) or not requested_targets:
        raise ValueError("targetChunkIds must be a non-empty list.")
    if any(not isinstance(value, str) for value in requested_targets):
        raise ValueError("targetChunkIds must contain strings.")
    normalized_targets = [value.strip() for value in requested_targets]
    if any(not value for value in normalized_targets) or len(normalized_targets) != len(set(normalized_targets)):
        raise ValueError("targetChunkIds must contain unique non-empty ids.")
    if len(normalized_targets) > 12:
        raise ValueError("targetChunkIds exceeds the RateAudit hotspot limit.")
    if any(len(value) > 160 for value in normalized_targets):
        raise ValueError("targetChunkIds contains an invalid id.")

    for key in (
        "dimensionId",
        "recommendedPromptId",
        "compareRevision",
        "scopeDigest",
        "sourceSha256",
        "planDigest",
    ):
        if not isinstance(request_payload.get(key), str):
            raise ValueError(f"{key} must be a string.")
        if not request_payload[key].strip():
            raise ValueError(f"{key} is required.")
    if not isinstance(request_payload.get("formatDigest"), str):
        raise ValueError("formatDigest must be a string.")
    sha256_pattern = re.compile(r"^[0-9a-f]{64}$")
    for key in ("sourceSha256", "scopeDigest", "planDigest"):
        if not sha256_pattern.fullmatch(str(request_payload.get(key, "") or "")):
            raise ValueError(f"{key} must be a lowercase SHA-256 digest.")
    requested_format_digest = str(request_payload.get("formatDigest", "") or "")
    if requested_format_digest and not sha256_pattern.fullmatch(requested_format_digest):
        raise ValueError("formatDigest must be empty or a lowercase SHA-256 digest.")
    identifier_pattern = re.compile(r"^[a-z0-9][a-z0-9_-]{0,63}$")
    for key in ("dimensionId", "recommendedPromptId"):
        if not identifier_pattern.fullmatch(str(request_payload.get(key, "") or "")):
            raise ValueError(f"{key} is invalid.")
    if len(str(request_payload.get("compareRevision", "") or "")) > 200:
        raise ValueError("compareRevision is invalid.")

    try:
        report = get_document_rate_audit(str(normalized_source), str(normalized_output))
    except ValueError as exc:
        raise StaleRateAuditStrategyPlanError(["strategy_recompute_failed"]) from exc
    binding = report.get("strategyBinding") if isinstance(report.get("strategyBinding"), dict) else {}
    mismatch_codes: list[str] = []
    if not bool(binding.get("ready")):
        mismatch_codes.append("strategy_not_ready")
        blocked_reason = str(binding.get("blockedReason", "") or "").strip()
        if blocked_reason:
            mismatch_codes.append(blocked_reason)

    comparisons = (
        ("compareRevision", "compare_revision_mismatch"),
        ("sourceSha256", "source_sha256_mismatch"),
        ("scopeDigest", "scope_digest_mismatch"),
        ("formatDigest", "format_digest_mismatch"),
        ("dimensionId", "dimension_mismatch"),
        ("recommendedPromptId", "prompt_mismatch"),
        ("planDigest", "plan_digest_mismatch"),
    )
    for key, code in comparisons:
        if str(request_payload.get(key, "") or "") != str(binding.get(key, "") or ""):
            mismatch_codes.append(code)
    if normalized_targets != list(binding.get("targetChunkIds", []) or []):
        mismatch_codes.append("target_chunks_mismatch")
    if mismatch_codes:
        raise StaleRateAuditStrategyPlanError(mismatch_codes)
    return report, binding


def get_document_rate_audit(source_path: str, output_path: str | None = None) -> dict[str, Any]:
    """Build an offline, explainable source-to-round risk trajectory.

    The report deliberately exposes heuristic writing signals rather than an
    invented detector probability.  Explicit output paths must stay under the
    generated ``finish`` tree even when this function is called outside Flask.
    """

    normalized_source = normalize_path(Path(source_path))
    if not is_path_under(normalized_source, ROOT_DIR / "origin"):
        raise ValueError("Source file must stay under the origin workspace directory.")
    if not normalized_source.exists() or not normalized_source.is_file():
        raise ValueError(f"Source file does not exist: {normalized_source}")

    source_text_path, _ = ensure_round_input_text(normalized_source)
    source_text = _read_rate_audit_text(source_text_path, label="Source text")
    history = get_document_history(str(normalized_source))
    all_rounds = [item for item in history.get("rounds", []) if isinstance(item, dict)]

    explicit_output: Path | None = None
    if str(output_path or "").strip():
        explicit_output = normalize_path(Path(str(output_path)))
        if not is_path_under(explicit_output, ROOT_DIR / "finish"):
            raise ValueError("Rate-audit output must stay under the finish workspace directory.")
        if not explicit_output.exists() or not explicit_output.is_file():
            raise ValueError(f"Rate-audit output does not exist: {explicit_output}")

    selected_round: dict[str, Any] | None = None
    if explicit_output is not None:
        for item in all_rounds:
            raw_path = str(item.get("outputPath", "") or "").strip()
            if raw_path and normalize_path(Path(raw_path)) == explicit_output:
                selected_round = item
                break
    elif all_rounds:
        selected_round = max(
            all_rounds,
            key=lambda item: (
                str(item.get("timestamp", "") or ""),
                int(item.get("round", 0) or 0),
            ),
        )

    # Every round is captured through the same short per-output lock.  The
    # resulting in-memory snapshot is then reused for trajectory text,
    # effective chunks, DOCX contract input, and strategy-plan hashes.
    materialized_outputs: dict[str, dict[str, Any]] = {}

    def materialize(candidate: Path, compare_payload: dict[str, Any] | None = None) -> dict[str, Any]:
        del compare_payload  # compare must come from the captured snapshot.
        cache_key = str(candidate.resolve())
        cached = materialized_outputs.get(cache_key)
        if cached is not None:
            return cached
        snapshot = read_round_artifact_snapshot(candidate, include_internal=True)
        payload = _materialize_rate_audit_output(candidate, artifact_snapshot=snapshot)
        materialized_outputs[cache_key] = payload
        return payload

    selected_compare_payload: dict[str, Any] | None = None
    if explicit_output is not None and selected_round is None:
        explicit_materialized = materialize(explicit_output)
        raw_selected_compare = explicit_materialized.get("comparePayload")
        selected_compare_payload = raw_selected_compare if isinstance(raw_selected_compare, dict) else None
        expected_doc_id = str(history.get("docId", "") or "").strip()
        compare_doc_id = str((selected_compare_payload or {}).get("docId", "") or "").strip()
        if not expected_doc_id or compare_doc_id != expected_doc_id:
            raise ValueError("Rate-audit output does not belong to the selected source document.")

    route_rounds = (
        [item for item in all_rounds if _same_rate_audit_route(item, selected_round)]
        if selected_round is not None
        else []
    )
    route_rounds.sort(key=lambda item: (int(item.get("round", 0) or 0), str(item.get("timestamp", "") or "")))

    stages: list[dict[str, Any]] = []

    stage_path_ids: dict[str, str] = {}
    used_stage_ids: set[str] = set()
    for index, item in enumerate(route_rounds):
        raw_path = str(item.get("outputPath", "") or "").strip()
        if not raw_path:
            continue
        candidate = normalize_path(Path(raw_path))
        if not is_path_under(candidate, ROOT_DIR / "finish") or not candidate.exists() or not candidate.is_file():
            continue
        round_number = int(item.get("round", 0) or 0)
        base_stage_id = f"round-{round_number}" if round_number > 0 else f"result-{index + 1}"
        stage_id = base_stage_id
        suffix = 2
        while stage_id in used_stage_ids:
            stage_id = f"{base_stage_id}-{suffix}"
            suffix += 1
        used_stage_ids.add(stage_id)
        stage_path_ids[str(candidate)] = stage_id
        stages.append(
            {
                "id": stage_id,
                "label": f"第 {round_number} 轮" if round_number > 0 else f"结果 {index + 1}",
                "round": round_number if round_number > 0 else None,
                "text": str(materialize(candidate).get("text", "") or ""),
            }
        )

    selected_output = explicit_output
    if selected_output is None and selected_round is not None:
        raw_selected_path = str(selected_round.get("outputPath", "") or "").strip()
        if raw_selected_path:
            candidate = normalize_path(Path(raw_selected_path))
            if is_path_under(candidate, ROOT_DIR / "finish") and candidate.exists() and candidate.is_file():
                selected_output = candidate

    current_stage_id = "source"
    selected_materialized: dict[str, Any] | None = None
    if selected_output is not None:
        selected_materialized = materialize(selected_output, selected_compare_payload)
        selected_key = str(selected_output)
        current_stage_id = stage_path_ids.get(selected_key, "current")
        if current_stage_id == "current":
            stages.append(
                {
                    "id": "current",
                    "label": "当前结果",
                    "round": None,
                    "text": str(selected_materialized.get("text", "") or ""),
                }
            )

    current_chunks = (
        selected_materialized.get("chunks")
        if isinstance(selected_materialized, dict) and isinstance(selected_materialized.get("chunks"), list)
        else None
    )

    selected_body_map = None
    if isinstance(selected_materialized, dict):
        selected_snapshot = selected_materialized.get("artifactSnapshot")
        selected_internal = (
            selected_snapshot.get("_internal")
            if isinstance(selected_snapshot, dict) and isinstance(selected_snapshot.get("_internal"), dict)
            else None
        )
        captured_body_map_payload = (
            selected_internal.get("bodyMapPayload")
            if isinstance(selected_internal, dict)
            else None
        )
        if captured_body_map_payload is not None:
            selected_body_map = docx_body_map_from_payload(captured_body_map_payload)
            if selected_body_map is None:
                raise RoundArtifactSnapshotError(
                    "round_snapshot_body_map_corrupt",
                    "Captured DOCX body map could not be reconstructed for RateAudit.",
                )
    contract_candidate_texts: list[str] | None = None
    if selected_body_map is not None:
        effective_paragraphs = selected_materialized.get("paragraphs") if isinstance(selected_materialized, dict) else None
        if isinstance(effective_paragraphs, list) and len(effective_paragraphs) == len(selected_body_map.units):
            selected_body_map = update_docx_body_map_texts(
                selected_body_map,
                [str(item) for item in effective_paragraphs],
                round_number=selected_body_map.round_number,
            )
        contract_candidate_texts = selected_body_map.current_texts()
    elif selected_output is not None and normalized_source.suffix.lower() == ".docx":
        effective_paragraphs = selected_materialized.get("paragraphs") if isinstance(selected_materialized, dict) else None
        contract_candidate_texts = (
            [str(item) for item in effective_paragraphs]
            if isinstance(effective_paragraphs, list)
            else _split_text_into_blocks(_read_rate_audit_text(selected_output, label="Current output"))
        )
    content_contract = build_document_edit_contract(
        normalized_source,
        body_map=selected_body_map,
        candidate_texts=contract_candidate_texts,
        stage="rate_audit",
        report_path=get_document_edit_contract_path(normalized_source),
    )

    route_metadata = selected_round or selected_compare_payload or {}
    prompt_profile = str(route_metadata.get("promptProfile", DEFAULT_PROMPT_PROFILE) or DEFAULT_PROMPT_PROFILE)
    raw_prompt_sequence = route_metadata.get("promptSequence")
    if not isinstance(raw_prompt_sequence, list) or not raw_prompt_sequence:
        try:
            from app_config import load_app_config

            active_config = load_app_config()
            prompt_profile = str(active_config.get("promptProfile", prompt_profile) or prompt_profile)
            raw_prompt_sequence = active_config.get("promptSequence")
        except Exception:
            raw_prompt_sequence = None
    prompt_sequence = normalize_prompt_sequence(prompt_profile, raw_prompt_sequence)
    selected_round_number = int(route_metadata.get("round", 0) or 0)
    current_prompt_id = ""
    next_prompt_id = ""
    if selected_round_number > 0:
        try:
            current_prompt_id = get_prompt_id_for_round(prompt_profile, selected_round_number, prompt_sequence)
        except ValueError:
            current_prompt_id = ""
        if selected_round_number < len(prompt_sequence):
            next_prompt_id = str(prompt_sequence[selected_round_number])
    elif current_stage_id == "source" and prompt_sequence:
        next_prompt_id = str(prompt_sequence[0])

    report = build_rate_audit_report(
        source_text=source_text,
        stages=stages,
        current_chunks=current_chunks,
        source_path=str(normalized_source),
        current_output_path=str(selected_output or ""),
        current_stage_id=current_stage_id,
        content_contract=content_contract,
        current_prompt_id=current_prompt_id,
        next_prompt_id=next_prompt_id,
    )
    report["effectiveText"] = {
        "source": str((selected_materialized or {}).get("source", "source_document")),
        "reviewDecisionCount": len((selected_materialized or {}).get("decisions", {}) or {}),
        "reviewDecisionsPath": str((selected_materialized or {}).get("decisionsPath", "") or ""),
        "chunkCount": len(current_chunks or []),
    }
    report["strategyBinding"] = _build_rate_audit_strategy_binding(
        report,
        source_path=normalized_source,
        output_path=selected_output,
        materialized=selected_materialized,
    )
    _materialize_rate_audit_plateau_state(report, report["strategyBinding"])
    return report


def list_document_histories() -> dict[str, Any]:
    indexed_documents = _list_indexed_history_documents()
    if indexed_documents is not None:
        indexed_records = {
            str(document.get("docId", "")): _indexed_history_document_to_record_entry(document)
            for document in indexed_documents
            if isinstance(document, dict) and str(document.get("docId", "")).strip()
        }
        items = [
            _record_entry_to_history(doc_id, entry, all_records=indexed_records)
            for doc_id, entry in indexed_records.items()
        ]
        items.sort(key=lambda item: (item.get("lastTimestamp", ""), item.get("docId", "")), reverse=True)
        return {
            "items": items,
            "total": len(items),
        }

    records = list_records()
    items = [
        _record_entry_to_history(doc_id, entry, all_records=records)
        for doc_id, entry in records.items()
        if isinstance(entry, dict)
    ]
    items.sort(key=lambda item: (item.get("lastTimestamp", ""), item.get("docId", "")), reverse=True)
    return {
        "items": items,
        "total": len(items),
    }


def query_history_artifact_governance(filters: dict[str, Any] | None = None) -> dict[str, Any]:
    normalized_filters = filters if isinstance(filters, dict) else {}
    try:
        from fyadr_history_db import query_history_artifacts_from_index

        result = query_history_artifacts_from_index(normalized_filters)
    except Exception as exc:
        return {
            "ok": False,
            "source": "sqlite",
            "error": str(exc),
            "filters": normalized_filters,
            "items": [],
            "total": 0,
            "limit": 0,
            "offset": 0,
            "hasMore": False,
            "stats": _empty_artifact_stats(),
        }
    if result is None:
        return {
            "ok": False,
            "source": "sqlite",
            "error": "SQLite history index is unavailable or stale.",
            "filters": normalized_filters,
            "items": [],
            "total": 0,
            "limit": 0,
            "offset": 0,
            "hasMore": False,
            "stats": _empty_artifact_stats(),
        }
    return result


def check_history_database_governance() -> dict[str, Any]:
    try:
        from fyadr_records import check_history_index

        return check_history_index(strict=False)
    except Exception as exc:
        return {
            "ok": False,
            "error": str(exc),
            "issueCount": 1,
            "errorCount": 1,
            "warningCount": 0,
            "repairableIssueCount": 1,
            "issues": [
                {
                    "code": "history_index_check_failed",
                    "severity": "error",
                    "message": str(exc),
                    "repairable": True,
                    "recommendedAction": "history-db-repair",
                }
            ],
        }


def repair_history_database_governance() -> dict[str, Any]:
    try:
        from fyadr_records import repair_history_index

        return repair_history_index(strict=True)
    except Exception as exc:
        return {
            "ok": False,
            "error": str(exc),
            "before": check_history_database_governance(),
            "after": {
                "ok": False,
                "issueCount": 1,
                "errorCount": 1,
                "warningCount": 0,
                "repairableIssueCount": 1,
                "issues": [
                    {
                        "code": "history_index_repair_failed",
                        "severity": "error",
                        "message": str(exc),
                        "repairable": True,
                        "recommendedAction": "history-db-repair",
                    }
                ],
            },
        }


def list_history_database_backups(validate: bool = False) -> dict[str, Any]:
    from fyadr_records import list_history_index_backups

    return list_history_index_backups(validate=validate)


def backup_history_database_governance(reason: str = "manual", keep: int = 12) -> dict[str, Any]:
    from fyadr_records import backup_history_index

    return backup_history_index(reason=reason, keep=keep)


def compact_history_database_governance(create_backup: bool = True, keep: int = 12) -> dict[str, Any]:
    from fyadr_records import compact_history_index

    return compact_history_index(create_backup=create_backup, keep=keep)


def recover_history_database_governance(backup_path: str | None = None, keep: int = 12) -> dict[str, Any]:
    from fyadr_records import recover_history_index

    return recover_history_index(backup_path=backup_path, keep=keep)


def get_history_database_maintenance_summary() -> dict[str, Any]:
    from fyadr_records import get_history_index_maintenance_summary

    return get_history_index_maintenance_summary()


def ensure_history_database_ready(reason: str = "app", max_age_seconds: float = 30, compact: bool = True) -> dict[str, Any]:
    from fyadr_records import ensure_history_governance_ready

    return ensure_history_governance_ready(reason=reason, max_age_seconds=max_age_seconds, compact=compact)


HISTORY_ORPHAN_SCAN_DIRS: dict[str, Path] = {
    "sources": ROOT_DIR / "origin",
    "intermediate": ROOT_DIR / "finish" / "intermediate",
    "exports": ROOT_DIR / "finish" / "web_exports",
    "reports": ROOT_DIR / "finish" / "detection_reports",
}
HISTORY_ORPHAN_EXCLUDED_FILENAMES: set[str] = set()
ROUND_ARTIFACT_PATTERN = re.compile(r"_round\d+(?:[_.]|$)")


def _empty_artifact_stats() -> dict[str, Any]:
    return {
        "total": 0,
        "existing": 0,
        "intermediate": 0,
        "exports": 0,
        "reports": 0,
        "sources": 0,
        "external": 0,
        "missing": 0,
        "bytes": 0,
    }


def _path_to_workspace_display(path: Path) -> str:
    normalized_path = normalize_path(path)
    try:
        return str(normalized_path.relative_to(ROOT_DIR)).replace("\\", "/")
    except ValueError:
        return str(normalized_path)


def _history_artifact_kind(path: Path) -> str:
    normalized_path = normalize_path(path)
    try:
        relative = normalized_path.relative_to(ROOT_DIR)
    except ValueError:
        return "external"
    parts = relative.parts
    if parts and parts[0] == "origin":
        return "sources"
    if len(parts) < 2 or parts[0] != "finish":
        return "external"
    if parts[1] == "web_exports":
        return "exports"
    if parts[1] == "detection_reports":
        return "reports"
    if normalized_path.name.endswith((".audit.json", ".ooxml_audit.json", ".text_integrity.json", ".guard.json", "_validation.json", "_format_preflight.json")):
        return "reports"
    return "intermediate"


def _is_cleanable_history_artifact(path: Path) -> bool:
    normalized_path = normalize_path(path)
    if normalized_path.name in HISTORY_ORPHAN_EXCLUDED_FILENAMES:
        return False
    if not normalized_path.exists() or not normalized_path.is_file():
        return False
    if is_path_under(normalized_path, HISTORY_ORPHAN_SCAN_DIRS["sources"]):
        return True
    if is_path_under(normalized_path, HISTORY_ORPHAN_SCAN_DIRS["exports"]):
        return True
    if is_path_under(normalized_path, HISTORY_ORPHAN_SCAN_DIRS["reports"]):
        return normalized_path.suffix.lower() in {".pdf", ".json", ".txt"}
    if not is_path_under(normalized_path, HISTORY_ORPHAN_SCAN_DIRS["intermediate"]):
        return False

    name = normalized_path.name
    if ROUND_ARTIFACT_PATTERN.search(name):
        return True
    if name.endswith(("_format_preflight.json", "_review_decisions.json", ".ooxml_audit.json", ".text_integrity.json")):
        return True
    if ("_export" in name or "_reviewed_export" in name) and normalized_path.suffix.lower() in {".docx", ".txt", ".json"}:
        return True
    return False


def _normalize_protected_paths(protected_paths: object | None) -> set[Path]:
    raw_items: list[Any]
    if protected_paths is None:
        raw_items = []
    elif isinstance(protected_paths, list):
        raw_items = protected_paths
    else:
        raw_items = [protected_paths]

    normalized_paths: set[Path] = set()
    for item in raw_items:
        value = str(item or "").strip()
        if not value:
            continue
        try:
            normalized_paths.add(normalize_path(Path(value)))
        except Exception:
            continue
    return normalized_paths


def _expand_certified_export_bundle_paths(paths: set[Path]) -> set[Path]:
    expanded = set(paths)
    candidate_manifests: set[Path] = set()
    for path in paths:
        if path.suffix.lower() == ".docx":
            candidate_manifests.add(path.with_suffix(".evidence.json"))
        elif path.name.endswith(".evidence.json"):
            candidate_manifests.add(path)

    allowed_bundle_roots = tuple(HISTORY_ORPHAN_SCAN_DIRS.values())
    for manifest_path in candidate_manifests:
        if not any(is_path_under(manifest_path, root) for root in allowed_bundle_roots):
            continue
        if not manifest_path.exists() or not manifest_path.is_file():
            continue
        try:
            payload = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if not isinstance(payload, dict) or str(payload.get("status", "")) != "passed":
            continue
        raw_artifact_path = str(payload.get("artifactPath", "") or "").strip()
        if not raw_artifact_path:
            continue
        artifact_path = normalize_path(Path(raw_artifact_path))
        if artifact_path not in paths and manifest_path not in paths:
            continue
        expanded.add(manifest_path)
        if any(is_path_under(artifact_path, root) for root in allowed_bundle_roots):
            expanded.add(artifact_path)
        reports = payload.get("reports")
        if isinstance(reports, dict):
            for raw_report_path in reports.values():
                value = str(raw_report_path or "").strip()
                if not value:
                    continue
                report_path = normalize_path(Path(value))
                if any(is_path_under(report_path, root) for root in allowed_bundle_roots):
                    expanded.add(report_path)
    return expanded


def _collect_referenced_history_artifacts(protected_paths: object | None = None) -> set[Path]:
    referenced_paths = _normalize_protected_paths(protected_paths)
    referenced_paths = _expand_certified_export_bundle_paths(referenced_paths)
    indexed_paths = list_referenced_history_artifact_paths()
    if indexed_paths is not None:
        for raw_path in indexed_paths:
            try:
                referenced_paths.add(normalize_path(Path(raw_path)))
            except Exception:
                continue
        return _expand_certified_export_bundle_paths(referenced_paths)

    records = list_records()
    for entry in records.values():
        if not isinstance(entry, dict):
            continue
        origin_path = entry.get("origin_path")
        if isinstance(origin_path, str) and origin_path.strip():
            try:
                referenced_paths.add(normalize_path(Path(origin_path)))
            except Exception:
                pass
        rounds = entry.get("rounds")
        if not isinstance(rounds, list):
            continue
        normalized_rounds = [item for item in rounds if isinstance(item, dict)]
        referenced_paths.update(_collect_round_file_paths(normalized_rounds))
        for round_item in normalized_rounds:
            compare_path = round_item.get("compare_path")
            if isinstance(compare_path, str) and compare_path.strip():
                try:
                    normalized_compare_path = normalize_path(Path(compare_path))
                    referenced_paths.add(normalized_compare_path.with_name(f"{normalized_compare_path.stem}_review_decisions.json"))
                except Exception:
                    continue
    return _expand_certified_export_bundle_paths(referenced_paths)


def _iter_cleanable_history_artifacts() -> list[Path]:
    candidates: list[Path] = []
    for root_path in HISTORY_ORPHAN_SCAN_DIRS.values():
        if not root_path.exists():
            continue
        for path in root_path.rglob("*"):
            try:
                normalized_path = normalize_path(path)
                if _is_cleanable_history_artifact(normalized_path):
                    candidates.append(normalized_path)
            except OSError:
                continue
    return sorted(set(candidates), key=lambda item: str(item).lower())


def _artifact_file_entry(path: Path) -> dict[str, Any]:
    normalized_path = normalize_path(path)
    kind = _history_artifact_kind(normalized_path)
    size = 0
    modified_at = ""
    try:
        stat = normalized_path.stat()
        size = stat.st_size
        modified_at = datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat().replace("+00:00", "Z")
    except OSError:
        pass
    return {
        "path": str(normalized_path),
        "relativePath": _path_to_workspace_display(normalized_path),
        "kind": kind,
        "bytes": size,
        "modifiedAt": modified_at,
    }


def _build_stats_from_file_entries(entries: list[dict[str, Any]]) -> dict[str, Any]:
    stats = _empty_artifact_stats()
    for entry in entries:
        kind = str(entry.get("kind", "external"))
        size = int(entry.get("bytes", 0) or 0)
        stats["total"] += 1
        stats["existing"] += 1
        stats["bytes"] += max(0, size)
        if kind in {"sources", "intermediate", "exports", "reports"}:
            stats[kind] += 1
        else:
            stats["external"] += 1
    return stats


def _build_kind_stats(entries: list[dict[str, Any]]) -> dict[str, dict[str, int]]:
    kind_stats = {
        "sources": {"files": 0, "bytes": 0},
        "intermediate": {"files": 0, "bytes": 0},
        "exports": {"files": 0, "bytes": 0},
        "reports": {"files": 0, "bytes": 0},
    }
    for entry in entries:
        kind = str(entry.get("kind", ""))
        if kind not in kind_stats:
            continue
        kind_stats[kind]["files"] += 1
        kind_stats[kind]["bytes"] += int(entry.get("bytes", 0) or 0)
    return kind_stats


def _scan_history_orphan_entries(protected_paths: object | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], set[Path]]:
    referenced_paths = _collect_referenced_history_artifacts(protected_paths)
    candidates = _iter_cleanable_history_artifacts()
    referenced_files = [
        _artifact_file_entry(path)
        for path in candidates
        if path in referenced_paths
    ]
    orphan_files = [
        _artifact_file_entry(path)
        for path in candidates
        if path not in referenced_paths
    ]
    orphan_files.sort(key=lambda item: (str(item.get("kind", "")), str(item.get("relativePath", "")).lower()))
    return referenced_files, orphan_files, referenced_paths


def scan_history_orphan_artifacts(protected_paths: object | None = None) -> dict[str, Any]:
    referenced_files, orphan_files, referenced_paths = _scan_history_orphan_entries(protected_paths)
    protected = sorted(
        _path_to_workspace_display(path)
        for path in referenced_paths
        if any(is_path_under(path, root_path) for root_path in HISTORY_ORPHAN_SCAN_DIRS.values())
    )
    return {
        "scannedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "rootDir": str(ROOT_DIR),
        "scanDirs": {
            kind: _path_to_workspace_display(path)
            for kind, path in HISTORY_ORPHAN_SCAN_DIRS.items()
        },
        "protectedPaths": protected,
        "referencedStats": _build_stats_from_file_entries(referenced_files),
        "orphanStats": _build_stats_from_file_entries(orphan_files),
        "orphanKindStats": _build_kind_stats(orphan_files),
        "orphanFiles": orphan_files[:200],
        "hasMore": len(orphan_files) > 200,
        "totalOrphanFiles": len(orphan_files),
    }


def delete_history_orphan_artifacts(protected_paths: object | None = None) -> dict[str, Any]:
    before_scan = scan_history_orphan_artifacts(protected_paths)
    _, orphan_files, _ = _scan_history_orphan_entries(protected_paths)
    deleted_entries: list[dict[str, Any]] = []
    failed_files: list[dict[str, str]] = []

    for item in orphan_files:
        if not isinstance(item, dict):
            continue
        raw_path = str(item.get("path", "")).strip()
        if not raw_path:
            continue
        path = normalize_path(Path(raw_path))
        if not _is_cleanable_history_artifact(path):
            continue
        entry = _artifact_file_entry(path)
        try:
            path.unlink()
            deleted_entries.append(entry)
        except OSError as exc:
            failed_files.append({
                "path": entry["relativePath"],
                "message": str(exc),
            })

    after_scan = scan_history_orphan_artifacts(protected_paths)
    return {
        "deletedFiles": [str(item["relativePath"]) for item in deleted_entries],
        "deletedFileStats": _build_stats_from_file_entries(deleted_entries),
        "failedFiles": failed_files,
        "before": before_scan,
        "after": after_scan,
    }


def reset_round_progress(
    source_path: str,
    prompt_profile: str,
    round_number: int,
    prompt_sequence: object | None = None,
) -> dict[str, Any]:
    normalized_profile = normalize_prompt_profile(prompt_profile)
    normalized_prompt_sequence = normalize_prompt_sequence(normalized_profile, prompt_sequence)
    normalized_source = normalize_path(Path(source_path))
    try:
        relative_doc_id = normalized_source.relative_to(ROOT_DIR)
        doc_id = normalize_doc_id(str(relative_doc_id).replace("\\", "/"))
    except ValueError:
        doc_id = normalize_doc_id(str(normalized_source))
    target_round = int(round_number)
    artifact_stem = build_round_artifact_stem(
        doc_id,
        normalized_profile,
        normalized_prompt_sequence,
    )
    intermediate_dir = ROOT_DIR / "finish" / "intermediate"
    output_path = intermediate_dir / f"{artifact_stem}_round{target_round}.txt"
    manifest_path = intermediate_dir / f"{artifact_stem}_round{target_round}_manifest.json"
    deleted_files: list[str] = []

    try:
        records_payload = delete_rounds(
            doc_id,
            target_round,
            prompt_profile=normalized_profile,
            prompt_sequence=normalized_prompt_sequence if is_prompt_sequence_customizable(normalized_profile) else None,
        )
    except ValueError as exc:
        message = str(exc)
        if "Document record not found" not in message and "No rounds found" not in message:
            raise
        records_payload = {
            "docId": doc_id,
            "deletedRounds": [],
            "deletedFiles": [],
            "remainingRounds": [],
        }

    candidate_paths = {
        output_path,
        intermediate_dir / f"{artifact_stem}_round{target_round}_input.txt",
        output_path.with_name(f"{output_path.stem}_checkpoint.json"),
        output_path.with_name(f"{output_path.stem}_compare.json"),
        output_path.with_name(f"{output_path.stem}_quality.json"),
        output_path.with_name(f"{output_path.stem}_bodymap.json"),
        intermediate_dir / f"{artifact_stem}_round{target_round}_body_map.json",
        intermediate_dir / f"{artifact_stem}_round{target_round}_validation.json",
        manifest_path,
    }

    for candidate_path in candidate_paths:
        try:
            normalized_path = normalize_path(candidate_path)
            if normalized_path.exists():
                normalized_path.unlink()
                deleted_files.append(str(normalized_path))
        except Exception:
            continue

    return {
        "docId": doc_id,
        "round": target_round,
        "promptProfile": normalized_profile,
        "promptSequence": normalized_prompt_sequence,
        "deletedRounds": records_payload.get("deletedRounds", []),
        "deletedFiles": [*records_payload.get("deletedFiles", []), *deleted_files],
    }


def _build_checkpoint_resume_details(
    *,
    effective_round: int,
    chunk_ids: Any,
    chunk_outputs: Any,
    completed_chunks: int,
    total_chunks: int,
    last_error: str,
) -> dict[str, Any]:
    ordered_chunk_ids = [str(item) for item in chunk_ids if isinstance(item, str)] if isinstance(chunk_ids, list) else []
    completed_chunk_ids = set(chunk_outputs) if isinstance(chunk_outputs, dict) else set()
    pending_chunk_ids = [chunk_id for chunk_id in ordered_chunk_ids if chunk_id not in completed_chunk_ids]
    next_chunk_id = pending_chunk_ids[0] if pending_chunk_ids else ""
    next_chunk_index = (ordered_chunk_ids.index(next_chunk_id) + 1) if next_chunk_id in ordered_chunk_ids else 0
    failed_chunk_id = ""
    failed_match = re.search(r"Chunk\s+([A-Za-z0-9_.:-]+)\s+failed", last_error)
    if failed_match:
        failed_chunk_id = failed_match.group(1)

    if total_chunks and completed_chunks >= total_chunks:
        resume_stage = "finalize_output"
        resume_action_label = "继续收尾"
        resume_explanation = f"第 {effective_round} 轮所有分块都已落盘，继续时会直接进入合并、Diff 和记录写入阶段，不会重跑 100% 已完成的分块。"
    elif next_chunk_id:
        resume_stage = "continue_chunks"
        resume_action_label = f"从第 {next_chunk_index}/{total_chunks or '?'} 块继续"
        resume_explanation = f"继续时会复用已完成分块，并从 {next_chunk_id} 开始处理；不会从第一块重跑。"
    else:
        resume_stage = "inspect_checkpoint"
        resume_action_label = "检查断点"
        resume_explanation = "断点存在，但无法可靠判断下一块；建议先刷新状态，仍异常再放弃本轮进度。"

    return {
        "remainingChunks": max(0, total_chunks - completed_chunks) if total_chunks else 0,
        "nextChunkId": next_chunk_id,
        "nextChunkIndex": next_chunk_index,
        "failedChunkId": failed_chunk_id,
        "resumeStage": resume_stage,
        "resumeActionLabel": resume_action_label,
        "resumeExplanation": resume_explanation,
    }


def get_round_progress_status(
    source_path: str,
    prompt_profile: str,
    round_number: int | None = None,
    prompt_sequence: object | None = None,
) -> dict[str, Any]:
    normalized_profile = normalize_prompt_profile(prompt_profile)
    normalized_prompt_sequence = normalize_prompt_sequence(normalized_profile, prompt_sequence)
    normalized_source = normalize_path(Path(source_path))
    effective_round = int(round_number or 0)
    if effective_round <= 0:
        status = get_document_status(
            str(normalized_source),
            prompt_profile=normalized_profile,
            prompt_sequence=normalized_prompt_sequence,
        )
        effective_round = int(status.get("nextRound", 0) or 0)
    if effective_round <= 0:
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": None,
            "checkpointExists": False,
            "canResume": False,
            "completedChunks": 0,
            "totalChunks": 0,
            "progressPercent": 0,
            "checkpointPath": "",
            "lastError": "",
            "updatedAt": "",
            "validationEventCount": 0,
            "message": "当前没有待执行轮次。",
        }

    context = build_round_context(
        normalized_source,
        round_number=effective_round,
        prompt_profile=normalized_profile,
        prompt_sequence=normalized_prompt_sequence,
    )
    checkpoint_path = get_round_checkpoint_path(context.output_text_path)
    if not checkpoint_path.exists():
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": effective_round,
            "checkpointExists": False,
            "canResume": False,
            "completedChunks": 0,
            "totalChunks": 0,
            "progressPercent": 0,
            "checkpointPath": str(checkpoint_path),
            "lastError": "",
            "updatedAt": "",
            "validationEventCount": 0,
            "message": "当前轮次暂无断点。",
        }

    try:
        payload = json.loads(checkpoint_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": effective_round,
            "checkpointExists": True,
            "canResume": False,
            "completedChunks": 0,
            "totalChunks": 0,
            "progressPercent": 0,
            "checkpointPath": str(checkpoint_path),
            "lastError": safe_public_error_message(exc),
            "updatedAt": "",
            "validationEventCount": 0,
            "message": "断点文件无法读取，可放弃本轮进度后重新开始。",
        }
    if not isinstance(payload, dict):
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": effective_round,
            "checkpointExists": True,
            "canResume": False,
            "completedChunks": 0,
            "totalChunks": 0,
            "progressPercent": 0,
            "checkpointPath": str(checkpoint_path),
            "lastError": "Checkpoint payload is not an object.",
            "updatedAt": "",
            "validationEventCount": 0,
            "message": "断点文件格式异常，可放弃本轮进度后重新开始。",
        }
    if payload.get("completed") is True:
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": effective_round,
            "checkpointExists": False,
            "canResume": False,
            "completedChunks": 0,
            "totalChunks": 0,
            "progressPercent": 0,
            "checkpointPath": str(checkpoint_path),
            "lastError": "",
            "updatedAt": str(payload.get("updated_at", "") or ""),
            "validationEventCount": 0,
            "message": "当前轮次已完成；残留断点标记无需处理。",
        }
    checkpoint_last_error, checkpoint_last_error_details = sanitize_persisted_error(
        payload.get("last_error"),
        payload.get("last_error_details"),
    )
    try:
        checkpoint_metadata: dict[str, object] | None = None
        if effective_round > 1:
            parent_snapshot = _capture_downstream_parent_snapshot(
                context,
                expected_previous_compare_revision=None,
                require_revision=False,
                include_internal=True,
            )
            internal = (parent_snapshot or {}).get("_internal")
            if not isinstance(internal, dict) or not isinstance(internal.get("effectiveText"), str):
                raise ValueError("Parent snapshot has no effective checkpoint input.")
            input_text = str(internal["effectiveText"])
            checkpoint_metadata = {
                "parent_input_provenance": {
                    "parentOutputPath": str(context.parent_output_path or ""),
                    "parentCompareRevision": str((parent_snapshot or {}).get("compareRevision", "") or ""),
                    "parentContentRevision": str((parent_snapshot or {}).get("contentRevision", "") or ""),
                    "parentReviewRevision": str((parent_snapshot or {}).get("reviewRevision", "") or ""),
                    "effectiveInputSha256": str((parent_snapshot or {}).get("effectiveTextSha256", "") or ""),
                    "materializationSource": str((parent_snapshot or {}).get("materializationSource", "") or ""),
                    "parentArtifactSnapshotDigest": str((parent_snapshot or {}).get("artifactSnapshotDigest", "") or ""),
                }
            }
        else:
            input_text = context.input_text_path.read_text(encoding="utf-8")
        chunk_metric = get_chunk_metric(normalized_profile, normalized_prompt_sequence)
        manifest = build_manifest(input_text, chunk_limit=DEFAULT_CHUNK_LIMIT, chunk_metric=chunk_metric)
        status_style_profile = build_global_style_profile_from_texts(
            [str(chunk.text) for chunk in manifest.chunks]
        )
        checkpoint_metadata = {
            **(checkpoint_metadata or {}),
            "style_card_version": STYLE_CARD_VERSION,
            "global_style_profile_sha256": _sha256_json(status_style_profile),
            "source_pattern_profile_sha256": str(
                (status_style_profile.get("documentPatternBaseline") or {}).get("profileSha256", "")
                if isinstance(status_style_profile.get("documentPatternBaseline"), dict)
                else ""
            ),
        }
        prompt_text = load_prompt(normalized_profile, effective_round, normalized_prompt_sequence)
        checkpoint_signature = _build_checkpoint_signature(
            doc_id=context.doc_id,
            round_number=effective_round,
            prompt_profile=normalized_profile,
            prompt_sequence=normalized_prompt_sequence,
            input_path=context.input_text_path,
            output_path=context.output_text_path,
            manifest_path=context.manifest_path,
            manifest_chunk_ids=[chunk.chunk_id for chunk in manifest.chunks],
            input_sha256=_sha256_text(input_text),
            prompt_sha256=_sha256_text(prompt_text),
            chunk_limit=DEFAULT_CHUNK_LIMIT,
            chunk_metric=chunk_metric,
            checkpoint_metadata=checkpoint_metadata,
        )
    except Exception as exc:
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": effective_round,
            "checkpointExists": True,
            "canResume": False,
            "completedChunks": 0,
            "totalChunks": 0,
            "progressPercent": 0,
            "checkpointPath": str(checkpoint_path),
            "lastError": safe_public_error_message(exc),
            "updatedAt": str(payload.get("updated_at", "") or ""),
            "validationEventCount": 0,
            "resumeStage": "inspect_checkpoint",
            "resumeActionLabel": "检查断点",
            "resumeExplanation": "断点存在，但当前正文状态无法校验；为避免误续跑，请刷新历史文档后再继续。",
            "message": "断点存在，但无法确认是否属于当前正文。",
        }
    if not _is_checkpoint_compatible(payload, checkpoint_signature):
        return {
            "sourcePath": str(normalized_source),
            "promptProfile": normalized_profile,
            "promptSequence": normalized_prompt_sequence,
            "round": effective_round,
            "checkpointExists": True,
            "canResume": False,
            "completedChunks": int(payload.get("completed_chunk_count", 0) or 0),
            "totalChunks": len(payload.get("chunk_ids", [])) if isinstance(payload.get("chunk_ids"), list) else 0,
            "progressPercent": 0,
            "checkpointPath": str(checkpoint_path),
            "lastError": checkpoint_last_error,
            "lastErrorDetails": checkpoint_last_error_details,
            "updatedAt": str(payload.get("updated_at", "") or ""),
            "validationEventCount": 0,
            "resumeStage": "inspect_checkpoint",
            "resumeActionLabel": "断点不匹配",
            "resumeExplanation": "该断点与当前正文、Prompt 或分块清单不一致，继续会重新生成；请确认历史文档和自定义流程一致。",
            "message": "断点与当前正文不匹配，已阻止假续跑提示。",
        }
    chunk_outputs = payload.get("chunk_outputs")
    chunk_ids = payload.get("chunk_ids")
    validation_events = payload.get("validation_events")
    completed_chunks = int(payload.get("completed_chunk_count", 0) or 0)
    if isinstance(chunk_outputs, dict):
        completed_chunks = max(completed_chunks, len(chunk_outputs))
    total_chunks = len(chunk_ids) if isinstance(chunk_ids, list) else 0
    progress_percent = round((completed_chunks / total_chunks) * 100) if total_chunks else 0
    last_error = checkpoint_last_error
    last_error_details = checkpoint_last_error_details
    updated_at = str(payload.get("updated_at", "") or "").strip()
    validation_event_count = len(validation_events) if isinstance(validation_events, list) else 0
    resume_details = _build_checkpoint_resume_details(
        effective_round=effective_round,
        chunk_ids=chunk_ids,
        chunk_outputs=chunk_outputs,
        completed_chunks=completed_chunks,
        total_chunks=total_chunks,
        last_error=last_error,
    )
    return {
        "sourcePath": str(normalized_source),
        "promptProfile": normalized_profile,
        "promptSequence": normalized_prompt_sequence,
        "round": effective_round,
        "checkpointExists": True,
        "canResume": True,
        "completedChunks": completed_chunks,
        "totalChunks": total_chunks,
        "progressPercent": progress_percent,
        "checkpointPath": str(checkpoint_path),
        "lastError": last_error,
        "lastErrorDetails": last_error_details,
        "updatedAt": updated_at,
        "validationEventCount": validation_event_count,
        **resume_details,
        "message": (
            f"发现第 {effective_round} 轮断点：已完成 {completed_chunks}/{total_chunks} 块，{resume_details['resumeActionLabel']}。"
            if total_chunks
            else f"发现第 {effective_round} 轮断点。"
        ),
    }


def delete_document_history(
    doc_id: str,
    from_round: int | None = None,
    prompt_profile: str | None = None,
    prompt_sequence: object | None = None,
    mode: str | None = None,
) -> dict[str, Any]:
    normalized_doc_id = normalize_doc_id(doc_id)
    if from_round is None:
        return delete_document(normalized_doc_id, mode=mode)
    normalized_profile = normalize_prompt_profile(prompt_profile) if prompt_profile is not None else None
    return delete_rounds(
        normalized_doc_id,
        from_round,
        prompt_profile=normalized_profile,
        prompt_sequence=normalize_prompt_sequence(normalized_profile, prompt_sequence) if normalized_profile and is_prompt_sequence_customizable(normalized_profile) else None,
        mode=mode,
    )


def preview_document_history_delete(
    doc_id: str,
    from_round: int | None = None,
    prompt_profile: str | None = None,
    prompt_sequence: object | None = None,
    mode: str | None = None,
) -> dict[str, Any]:
    normalized_doc_id = normalize_doc_id(doc_id)
    normalized_profile = normalize_prompt_profile(prompt_profile) if prompt_profile is not None else None
    return preview_delete_document(
        normalized_doc_id,
        from_round,
        prompt_profile=normalized_profile,
        prompt_sequence=normalize_prompt_sequence(normalized_profile, prompt_sequence) if normalized_profile and is_prompt_sequence_customizable(normalized_profile) else None,
        mode=mode,
    )


def _capture_downstream_parent_snapshot(
    context: Any,
    *,
    expected_previous_compare_revision: str | None,
    expected_parent_input_binding: dict[str, Any] | None = None,
    require_revision: bool,
    require_complete_binding: bool = False,
    include_internal: bool,
) -> dict[str, Any] | None:
    expected_binding = _normalize_expected_parent_input_binding(
        expected_previous_compare_revision,
        expected_parent_input_binding,
    )
    if int(getattr(context, "round_number", 0) or 0) <= 1:
        if expected_binding:
            raise StaleRoundInputError(
                "",
                ["parent_round_missing"],
                "A parent input binding was supplied, but the next run no longer has a parent round.",
            )
        return None
    parent_output_path = getattr(context, "parent_output_path", None)
    if not isinstance(parent_output_path, Path):
        raise StaleRoundInputError("", ["parent_output_missing"], "Server-derived parent output is unavailable.")
    requested_revision = expected_binding.get("compareRevision", "")
    if require_complete_binding:
        missing_fields = [
            field
            for field, _mismatch_code in PARENT_INPUT_BINDING_FIELDS
            if not expected_binding.get(field)
        ]
        if missing_fields:
            raise RoundInputRevisionRequiredError(
                "A complete previous-round input binding is required for downstream rounds. "
                f"Missing: {', '.join(missing_fields)}."
            )
    elif require_revision and not requested_revision:
        raise RoundInputRevisionRequiredError("expectedPreviousCompareRevision is required for downstream rounds.")
    for field, value in expected_binding.items():
        if len(value) > 200:
            raise ValueError(f"Expected parent input binding field {field} is invalid.")
    try:
        snapshot = read_round_artifact_snapshot(
            parent_output_path,
            include_internal=include_internal,
        )
    except RoundArtifactSnapshotError as exc:
        raise StaleRoundInputError(
            "",
            [exc.code],
            f"Parent round snapshot is not safe to use: {exc}",
        ) from exc
    current_revision = str(snapshot.get("compareRevision", "") or "").strip()
    if not current_revision:
        raise StaleRoundInputError("", ["parent_compare_revision_missing"])
    current_binding = _round_snapshot_parent_input_binding(snapshot)
    mismatch_codes = [
        mismatch_code
        for field, mismatch_code in PARENT_INPUT_BINDING_FIELDS
        if expected_binding.get(field)
        and expected_binding.get(field) != current_binding.get(field)
    ]
    if mismatch_codes:
        raise StaleRoundInputError(
            current_revision,
            mismatch_codes,
            "Downstream round approval no longer matches the current parent snapshot.",
        )
    return snapshot


def preflight_run_round_input(
    source_path: str,
    model_config: dict[str, Any],
    *,
    expected_previous_compare_revision: str | None = None,
    expected_parent_input_binding: dict[str, Any] | None = None,
    require_revision: bool = True,
) -> dict[str, Any]:
    """Bind a run request to one complete parent generation before HTTP 202."""

    prompt_profile = normalize_prompt_profile(model_config.get("promptProfile", DEFAULT_PROMPT_PROFILE))
    prompt_sequence = normalize_prompt_sequence(prompt_profile, model_config.get("promptSequence"))
    status = get_document_status(source_path, prompt_profile=prompt_profile, prompt_sequence=prompt_sequence)
    if bool(status.get("isComplete")):
        raise ValueError("Document already completed all selected rounds.")
    effective_round = int(status.get("nextRound") or 1)
    context = build_round_context(
        source_path,
        round_number=effective_round,
        prompt_profile=prompt_profile,
        prompt_sequence=prompt_sequence,
    )
    snapshot = _capture_downstream_parent_snapshot(
        context,
        expected_previous_compare_revision=expected_previous_compare_revision,
        expected_parent_input_binding=expected_parent_input_binding,
        require_revision=require_revision,
        require_complete_binding=require_revision,
        include_internal=False,
    )
    parent_input_binding = _round_snapshot_parent_input_binding(snapshot)
    return {
        "round": effective_round,
        "parentOutputPath": str(getattr(context, "parent_output_path", "") or ""),
        "parentCompareRevision": str((snapshot or {}).get("compareRevision", "") or ""),
        "parentContentRevision": str((snapshot or {}).get("contentRevision", "") or ""),
        "parentReviewRevision": str((snapshot or {}).get("reviewRevision", "") or ""),
        "parentArtifactSnapshotDigest": str((snapshot or {}).get("artifactSnapshotDigest", "") or ""),
        "effectiveInputSha256": str((snapshot or {}).get("effectiveTextSha256", "") or ""),
        "materializationSource": str((snapshot or {}).get("materializationSource", "") or ""),
        "parentInputBinding": parent_input_binding,
    }


def run_round_for_app(
    source_path: str,
    model_config: dict[str, Any],
    round_number: int | None = None,
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
    cancel_check: Callable[[], bool] | None = None,
    expected_previous_compare_revision: str | None = None,
    expected_parent_input_binding: dict[str, Any] | None = None,
) -> dict[str, Any]:
    from round_helper import run_document_round

    prompt_profile = normalize_prompt_profile(model_config.get("promptProfile", DEFAULT_PROMPT_PROFILE))
    prompt_sequence = normalize_prompt_sequence(prompt_profile, model_config.get("promptSequence"))
    status = get_document_status(source_path, prompt_profile=prompt_profile, prompt_sequence=prompt_sequence)
    max_rounds = get_max_rounds(prompt_profile, prompt_sequence)
    if bool(status.get("isComplete")):
        raise ValueError(f"Document already completed all {max_rounds} rounds.")
    effective_round = int(round_number or status.get("nextRound") or 1)
    if effective_round < 1 or effective_round > max_rounds:
        raise ValueError(f"Round {effective_round} is outside the selected {max_rounds} round prompt workflow.")
    effective_model_config = _resolve_round_model_config(model_config, prompt_profile, effective_round)
    context = build_round_context(
        source_path,
        round_number=effective_round,
        prompt_profile=prompt_profile,
        prompt_sequence=prompt_sequence,
    )
    parent_artifact_snapshot = _capture_downstream_parent_snapshot(
        context,
        expected_previous_compare_revision=expected_previous_compare_revision,
        expected_parent_input_binding=expected_parent_input_binding,
        require_revision=(
            expected_previous_compare_revision is not None
            or expected_parent_input_binding is not None
        ),
        require_complete_binding=expected_parent_input_binding is not None,
        include_internal=True,
    )

    base_url = str(effective_model_config.get("baseUrl", "")).strip()
    api_key = str(effective_model_config.get("apiKey", "")).strip()
    model = str(effective_model_config.get("model", "")).strip()
    api_type = str(effective_model_config.get("apiType", "chat_completions")).strip()
    streaming = normalize_streaming(
        effective_model_config.get("streaming", effective_model_config.get("stream"))
    )
    temperature = float(effective_model_config.get("temperature", model_config.get("temperature", 0.7)) or 0.7)
    configured_timeout_seconds = _coerce_int_config(
        effective_model_config.get("requestTimeoutSeconds", DEFAULT_REQUEST_TIMEOUT_SECONDS),
        default=DEFAULT_REQUEST_TIMEOUT_SECONDS,
        minimum=30,
        maximum=3600,
    )
    request_timeout_seconds = _coerce_rewrite_timeout_seconds(configured_timeout_seconds)
    configured_max_retries = _coerce_int_config(
        effective_model_config.get("maxRetries", DEFAULT_MAX_RETRIES),
        default=DEFAULT_MAX_RETRIES,
        minimum=0,
        maximum=10,
    )
    max_retries = max(configured_max_retries, MIN_REWRITE_TRANSIENT_RETRIES)
    rewrite_concurrency = _coerce_int_config(
        model_config.get("rewriteConcurrency", DEFAULT_REWRITE_CONCURRENCY),
        default=DEFAULT_REWRITE_CONCURRENCY,
        minimum=1,
        maximum=MAX_REWRITE_CONCURRENCY,
    )

    if not base_url or not api_key or not model:
        raise ValueError("Model configuration is incomplete.")

    if progress_callback is not None:
        progress_callback({
            "phase": "model-selected",
            "round": effective_round,
            "roundModel": {
                "providerId": str(effective_model_config.get("providerId", "")),
                "providerName": str(effective_model_config.get("providerName", "")),
                "baseUrl": base_url,
                "model": model,
                "apiType": api_type,
                "temperature": temperature,
                "rateLimitWindowMinutes": _coerce_rate_window_minutes(effective_model_config),
                "rateLimitMaxRequests": _coerce_rate_max_requests(effective_model_config),
                "routeSource": str(effective_model_config.get("routeSource", "default")),
            },
            "concurrency": rewrite_concurrency,
            "configuredConcurrency": rewrite_concurrency,
            "requestTimeoutSeconds": request_timeout_seconds,
            "configuredRequestTimeoutSeconds": configured_timeout_seconds,
            "maxRetries": max_retries,
            "configuredMaxRetries": configured_max_retries,
        })

    def ensure_not_cancelled() -> None:
        if cancel_check is not None and cancel_check():
            raise RuntimeError("Run was interrupted by user. Completed chunks are kept; click continue to resume this round.")

    rate_limiter = _build_provider_rate_limiter(effective_model_config)

    def transform(_: str, prompt_input: str, __: int, chunk_id: str) -> str:
        ensure_not_cancelled()
        _wait_for_provider_cooldown(effective_model_config, cancel_check=cancel_check)
        ensure_not_cancelled()
        rate_limiter()
        ensure_not_cancelled()

        def emit_retry_wait(event: dict[str, object]) -> None:
            if progress_callback is None:
                return
            safe_retry = safe_retry_progress(event)
            progress_callback(
                {
                    "phase": "provider-retry-wait",
                    "round": effective_round,
                    "chunkId": chunk_id,
                    **safe_retry,
                    "attempts": safe_retry.get("attempt"),
                    "concurrency": rewrite_concurrency,
                    "configuredConcurrency": rewrite_concurrency,
                }
            )
        last_stream_emit_at = {"t": 0.0}
        last_stream_event_count = {"n": 0}

        def emit_stream_delta(event: dict[str, object]) -> None:
            if progress_callback is None:
                return
            try:
                upstream_event_count = max(0, int(event.get("eventCount", 0) or 0))
            except (TypeError, ValueError):
                upstream_event_count = 0
            event_count = max(last_stream_event_count["n"] + 1, upstream_event_count)
            done = event.get("done") is True
            try:
                final_text_chars = max(0, int(event.get("finalTextChars", 0) or 0)) if done else 0
            except (TypeError, ValueError):
                final_text_chars = 0
            now = time.monotonic()
            # Transport events are metadata-only.  Never forward delta/text,
            # provider payloads, endpoints, or previews into progress/UI.
            if (
                last_stream_event_count["n"] > 0
                and event_count - last_stream_event_count["n"] < 4
                and now - last_stream_emit_at["t"] < 0.18
                and not done
            ):
                return
            last_stream_emit_at["t"] = now
            last_stream_event_count["n"] = event_count
            progress_callback(
                {
                    "phase": "provider-stream",
                    "round": effective_round,
                    "chunkId": chunk_id,
                    "streamEventCount": event_count,
                    "streamDone": done,
                    "finalTextChars": final_text_chars,
                    "reasoningSuppressed": True,
                    "providerContentStored": False,
                    "concurrency": rewrite_concurrency,
                    "configuredConcurrency": rewrite_concurrency,
                }
            )

        try:
            output = llm_completion(
                prompt_input,
                model=model,
                api_key=api_key,
                base_url=base_url,
                api_type=api_type,
                temperature=temperature,
                timeout=request_timeout_seconds,
                max_retries=max_retries,
                retry_backoff_seconds=REWRITE_RETRY_BACKOFF_SECONDS,
                retry_callback=emit_retry_wait,
                stream=streaming,
                stream_callback=emit_stream_delta if streaming else None,
            )
        except LLMRequestError as exc:
            _register_provider_failure(effective_model_config, exc)
            raise
        _register_provider_success(effective_model_config)
        return output

    checkpoint_metadata = {
        "base_url": base_url,
        "model": model,
        "api_type": api_type,
        "temperature": temperature,
        "prompt_profile": prompt_profile,
        "prompt_sequence": prompt_sequence,
        "round_model_key": _round_model_key(prompt_profile, effective_round),
        "round_model_provider_id": str(effective_model_config.get("providerId", "")),
        "round_model_provider": str(effective_model_config.get("providerName", "")),
        "round_model_route_source": str(effective_model_config.get("routeSource", "default")),
        "request_timeout_seconds": request_timeout_seconds,
        "configured_request_timeout_seconds": configured_timeout_seconds,
        "max_retries": max_retries,
        "configured_max_retries": configured_max_retries,
        "retry_backoff_seconds": REWRITE_RETRY_BACKOFF_SECONDS,
        "rewrite_concurrency": rewrite_concurrency,
        "rate_limit_window_minutes": _coerce_rate_window_minutes(effective_model_config),
        "rate_limit_max_requests": _coerce_rate_max_requests(effective_model_config),
    }

    result = run_document_round(
        source_path,
        transform=transform,
        round_number=effective_round,
        prompt_profile=prompt_profile,
        prompt_sequence=prompt_sequence,
        progress_callback=progress_callback or emit_progress_event,
        checkpoint_metadata=checkpoint_metadata,
        cancel_check=cancel_check,
        max_concurrency=rewrite_concurrency,
        parent_artifact_snapshot=parent_artifact_snapshot,
    )
    return {
        "round": int(result["round"]),
        "outputPath": str(result["output_path"]),
        "manifestPath": str(result["manifest_path"]),
        "comparePath": str(result["compare_path"]),
        "qualityPath": str(result.get("quality_path", "")),
        "bodyMapPath": str(result.get("body_map_path", "")),
        "validationPath": str(result.get("validation_path", "")),
        "editContractPath": str(result.get("edit_contract_path", "")),
        "editContract": result.get("edit_contract", {}),
        "chunkLimit": int(result["chunk_limit"]),
        "inputSegmentCount": int(result["input_segment_count"]),
        "outputSegmentCount": int(result["output_segment_count"]),
        "paragraphCount": int(result["paragraph_count"]),
        "roundModel": {
            "round": effective_round,
            "providerId": str(effective_model_config.get("providerId", "")),
            "providerName": str(effective_model_config.get("providerName", "")),
            "baseUrl": base_url,
            "model": model,
            "apiType": api_type,
            "temperature": temperature,
            "rateLimitWindowMinutes": _coerce_rate_window_minutes(effective_model_config),
            "rateLimitMaxRequests": _coerce_rate_max_requests(effective_model_config),
            "routeSource": str(effective_model_config.get("routeSource", "default")),
        },
        "promptSequence": prompt_sequence,
        "docEntry": result["doc_entry"],
        "roundContext": result["round_context"],
        "qualitySummary": result.get("quality_summary", {}),
        "runAudit": result.get("run_audit", {}),
        "inputProvenance": result.get("input_provenance", {}),
    }


def test_model_connection(model_config: dict[str, Any]) -> dict[str, Any]:
    base_url = str(model_config.get("baseUrl", "")).strip()
    api_key = str(model_config.get("apiKey", "")).strip()
    model = str(model_config.get("model", "")).strip()
    api_type = str(model_config.get("apiType", "chat_completions")).strip()
    request_timeout_seconds = _coerce_int_config(
        model_config.get("requestTimeoutSeconds", DEFAULT_REQUEST_TIMEOUT_SECONDS),
        default=DEFAULT_REQUEST_TIMEOUT_SECONDS,
        minimum=30,
        maximum=3600,
    )
    max_retries = _coerce_int_config(
        model_config.get("maxRetries", DEFAULT_MAX_RETRIES),
        default=DEFAULT_MAX_RETRIES,
        minimum=0,
        maximum=10,
    )

    if not base_url or not api_key or not model:
        raise ValueError("Model configuration is incomplete.")

    result = test_llm_connection(
        model=model,
        api_key=api_key,
        base_url=base_url,
        api_type=api_type,
        timeout=request_timeout_seconds,
        max_retries=max_retries,
    )
    return {
        "ok": True,
        "message": "接口连通性测试成功。",
        **result,
    }


def list_available_models(model_config: dict[str, Any]) -> dict[str, Any]:
    base_url = str(model_config.get("baseUrl", "")).strip()
    api_key = str(model_config.get("apiKey", "")).strip()
    request_timeout_seconds = _coerce_int_config(
        model_config.get("requestTimeoutSeconds", DEFAULT_REQUEST_TIMEOUT_SECONDS),
        default=DEFAULT_REQUEST_TIMEOUT_SECONDS,
        minimum=30,
        maximum=3600,
    )
    max_retries = _coerce_int_config(
        model_config.get("maxRetries", DEFAULT_MAX_RETRIES),
        default=DEFAULT_MAX_RETRIES,
        minimum=0,
        maximum=10,
    )

    if not base_url:
        raise ValueError("baseUrl is required before loading models.")
    if not api_key:
        raise ValueError("apiKey is required before loading models.")

    result = list_llm_models(
        api_key=api_key,
        base_url=base_url,
        timeout=request_timeout_seconds,
        max_retries=max_retries,
    )
    return {
        "ok": True,
        "message": "Model catalog loaded successfully.",
        **result,
    }


def _round_snapshot_export_evidence(snapshot: dict[str, Any]) -> dict[str, Any]:
    """Return non-secret revision evidence for the exact exported generation."""

    return {
        "outputPath": str(snapshot.get("outputPath", "") or ""),
        "docId": str(snapshot.get("docId", "") or ""),
        "round": int(snapshot.get("round", 0) or 0),
        "roundSnapshotVersion": int(snapshot.get("version", 0) or 0),
        "compareRevision": str(snapshot.get("compareRevision", "") or ""),
        "reviewRevision": str(snapshot.get("reviewRevision", "") or ""),
        "contentRevision": str(snapshot.get("contentRevision", "") or ""),
        "artifactSnapshotDigest": str(snapshot.get("artifactSnapshotDigest", "") or ""),
        "effectiveTextSha256": str(snapshot.get("effectiveTextSha256", "") or ""),
        "capturedOutputSha256": str(snapshot.get("outputSha256", "") or ""),
        "capturedBodyMapSha256": str(snapshot.get("bodyMapSha256", "") or ""),
        "capturedManifestSha256": str(snapshot.get("manifestSha256", "") or ""),
    }


def _validate_export_snapshot_preconditions(
    snapshot: dict[str, Any] | None,
    *,
    expected_doc_id: str | None,
    expected_round: int | None,
    expected_compare_revision: str | None,
    expected_content_revision: str | None,
    expected_artifact_snapshot_digest: str | None,
) -> None:
    expectations = {
        "docId": expected_doc_id,
        "round": expected_round,
        "compareRevision": expected_compare_revision,
        "contentRevision": expected_content_revision,
        "artifactSnapshotDigest": expected_artifact_snapshot_digest,
    }
    if not any(value is not None for value in expectations.values()):
        return
    if snapshot is None:
        raise RoundArtifactSnapshotError(
            "round_snapshot_precondition_failed",
            "Export preconditions require a revision-consistent round snapshot.",
        )
    mismatches = [
        key
        for key, expected in expectations.items()
        if expected is not None and snapshot.get(key) != expected
    ]
    if mismatches:
        raise RoundArtifactSnapshotError(
            "round_snapshot_precondition_failed",
            "The selected round changed before export; refresh and confirm the current revision.",
            details={"mismatchFields": mismatches},
        )


def export_round_output(
    output_path: str,
    export_path: str,
    target_format: str,
    format_mode: str | None = None,
    *,
    expected_doc_id: str | None = None,
    expected_round: int | None = None,
    expected_compare_revision: str | None = None,
    expected_content_revision: str | None = None,
    expected_artifact_snapshot_digest: str | None = None,
) -> dict[str, Any]:
    normalized_target_format = str(target_format or "").strip().lower()
    if normalized_target_format not in {"txt", "docx"}:
        raise ValueError(f"Unsupported export format: {target_format}")
    normalized_output_path = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    if not normalized_output_path.exists():
        raise ValueError(f"Output file does not exist: {normalized_output_path}")
    if not normalized_output_path.is_file():
        raise ValueError(f"Output path is not a file: {normalized_output_path}")
    normalized_export_path = _resolve_api_path(export_path, allowed_roots=API_EXPORT_ALLOWED_ROOTS, label="Export path")
    normalized_export_path.parent.mkdir(parents=True, exist_ok=True)
    artifact_snapshot: dict[str, Any] | None = None
    certification: str | None = None
    try:
        # The lock is held only while immutable bytes/payloads are captured.
        # TXT writing and the substantially longer DOCX rebuild/audit happen
        # after it is released and use this in-memory generation exclusively.
        artifact_snapshot = read_round_artifact_snapshot(
            normalized_output_path,
            include_internal=True,
        )
    except RoundArtifactSnapshotError as exc:
        if normalized_target_format == "docx" and exc.code.startswith("round_snapshot_body_map_"):
            message = "DOCX 导出已阻断：冻结 body map 缺失、损坏或与当前轮次不一致。"
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="provenance",
                    label="Word 来源",
                    message=message,
                    issue_count=1,
                    samples=[{"code": exc.code, "message": str(exc)}],
                ),
            ) from exc
        legacy_plain_output = bool(
            exc.code == "round_snapshot_compare_missing"
            and not _has_fyadr_round_or_docx_provenance(normalized_output_path)
        )
        if not legacy_plain_output:
            raise
        certification = PLAIN_EXPORT_CERTIFICATION

    _validate_export_snapshot_preconditions(
        artifact_snapshot,
        expected_doc_id=expected_doc_id,
        expected_round=expected_round,
        expected_compare_revision=expected_compare_revision,
        expected_content_revision=expected_content_revision,
        expected_artifact_snapshot_digest=expected_artifact_snapshot_digest,
    )

    review_decisions = (
        copy.deepcopy((artifact_snapshot.get("review") or {}).get("decisions"))
        if isinstance(artifact_snapshot, dict)
        else {}
    )
    if not isinstance(review_decisions, dict):
        raise RoundArtifactSnapshotError(
            "round_snapshot_review_corrupt",
            "Captured review decisions are invalid for export.",
        )
    if normalized_target_format == "txt":
        result = _export_txt_round(
            normalized_output_path,
            normalized_export_path,
            review_decisions,
            artifact_snapshot=artifact_snapshot,
            certification=certification,
        )
    else:
        result = _export_docx_round(
            normalized_output_path,
            normalized_export_path,
            review_decisions,
            format_mode=format_mode,
            artifact_snapshot=artifact_snapshot,
            certification=certification,
        )
    result.setdefault("outputPath", str(normalized_output_path.resolve()))
    return result


def _export_txt_round(
    normalized_output_path: Path,
    normalized_export_path: Path,
    review_decisions: dict[str, Any],
    *,
    artifact_snapshot: dict[str, Any] | None = None,
    certification: str | None = None,
) -> dict[str, Any]:
    """Write the plain-text export, applying review decisions when present."""

    result = {
        "format": "txt",
        "path": str(normalized_export_path),
        "evidenceVersion": 1,
        "overallStatus": "passed",
        "sourceKind": "plain_text",
        "contentContractStatus": "not_applicable",
        "formatLockStatus": "not_applicable",
        "checksPerformed": ["text_export"],
    }
    if certification:
        result["certification"] = certification

    if artifact_snapshot is not None:
        internal = artifact_snapshot.get("_internal")
        blocks = internal.get("effectiveParagraphs") if isinstance(internal, dict) else None
        if not isinstance(blocks, list) or not blocks or not all(isinstance(item, str) for item in blocks):
            raise RoundArtifactSnapshotError(
                "round_snapshot_materialization_failed",
                "Captured effective paragraphs are invalid for TXT export.",
            )
        normalized_export_path.write_text("\n\n".join(blocks), encoding="utf-8")
        result.update(_round_snapshot_export_evidence(artifact_snapshot))
        return result

    if review_decisions:
        blocks = _build_reviewed_paragraphs_from_compare(normalized_output_path, review_decisions)
        normalized_export_path.write_text("\n\n".join(blocks), encoding="utf-8")
        return result

    compare_payload = _load_compare_payload_for_output(normalized_output_path)
    if compare_payload is not None:
        blocks, _paragraph_source = _read_output_paragraphs_for_export(normalized_output_path)
        normalized_export_path.write_text("\n\n".join(blocks), encoding="utf-8")
    else:
        shutil.copyfile(normalized_output_path, normalized_export_path)
    return result


def _resolve_docx_format_mode() -> str:
    """Return the only product-safe DOCX export mode.

    Older builds persisted ``school_rules`` and accepted an explicit export
    override that rewrote typography and paragraph layout. The product now
    removes that branch entirely: the legacy value is only migrated, while the
    source DOCX remains the sole formatting truth.
    """

    return "preserve_original"


def _sha256_export_artifact(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_export_evidence_json_atomically(path: Path, payload: dict[str, Any]) -> None:
    persisted_payload = payload
    if isinstance(payload.get("chunks"), list) and (
        "docId" in payload or "promptProfile" in payload or "validationEvents" in payload
    ):
        persisted_payload = _normalize_compare_failed_attempts(copy.deepcopy(payload))
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(
        f".{path.name}.{threading.get_ident()}.{time.monotonic_ns()}.tmp"
    )
    try:
        temporary_path.write_text(
            json.dumps(persisted_payload, ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        temporary_path.replace(path)
    finally:
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass


def _replace_staging_path_references(value: Any, *, staging_path: str, committed_path: str) -> Any:
    if isinstance(value, dict):
        return {
            str(key): _replace_staging_path_references(
                item,
                staging_path=staging_path,
                committed_path=committed_path,
            )
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [
            _replace_staging_path_references(
                item,
                staging_path=staging_path,
                committed_path=committed_path,
            )
            for item in value
        ]
    if isinstance(value, str):
        return value.replace(staging_path, committed_path)
    return value


def _prepare_docx_export_evidence_bundle(
    result: dict[str, Any],
    *,
    staging_export_path: Path,
    committed_export_path: Path,
    artifact_sha256: str,
    attempt_id: str,
) -> Path:
    """Finalize reports only after the immutable DOCX has been published."""

    staging_path = str(staging_export_path.resolve())
    committed_path = str(committed_export_path.resolve())
    if not committed_export_path.exists() or not committed_export_path.is_file():
        raise RuntimeError("Certified DOCX artifact is missing before evidence publication.")
    if _sha256_export_artifact(committed_export_path) != artifact_sha256:
        raise RuntimeError("Certified DOCX artifact hash changed before evidence publication.")

    checks_performed = {
        str(item).strip()
        for item in (result.get("checksPerformed") or [])
        if str(item).strip()
    }
    source_kind = str(result.get("sourceKind", "") or "")
    source_sha256 = str(result.get("sourceSha256", "") or "").strip().lower()
    provenance_source_path = str(result.get("provenanceSourcePath", "") or "").strip()
    required_checks = {"document_generation"}
    required_report_keys: set[str] = set()
    if source_kind == "original_docx":
        required_checks.update(
            {
                "pre_export_guard",
                "content_contract",
                "text_integrity",
                "protected_text_audit",
                "ooxml_integrity",
                "format_lock",
                "post_export_contract",
                "source_generation_anchor",
            }
        )
        required_report_keys.update(
            {
                "guardPath",
                "auditPath",
                "ooxmlAuditPath",
                "formatLockPath",
                "contentContractPath",
            }
        )
        if result.get("contentContractStatus") != "passed" or result.get("formatLockStatus") != "passed":
            raise RuntimeError("Original-DOCX evidence statuses are incomplete.")
        if result.get("semanticRangeTopologyValid") is not True:
            raise RuntimeError("Original-DOCX evidence has invalid semantic range topology.")
        if int(result.get("editableSemanticRangeAnchorUnitCount", 0) or 0) != 0:
            raise RuntimeError("Original-DOCX evidence contains an editable semantic range anchor.")
        if int(result.get("editableSemanticRangeCoveredUnitCount", 0) or 0) != 0:
            raise RuntimeError("Original-DOCX evidence contains editable text inside a semantic range.")
        if int(result.get("editableSemanticPointReferenceUnitCount", 0) or 0) != 0:
            raise RuntimeError("Original-DOCX evidence contains an editable semantic point reference.")
        if re.fullmatch(r"[0-9a-f]{64}", source_sha256) is None:
            raise RuntimeError("Original-DOCX evidence has no valid source SHA-256 binding.")
        if not provenance_source_path:
            raise RuntimeError("Original-DOCX evidence has no provenance source path.")
    elif source_kind == "generated_docx":
        if result.get("contentContractStatus") != "not_applicable" or result.get("formatLockStatus") != "not_applicable":
            raise RuntimeError("Generated-DOCX evidence applicability is inconsistent.")
    else:
        raise RuntimeError("DOCX export source kind is missing or unsupported.")
    missing_checks = sorted(required_checks - checks_performed)
    if missing_checks:
        raise RuntimeError("DOCX evidence omitted required checks: " + ", ".join(missing_checks))

    report_keys = (
        "guardPath",
        "auditPath",
        "ooxmlAuditPath",
        "formatLockPath",
        "contentContractPath",
    )
    published_reports: dict[str, str] = {}
    expected_report_outcome = {
        "guardPath": "ok",
        "auditPath": "ok",
        "ooxmlAuditPath": "ok",
        "formatLockPath": "ok",
        "contentContractPath": "ready",
    }
    for key in report_keys:
        raw_path = str(result.get(key, "") or "").strip()
        if not raw_path:
            if key in required_report_keys:
                raise RuntimeError(f"DOCX evidence report path is missing: {key}")
            continue
        report_path = normalize_path(Path(raw_path))
        if not report_path.exists() or not report_path.is_file():
            raise RuntimeError(f"DOCX evidence report does not exist: {key}")
        try:
            payload = json.loads(report_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuntimeError(f"DOCX evidence report is unreadable: {key}") from exc
        if not isinstance(payload, dict):
            raise RuntimeError(f"DOCX evidence report is not an object: {key}")
        outcome_field = expected_report_outcome.get(key)
        if outcome_field and not bool(payload.get(outcome_field)):
            raise RuntimeError(f"DOCX evidence report did not pass: {key}")
        if source_kind == "original_docx" and key in {
            "auditPath",
            "ooxmlAuditPath",
            "formatLockPath",
            "contentContractPath",
        }:
            if str(payload.get("sourceSha256", "") or "").strip().lower() != source_sha256:
                raise RuntimeError(f"DOCX evidence source SHA-256 mismatch: {key}")
            if str(payload.get("expectedSourceSha256", "") or "").strip().lower() != source_sha256:
                raise RuntimeError(f"DOCX evidence expected source SHA-256 mismatch: {key}")
            if payload.get("sourceGenerationStable") is not True:
                raise RuntimeError(f"DOCX evidence source generation is unstable: {key}")
            if str(payload.get("provenanceSourcePath", "") or "").strip() != provenance_source_path:
                raise RuntimeError(f"DOCX evidence provenance source mismatch: {key}")
        finalized_payload = _replace_staging_path_references(
            payload,
            staging_path=staging_path,
            committed_path=committed_path,
        )
        if not isinstance(finalized_payload, dict):
            continue
        finalized_payload["exportAttemptId"] = attempt_id
        finalized_payload["publishedArtifactPath"] = committed_path
        finalized_payload["artifactSha256"] = artifact_sha256
        if "exportPath" in finalized_payload:
            finalized_payload["exportPath"] = committed_path
        if "exportSha256" in finalized_payload:
            finalized_payload["exportSha256"] = artifact_sha256
        _write_export_evidence_json_atomically(report_path, finalized_payload)
        published_reports[key] = str(report_path.resolve())

    manifest_path = committed_export_path.with_suffix(".evidence.json")
    manifest = {
        "version": 1,
        "status": "passed",
        "createdAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "exportAttemptId": attempt_id,
        "artifactPath": committed_path,
        "artifactSha256": artifact_sha256,
        "format": "docx",
        "overallStatus": str(result.get("overallStatus", "passed") or "passed"),
        "sourceKind": source_kind,
        "contentContractStatus": str(result.get("contentContractStatus", "unknown") or "unknown"),
        "formatLockStatus": str(result.get("formatLockStatus", "unknown") or "unknown"),
        "checksPerformed": sorted(checks_performed),
        "reports": published_reports,
        "roundArtifactSnapshot": {
            "version": int(result.get("roundSnapshotVersion", 0) or 0),
            "outputPath": str(result.get("outputPath", "") or ""),
            "docId": str(result.get("docId", "") or ""),
            "round": int(result.get("round", 0) or 0),
            "compareRevision": str(result.get("compareRevision", "") or ""),
            "reviewRevision": str(result.get("reviewRevision", "") or ""),
            "contentRevision": str(result.get("contentRevision", "") or ""),
            "artifactSnapshotDigest": str(result.get("artifactSnapshotDigest", "") or ""),
            "effectiveTextSha256": str(result.get("effectiveTextSha256", "") or ""),
            "outputSha256": str(result.get("capturedOutputSha256", "") or ""),
            "bodyMapSha256": str(result.get("capturedBodyMapSha256", "") or ""),
            "manifestSha256": str(result.get("capturedManifestSha256", "") or ""),
        },
    }
    certification = str(result.get("certification", "") or "").strip()
    if certification:
        manifest["certification"] = certification
    if source_kind == "original_docx":
        manifest["sourceSha256"] = source_sha256
        manifest["provenanceSourcePath"] = provenance_source_path
        manifest["semanticBoundarySummary"] = {
            "rangeCount": int(result.get("semanticRangeCount", 0) or 0),
            "rangeTopologyValid": result.get("semanticRangeTopologyValid") is True,
            "rangeIssueCount": int(result.get("semanticRangeIssueCount", 0) or 0),
            "rangeAnchorUnitCount": int(result.get("semanticRangeAnchorUnitCount", 0) or 0),
            "protectedRangeAnchorUnitCount": int(
                result.get("protectedSemanticRangeAnchorUnitCount", 0) or 0
            ),
            "editableRangeAnchorUnitCount": int(
                result.get("editableSemanticRangeAnchorUnitCount", 0) or 0
            ),
            "rangeCoveredUnitCount": int(
                result.get("semanticRangeCoveredUnitCount", 0) or 0
            ),
            "protectedRangeCoveredUnitCount": int(
                result.get("protectedSemanticRangeCoveredUnitCount", 0) or 0
            ),
            "editableRangeCoveredUnitCount": int(
                result.get("editableSemanticRangeCoveredUnitCount", 0) or 0
            ),
            "pointReferenceUnitCount": int(
                result.get("semanticPointReferenceUnitCount", 0) or 0
            ),
            "protectedPointReferenceUnitCount": int(
                result.get("protectedSemanticPointReferenceUnitCount", 0) or 0
            ),
            "editablePointReferenceUnitCount": int(
                result.get("editableSemanticPointReferenceUnitCount", 0) or 0
            ),
        }
    _write_export_evidence_json_atomically(manifest_path, manifest)
    return manifest_path


def _cleanup_unpublished_docx_evidence_bundle(
    result: dict[str, Any],
    *,
    published_export_path: Path,
) -> None:
    candidate_paths = {
        published_export_path.with_suffix(".evidence.json"),
    }
    for key in (
        "guardPath",
        "auditPath",
        "ooxmlAuditPath",
        "formatLockPath",
        "contentContractPath",
    ):
        value = str(result.get(key, "") or "").strip()
        if value:
            candidate_paths.add(normalize_path(Path(value)))
    for path in candidate_paths:
        if not is_path_under(path, ROOT_DIR / "finish"):
            continue
        try:
            path.unlink(missing_ok=True)
        except OSError:
            continue


def _resolve_docx_origin_bundle_for_export(
    output_path: Path,
    body_map: Any,
) -> tuple[Path, Path] | None:
    origin_bundle = _find_origin_docx_for_output(output_path)
    if origin_bundle is not None:
        return origin_bundle

    declared_origin = _find_declared_origin_docx_for_output(output_path)
    body_map_path = _find_body_map_path_for_output(output_path)
    body_map_artifact_exists = bool(body_map_path is not None and body_map_path.exists())
    if body_map is None:
        if declared_origin is not None or body_map_artifact_exists:
            message = "DOCX 导出已阻断：输出带有 Word 来源痕迹，但冻结 body map 缺失或损坏，不能降级为新建 Word。"
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="provenance",
                    label="Word 来源",
                    message=message,
                    issue_count=1,
                    samples=[
                        {
                            "code": "docx_body_map_missing_or_invalid",
                            "message": "无法证明当前结果仍只对应原 Word 的冻结正文范围。",
                        }
                    ],
                ),
            )
        return None

    source_docx_path = normalize_path(Path(str(getattr(body_map, "source_path", "") or "")))
    source_in_workspace = any(
        is_path_under(source_docx_path, root)
        for root in (ROOT_DIR / "origin", ROOT_DIR / "finish")
    )
    if source_docx_path.suffix.lower() != ".docx" or not source_in_workspace:
        message = "DOCX 导出已阻断：body map 声明的源 Word 路径无效或超出原文目录。"
        raise ExportRoundError(
            message,
            _build_export_failure(
                stage="provenance",
                label="Word 来源",
                message=message,
                issue_count=1,
                samples=[{"code": "docx_body_map_source_invalid", "message": message}],
            ),
        )
    if declared_origin is not None and normalize_path(declared_origin) != source_docx_path:
        message = "DOCX 导出已阻断：body map 与历史记录声明了不同的源 Word。"
        raise ExportRoundError(
            message,
            _build_export_failure(
                stage="provenance",
                label="Word 来源",
                message=message,
                issue_count=1,
                samples=[{"code": "docx_provenance_mismatch", "message": message}],
            ),
        )
    if not source_docx_path.exists() or not source_docx_path.is_file():
        message = "DOCX 导出已阻断：body map 对应的源 Word 已不存在。"
        raise ExportRoundError(
            message,
            _build_export_failure(
                stage="provenance",
                label="Word 来源",
                message=message,
                issue_count=1,
                samples=[{"code": "docx_source_missing", "message": message}],
            ),
        )

    snapshot_path = normalize_path(Path(str(getattr(body_map, "snapshot_path", "") or "")))
    canonical_snapshot_path = get_docx_snapshot_path(source_docx_path)
    if snapshot_path != canonical_snapshot_path or not is_path_under(snapshot_path, ROOT_DIR / "finish"):
        message = "DOCX 导出已阻断：body map 的 Word 快照路径不是当前源文档的规范冻结快照。"
        raise ExportRoundError(
            message,
            _build_export_failure(
                stage="provenance",
                label="Word 来源",
                message=message,
                issue_count=1,
                samples=[{"code": "docx_snapshot_path_invalid", "message": message}],
            ),
        )

    ensure_docx_processing_assets(source_docx_path, snapshot_path=snapshot_path)
    validation_report = validate_docx_body_map(
        body_map,
        source_path=source_docx_path,
        snapshot_path=snapshot_path,
    )
    if not bool(validation_report.get("ok")):
        message = "DOCX 导出已阻断：孤立轮次的 body map 未通过源 Word 权威校验，不能降级为普通 Word。"
        raise ExportRoundError(
            message,
            _build_export_failure(
                stage="provenance",
                label="Word 来源",
                message=message,
                report=validation_report,
                issue_count=len(validation_report.get("blockingIssues", []) or []),
                warning_count=len(validation_report.get("warnings", []) or []),
                sample_keys=("blockingIssues", "warnings"),
            ),
        )
    return source_docx_path, snapshot_path


_DOCX_SOURCE_ANCHOR_LOCK = threading.Lock()


def _windows_file_change_time_ns(descriptor: int) -> int:
    """Return the NTFS/ReFS change timestamp for one open Windows file.

    Python 3.11 exposes Windows ``st_ctime`` as creation time, so it cannot
    detect an A -> B -> A in-place replacement whose bytes and mtime are
    restored.  ``FILE_BASIC_INFO.ChangeTime`` is the kernel-maintained
    metadata-change clock and closes that generation-seal gap.
    """

    if os.name != "nt":
        return 0
    import ctypes
    import msvcrt  # type: ignore[import-not-found]

    class FileBasicInfo(ctypes.Structure):
        _fields_ = [
            ("creation_time", ctypes.c_longlong),
            ("last_access_time", ctypes.c_longlong),
            ("last_write_time", ctypes.c_longlong),
            ("change_time", ctypes.c_longlong),
            ("file_attributes", ctypes.c_ulong),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    get_file_information = kernel32.GetFileInformationByHandleEx
    get_file_information.argtypes = [ctypes.c_void_p, ctypes.c_int, ctypes.c_void_p, ctypes.c_ulong]
    get_file_information.restype = ctypes.c_int
    info = FileBasicInfo()
    raw_handle = msvcrt.get_osfhandle(descriptor)
    if raw_handle == -1:
        raise OSError("Invalid Windows file handle while sealing a DOCX source generation.")
    succeeded = get_file_information(
        ctypes.c_void_p(raw_handle),
        0,  # FILE_INFO_BY_HANDLE_CLASS.FileBasicInfo
        ctypes.byref(info),
        ctypes.sizeof(info),
    )
    if not succeeded:
        raise ctypes.WinError(ctypes.get_last_error())
    # Windows stores this value in 100 ns ticks since 1601.  Its epoch is not
    # relevant; only exact equality within one export generation is compared.
    return int(info.change_time) * 100


def _docx_source_generation_stat_from_result(
    stat: os.stat_result,
    *,
    descriptor: int | None = None,
) -> dict[str, int]:
    result = {
        "device": int(stat.st_dev),
        "inode": int(stat.st_ino),
        "size": int(stat.st_size),
        "mtimeNs": int(stat.st_mtime_ns),
        "ctimeNs": int(stat.st_ctime_ns),
    }
    if os.name == "nt":
        if descriptor is None:
            raise OSError("Windows DOCX generation stats require an open file descriptor.")
        result["changeTimeNs"] = _windows_file_change_time_ns(descriptor)
    return result


def _docx_source_generation_stat(path: Path) -> dict[str, int]:
    with path.open("rb") as handle:
        return _docx_source_generation_stat_from_result(
            os.fstat(handle.fileno()),
            descriptor=handle.fileno(),
        )


def _raise_docx_source_generation_error(stage: str, message: str, *, codes: list[str]) -> None:
    raise ExportRoundError(
        message,
        _build_export_failure(
            stage="source-anchor",
            label="Word 源文件代际",
            message=message,
            issue_count=max(1, len(codes)),
            samples=[{"code": code, "message": message, "stage": stage} for code in codes[:8]],
        ),
    )


def _assert_docx_source_generation_seal(seal: dict[str, Any], *, stage: str) -> None:
    source_path = normalize_path(Path(str(seal.get("sourcePath", "") or "")))
    expected_sha256 = str(seal.get("sourceSha256", "") or "")
    expected_stat = seal.get("sourceStat")
    mismatch_codes: list[str] = []
    try:
        stat_before = _docx_source_generation_stat(source_path)
        current_sha256 = _sha256_file_bytes(source_path)
        stat_after = _docx_source_generation_stat(source_path)
    except OSError:
        stat_before = {}
        stat_after = {}
        current_sha256 = ""
        mismatch_codes.append("source_path_unreadable")
    if not isinstance(expected_stat, dict) or stat_before != expected_stat or stat_after != expected_stat:
        mismatch_codes.append("source_generation_stat_mismatch")
    if not expected_sha256 or current_sha256 != expected_sha256:
        mismatch_codes.append("source_generation_sha256_mismatch")
    if mismatch_codes:
        _raise_docx_source_generation_error(
            stage,
            "DOCX 导出已阻断：源 Word 在本次导出期间发生过代际替换。",
            codes=sorted(set(mismatch_codes)),
        )


def _capture_docx_export_source_anchor(
    source_path: Path,
    *,
    snapshot_path: Path,
    body_map: Any,
) -> dict[str, Any]:
    normalized_source = source_path.resolve()
    cached_snapshot = _load_docx_snapshot(snapshot_path)
    body_map_sha256 = str(getattr(body_map, "source_sha256", "") or "").strip().lower()
    snapshot_sha256 = str(getattr(cached_snapshot, "source_sha256", "") or "").strip().lower()
    if not body_map_sha256 or not snapshot_sha256 or body_map_sha256 != snapshot_sha256:
        _raise_docx_source_generation_error(
            "capture",
            "DOCX 导出已阻断：body map 与快照没有同一份源 Word 哈希。",
            codes=["source_anchor_claim_mismatch"],
        )
    expected_sha256 = body_map_sha256

    anchor_root = ROOT_DIR / "finish" / "intermediate" / "docx_source_anchors"
    anchor_root.mkdir(parents=True, exist_ok=True)
    temporary_path = anchor_root / (
        f".{expected_sha256}.{threading.get_ident()}.{time.monotonic_ns()}.tmp.docx"
    )
    anchor_path = anchor_root / f"{expected_sha256}.docx"
    source_stat_before: dict[str, int] = {}
    source_stat_after: dict[str, int] = {}
    source_handle_stat_before: dict[str, int] = {}
    source_handle_stat_after: dict[str, int] = {}
    captured_sha256 = ""
    try:
        with normalized_source.open("rb") as source_handle, temporary_path.open("xb") as anchor_handle:
            source_stat_before = _docx_source_generation_stat(normalized_source)
            source_handle_stat_before = _docx_source_generation_stat_from_result(
                os.fstat(source_handle.fileno()),
                descriptor=source_handle.fileno(),
            )
            digest = hashlib.sha256()
            for block in iter(lambda: source_handle.read(1024 * 1024), b""):
                digest.update(block)
                anchor_handle.write(block)
            anchor_handle.flush()
            captured_sha256 = digest.hexdigest()
            source_handle_stat_after = _docx_source_generation_stat_from_result(
                os.fstat(source_handle.fileno()),
                descriptor=source_handle.fileno(),
            )
            source_stat_after = _docx_source_generation_stat(normalized_source)
        if (
            source_stat_before != source_stat_after
            or source_handle_stat_before != source_handle_stat_after
            or source_handle_stat_before != source_stat_before
        ):
            _raise_docx_source_generation_error(
                "capture",
                "DOCX 导出已阻断：捕获源 Word 时文件代际发生变化。",
                codes=["source_generation_changed_during_capture"],
            )
        if captured_sha256 != expected_sha256:
            _raise_docx_source_generation_error(
                "capture",
                "DOCX 导出已阻断：捕获到的源 Word 与冻结快照哈希不一致。",
                codes=["source_anchor_sha256_mismatch"],
            )
        with _DOCX_SOURCE_ANCHOR_LOCK:
            if anchor_path.exists():
                if _sha256_file_bytes(anchor_path) != expected_sha256:
                    _raise_docx_source_generation_error(
                        "capture",
                        "DOCX 导出已阻断：内部源文件 anchor 已损坏。",
                        codes=["source_anchor_cache_corrupt"],
                    )
            else:
                temporary_path.replace(anchor_path)
                anchor_path.chmod(0o444)
            anchor_extracted_path = anchor_root / f"{expected_sha256}.extracted.txt"
            anchor_snapshot_path = anchor_root / f"{expected_sha256}.snapshot.json"
            anchor_scope_path = anchor_root / f"{expected_sha256}.scope.json"
            _, _, anchor_snapshot = ensure_docx_processing_assets(
                anchor_path,
                extracted_path=anchor_extracted_path,
                snapshot_path=anchor_snapshot_path,
                scope_diagnostics_path=anchor_scope_path,
            )
        if anchor_snapshot.source_sha256 != expected_sha256:
            _raise_docx_source_generation_error(
                "capture",
                "DOCX 导出已阻断：内部源文件 anchor 快照哈希不一致。",
                codes=["source_anchor_snapshot_mismatch"],
            )
        anchored_cached_snapshot = copy.deepcopy(cached_snapshot)
        anchored_cached_snapshot.source_path = str(anchor_path.resolve())
        anchored_cached_snapshot.source_size = int(anchor_snapshot.source_size)
        anchored_cached_snapshot.source_mtime_ns = int(anchor_snapshot.source_mtime_ns)
        anchored_cached_snapshot.source_sha256 = expected_sha256
        cached_snapshot_derivation, _ = verify_docx_snapshot_derivation(
            anchored_cached_snapshot,
            anchor_path,
        )
        seal = {
            "sourcePath": str(normalized_source),
            "sourceSha256": expected_sha256,
            "sourceStat": dict(source_stat_before),
        }
        _assert_docx_source_generation_seal(seal, stage="after_capture")
        return {
            "anchorPath": anchor_path.resolve(),
            "snapshotPath": anchor_snapshot_path.resolve(),
            "extractedPath": anchor_extracted_path.resolve(),
            "scopePath": anchor_scope_path.resolve(),
            "snapshot": anchor_snapshot,
            "cachedSnapshotDerivation": cached_snapshot_derivation,
            "sourceSha256": expected_sha256,
            "provenanceSourcePath": normalized_source,
            "seal": seal,
        }
    finally:
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass


def _retag_docx_body_map_for_source_anchor(body_map: Any, anchor: dict[str, Any]) -> Any:
    if body_map is None:
        return None
    anchored_body_map = copy.deepcopy(body_map)
    anchor_snapshot = anchor.get("snapshot")
    anchored_body_map.source_path = str(anchor["anchorPath"])
    anchored_body_map.source_size = int(getattr(anchor_snapshot, "source_size", 0) or 0)
    anchored_body_map.source_mtime_ns = int(getattr(anchor_snapshot, "source_mtime_ns", 0) or 0)
    anchored_body_map.source_sha256 = str(anchor["sourceSha256"])
    anchored_body_map.snapshot_path = str(anchor["snapshotPath"])
    anchored_body_map.snapshot_version = int(getattr(anchor_snapshot, "version", 0) or 0)
    return anchored_body_map


def _export_docx_round(
    normalized_output_path: Path,
    normalized_export_path: Path,
    review_decisions: dict[str, Any],
    *,
    format_mode: str | None = None,
    artifact_snapshot: dict[str, Any] | None = None,
    certification: str | None = None,
) -> dict[str, Any]:
    """Build and audit a DOCX in staging, then atomically publish it.

    The final path is never used as the working file.  A failed post-build
    audit therefore cannot leave an uncertified document behind or overwrite a
    previously successful export.  Diagnostic JSON reports intentionally keep
    their stable final-path names so a blocked export still has actionable
    evidence after the staging DOCX is removed.
    """

    attempt_id = f"{int(time.time() * 1000)}-{threading.get_ident()}-{time.monotonic_ns()}"
    artifact_stem = truncate_utf8_filename_component(
        normalized_export_path.stem,
        max_bytes=96,
        fallback="paper",
    )
    staging_export_path = normalized_export_path.with_name(
        f".{artifact_stem}.{attempt_id}.tmp"
        f"{normalized_export_path.suffix}"
    )
    published_export_path = normalized_export_path.with_name(
        f"{artifact_stem}.{attempt_id}{normalized_export_path.suffix}"
    )
    try:
        result = _build_and_audit_docx_round(
            normalized_output_path,
            staging_export_path,
            review_decisions,
            committed_export_path=published_export_path,
            evidence_export_path=published_export_path,
            format_mode=format_mode,
            artifact_snapshot=artifact_snapshot,
        )
        if artifact_snapshot is not None:
            result.update(_round_snapshot_export_evidence(artifact_snapshot))
        if certification:
            result["certification"] = certification
        source_generation_seal = result.pop("_sourceGenerationSeal", None)
        if str(result.get("sourceKind", "") or "") == "original_docx":
            if not isinstance(source_generation_seal, dict):
                _raise_docx_source_generation_error(
                    "before_publish",
                    "DOCX 导出已阻断：发布前缺少源 Word 代际封条。",
                    codes=["source_generation_seal_missing"],
                )
            _assert_docx_source_generation_seal(
                source_generation_seal,
                stage="before_publish",
            )
        artifact_sha256 = _sha256_export_artifact(staging_export_path)
        staging_export_path.replace(published_export_path)
        try:
            evidence_manifest_path = _prepare_docx_export_evidence_bundle(
                result,
                staging_export_path=staging_export_path,
                committed_export_path=published_export_path,
                artifact_sha256=artifact_sha256,
                attempt_id=attempt_id,
            )
        except Exception:
            published_export_path.unlink(missing_ok=True)
            _cleanup_unpublished_docx_evidence_bundle(
                result,
                published_export_path=published_export_path,
            )
            raise

        # Keep the requested fixed name only as a backwards-compatible latest
        # alias.  It is deliberately not the certified artifact returned to
        # callers; evidence and downloads bind to the immutable attempt path.
        latest_alias_temp = normalized_export_path.with_name(
            f".{artifact_stem}.{attempt_id}.latest.tmp"
        )
        latest_alias_warning = ""
        try:
            shutil.copyfile(published_export_path, latest_alias_temp)
            latest_alias_temp.replace(normalized_export_path)
        except OSError as exc:
            latest_alias_warning = f"latest_alias_update_failed:{type(exc).__name__}"
        finally:
            try:
                latest_alias_temp.unlink(missing_ok=True)
            except OSError:
                pass

        result["path"] = str(published_export_path)
        result["latestAliasPath"] = str(normalized_export_path)
        result["latestAliasWarning"] = latest_alias_warning
        result["exportAttemptId"] = attempt_id
        result["artifactSha256"] = artifact_sha256
        result["evidenceManifestPath"] = str(evidence_manifest_path)
        return result
    finally:
        try:
            staging_export_path.unlink(missing_ok=True)
        except OSError:
            pass


def _build_and_audit_docx_round(
    normalized_output_path: Path,
    normalized_export_path: Path,
    review_decisions: dict[str, Any],
    *,
    committed_export_path: Path,
    evidence_export_path: Path,
    format_mode: str | None = None,
    artifact_snapshot: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Rebuild a DOCX export, guarding layout, format rules, and protected content."""

    layout_mode = "plain_text_docx"
    paragraph_source = "text"
    validation_path = _find_validation_path_for_output(normalized_output_path)
    captured_effective_paragraphs: list[str] | None = None
    if artifact_snapshot is not None:
        if str(artifact_snapshot.get("outputPath", "") or "") != str(normalized_output_path.resolve()):
            raise RoundArtifactSnapshotError(
                "round_snapshot_output_mismatch",
                "Captured export snapshot belongs to another round output.",
            )
        internal = artifact_snapshot.get("_internal")
        if not isinstance(internal, dict):
            raise RoundArtifactSnapshotError(
                "round_snapshot_capture_failed",
                "Internal round snapshot payload is unavailable for DOCX export.",
            )
        compare_payload = internal.get("comparePayload")
        raw_effective_paragraphs = internal.get("effectiveParagraphs")
        raw_body_map_payload = internal.get("bodyMapPayload")
        snapshot_review = artifact_snapshot.get("review")
        snapshot_decisions = snapshot_review.get("decisions") if isinstance(snapshot_review, dict) else None
        if (
            not isinstance(compare_payload, dict)
            or not isinstance(raw_effective_paragraphs, list)
            or not raw_effective_paragraphs
            or not all(isinstance(item, str) for item in raw_effective_paragraphs)
            or not isinstance(snapshot_decisions, dict)
        ):
            raise RoundArtifactSnapshotError(
                "round_snapshot_materialization_failed",
                "Captured round generation is incomplete for DOCX export.",
            )
        captured_effective_paragraphs = list(raw_effective_paragraphs)
        review_decisions = copy.deepcopy(snapshot_decisions)
        body_map = docx_body_map_from_payload(raw_body_map_payload) if raw_body_map_payload is not None else None
        if raw_body_map_payload is not None and body_map is None:
            raise RoundArtifactSnapshotError(
                "round_snapshot_body_map_corrupt",
                "Captured DOCX body map could not be reconstructed for export.",
            )
    else:
        body_map = _load_body_map_for_output(normalized_output_path)
        compare_payload = _load_compare_payload_for_output(normalized_output_path)
    origin_docx_bundle = _resolve_docx_origin_bundle_for_export(normalized_output_path, body_map)
    source_anchor: dict[str, Any] | None = None
    provenance_source_path: Path | None = None
    if origin_docx_bundle is not None:
        provenance_source_path = origin_docx_bundle[0]
        source_anchor = _capture_docx_export_source_anchor(
            origin_docx_bundle[0],
            snapshot_path=origin_docx_bundle[1],
            body_map=body_map,
        )
        cached_snapshot_derivation = source_anchor.get("cachedSnapshotDerivation")
        if not isinstance(cached_snapshot_derivation, dict) or not bool(
            cached_snapshot_derivation.get("ok")
        ):
            guard_report_path = evidence_export_path.with_suffix(".guard.json")
            guard_report = {
                "ok": False,
                "mode": "docx-export",
                "sourcePath": str(origin_docx_bundle[0].resolve()),
                "snapshotPath": str(origin_docx_bundle[1].resolve()),
                "blockingIssueCount": 1,
                "warningCount": 0,
                "blockingIssues": [
                    {
                        "code": "body_map_snapshot_authority_mismatch",
                        "severity": "error",
                        "message": "冻结的 Word 正文范围与内容寻址源文件重新派生的权威范围不一致。",
                        "cachedDigest": str(
                            (cached_snapshot_derivation or {}).get("cachedDigest", "")
                        ),
                        "authoritativeDigest": str(
                            (cached_snapshot_derivation or {}).get("authoritativeDigest", "")
                        ),
                        "mismatchUnitIndexes": list(
                            (cached_snapshot_derivation or {}).get("mismatchUnitIndexes", [])
                        )[:40],
                    }
                ],
                "warnings": [],
                "reportPath": str(guard_report_path.resolve()),
            }
            guard_report_path.parent.mkdir(parents=True, exist_ok=True)
            guard_report_path.write_text(
                json.dumps(guard_report, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            message = "DOCX 导出已阻断：冻结的 Word 正文范围不是从本次捕获的源文件代际派生。"
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="guard",
                    label="导出前保护",
                    message=message,
                    report=guard_report,
                    report_path=guard_report_path,
                    issue_count=1,
                    sample_keys=("blockingIssues",),
                ),
            )
        origin_docx_bundle = (
            Path(source_anchor["anchorPath"]),
            Path(source_anchor["snapshotPath"]),
        )
        body_map = _retag_docx_body_map_for_source_anchor(body_map, source_anchor)
    guard_report: dict[str, Any] | None = None
    content_contract_report: dict[str, Any] | None = None
    content_contract_path = get_export_edit_contract_path(evidence_export_path)
    contract_candidate_texts: list[str] | None = None
    contract_body_map: Any = None
    expected_text_targets: list[tuple[dict[str, Any], str]] = []
    text_integrity_performed = False
    post_export_contract_performed = False
    # ``format_mode`` remains in the Python signature for compatibility with
    # older callers, but it can no longer select a formatting-mutating path.
    # This is intentionally fail-closed even for an explicit ``school_rules``
    # value so Web, CLI and tests share the same content-only guarantee.
    _requested_format_mode = str(format_mode or "").strip().lower()
    resolved_format_mode = "preserve_original" if origin_docx_bundle is not None else "generated_default"
    preserve_original_format = True
    if origin_docx_bundle is not None:
        source_docx_path, snapshot_path = origin_docx_bundle
        ensure_docx_processing_assets(
            source_docx_path,
            snapshot_path=snapshot_path,
        )
        if body_map is not None:
            if captured_effective_paragraphs is not None:
                if len(captured_effective_paragraphs) != len(body_map.units):
                    raise RoundArtifactSnapshotError(
                        "round_snapshot_body_map_mismatch",
                        "Captured effective paragraphs do not match the captured DOCX body map.",
                        details={
                            "paragraphCount": len(captured_effective_paragraphs),
                            "bodyMapUnitCount": len(body_map.units),
                        },
                    )
                body_map = update_docx_body_map_texts(
                    body_map,
                    captured_effective_paragraphs,
                    round_number=body_map.round_number,
                )
                paragraph_source = "artifact_snapshot_review_materialized"
            elif review_decisions:
                reviewed_paragraphs = _build_reviewed_paragraphs_from_compare(normalized_output_path, review_decisions)
                reviewed_paragraphs = _ensure_reviewed_paragraph_count_for_body_map(
                    normalized_output_path,
                    compare_payload or {},
                    reviewed_paragraphs,
                    len(body_map.units),
                    review_decisions,
                )
                body_map = update_docx_body_map_texts(body_map, reviewed_paragraphs, round_number=body_map.round_number)
                paragraph_source = "body_map_review_decisions"
            elif compare_payload is not None:
                compare_paragraphs = _build_paragraphs_from_compare_payload(normalized_output_path, compare_payload)
                if compare_paragraphs is not None:
                    compare_paragraphs = _ensure_reviewed_paragraph_count_for_body_map(
                        normalized_output_path,
                        compare_payload,
                        compare_paragraphs,
                        len(body_map.units),
                        {},
                    )
                    body_map = update_docx_body_map_texts(body_map, compare_paragraphs, round_number=body_map.round_number)
                    paragraph_source = "body_map_compare"
            guard_report = run_docx_pre_export_guard(
                body_map.current_texts(),
                source_path=source_docx_path,
                snapshot_path=snapshot_path,
                export_path=committed_export_path,
                output_path=normalized_output_path,
                body_map=body_map,
                compare_payload=compare_payload,
                mode="docx-export",
                paragraph_source=paragraph_source,
                report_path=evidence_export_path.with_suffix(".guard.json"),
            )
            if not bool(guard_report.get("ok")):
                message = summarize_docx_export_guard_failure(guard_report, label="导出")
                raise ExportRoundError(
                    message,
                    _build_export_failure(
                        stage="guard",
                        label="导出前保护",
                        message=message,
                        report=guard_report,
                        issue_count=int(guard_report.get("blockingIssueCount", 0) or 0),
                        warning_count=int(guard_report.get("warningCount", 0) or 0),
                        sample_keys=("blockingIssues", "warnings"),
                    ),
                )
            contract_candidate_texts = body_map.current_texts()
            contract_body_map = body_map
            content_contract_report = build_document_edit_contract(
                source_docx_path,
                snapshot_path=snapshot_path,
                body_map=body_map,
                candidate_texts=contract_candidate_texts,
                stage="pre_export",
                expected_source_sha256=str((source_anchor or {}).get("sourceSha256", "") or "") or None,
                provenance_source_path=provenance_source_path,
                report_path=content_contract_path,
            )
            try:
                assert_document_edit_contract_ready(content_contract_report, label="DOCX 导出")
            except ValueError as exc:
                message = str(exc)
                raise ExportRoundError(
                    message,
                    _build_export_failure(
                        stage="content-contract",
                        label="正文与格式契约",
                        message=message,
                        report=content_contract_report,
                        sample_keys=("issues",),
                    ),
                ) from exc
            expected_text_targets = _build_docx_text_targets_from_body_map(body_map)
            if source_anchor is not None:
                _assert_docx_source_generation_seal(source_anchor["seal"], stage="before_rebuild")
            try:
                rebuild_docx_from_body_map_units(
                    body_map.units,
                    source_path=source_docx_path,
                    export_path=normalized_export_path,
                    preserve_format=preserve_original_format,
                )
            except ValueError as exc:
                raise _export_text_integrity_error(str(exc)) from exc
            layout_mode = "body-map-roundtrip"
            if captured_effective_paragraphs is None and not review_decisions:
                paragraph_source = "body_map"
        else:
            if captured_effective_paragraphs is not None:
                blocks = list(captured_effective_paragraphs)
                paragraph_source = "artifact_snapshot_review_materialized"
            elif review_decisions:
                blocks = _build_reviewed_paragraphs_from_compare(normalized_output_path, review_decisions)
                paragraph_source = "compare_review_decisions"
            else:
                blocks, paragraph_source = _read_output_paragraphs_for_export(normalized_output_path)
            snapshot = _load_docx_snapshot(snapshot_path)
            expected_paragraph_count = (
                snapshot.editable_unit_count
                if snapshot is not None
                else None
            )
            if expected_paragraph_count is not None and len(blocks) != expected_paragraph_count:
                message = (
                    "DOCX 导出已拦截：当前轮次正文段落数与原始 Word 快照不一致。"
                    f" 预期 {expected_paragraph_count} 段，实际 {len(blocks)} 段。"
                    " 请重新执行当前轮次，或回滚后重跑，避免导出成错位排版。"
                )
                raise ExportRoundError(
                    message,
                    _build_export_failure(
                        stage="paragraph-count",
                        label="正文段落",
                        message=message,
                        issue_count=1,
                        samples=[
                            {
                                "code": "paragraph_count_mismatch",
                                "message": "正文段落数与原始 Word 快照不一致",
                                "sample": f"expected={expected_paragraph_count}; actual={len(blocks)}",
                            }
                        ],
                    ),
                )
            guard_report = run_docx_pre_export_guard(
                blocks,
                source_path=source_docx_path,
                snapshot_path=snapshot_path,
                export_path=committed_export_path,
                output_path=normalized_output_path,
                compare_payload=compare_payload,
                mode="docx-export",
                paragraph_source=paragraph_source,
                report_path=evidence_export_path.with_suffix(".guard.json"),
            )
            if not bool(guard_report.get("ok")):
                message = summarize_docx_export_guard_failure(guard_report, label="导出")
                raise ExportRoundError(
                    message,
                    _build_export_failure(
                        stage="guard",
                        label="导出前保护",
                        message=message,
                        report=guard_report,
                        issue_count=int(guard_report.get("blockingIssueCount", 0) or 0),
                        warning_count=int(guard_report.get("warningCount", 0) or 0),
                        sample_keys=("blockingIssues", "warnings"),
                    ),
                )
            contract_candidate_texts = list(blocks)
            content_contract_report = build_document_edit_contract(
                source_docx_path,
                snapshot_path=snapshot_path,
                candidate_texts=contract_candidate_texts,
                stage="pre_export",
                expected_source_sha256=str((source_anchor or {}).get("sourceSha256", "") or "") or None,
                provenance_source_path=provenance_source_path,
                report_path=content_contract_path,
            )
            try:
                assert_document_edit_contract_ready(content_contract_report, label="DOCX 导出")
            except ValueError as exc:
                message = str(exc)
                raise ExportRoundError(
                    message,
                    _build_export_failure(
                        stage="content-contract",
                        label="正文与格式契约",
                        message=message,
                        report=content_contract_report,
                        sample_keys=("issues",),
                    ),
                ) from exc
            expected_text_targets = _build_docx_text_targets_from_snapshot(snapshot_path, blocks)
            if source_anchor is not None:
                _assert_docx_source_generation_seal(source_anchor["seal"], stage="before_rebuild")
            try:
                rebuild_docx_from_snapshot(
                    blocks,
                    source_path=source_docx_path,
                    snapshot_path=snapshot_path,
                    export_path=normalized_export_path,
                    preserve_format=preserve_original_format,
                )
            except ValueError as exc:
                raise _export_text_integrity_error(str(exc)) from exc
            layout_mode = "snapshot-compare-reflow" if paragraph_source.startswith("compare") else "snapshot-roundtrip"
    else:
        if captured_effective_paragraphs is not None:
            blocks = list(captured_effective_paragraphs)
            paragraph_source = "artifact_snapshot_review_materialized"
        elif review_decisions:
            blocks = _build_reviewed_paragraphs_from_compare(normalized_output_path, review_decisions)
            paragraph_source = "compare_review_decisions"
        else:
            blocks, paragraph_source = _read_output_paragraphs_for_export(normalized_output_path)
        write_docx_text(blocks, normalized_export_path)
    if source_anchor is not None:
        _assert_docx_source_generation_seal(source_anchor["seal"], stage="after_rebuild")
    # Source DOCX exports are complete after snapshot reconstruction.  The
    # immutable source layout is audited below by the dedicated format-lock and
    # OOXML checks; no school-rule parser or style rewrite runs here.
    if expected_text_targets:
        text_integrity_performed = True
        text_integrity_report = _audit_exported_docx_text_targets(normalized_export_path, expected_text_targets)
        if not bool(text_integrity_report.get("ok")):
            text_integrity_report_path = evidence_export_path.with_suffix(".text_integrity.json")
            text_integrity_report["reportPath"] = str(text_integrity_report_path)
            text_integrity_report_path.parent.mkdir(parents=True, exist_ok=True)
            text_integrity_report_path.write_text(
                json.dumps(text_integrity_report, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            issue_count = int(text_integrity_report.get("issueCount", 0) or 0)
            message = (
                "DOCX 导出已拦截：导出后的正文文本与本轮结果不一致。"
                f" 共 {issue_count} 个问题，报告：{text_integrity_report_path}"
            )
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="text-integrity",
                    label="正文文本",
                    message=message,
                    report=text_integrity_report,
                    report_path=text_integrity_report_path,
                    issue_count=issue_count,
                    sample_keys=("issues",),
                ),
            )
    audit_report_path = get_docx_audit_report_path(evidence_export_path)
    audit_report: dict[str, Any] | None = None
    ooxml_audit_report_path = get_docx_ooxml_audit_report_path(evidence_export_path)
    ooxml_audit_report: dict[str, Any] | None = None
    if origin_docx_bundle is not None:
        audit_report = audit_docx_export(
            normalized_export_path,
            source_path=origin_docx_bundle[0],
            snapshot_path=origin_docx_bundle[1],
            expected_source_sha256=str((source_anchor or {}).get("sourceSha256", "") or "") or None,
            provenance_source_path=provenance_source_path,
            report_path=audit_report_path,
        )
        if not bool(audit_report.get("ok")):
            issue_count = int(audit_report.get("issueCount", 0) or 0)
            message = (
                "DOCX 导出已拦截：审计发现保护区内容发生变化。"
                f" 共 {issue_count} 个问题，报告：{audit_report_path}"
            )
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="audit",
                    label="保护区审计",
                    message=message,
                    report=audit_report,
                    report_path=audit_report_path,
                    issue_count=issue_count,
                    sample_keys=("issues",),
                ),
            )
    if origin_docx_bundle is not None:
        ooxml_audit_report = audit_docx_ooxml_integrity(
            normalized_export_path,
            source_path=origin_docx_bundle[0],
            snapshot_path=origin_docx_bundle[1],
            expected_source_sha256=str((source_anchor or {}).get("sourceSha256", "") or "") or None,
            provenance_source_path=provenance_source_path,
            report_path=ooxml_audit_report_path,
        )
        if not bool(ooxml_audit_report.get("ok")):
            issue_count = int(ooxml_audit_report.get("issueCount", 0) or 0)
            message = (
                "DOCX export blocked: OOXML integrity audit found protected structure changes. "
                f"Issues: {issue_count}; report: {ooxml_audit_report_path}"
            )
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="ooxml",
                    label="Word 结构",
                    message=message,
                    report=ooxml_audit_report,
                    report_path=ooxml_audit_report_path,
                    issue_count=issue_count,
                    sample_keys=("issues",),
                ),
            )
    format_lock_report: dict[str, Any] | None = None
    if origin_docx_bundle is not None and preserve_original_format:
        # Fidelity-lock mode is a hard contract: all paragraph wrappers and
        # formatting, tables, and section properties must match the source after
        # excluding only editable w:t payloads. Never silently bypass this audit.
        format_lock_report_path = get_docx_format_lock_report_path(evidence_export_path)
        format_lock_report = audit_docx_format_lock(
            normalized_export_path,
            source_path=origin_docx_bundle[0],
            snapshot_path=origin_docx_bundle[1],
            expected_source_sha256=str((source_anchor or {}).get("sourceSha256", "") or "") or None,
            provenance_source_path=provenance_source_path,
            report_path=format_lock_report_path,
        )
        if not bool(format_lock_report.get("ok")):
            issue_count = int(format_lock_report.get("issueCount", 0) or 0)
            message = (
                "DOCX 导出已拦截：保真审计发现段落、表格或页面结构发生格式漂移。"
                f" 共 {issue_count} 个问题，报告：{format_lock_report_path}"
            )
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="format-lock",
                    label="格式保真",
                    message=message,
                    report=format_lock_report,
                    report_path=format_lock_report_path,
                    issue_count=issue_count,
                    sample_keys=("issues",),
                ),
            )
    if origin_docx_bundle is not None:
        content_contract_report = build_document_edit_contract(
            origin_docx_bundle[0],
            snapshot_path=origin_docx_bundle[1],
            body_map=contract_body_map,
            candidate_texts=contract_candidate_texts,
            stage="post_export",
            export_path=normalized_export_path,
            expected_source_sha256=str((source_anchor or {}).get("sourceSha256", "") or "") or None,
            provenance_source_path=provenance_source_path,
            export_evidence={
                "textIntegrityOk": True,
                "protectedTextAuditOk": bool(audit_report and audit_report.get("ok")),
                "ooxmlIntegrityOk": bool(ooxml_audit_report and ooxml_audit_report.get("ok")),
                "formatLockAuditOk": bool(format_lock_report and format_lock_report.get("ok")),
                "auditPath": str(audit_report_path),
                "ooxmlAuditPath": str(ooxml_audit_report_path),
                "formatLockPath": str((format_lock_report or {}).get("reportPath", "")),
            },
            report_path=content_contract_path,
        )
        post_export_contract_performed = True
        try:
            assert_document_edit_contract_ready(content_contract_report, label="DOCX 导出终检")
        except ValueError as exc:
            message = str(exc)
            raise ExportRoundError(
                message,
                _build_export_failure(
                    stage="content-contract-final",
                    label="正文与格式终检",
                    message=message,
                    report=content_contract_report,
                    sample_keys=("issues",),
                ),
            ) from exc
    checks_performed = ["document_generation"]
    if guard_report is not None:
        checks_performed.append("pre_export_guard")
    if content_contract_report is not None and origin_docx_bundle is not None:
        checks_performed.append("content_contract")
    if text_integrity_performed:
        checks_performed.append("text_integrity")
    if audit_report is not None:
        checks_performed.append("protected_text_audit")
    if ooxml_audit_report is not None:
        checks_performed.append("ooxml_integrity")
    if format_lock_report is not None:
        checks_performed.append("format_lock")
    if post_export_contract_performed:
        checks_performed.append("post_export_contract")
    if source_anchor is not None:
        _assert_docx_source_generation_seal(source_anchor["seal"], stage="after_post_contract")
        checks_performed.append("source_generation_anchor")

    return {
        "format": "docx",
        "path": str(committed_export_path),
        "evidenceVersion": 1,
        "overallStatus": "passed",
        "sourceKind": "original_docx" if origin_docx_bundle is not None else "generated_docx",
        "contentContractStatus": "passed" if origin_docx_bundle is not None else "not_applicable",
        "formatLockStatus": "passed" if format_lock_report is not None else "not_applicable",
        "checksPerformed": checks_performed,
        "layoutMode": layout_mode,
        "paragraphSource": paragraph_source,
        "formatMode": resolved_format_mode,
        "sourceSha256": str((source_anchor or {}).get("sourceSha256", "") or ""),
        "provenanceSourcePath": str(provenance_source_path or ""),
        "sourceAnchorPath": str((source_anchor or {}).get("anchorPath", "") or ""),
        "_sourceGenerationSeal": copy.deepcopy((source_anchor or {}).get("seal")),
        "formatScope": "editable_body_only" if origin_docx_bundle is not None else "generated_content_only",
        "validationPath": str(validation_path) if validation_path is not None else "",
        "auditPath": str(audit_report_path) if audit_report is not None else "",
        "auditIssueCount": int(audit_report.get("issueCount", 0) or 0) if audit_report is not None else 0,
        "ooxmlAuditPath": str(ooxml_audit_report_path) if ooxml_audit_report is not None else "",
        "ooxmlAuditIssueCount": int(ooxml_audit_report.get("issueCount", 0) or 0) if ooxml_audit_report is not None else 0,
        "formatLockPath": str(format_lock_report.get("reportPath", "")) if format_lock_report is not None else "",
        "formatLockIssueCount": int(format_lock_report.get("issueCount", 0) or 0) if format_lock_report is not None else 0,
        "formatLockEditableChecked": int(format_lock_report.get("editableChecked", 0) or 0) if format_lock_report is not None else 0,
        "contentContractPath": str((content_contract_report or {}).get("reportPath", "")),
        "contentContractReady": bool((content_contract_report or {}).get("ready", False)),
        "contentContractIssueCount": int((content_contract_report or {}).get("issueCount", 0) or 0),
        "editableUnitCount": int((content_contract_report or {}).get("editableUnitCount", 0) or 0),
        "protectedUnitCount": int((content_contract_report or {}).get("protectedUnitCount", 0) or 0),
        "protectedHeadingCount": int((content_contract_report or {}).get("protectedHeadingCount", 0) or 0),
        "editableHeadingCount": int((content_contract_report or {}).get("editableHeadingCount", 0) or 0),
        "semanticRangeCount": int(
            (content_contract_report or {}).get("semanticRangeCount", 0) or 0
        ),
        "semanticRangeTopologyValid": bool(
            (content_contract_report or {}).get("semanticRangeTopologyValid", False)
        ),
        "semanticRangeIssueCount": int(
            (content_contract_report or {}).get("semanticRangeIssueCount", 0) or 0
        ),
        "semanticRangeAnchorUnitCount": int(
            (content_contract_report or {}).get("semanticRangeAnchorUnitCount", 0) or 0
        ),
        "protectedSemanticRangeAnchorUnitCount": int(
            (content_contract_report or {}).get("protectedSemanticRangeAnchorUnitCount", 0) or 0
        ),
        "editableSemanticRangeAnchorUnitCount": int(
            (content_contract_report or {}).get("editableSemanticRangeAnchorUnitCount", 0) or 0
        ),
        "semanticRangeCoveredUnitCount": int(
            (content_contract_report or {}).get("semanticRangeCoveredUnitCount", 0) or 0
        ),
        "protectedSemanticRangeCoveredUnitCount": int(
            (content_contract_report or {}).get("protectedSemanticRangeCoveredUnitCount", 0) or 0
        ),
        "editableSemanticRangeCoveredUnitCount": int(
            (content_contract_report or {}).get("editableSemanticRangeCoveredUnitCount", 0) or 0
        ),
        "semanticPointReferenceUnitCount": int(
            (content_contract_report or {}).get("semanticPointReferenceUnitCount", 0) or 0
        ),
        "protectedSemanticPointReferenceUnitCount": int(
            (content_contract_report or {}).get("protectedSemanticPointReferenceUnitCount", 0) or 0
        ),
        "editableSemanticPointReferenceUnitCount": int(
            (content_contract_report or {}).get("editableSemanticPointReferenceUnitCount", 0) or 0
        ),
        "modelInputMatchesEditableUnits": bool((content_contract_report or {}).get("modelInputMatchesEditableUnits", False)),
        "guardPath": str(guard_report.get("reportPath", "")) if guard_report is not None else "",
        "guardIssueCount": int(guard_report.get("blockingIssueCount", 0) or 0) if guard_report is not None else 0,
        "guardWarningCount": int(guard_report.get("warningCount", 0) or 0) if guard_report is not None else 0,
        "guardIssueSamples": _collect_issue_samples(guard_report, keys=("blockingIssues", "warnings")),
        "auditIssueSamples": _collect_issue_samples(audit_report, keys=("issues",)),
        "ooxmlAuditIssueSamples": _collect_issue_samples(ooxml_audit_report, keys=("issues",)),
    }


def _build_docx_text_targets_from_body_map(body_map: Any) -> list[tuple[dict[str, Any], str]]:
    targets: list[tuple[dict[str, Any], str]] = []
    for unit in getattr(body_map, "units", []) or []:
        target = getattr(unit, "target", {})
        if isinstance(target, dict):
            expected_text = _restore_rewritten_text_with_source_whitespace(
                str(getattr(unit, "current_text", "")),
                original_text=str(getattr(unit, "original_text", "")),
                leading_whitespace=str(getattr(unit, "leading_whitespace", "")),
                trailing_whitespace=str(getattr(unit, "trailing_whitespace", "")),
            )
            targets.append((dict(target), expected_text))
    return targets


def _build_docx_text_targets_from_snapshot(snapshot_path: Path, rewritten_paragraphs: list[str]) -> list[tuple[dict[str, Any], str]]:
    snapshot = _load_docx_snapshot(snapshot_path)
    if snapshot is None:
        return []
    return [
        (
            dict(unit.target),
            _restore_rewritten_text_with_source_whitespace(
                str(rewritten_text),
                original_text=unit.text,
                leading_whitespace=unit.leading_whitespace,
                trailing_whitespace=unit.trailing_whitespace,
            ),
        )
        for unit, rewritten_text in zip(snapshot.editable_units(), rewritten_paragraphs)
    ]


def _audit_exported_docx_text_targets(export_path: Path, expected_targets: list[tuple[dict[str, Any], str]]) -> dict[str, Any]:
    document = Document(str(export_path.resolve()))
    issues: list[dict[str, Any]] = []
    checks: list[dict[str, Any]] = []
    for index, (target, expected_text) in enumerate(expected_targets):
        try:
            paragraph = _resolve_target_paragraph(document, target)
            actual_text = paragraph.text
        except Exception as exc:
            issues.append({
                "code": "docx_text_target_unreadable",
                "message": f"无法读取导出后的正文目标：{type(exc).__name__}",
                "location": json.dumps(target, ensure_ascii=False, sort_keys=True),
            })
            continue
        expected = str(expected_text)
        check = {
            "index": index,
            "location": json.dumps(target, ensure_ascii=False, sort_keys=True),
            "expectedPreview": _truncate_issue_text(expected, 120),
            "actualPreview": _truncate_issue_text(actual_text, 120),
        }
        checks.append(check)
        if actual_text != expected:
            issues.append({
                "code": "docx_exported_text_changed",
                "message": "导出后的正文文本与本轮结果不一致。",
                "location": check["location"],
                "sample": f"expected={check['expectedPreview']}; actual={check['actualPreview']}",
            })
    return {
        "ok": not issues,
        "issueCount": len(issues),
        "checked": len(checks),
        "issues": issues[:20],
        "sampleChecks": checks[:8],
    }


def _export_text_integrity_error(message: str) -> ExportRoundError:
    return ExportRoundError(
        message,
        _build_export_failure(
            stage="text-integrity",
            label="正文文本",
            message=message,
            issue_count=1,
            samples=[{
                "code": "docx_text_integrity_failed",
                "message": _truncate_issue_text(message, 240),
            }],
        ),
    )


def read_output_text(output_path: str, max_chars: int | None = None) -> dict[str, Any]:
    normalized_output_path = _resolve_api_path(output_path, allowed_roots=API_READ_ALLOWED_ROOTS, label="Output path")
    text = normalized_output_path.read_text(encoding="utf-8")
    total_chars = len(text)
    normalized_limit = max_chars if isinstance(max_chars, int) and max_chars > 0 else None
    truncated = normalized_limit is not None and total_chars > normalized_limit
    preview_text = text
    if truncated and normalized_limit is not None:
        preview_text = (
            text[:normalized_limit].rstrip()
            + "\n\n[预览已截断，导出文件可查看完整内容]"
        )
    return {
        "path": str(normalized_output_path),
        "text": preview_text,
        "truncated": truncated,
        "totalChars": total_chars,
        "previewChars": len(preview_text),
    }


def read_round_compare(output_path: str, *, include_revision: bool = False) -> dict[str, Any]:
    normalized_output_path = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    compare_path = _find_compare_path_for_output(normalized_output_path)
    if not compare_path.exists():
        raise ValueError(f"Compare data not found for output: {normalized_output_path}")
    raw_payload = json.loads(compare_path.read_text(encoding="utf-8"))
    if not isinstance(raw_payload, dict):
        raise ValueError(f"Invalid compare data payload: {compare_path}")
    payload = _normalize_compare_failed_attempts(copy.deepcopy(raw_payload))
    if not _compare_payload_is_usable(normalized_output_path, payload):
        raise ValueError("本轮结果不完整：输出或 Diff 数据为空，不能视为已完成。")
    if include_revision:
        # Response-only concurrency token. Do not persist this derived field in
        # legacy compare files that predate updatedAt.
        payload["compareRevision"] = _compare_revision_for_review(payload, compare_path)
    return payload


def _find_review_decisions_path_for_output(output_path: Path) -> Path:
    compare_path = _find_compare_path_for_output(output_path)
    return compare_path.with_name(f"{compare_path.stem}_review_decisions.json")


def _compare_revision_for_review(compare_payload: dict[str, Any], compare_path: Path) -> str:
    revision = str(compare_payload.get("updatedAt", "") or "").strip()
    if revision:
        return revision
    try:
        return f"sha256:{_sha256_file_bytes(compare_path)}"
    except OSError:
        return ""


def _round_snapshot_fail(
    code: str,
    message: str,
    *,
    retryable: bool = False,
    details: dict[str, Any] | None = None,
) -> None:
    raise RoundArtifactSnapshotError(
        code,
        message,
        retryable=retryable,
        details=details,
    )


def _round_snapshot_sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _round_snapshot_normalize_newlines(value: str) -> str:
    """Compare text artifacts independent of the host text-mode newline convention."""
    return value.replace("\r\n", "\n").replace("\r", "\n")


def _round_snapshot_digest(parts: list[tuple[str, bytes | None]]) -> str:
    """Hash named, length-delimited artifact bytes without concatenation ambiguity."""

    digest = hashlib.sha256()
    digest.update(b"fyadr:round-artifact-snapshot:v1\x00")
    for name, payload in parts:
        encoded_name = str(name).encode("utf-8")
        digest.update(len(encoded_name).to_bytes(4, "big"))
        digest.update(encoded_name)
        if payload is None:
            digest.update(b"\x00")
            continue
        digest.update(b"\x01")
        digest.update(len(payload).to_bytes(8, "big"))
        digest.update(payload)
    return digest.hexdigest()


def _read_round_snapshot_artifact_bytes(
    path: Path,
    *,
    missing_code: str,
    unreadable_code: str,
    label: str,
) -> bytes:
    if not path.exists() or not path.is_file():
        _round_snapshot_fail(missing_code, f"{label} is missing: {path}")
    try:
        return path.read_bytes()
    except OSError as exc:
        _round_snapshot_fail(
            unreadable_code,
            f"{label} is unreadable: {path}",
            details={"errorType": type(exc).__name__},
        )
    raise AssertionError("unreachable")


def _parse_round_snapshot_json_object(
    payload: bytes,
    *,
    code: str,
    label: str,
) -> dict[str, Any]:
    try:
        decoded = payload.decode("utf-8")
        value = json.loads(decoded)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        _round_snapshot_fail(
            code,
            f"{label} is not valid UTF-8 JSON.",
            details={"errorType": type(exc).__name__},
        )
    if not isinstance(value, dict):
        _round_snapshot_fail(code, f"{label} must contain a JSON object.")
    return value


def _round_snapshot_declared_path_matches(
    declared_path: Any,
    expected_path: Path,
    *,
    code: str,
    label: str,
    allow_missing: bool = False,
) -> None:
    if declared_path is None and allow_missing:
        return
    if not isinstance(declared_path, str):
        _round_snapshot_fail(code, f"{label} outputPath must be a string.")
    raw_path = declared_path.strip()
    if not raw_path:
        _round_snapshot_fail(code, f"{label} does not declare outputPath.")
    try:
        resolved = normalize_path(Path(raw_path)).resolve()
    except (OSError, RuntimeError, ValueError) as exc:
        _round_snapshot_fail(
            code,
            f"{label} declares an invalid outputPath.",
            details={"errorType": type(exc).__name__},
        )
    if resolved != expected_path.resolve():
        _round_snapshot_fail(
            code,
            f"{label} belongs to another round output.",
            details={"declaredOutputPath": raw_path},
        )


def _normalize_round_snapshot_review_decision(
    decision: Any,
    *,
    chunk_id: str,
) -> str | dict[str, Any]:
    if isinstance(decision, str):
        if decision not in {"source", "rewrite", "source_confirmed", "rewrite_confirmed"}:
            _round_snapshot_fail(
                "round_snapshot_review_decision_invalid",
                f"Review decision for chunk {chunk_id} is invalid.",
                details={"chunkId": chunk_id},
            )
        return decision
    if not isinstance(decision, dict) or str(decision.get("mode", "")).strip().lower() != "custom":
        _round_snapshot_fail(
            "round_snapshot_review_decision_invalid",
            f"Review decision for chunk {chunk_id} has an invalid shape.",
            details={"chunkId": chunk_id},
        )
    text = decision.get("text")
    if not isinstance(text, str) or not text.strip():
        _round_snapshot_fail(
            "round_snapshot_review_decision_invalid",
            f"Custom review decision for chunk {chunk_id} has no text.",
            details={"chunkId": chunk_id},
        )
    if "source" in decision and not isinstance(decision.get("source"), str):
        _round_snapshot_fail(
            "round_snapshot_review_decision_invalid",
            f"Custom review decision source for chunk {chunk_id} is invalid.",
            details={"chunkId": chunk_id},
        )
    if "confirmed" in decision and not isinstance(decision.get("confirmed"), bool):
        _round_snapshot_fail(
            "round_snapshot_review_decision_invalid",
            f"Custom review confirmation for chunk {chunk_id} is invalid.",
            details={"chunkId": chunk_id},
        )
    if "attempt" in decision and (
        isinstance(decision.get("attempt"), bool)
        or (
            decision.get("attempt") is not None
            and not isinstance(decision.get("attempt"), int)
        )
    ):
        _round_snapshot_fail(
            "round_snapshot_review_decision_invalid",
            f"Custom review attempt for chunk {chunk_id} is invalid.",
            details={"chunkId": chunk_id},
        )
    if "error" in decision and not isinstance(decision.get("error"), str):
        _round_snapshot_fail(
            "round_snapshot_review_decision_invalid",
            f"Custom review error evidence for chunk {chunk_id} is invalid.",
            details={"chunkId": chunk_id},
        )
    return _normalize_review_decision_value(decision)


def _parse_round_snapshot_manifest(payload: dict[str, Any]) -> ChunkManifest:
    try:
        raw_paragraphs = payload["paragraphs"]
        raw_chunks = payload["chunks"]
        if not isinstance(raw_paragraphs, list) or not isinstance(raw_chunks, list):
            raise TypeError("manifest lists are missing")
        chunk_metric = str(payload.get("chunk_metric", "char"))
        if chunk_metric not in {"char", "word"}:
            raise ValueError("chunk metric is invalid")
        manifest = ChunkManifest(
            chunk_limit=int(payload["chunk_limit"]),
            chunk_metric=chunk_metric,  # type: ignore[arg-type]
            paragraph_count=int(payload["paragraph_count"]),
            chunk_count=int(payload["chunk_count"]),
            paragraphs=[ParagraphManifest(**item) for item in raw_paragraphs],
            chunks=[Chunk(**item) for item in raw_chunks],
        )
    except (KeyError, TypeError, ValueError) as exc:
        _round_snapshot_fail(
            "round_snapshot_manifest_corrupt",
            "Round manifest has an invalid schema.",
            details={"errorType": type(exc).__name__},
        )
    if (
        manifest.paragraph_count <= 0
        or manifest.chunk_count <= 0
        or manifest.paragraph_count != len(manifest.paragraphs)
        or manifest.chunk_count != len(manifest.chunks)
    ):
        _round_snapshot_fail(
            "round_snapshot_manifest_corrupt",
            "Round manifest counts do not match its captured entries.",
        )
    return manifest


def _resolve_round_snapshot_manifest_path(
    output_path: Path,
    compare_payload: dict[str, Any],
) -> tuple[Path | None, bool]:
    declared_manifest_path = compare_payload.get("manifestPath", "")
    if declared_manifest_path is not None and not isinstance(declared_manifest_path, str):
        _round_snapshot_fail(
            "round_snapshot_manifest_path_invalid",
            "Round compare manifestPath must be a string.",
        )
    raw_manifest_path = str(declared_manifest_path or "").strip()
    explicitly_declared = bool(raw_manifest_path)
    manifest_path: Path | None
    if raw_manifest_path:
        manifest_path = normalize_path(Path(raw_manifest_path)).resolve()
    else:
        manifest_path = _find_manifest_path_for_output(output_path)
        manifest_path = manifest_path.resolve() if manifest_path is not None else None
    if manifest_path is not None and not is_path_under(manifest_path, ROOT_DIR / "finish"):
        _round_snapshot_fail(
            "round_snapshot_manifest_path_invalid",
            "Round manifest must stay under the finish workspace directory.",
        )
    return manifest_path, explicitly_declared


def _capture_round_artifact_snapshot_unlocked(
    normalized_output_path: Path,
    *,
    max_preview_chars: int | None,
    include_internal: bool,
) -> dict[str, Any]:
    output_bytes = _read_round_snapshot_artifact_bytes(
        normalized_output_path,
        missing_code="round_snapshot_output_missing",
        unreadable_code="round_snapshot_output_unreadable",
        label="Round output",
    )
    try:
        raw_output_text = output_bytes.decode("utf-8")
    except UnicodeDecodeError as exc:
        _round_snapshot_fail(
            "round_snapshot_output_corrupt",
            "Round output is not valid UTF-8 text.",
            details={"errorType": type(exc).__name__},
        )
    if not raw_output_text.strip():
        _round_snapshot_fail("round_snapshot_output_empty", "Round output is empty.")

    compare_path = _find_compare_path_for_output(normalized_output_path).resolve()
    if not is_path_under(compare_path, ROOT_DIR / "finish"):
        _round_snapshot_fail(
            "round_snapshot_compare_path_invalid",
            "Round compare data must stay under the finish workspace directory.",
        )
    compare_bytes = _read_round_snapshot_artifact_bytes(
        compare_path,
        missing_code="round_snapshot_compare_missing",
        unreadable_code="round_snapshot_compare_unreadable",
        label="Round compare data",
    )
    raw_compare_payload = _parse_round_snapshot_json_object(
        compare_bytes,
        code="round_snapshot_compare_corrupt",
        label="Round compare data",
    )
    compare_payload = _normalize_compare_failed_attempts(copy.deepcopy(raw_compare_payload))
    raw_compare_output_path = compare_payload.get("outputPath")
    compare_relocated = False
    if isinstance(raw_compare_output_path, str) and raw_compare_output_path.strip():
        try:
            compare_relocated = normalize_path(Path(raw_compare_output_path)).resolve() != normalized_output_path.resolve()
        except (OSError, RuntimeError, ValueError):
            compare_relocated = False
            _round_snapshot_declared_path_matches(
                raw_compare_output_path,
                normalized_output_path,
                code="round_snapshot_compare_output_mismatch",
                label="Round compare data",
            )
    if compare_relocated:
        # Legacy orphan-recovery workflows may move a complete round bundle
        # without rewriting its compare metadata.  Only accept that relocation
        # when every declared manifest dependency has a local sibling; never
        # follow the old path because it is protected by a different lock and
        # could otherwise reintroduce a mixed-generation read.
        raw_declared_manifest = str(compare_payload.get("manifestPath", "") or "").strip()
        local_manifest_path = _find_manifest_path_for_output(
            normalized_output_path,
            include_records=False,
        )
        if raw_declared_manifest and local_manifest_path is None:
            _round_snapshot_fail(
                "round_snapshot_compare_output_mismatch",
                "Relocated round compare has no local manifest sibling.",
            )
        compare_payload["outputPath"] = str(normalized_output_path)
        if local_manifest_path is not None:
            compare_payload["manifestPath"] = str(local_manifest_path.resolve())
    else:
        _round_snapshot_declared_path_matches(
            raw_compare_output_path,
            normalized_output_path,
            code="round_snapshot_compare_output_mismatch",
            label="Round compare data",
            allow_missing="outputPath" not in compare_payload,
        )
    legacy_record_context: tuple[dict[str, Any], dict[str, Any]] | None = None
    raw_doc_id = compare_payload.get("docId")
    doc_id = raw_doc_id.strip() if isinstance(raw_doc_id, str) else ""
    if not doc_id and "docId" not in compare_payload:
        legacy_record_context = _find_record_context_for_output(normalized_output_path)
        if legacy_record_context is not None:
            legacy_entry, _legacy_round_item = legacy_record_context
            legacy_origin = str(legacy_entry.get("origin_path", "") or "").strip()
            if legacy_origin:
                doc_id = normalize_doc_id(legacy_origin)
                compare_payload["docId"] = doc_id
    if not doc_id:
        _round_snapshot_fail("round_snapshot_compare_identity_invalid", "Round compare data has no docId.")
    raw_round = compare_payload.get("round")
    if raw_round is None and "round" not in compare_payload:
        if legacy_record_context is None:
            legacy_record_context = _find_record_context_for_output(normalized_output_path)
        if legacy_record_context is not None:
            _legacy_entry, legacy_round_item = legacy_record_context
            raw_round = legacy_round_item.get("round")
            compare_payload["round"] = raw_round
    if isinstance(raw_round, bool) or not isinstance(raw_round, int):
        _round_snapshot_fail("round_snapshot_compare_identity_invalid", "Round compare data has an invalid round.")
    round_number = raw_round
    if round_number <= 0:
        _round_snapshot_fail("round_snapshot_compare_identity_invalid", "Round compare data has an invalid round.")

    raw_chunks = compare_payload.get("chunks")
    if not isinstance(raw_chunks, list) or not raw_chunks:
        _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare data has no chunks.")
    raw_expected_chunk_count = compare_payload.get("chunkCount")
    raw_expected_paragraph_count = compare_payload.get("paragraphCount")
    expected_chunk_count = (
        raw_expected_chunk_count
        if isinstance(raw_expected_chunk_count, int) and not isinstance(raw_expected_chunk_count, bool)
        else None
    )
    expected_paragraph_count = (
        raw_expected_paragraph_count
        if isinstance(raw_expected_paragraph_count, int) and not isinstance(raw_expected_paragraph_count, bool)
        else None
    )
    if expected_chunk_count is None or expected_chunk_count <= 0 or expected_chunk_count != len(raw_chunks):
        _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare chunkCount is inconsistent.")
    if expected_paragraph_count is None and "paragraphCount" not in compare_payload:
        legacy_paragraph_indices = [
            item.get("paragraphIndex")
            for item in raw_chunks
            if isinstance(item, dict)
        ]
        if (
            legacy_paragraph_indices
            and all(
                isinstance(value, int) and not isinstance(value, bool) and value >= 0
                for value in legacy_paragraph_indices
            )
        ):
            expected_paragraph_count = max(legacy_paragraph_indices) + 1
            compare_payload["paragraphCount"] = expected_paragraph_count
    if expected_paragraph_count is None or expected_paragraph_count <= 0:
        _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare paragraphCount is invalid.")
    valid_chunk_ids: set[str] = set()
    for item in raw_chunks:
        if not isinstance(item, dict):
            _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare contains an invalid chunk.")
        raw_chunk_id = item.get("chunkId")
        chunk_id = raw_chunk_id.strip() if isinstance(raw_chunk_id, str) else ""
        if not chunk_id or chunk_id in valid_chunk_ids:
            _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare chunk ids are missing or duplicated.")
        if not isinstance(item.get("inputText"), str) or not isinstance(item.get("outputText"), str):
            _round_snapshot_fail(
                "round_snapshot_compare_corrupt",
                f"Round compare chunk {chunk_id} has invalid text fields.",
                details={"chunkId": chunk_id},
            )
        valid_chunk_ids.add(chunk_id)

    compare_sha256 = _round_snapshot_sha256(compare_bytes)
    raw_compare_updated_at = compare_payload.get("updatedAt", "")
    if raw_compare_updated_at is not None and not isinstance(raw_compare_updated_at, str):
        _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare updatedAt must be a string.")
    compare_revision = str(raw_compare_updated_at or "").strip() or f"sha256:{compare_sha256}"
    raw_compare_review_updated_at = compare_payload.get("reviewUpdatedAt", "")
    if raw_compare_review_updated_at is not None and not isinstance(raw_compare_review_updated_at, str):
        _round_snapshot_fail("round_snapshot_compare_corrupt", "Round compare reviewUpdatedAt must be a string.")
    compare_review_updated_at = str(raw_compare_review_updated_at or "").strip()

    review_path = _find_review_decisions_path_for_output(normalized_output_path).resolve()
    if not is_path_under(review_path, ROOT_DIR / "finish"):
        _round_snapshot_fail(
            "round_snapshot_review_path_invalid",
            "Review decisions must stay under the finish workspace directory.",
        )
    review_exists = review_path.exists()
    if compare_review_updated_at and not review_exists:
        _round_snapshot_fail(
            "round_snapshot_review_missing",
            "Round compare declares review decisions, but the sidecar is missing.",
        )

    review_bytes: bytes | None = None
    review_payload: dict[str, Any] | None = None
    decisions: dict[str, Any] = {}
    review_updated_at = ""
    review_base_compare_revision = ""
    review_link_status = "none"
    if review_exists:
        review_bytes = _read_round_snapshot_artifact_bytes(
            review_path,
            missing_code="round_snapshot_review_missing",
            unreadable_code="round_snapshot_review_unreadable",
            label="Review decision sidecar",
        )
        review_payload = _parse_round_snapshot_json_object(
            review_bytes,
            code="round_snapshot_review_corrupt",
            label="Review decision sidecar",
        )
        _round_snapshot_declared_path_matches(
            review_payload.get("outputPath"),
            normalized_output_path,
            code="round_snapshot_review_output_mismatch",
            label="Review decision sidecar",
        )
        if "docId" in review_payload:
            raw_review_doc_id = review_payload.get("docId")
            if not isinstance(raw_review_doc_id, str) or raw_review_doc_id.strip() != doc_id:
                _round_snapshot_fail(
                    "round_snapshot_review_identity_mismatch",
                    "Review decision sidecar belongs to another document.",
                )
        if "round" in review_payload:
            review_round = review_payload.get("round")
            if isinstance(review_round, bool) or not isinstance(review_round, int) or review_round != round_number:
                _round_snapshot_fail(
                    "round_snapshot_review_identity_mismatch",
                    "Review decision sidecar belongs to another round.",
                )
        raw_review_updated_at = review_payload.get("updatedAt")
        if not isinstance(raw_review_updated_at, str):
            _round_snapshot_fail(
                "round_snapshot_review_corrupt",
                "Review decision sidecar updatedAt must be a string.",
            )
        review_updated_at = raw_review_updated_at.strip()
        if not review_updated_at:
            _round_snapshot_fail(
                "round_snapshot_review_corrupt",
                "Review decision sidecar has no updatedAt revision.",
            )
        if compare_review_updated_at:
            if review_updated_at != compare_review_updated_at:
                _round_snapshot_fail(
                    "round_snapshot_review_link_mismatch",
                    "Review decision sidecar is not linked to compare.reviewUpdatedAt.",
                    details={
                        "compareReviewUpdatedAt": compare_review_updated_at,
                        "reviewUpdatedAt": review_updated_at,
                    },
                )
            review_link_status = "linked"
        else:
            review_link_status = "legacy_unversioned"
        for key in ("compareRevision", "reviewBaseCompareRevision"):
            if key in review_payload and not isinstance(review_payload.get(key), str):
                _round_snapshot_fail(
                    "round_snapshot_review_corrupt",
                    f"Review decision sidecar field {key} must be a string.",
                )
        review_base_compare_revision = str(
            review_payload.get(
                "reviewBaseCompareRevision",
                review_payload.get("compareRevision", review_updated_at),
            )
            or ""
        ).strip()
        raw_decisions = review_payload.get("decisions")
        if not isinstance(raw_decisions, dict):
            _round_snapshot_fail(
                "round_snapshot_review_corrupt",
                "Review decision sidecar decisions must be an object.",
            )
        unknown_chunk_ids = sorted(str(chunk_id) for chunk_id in raw_decisions if str(chunk_id) not in valid_chunk_ids)
        if unknown_chunk_ids:
            _round_snapshot_fail(
                "round_snapshot_review_chunk_mismatch",
                "Review decision sidecar references chunks outside the captured compare generation.",
                details={"chunkIds": unknown_chunk_ids[:12]},
            )
        decisions = {
            str(chunk_id): _normalize_round_snapshot_review_decision(decision, chunk_id=str(chunk_id))
            for chunk_id, decision in raw_decisions.items()
        }

    try:
        _assert_document_release_payload(compare_payload, decisions)
    except DocumentReleaseGateError as exc:
        _round_snapshot_fail(
            "round_snapshot_release_gate_failed",
            "Captured review materialization failed the document release gate.",
            details=exc.details,
        )

    manifest_path, manifest_explicitly_declared = _resolve_round_snapshot_manifest_path(
        normalized_output_path,
        compare_payload,
    )
    manifest_bytes: bytes | None = None
    manifest_payload: dict[str, Any] | None = None
    captured_manifest: ChunkManifest | None = None
    if manifest_path is not None and manifest_path.exists():
        manifest_bytes = _read_round_snapshot_artifact_bytes(
            manifest_path,
            missing_code="round_snapshot_manifest_missing",
            unreadable_code="round_snapshot_manifest_unreadable",
            label="Round manifest",
        )
        manifest_payload = _parse_round_snapshot_json_object(
            manifest_bytes,
            code="round_snapshot_manifest_corrupt",
            label="Round manifest",
        )
        captured_manifest = _parse_round_snapshot_manifest(manifest_payload)
    elif manifest_explicitly_declared:
        _round_snapshot_fail(
            "round_snapshot_manifest_missing",
            "Round compare declares a manifest, but it is missing.",
        )
    else:
        manifest_path = None

    paragraphs = _build_paragraphs_from_compare_payload(
        normalized_output_path,
        compare_payload,
        decisions,
        captured_manifest=captured_manifest,
        allow_manifest_disk_read=False,
    )
    if paragraphs is None or len(paragraphs) != expected_paragraph_count:
        _round_snapshot_fail(
            "round_snapshot_materialization_failed",
            "Captured compare and review decisions cannot restore every round paragraph.",
        )
    effective_chunks: list[dict[str, Any]] = []
    for index, item in enumerate(raw_chunks):
        chunk_id = str(item.get("chunkId", "") or "").strip()
        input_text = str(item.get("inputText", "") or "")
        selected_text = _select_review_text(
            input_text,
            item.get("outputText", ""),
            decisions.get(chunk_id),
            chunk=item,
        )
        effective_chunk_text = normalize_chunk_output(input_text, selected_text)
        if not effective_chunk_text.strip():
            _round_snapshot_fail(
                "round_snapshot_materialization_failed",
                f"Captured chunk {chunk_id} materialized to empty text.",
                details={"chunkId": chunk_id},
            )
        try:
            paragraph_index = int(item.get("paragraphIndex", index))
            chunk_index = int(item.get("chunkIndex", 0))
        except (TypeError, ValueError):
            _round_snapshot_fail(
                "round_snapshot_compare_corrupt",
                f"Round compare chunk {chunk_id} has an invalid position.",
                details={"chunkId": chunk_id},
            )
        effective_chunks.append(
            {
                "chunkId": chunk_id,
                "paragraphIndex": paragraph_index,
                "chunkIndex": chunk_index,
                "text": effective_chunk_text,
            }
        )
    effective_text = "\n\n".join(paragraphs)
    if not effective_text.strip():
        _round_snapshot_fail(
            "round_snapshot_materialization_failed",
            "Captured compare and review decisions materialized to empty text.",
        )

    declared_docx_source = _find_declared_origin_docx_for_output(normalized_output_path)
    body_map_path = _find_body_map_path_for_output(normalized_output_path)
    body_map_path = body_map_path.resolve() if body_map_path is not None else None
    if body_map_path is not None and not is_path_under(body_map_path, ROOT_DIR / "finish"):
        _round_snapshot_fail(
            "round_snapshot_body_map_path_invalid",
            "DOCX body map must stay under the finish workspace directory.",
        )
    body_map_bytes: bytes | None = None
    body_map_payload: dict[str, Any] | None = None
    body_map_current_texts: list[str] | None = None
    if body_map_path is not None and body_map_path.exists():
        body_map_bytes = _read_round_snapshot_artifact_bytes(
            body_map_path,
            missing_code="round_snapshot_body_map_missing",
            unreadable_code="round_snapshot_body_map_unreadable",
            label="DOCX body map",
        )
        body_map_payload = _parse_round_snapshot_json_object(
            body_map_bytes,
            code="round_snapshot_body_map_corrupt",
            label="DOCX body map",
        )
        raw_units = body_map_payload.get("units")
        if not isinstance(raw_units, list) or not raw_units:
            _round_snapshot_fail(
                "round_snapshot_body_map_corrupt",
                "DOCX body map has no captured editable units.",
            )
        body_map_current_texts = []
        for raw_unit in raw_units:
            if not isinstance(raw_unit, dict) or not isinstance(raw_unit.get("current_text"), str):
                _round_snapshot_fail(
                    "round_snapshot_body_map_corrupt",
                    "DOCX body map contains an invalid editable unit.",
                )
            body_map_current_texts.append(str(raw_unit["current_text"]))
        try:
            editable_unit_count = int(body_map_payload.get("editable_unit_count", len(raw_units)))
        except (TypeError, ValueError):
            editable_unit_count = -1
        if editable_unit_count != len(raw_units):
            _round_snapshot_fail(
                "round_snapshot_body_map_corrupt",
                "DOCX body map editable-unit count is inconsistent.",
            )
    elif declared_docx_source is not None:
        _round_snapshot_fail(
            "round_snapshot_body_map_missing",
            "DOCX round declares Word provenance, but its body map is missing.",
        )
    else:
        body_map_path = None

    output_sha256 = _round_snapshot_sha256(output_bytes)
    effective_text_sha256 = hashlib.sha256(effective_text.encode("utf-8")).hexdigest()
    review_revision = _round_snapshot_sha256(
        review_bytes if review_bytes is not None else ROUND_ARTIFACT_SNAPSHOT_NO_REVIEW_SENTINEL
    )
    body_map_sha256 = _round_snapshot_sha256(body_map_bytes) if body_map_bytes is not None else None
    manifest_sha256 = _round_snapshot_sha256(manifest_bytes) if manifest_bytes is not None else None
    content_revision = _round_snapshot_digest(
        [
            ("compare", compare_bytes),
            (
                "review",
                review_bytes if review_bytes is not None else ROUND_ARTIFACT_SNAPSHOT_NO_REVIEW_SENTINEL,
            ),
        ]
    )
    artifact_snapshot_digest = _round_snapshot_digest(
        [
            ("contentRevision", content_revision.encode("ascii")),
            ("output", output_bytes),
            ("bodyMap", body_map_bytes),
            ("manifest", manifest_bytes),
        ]
    )

    normalized_preview_limit = (
        max_preview_chars
        if isinstance(max_preview_chars, int) and not isinstance(max_preview_chars, bool) and max_preview_chars > 0
        else None
    )
    preview_text = effective_text
    truncated = normalized_preview_limit is not None and len(effective_text) > normalized_preview_limit
    if truncated and normalized_preview_limit is not None:
        preview_text = (
            effective_text[:normalized_preview_limit].rstrip()
            + "\n\n[预览已截断，导出文件可查看完整内容]"
        )

    canonical_output_path = str(normalized_output_path)
    response_compare = copy.deepcopy(compare_payload)
    response_compare["outputPath"] = canonical_output_path
    response_compare["compareRevision"] = compare_revision
    response_review = {
        "path": str(review_path),
        "outputPath": canonical_output_path,
        "docId": doc_id,
        "round": round_number,
        "decisions": decisions,
        "updatedAt": review_updated_at,
        "compareRevision": compare_revision,
        "currentCompareRevision": compare_revision,
        "reviewBaseCompareRevision": review_base_compare_revision,
        "reviewLinkReady": True,
        "reviewLinkStatus": review_link_status,
    }
    response: dict[str, Any] = {
        "version": ROUND_ARTIFACT_SNAPSHOT_VERSION,
        "outputPath": canonical_output_path,
        "docId": doc_id,
        "round": round_number,
        "materializationSource": "review_materialized_compare",
        "compare": response_compare,
        "review": response_review,
        "effectivePreview": {
            "path": canonical_output_path,
            "text": preview_text,
            "truncated": truncated,
            "totalChars": len(effective_text),
            "previewChars": len(preview_text),
        },
        "compareRevision": compare_revision,
        "reviewRevision": review_revision,
        "contentRevision": content_revision,
        "artifactSnapshotDigest": artifact_snapshot_digest,
        "compareSha256": compare_sha256,
        "reviewSha256": _round_snapshot_sha256(review_bytes) if review_bytes is not None else None,
        "effectiveTextSha256": effective_text_sha256,
        "outputSha256": output_sha256,
        "bodyMapSha256": body_map_sha256,
        "manifestSha256": manifest_sha256,
        "rawOutputMatchesEffective": (
            _round_snapshot_normalize_newlines(raw_output_text)
            == _round_snapshot_normalize_newlines(effective_text)
        ),
        "bodyMapMatchesEffective": (
            body_map_current_texts == paragraphs
            if body_map_current_texts is not None
            else None
        ),
    }
    if include_internal:
        response["_internal"] = {
            "effectiveText": effective_text,
            "effectiveParagraphs": list(paragraphs),
            "effectiveChunks": effective_chunks,
            "rawOutputText": raw_output_text,
            "comparePayload": copy.deepcopy(compare_payload),
            "reviewPayload": copy.deepcopy(review_payload),
            "bodyMapPayload": copy.deepcopy(body_map_payload),
            "bodyMapCurrentTexts": copy.deepcopy(body_map_current_texts),
            "manifestPayload": copy.deepcopy(manifest_payload),
            "compareBytes": compare_bytes,
            "reviewBytes": review_bytes,
            "outputBytes": output_bytes,
            "bodyMapBytes": body_map_bytes,
            "manifestBytes": manifest_bytes,
            "outputPath": canonical_output_path,
            "comparePath": str(compare_path),
            "reviewPath": str(review_path),
            "bodyMapPath": str(body_map_path or ""),
            "manifestPath": str(manifest_path or ""),
        }
    return response


def read_round_artifact_snapshot(
    output_path: str | Path,
    *,
    max_preview_chars: int | None = None,
    include_internal: bool = False,
    lock_timeout_seconds: float = ROUND_ARTIFACT_SNAPSHOT_LOCK_TIMEOUT_SECONDS,
) -> dict[str, Any]:
    """Capture one revision-aware compare/review/output/body-map/manifest view.

    Compare plus review decisions are the canonical content. Raw output and
    body-map text are captured as materialized-cache evidence and may legally
    lag immediately after a review-only save.
    """

    normalized_output_path = _resolve_api_path(
        output_path,
        allowed_roots=API_OUTPUT_ALLOWED_ROOTS,
        label="Output path",
    )
    try:
        normalized_timeout = min(30.0, max(0.0, float(lock_timeout_seconds)))
    except (TypeError, ValueError) as exc:
        raise ValueError("lock_timeout_seconds must be a non-negative number.") from exc
    lock = get_output_rerun_lock(normalized_output_path)
    if not lock.acquire(timeout=normalized_timeout):
        _round_snapshot_fail(
            "round_snapshot_busy",
            "Round artifacts are currently being updated; retry the snapshot shortly.",
            retryable=True,
            details={
                "outputPath": str(normalized_output_path),
                "retryAfterMs": 800,
            },
        )
    try:
        try:
            return _capture_round_artifact_snapshot_unlocked(
                normalized_output_path,
                max_preview_chars=max_preview_chars,
                include_internal=bool(include_internal),
            )
        except RoundArtifactSnapshotError:
            raise
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            raise RoundArtifactSnapshotError(
                "round_snapshot_capture_failed",
                "Round artifacts could not be captured as one coherent snapshot.",
                details={"errorType": type(exc).__name__},
            ) from exc
    finally:
        lock.release()


def _load_review_decisions_unlocked(normalized_output_path: Path) -> dict[str, Any]:
    decisions_path = _find_review_decisions_path_for_output(normalized_output_path)
    updated_at = ""
    if not decisions_path.exists():
        payload: dict[str, Any] = {}
    else:
        try:
            loaded_payload = json.loads(decisions_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            loaded_payload = {}
        payload = loaded_payload if isinstance(loaded_payload, dict) else {}
        updated_at = str(payload.get("updatedAt", "") or "").strip()
    decisions = payload.get("decisions")
    if not isinstance(decisions, dict):
        decisions = {}
    normalized_decisions = {
        str(chunk_id): _normalize_review_decision_value(decision)
        for chunk_id, decision in decisions.items()
    }
    compare_path = _find_compare_path_for_output(normalized_output_path)
    compare_payload = _load_compare_payload_for_output(normalized_output_path)
    compare_revision = (
        _compare_revision_for_review(compare_payload, compare_path)
        if isinstance(compare_payload, dict)
        else ""
    )
    review_base_compare_revision = str(
        payload.get("reviewBaseCompareRevision", payload.get("compareRevision", updated_at)) or ""
    ).strip()
    compare_review_updated_at = str((compare_payload or {}).get("reviewUpdatedAt", "") or "").strip()
    return {
        "path": str(decisions_path),
        "decisions": normalized_decisions,
        "updatedAt": updated_at,
        # Compatibility alias: this is the token a new write must compare with,
        # not the historical revision at which the sidecar was created.
        "compareRevision": compare_revision,
        "currentCompareRevision": compare_revision,
        "reviewBaseCompareRevision": review_base_compare_revision,
        "reviewLinkReady": bool(
            not updated_at
            or not compare_review_updated_at
            or compare_review_updated_at == updated_at
        ),
    }


def load_review_decisions(output_path: str) -> dict[str, Any]:
    normalized_output_path = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    with get_output_rerun_lock(normalized_output_path):
        return _load_review_decisions_unlocked(normalized_output_path)


def _save_review_decisions_unlocked(
    output_path: str,
    decisions: dict[str, Any],
    *,
    expected_compare_revision: str | None = None,
    require_compare_revision: bool = False,
) -> dict[str, Any]:
    normalized_output_path = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    decisions_path = _find_review_decisions_path_for_output(normalized_output_path)
    compare_payload = read_round_compare(str(normalized_output_path))
    compare_path = _find_compare_path_for_output(normalized_output_path)
    current_compare_revision = _compare_revision_for_review(compare_payload, compare_path)
    requested_revision = str(expected_compare_revision or "").strip()
    if require_compare_revision and not requested_revision:
        raise ReviewRevisionRequiredError("expectedCompareRevision is required.")
    if requested_revision and requested_revision != current_compare_revision:
        raise StaleReviewDecisionsError(current_compare_revision)
    current_review_state = _load_review_decisions_unlocked(normalized_output_path)
    if current_review_state.get("reviewLinkReady") is False:
        raise InconsistentReviewStateError(current_compare_revision)
    valid_chunk_ids = {
        str(chunk.get("chunkId"))
        for chunk in compare_payload.get("chunks", [])
        if isinstance(chunk, dict) and str(chunk.get("chunkId", "")).strip()
    }
    normalized_decisions = {
        str(chunk_id): _normalize_review_decision_value(decision)
        for chunk_id, decision in decisions.items()
        if str(chunk_id) in valid_chunk_ids
    }
    _validate_docx_custom_review_format_anchors(
        normalized_output_path,
        compare_payload,
        normalized_decisions,
    )
    # Validate the complete staged materialization before mutating timestamps,
    # compare metadata, or the linked review sidecar.
    _assert_document_release_payload(compare_payload, normalized_decisions)
    # This field is not part of the authoritative compare protocol.  Older or
    # concurrent producers may nevertheless have persisted a materialized
    # document-delta assessment for a previous review selection.  A review
    # decision can change the effective text without changing any raw
    # candidate, so carrying that assessment into the new compare revision
    # would expose stale ``passed`` evidence to non-release consumers.  The
    # release gate above always performs its own canonical fresh assessment;
    # discard the optional cache instead of widening the persisted protocol.
    compare_payload.pop("materializedSourceRelativeDocumentDelta", None)
    compare_chunks = compare_payload.get("chunks")
    if isinstance(compare_chunks, list):
        for chunk in compare_chunks:
            if not isinstance(chunk, dict):
                continue
            chunk_id = str(chunk.get("chunkId", "") or "")
            decision = normalized_decisions.get(chunk_id)
            decision_confirmed = bool(
                decision in {"source_confirmed", "rewrite_confirmed"}
                if isinstance(decision, str)
                else isinstance(decision, dict) and decision.get("confirmed") is True
            )
            if decision_confirmed:
                chunk.pop("rateAuditStrategyReviewRequired", None)
    updated_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    compare_payload["updatedAt"] = updated_at
    compare_payload["reviewUpdatedAt"] = updated_at
    review_payload = {
        "outputPath": str(normalized_output_path),
        "updatedAt": updated_at,
        "compareRevision": updated_at,
        "reviewBaseCompareRevision": updated_at,
        "decisions": normalized_decisions,
    }
    next_payloads = (
        (
            decisions_path,
            json.dumps(review_payload, ensure_ascii=False, indent=2).encode("utf-8"),
        ),
        (
            compare_path,
            json.dumps(compare_payload, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8"),
        ),
    )
    originals: dict[Path, bytes | None] = {}
    try:
        for path, _ in next_payloads:
            originals[path] = path.read_bytes() if path.exists() else None
        # The compare revision is the commit marker and is deliberately last.
        for path, payload in next_payloads:
            _replace_file_bytes_atomically(path, payload)
        disk_compare = json.loads(compare_path.read_text(encoding="utf-8"))
        disk_review = json.loads(decisions_path.read_text(encoding="utf-8"))
        if disk_compare.get("updatedAt") != updated_at or disk_compare.get("reviewUpdatedAt") != updated_at:
            raise ValueError("Review decision compare revision failed its post-commit check.")
        if disk_review.get("compareRevision") != updated_at or disk_review.get("decisions") != normalized_decisions:
            raise ValueError("Review decision sidecar failed its post-commit check.")
    except Exception:
        for path, original in reversed(list(originals.items())):
            try:
                if original is None:
                    path.unlink(missing_ok=True)
                else:
                    _replace_file_bytes_atomically(path, original)
            except OSError:
                pass
        raise
    return {
        "path": str(decisions_path),
        "decisions": normalized_decisions,
        "updatedAt": updated_at,
        "compareRevision": updated_at,
        "currentCompareRevision": updated_at,
        "reviewBaseCompareRevision": updated_at,
        "reviewLinkReady": True,
    }


def save_review_decisions(
    output_path: str,
    decisions: dict[str, Any],
    *,
    expected_compare_revision: str | None = None,
    require_compare_revision: bool = False,
) -> dict[str, Any]:
    normalized_output = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    with get_output_rerun_lock(normalized_output):
        return _save_review_decisions_unlocked(
            str(normalized_output),
            decisions,
            expected_compare_revision=expected_compare_revision,
            require_compare_revision=require_compare_revision,
        )


def _build_transform_from_model_config(model_config: dict[str, Any]) -> tuple[Callable[[str, str, int, str], str], str]:
    base_url = str(model_config.get("baseUrl", "")).strip()
    api_key = str(model_config.get("apiKey", "")).strip()
    model = str(model_config.get("model", "")).strip()
    api_type = str(model_config.get("apiType", "chat_completions")).strip()
    streaming = normalize_streaming(model_config.get("streaming", model_config.get("stream")))
    temperature = float(model_config.get("temperature", 0.7) or 0.7)
    request_timeout_seconds = _coerce_rewrite_timeout_seconds(
        model_config.get("requestTimeoutSeconds", DEFAULT_REQUEST_TIMEOUT_SECONDS)
    )
    configured_max_retries = _coerce_int_config(
        model_config.get("maxRetries", DEFAULT_MAX_RETRIES),
        default=DEFAULT_MAX_RETRIES,
        minimum=0,
        maximum=10,
    )
    max_retries = max(configured_max_retries, MIN_REWRITE_TRANSIENT_RETRIES)

    if not base_url or not api_key or not model:
        raise ValueError("Model configuration is incomplete.")

    rate_limiter = _build_provider_rate_limiter(model_config)

    def transform(_chunk_text: str, prompt_input: str, _round: int, _chunk_id: str) -> str:
        _wait_for_provider_cooldown(model_config)
        rate_limiter()
        try:
            output = llm_completion(
                prompt_input,
                model=model,
                api_key=api_key,
                base_url=base_url,
                api_type=api_type,
                temperature=temperature,
                timeout=request_timeout_seconds,
                max_retries=max_retries,
                retry_backoff_seconds=REWRITE_RETRY_BACKOFF_SECONDS,
                stream=streaming,
            )
        except LLMRequestError as exc:
            _register_provider_failure(model_config, exc)
            raise
        _register_provider_success(model_config)
        return output

    return transform, "online"


RERUN_FEEDBACK_TEMPLATES: dict[str, dict[str, Any]] = {
    "academic_register_drift": {
        "tag": "academic-register-repair",
        "advice": "删除模型新引入的聊天式、随意式词组，改成准确、克制的学术书面语；原文已有表达只在不改变含义时调整。",
        "instruction": "- Academic-register repair: replace newly introduced colloquial wording with precise written academic language while preserving every claim and qualifier.",
    },
    "citation_missing": {
        "tag": "citation-repair",
        "advice": "补回原文所有引用标记，引用必须跟随对应观点，不要移动到无关句子。",
        "instruction": "- Citation repair: preserve each citation marker from the input and place it near the same claim or sentence logic.",
    },
    "over_expanded": {
        "tag": "control-expansion",
        "advice": "控制长度，删除新增背景、泛化解释、额外评价和原文没有的结论。",
        "instruction": "- Length control: stay close to the source length and remove unsupported background, explanations, emphasis, or conclusions.",
    },
    "over_compressed": {
        "tag": "restore-detail",
        "advice": "补回被压缩掉的限定条件、因果关系、对比关系、实验细节和关键术语。",
        "instruction": "- Detail restoration: restore qualifiers, causal links, contrasts, technical terms, and key experimental details.",
    },
    "machine_like_expression": {
        "tag": "de-template-expression",
        "advice": "降低模板句式密度，减少套路连接词，避免整齐划一的总分总表达。",
        "instruction": "- De-template: vary sentence openings and rhythm; avoid generic transitions, boilerplate academic phrasing, and uniform paragraph cadence.",
    },
    "machine_style_drift": {
        "tag": "machine-style-drift-repair",
        "advice": "修复本次改写新引入的高风险模板句、机械连接或口语表达，不改变原文事实与逻辑关系。",
        "instruction": "- Style-drift repair: remove only high-confidence style problems newly introduced by this rewrite; preserve the source's legitimate academic constructions.",
    },
    "template_phrase_drift": {
        "tag": "avoid-new-template-phrase",
        "advice": "不要沿用或新增“在……背景下、具有重要意义、提供支持”等泛化模板句，改成该段落自己的具体表达。",
        "instruction": "- Template drift repair: remove newly introduced boilerplate phrases and replace them with concrete wording grounded in this paragraph.",
    },
    "connector_overuse": {
        "tag": "connector-trim",
        "advice": "只删减重复、公式化且不承担逻辑作用的过渡语；必要的因果、转折和条件标记必须保留。",
        "instruction": "- Connector trim: remove only redundant formulaic transitions; preserve every connector needed to express cause, contrast, condition, or scope.",
    },
    "template_phrase_density": {
        "tag": "phrase-diversify",
        "advice": "替换高频模板短语，不要使用平台化、套壳式、过度规整的论文腔。",
        "instruction": "- Phrase diversity: replace repeated template phrases with context-specific wording grounded in this paragraph.",
    },
    "generic_sentence_rhythm": {
        "tag": "rhythm-vary",
        "advice": "检查连续同构开句，只在完整语义边界做局部调整；不要强造短句。",
        "instruction": "- Rhythm repair: adjust genuinely repeated sentence frames at complete semantic boundaries; do not manufacture fragments or chase a length ratio.",
    },
    "uniform_sentence_rhythm": {
        "tag": "rhythm-vary",
        "advice": "检查连续同构开句，只在完整语义边界做局部调整；不要强造短句。",
        "instruction": "- Rhythm repair: adjust genuinely repeated sentence frames at complete semantic boundaries; do not manufacture fragments or chase a length ratio.",
    },
    "dimension_direction_not_effective": {
        "tag": "dimension-targeted-repair",
        "advice": "按本轮诊断做局部修复：节奏问题只处理重复句模，衔接问题只处理冗余过渡语；不得为了指标破坏语义。",
        "instruction": "- Dimension-targeted repair: fix the round's primary perturbation dimension specifically; see the dimension-targeted line above for which dimension.",
    },
    "structure_template_concentration": {
        "tag": "structure-diversify",
        "advice": "连续多句表层框架相似时，局部调整重复主语或开句即可；不要强塞被动句、长定语或额外从句。",
        "instruction": "- Repeated-frame repair: adjust repeated subjects or openings only where natural; do not force passive voice, stacked modifiers, or extra subordinate clauses to satisfy a classifier.",
    },
    "paragraph_length_symmetry": {
        "tag": "paragraph-length-vary",
        "advice": "相邻自然段长度过于整齐时，只在各段内部调整信息密度与句式节奏；不要合并、拆分、重排或删除自然段。",
        "instruction": "- Paragraph length symmetry: vary information density and sentence rhythm inside each paragraph so adjacent paragraph lengths are less uniform; do not merge, split, reorder, or drop natural paragraphs.",
    },
}


def _unique_strings(items: list[str]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for item in items:
        clean = " ".join(str(item or "").split())
        if clean and clean not in seen:
            seen.add(clean)
            unique.append(clean)
    return unique


def _collect_rerun_feedback_templates(flags: list[str], reasons: list[Any]) -> tuple[list[str], list[str], list[str]]:
    codes = list(flags)
    for reason in reasons:
        if isinstance(reason, dict):
            code = str(reason.get("code", "")).strip()
            if code:
                codes.append(code)
    tags: list[str] = []
    advice: list[str] = []
    instructions: list[str] = []
    for code in codes:
        template = RERUN_FEEDBACK_TEMPLATES.get(code)
        if not template:
            continue
        tags.append(str(template["tag"]))
        advice.append(str(template["advice"]))
        instructions.append(str(template["instruction"]))
    return _unique_strings(tags), _unique_strings(advice), _unique_strings(instructions)


def _format_rerun_evidence(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    try:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))[:500]
    except TypeError:
        return str(value)[:500]


def _truncate_prompt_block(text: str, limit: int = 6000) -> str:
    normalized = str(text or "").strip()
    if len(normalized) <= limit:
        return normalized
    half = max(1, limit // 2)
    return f"{normalized[:half]}\n...[TRUNCATED]...\n{normalized[-half:]}"


def _build_rerun_issue_cards(chunk: dict[str, Any]) -> list[str]:
    quality = chunk.get("quality") if isinstance(chunk.get("quality"), dict) else {}
    flags = quality.get("flags") if isinstance(quality, dict) else []
    reasons = quality.get("reviewReasons") if isinstance(quality, dict) else []
    cards: list[str] = []

    if isinstance(reasons, list):
        for reason in reasons[:8]:
            if not isinstance(reason, dict):
                continue
            code = str(reason.get("code", "review")).strip() or "review"
            level = str(reason.get("level", "medium")).strip() or "medium"
            message = str(reason.get("message", "")).strip()
            evidence = _format_rerun_evidence(reason.get("evidence"))
            if not message and not evidence:
                continue
            card = f"- [{code} / {level}] {message or '需要局部修复。'}"
            if evidence:
                card += f"\n  Evidence: {evidence}"
            cards.append(card)

    if isinstance(flags, list):
        known_codes = {card.split("]", 1)[0].lstrip("- [").split(" / ", 1)[0] for card in cards}
        for flag in flags:
            code = str(flag).strip()
            if code and code not in known_codes:
                cards.append(f"- [{code} / medium] 该块被标记为需审阅，请优先修复这个具体问题。")

    return _unique_strings(cards)


def _build_compact_round_brief(
    prompt_profile: str,
    round_number: int,
    prompt_sequence: object | None = None,
) -> list[str]:
    normalized_profile = normalize_prompt_profile(prompt_profile)
    normalized_prompt_sequence = normalize_prompt_sequence(normalized_profile, prompt_sequence)
    final_round = round_number >= get_max_rounds(normalized_profile, normalized_prompt_sequence)
    current_prompt_id = get_prompt_id_for_round(normalized_profile, round_number, normalized_prompt_sequence)
    if current_prompt_id == "prewrite":
        return [
            "- 当前轮次是润色改写：只做轻量自然化，不要大幅扩写。",
            "- 字数尽量贴近原文，段落角色、事实、术语和结论保持不变。",
        ]
    if current_prompt_id == "round2" or final_round:
        return [
            "- 当前轮次偏向最终自然化：降低机械论文腔，但不能新增观点或改变逻辑。",
            "- 避免总分总模板、泛化价值判断和过于整齐的连接词节奏。",
        ]
    return [
        "- 当前轮次是规范改写：表达可以更自然、更具解释性，但不能偏离原文。",
        "- 控制长度，优先修正句式与节奏，不要写成总结、建议或答疑。",
    ]


def _build_targeted_validation_retry_note(
    validation_error: str,
    *,
    input_text: str = "",
    failed_output: str = "",
    final_recovery: bool = False,
) -> str:
    repair_steps = _build_validation_repair_steps(validation_error)
    scope_repair = build_factual_scope_repair_guard(input_text, failed_output)
    sections = [
        "[TARGETED VALIDATION RETRY]",
        "- The previous targeted rerun output failed local hard validation.",
        "- Regenerate the same chunk and fix the validation problem directly.",
        "- Do not make broader changes just to avoid the error; keep the rewrite conservative.",
    ]
    if repair_steps:
        sections.append(repair_steps)
    if scope_repair:
        sections.append(scope_repair)
    if final_recovery:
        sections.append(
            "[FINAL SAFE RECOVERY]\n"
            "- This is the last bounded recovery attempt. Safety overrides every style objective.\n"
            "- If any repair could alter a term, fact, qualifier, relation, citation, number, or paragraph boundary, reproduce [INPUT TEXT] verbatim.\n"
            "- Do not explain the fallback and do not add a label; return the exact input body only."
        )
    sections.append(f"- Validation error: {validation_error.strip()}")
    return "\n".join(sections)

def _structure_direction_ineffective(direction: dict[str, Any] | None) -> bool:
    """True when sentence_structure dual-check reports structureDirection.effective=False."""
    if not isinstance(direction, dict):
        return False
    structure_direction = direction.get("structureDirection")
    if not isinstance(structure_direction, dict):
        return False
    return structure_direction.get("effective") is False


def _build_dimension_rerun_guidance(chunk: dict[str, Any]) -> tuple[str, str]:
    """Build conservative, dimension-specific rerun guidance."""
    quality = chunk.get("quality") if isinstance(chunk.get("quality"), dict) else {}
    flags = quality.get("advisoryFlags") if isinstance(quality, dict) else []
    if "dimension_direction_not_effective" not in (flags or []):
        return "", ""
    round_dimension = quality.get("roundDimension") if isinstance(quality, dict) else {}
    primary_metric = str(round_dimension.get("primaryMetric", "")) if isinstance(round_dimension, dict) else ""
    secondary_metric = str(round_dimension.get("secondaryMetric", "")) if isinstance(round_dimension, dict) else ""
    direction = quality.get("dimensionDirection") if isinstance(quality, dict) else {}
    # Structure dual-check under sentence_structure: inject structure-diversify
    # without passive / long-的 quotas when structureDirection failed.
    if (
        primary_metric in ("burstinessRatio", "structureConcentration")
        or secondary_metric == "structureConcentration"
    ) and _structure_direction_ineffective(direction if isinstance(direction, dict) else None):
        return (
            "- Dimension-targeted repair (structure-diversify / 本轮表层框架): "
            "连续多句表层主语与开句框架过于集中；请局部换用对象、条件或结果切入，"
            "只在完整语义单元上调整。不要强塞被动句、长“的”定语链或额外从句，"
            "也不要为分散分类器标签改写事实关系。",
            "本轮结构修复：打散重复表层框架与开句，禁止被动/长定语配额。",
        )
    if primary_metric == "burstinessRatio":
        return (
            "- Dimension-targeted repair (本轮句法与节奏): 检查连续重复的主语、开句或句界。"
            "只在完整分句边界做局部合并或拆分；不得插入无信息短句，不得强塞被动句、长定语或额外从句，"
            "也不要改动事实关系来追求句长统计值。",
            "本轮节奏修复：处理真实存在的重复句模，保持完整语义单元，不追逐最长/最短句比例。",
        )
    if primary_metric == "connectorDensity":
        return (
            "- Dimension-targeted repair (本轮衔接与终稿): 公式化过渡语仍然密集，或本轮新增了同类信号。"
            "只删除重复且不承担逻辑作用的部分；因果、转折、递进、条件和范围关系所需的连接词必须保留。",
            "本轮衔接修复：减少冗余公式化过渡语，同时保持每一处真实逻辑关系清楚。",
        )
    if primary_metric == "templateDensity":
        return (
            "- Dimension-targeted repair (本轮模板与空泛表达): 只处理诊断命中的模板句、泛化收束、"
            "空泛填充和不承载信息的套语。必须改成原文已有的具体对象、动作与关系；"
            "不得新增背景、意义、机制、数据、案例或结论，也不得顺带追逐句长和连接词指标。",
            "本轮模板修复：删除空泛套话并恢复具体陈述，不扩写、不拔高、不顺带重构其他维度。",
        )
    if primary_metric == "structureConcentration":
        return (
            "- Dimension-targeted repair (structure-diversify / 本轮表层框架): "
            "连续多句表层主语与开句框架过于集中；请局部换用对象、条件或结果切入。"
            "不要强塞被动句、长“的”定语链或额外从句。",
            "本轮结构修复：打散重复表层框架与开句，禁止被动/长定语配额。",
        )
    return "", ""


def _dimension_converged(input_text: str, output_text: str, round_dimension: dict[str, Any]) -> tuple[bool, dict[str, Any]]:
    """Re-score a rerun output on its OWN failed dimension (the closure check).

    Closes the detector-in-the-loop: a chunk flagged
    dimension_direction_not_effective is rerun with a dimension-targeted
    instruction; after the rerun passes hard validation we re-run the SAME
    dimension assessment on the new output. Converged = the rerun actually
    moved the dimension toward lower-AI (or the input was already satisfied).
    Returns (converged, dimensionDirection).

    Structure dual-check: even when ok/satisfied look true from rhythm alone,
    structureDirection.effective=False means the structure sub-signal has not
    converged and the loop must keep targeting structure-diversify.
    """
    if not round_dimension:
        return True, {}
    quality = _build_chunk_quality(input_text, output_text, round_dimension=round_dimension)
    direction = quality.get("dimensionDirection") if isinstance(quality, dict) else {}
    if not isinstance(direction, dict):
        return True, {}
    ok = bool(direction.get("ok", True))
    satisfied = bool(direction.get("satisfied", False))
    if _structure_direction_ineffective(direction):
        return False, direction
    converged = ok or satisfied
    return converged, direction


def _build_dimension_converge_retry_note(
    round_dimension: dict[str, Any],
    direction: dict[str, Any],
    attempt: int,
) -> str:
    """Strengthened instruction when a dimension-targeted rerun still didn't converge."""
    primary_metric = str(round_dimension.get("primaryMetric", "")) if isinstance(round_dimension, dict) else ""
    prefix = f"[DIMENSION-CONVERGE RETRY {attempt}] 上一次定向重跑仍有同类问题；本次只做必要的局部修复。"
    if _structure_direction_ineffective(direction if isinstance(direction, dict) else None):
        return (
            f"{prefix} 表层框架仍高度集中（structure-diversify）。"
            "请只调整重复主语与开句切入点；禁止为分散结构类型强塞被动句、长“的”定语或额外从句。"
        )
    if primary_metric == "burstinessRatio":
        return (
            f"{prefix} 请检查连续重复的句模，并只在完整分句边界做一次局部调整。"
            "不要插入短碎句，不要强制被动/长定语结构，也不要为了抬高统计值改变事实或逻辑。"
        )
    if primary_metric == "connectorDensity":
        return (
            f"{prefix} 公式化过渡语风险仍未出现可测量下降。"
            "仅删除重复且不承担关系表达的部分；必要的因果、转折、条件与范围连接词必须保留。"
        )
    if primary_metric == "templateDensity":
        return (
            f"{prefix} 模板句、泛化收束或空泛填充仍未出现同维度下降。"
            "仅删除不承载信息的部分，并用原文已有的对象、动作和关系表达；禁止补充背景、意义、机制或结论。"
        )
    if primary_metric == "structureConcentration":
        return (
            f"{prefix} 表层框架仍高度集中（structure-diversify）。"
            "请只调整重复主语与开句切入点；禁止被动/长定语配额。"
        )
    return prefix


def _build_rerun_strategy_note(chunk: dict[str, Any], user_feedback: str = "") -> tuple[str, list[str], list[str]]:
    quality = chunk.get("quality") if isinstance(chunk.get("quality"), dict) else {}
    flags = quality.get("flags") if isinstance(quality, dict) else []
    advisory_flags = quality.get("advisoryFlags") if isinstance(quality, dict) else []
    # advisory flags carry dimension/structure signals (e.g. structure_template_concentration,
    # template_phrase_drift) that should also drive targeted rerun feedback even when a
    # higher-severity hard flag is also present — include them so their templates fire.
    normalized_flags = _unique_strings(
        [str(flag) for flag in (list(flags) + list(advisory_flags))]
        if isinstance(flags, list) and isinstance(advisory_flags, list)
        else [str(flag) for flag in flags] if isinstance(flags, list) else []
    )
    advice = quality.get("rewriteAdvice") if isinstance(quality, dict) else []
    advice_items = [str(item).strip() for item in advice if str(item).strip()] if isinstance(advice, list) else []
    reasons = quality.get("reviewReasons") if isinstance(quality, dict) else []
    normalized_reasons = reasons if isinstance(reasons, list) else []
    template_tags, template_advice, template_instructions = _collect_rerun_feedback_templates(normalized_flags, normalized_reasons)
    advice_items = _unique_strings([*template_advice, *advice_items])
    # 维度定向重跑闭环：按本轮主攻维度生成定向指令，优先于通用模板
    dim_instruction, dim_advice = _build_dimension_rerun_guidance(chunk)
    if dim_advice:
        advice_items = _unique_strings([dim_advice, *advice_items])
    dim_tag = "dimension-targeted-repair" if dim_instruction else ""
    reason_messages: list[str] = []
    if isinstance(normalized_reasons, list):
        for reason in normalized_reasons[:5]:
            if not isinstance(reason, dict):
                continue
            message = str(reason.get("message", "")).strip()
            code = str(reason.get("code", "")).strip()
            if message:
                reason_messages.append(f"{code}: {message}" if code else message)
    expansion_ratio = quality.get("expansionRatio") if isinstance(quality, dict) else None
    instructions = [
        "[TARGETED RERUN NOTE]",
        "- The previous rewrite was marked as needing review. Fix the listed problems directly; do not repeat the same rewrite style.",
        "- Preserve every @@FYADR_*@@ placeholder exactly once and in its original order.",
        "- Do not change facts, citations, figure/table references, formula numbers, numbers, units, or conclusions.",
    ]
    if expansion_ratio is not None:
        instructions.append(f"- Previous expansion ratio: {expansion_ratio}. Use it to control length and information density.")
    if reason_messages:
        instructions.append("- Review diagnosis:")
        instructions.extend(f"  * {message}" for message in reason_messages)
    if advice_items:
        instructions.append("- Built-in targeted repair feedback:")
        instructions.extend(f"  * {item}" for item in advice_items[:6])
    clean_feedback = " ".join(str(user_feedback or "").split())
    if clean_feedback:
        instructions.append("- Human reviewer feedback for this rerun:")
        instructions.append(f"  * {clean_feedback[:800]}")
    strategy_tags: list[str] = [*template_tags]
    if dim_tag:
        strategy_tags.insert(0, dim_tag)

    if dim_instruction:
        instructions.append("- Dimension-targeted instructions (本轮主攻维度定向纠偏，优先于通用策略):")
        instructions.append(dim_instruction)

    if template_instructions:
        instructions.append("- Strategy-specific instructions:")
        instructions.extend(template_instructions[:8])

    strategy_tags = _unique_strings(strategy_tags)
    if not strategy_tags:
        strategy_tags.append("general-polish")
        instructions.append("- Improve sentence rhythm without increasing risk to structure or meaning.")

    instructions.append("- Return only the rewritten chunk body text.")
    return "\n".join(instructions), strategy_tags, advice_items


def _build_targeted_rerun_prompt_input(
    *,
    protected_input_text: str,
    previous_output_text: str,
    prompt_profile: str,
    prompt_sequence: object | None = None,
    round_number: int,
    chunk_id: str,
    rerun_note: str,
    issue_cards: list[str],
    style_card: str | None,
    validation_retry_note: str | None = None,
) -> str:
    compact_brief = _build_compact_round_brief(prompt_profile, round_number, prompt_sequence)
    sections = [
        "[TARGETED CHUNK RERUN]",
        (
            "This is a local repair prompt for exactly one chunk. "
            "The repair cards below have higher priority than generic rewrite habits from the full-round prompt."
        ),
        f"[ROUND {round_number}]",
        f"[CHUNK {chunk_id}]",
        "[NON-NEGOTIABLE OUTPUT RULES]\n"
        "- Return only the rewritten chunk body text; no labels, no explanations, no Markdown.\n"
        "- Preserve facts, claims, conclusions, citations, figure/table references, formula numbers, numbers, units, file paths, and technical terms.\n"
        "- Preserve every @@FYADR_*@@ placeholder exactly once and in its original order; do not translate, delete, duplicate, split, or rename placeholders.\n"
        "- Preserve paragraph count and paragraph role; do not merge or split paragraphs.\n"
        "- Do not copy the previous output's bad rhythm, boilerplate phrases, or unsupported additions.",
        "[COMPACT ROUND STYLE BRIEF]\n" + "\n".join(compact_brief),
    ]

    language_guard = build_language_guard(protected_input_text)
    if language_guard:
        sections.append(language_guard)
    paragraph_guard = build_paragraph_guard(protected_input_text)
    if paragraph_guard:
        sections.append(paragraph_guard)
    relation_guard = build_factual_relation_guard(protected_input_text)
    if relation_guard:
        sections.append(relation_guard)

    if issue_cards:
        sections.append("[REPAIR CARDS]\n" + "\n".join(issue_cards[:10]))
    if style_card:
        sections.append(style_card.strip())
    if rerun_note:
        sections.append(rerun_note.strip())
    if validation_retry_note:
        sections.append(validation_retry_note.strip())

    sections.extend(
        [
            "[TARGETED REPAIR DIRECTION]\n"
            "- Generate one conservative repair for this chunk; do not produce alternatives.\n"
            "- Follow the repair cards and human feedback directly while preserving structure and facts.",
            "[BAD PREVIOUS OUTPUT - DO NOT COPY ITS STYLE]\n" + _truncate_prompt_block(previous_output_text, 5000),
            "[INPUT TEXT]\n" + protected_input_text,
        ]
    )
    return "\n\n".join(section for section in sections if section)

def _append_compare_validation_event(compare_payload: dict[str, Any], event: dict[str, Any]) -> None:
    events = compare_payload.get("validationEvents")
    if not isinstance(events, list):
        events = []
        compare_payload["validationEvents"] = events
    events.append(event)


def _set_text_free_rerun_prompt_metadata(
    target_chunk: dict[str, Any],
    *,
    user_feedback: object,
) -> None:
    for field in (
        "rerunPromptNote",
        "rerunUserFeedback",
        "rerunIssueCards",
        "rerunStyleCard",
        "rerunGlobalStyleProfile",
    ):
        target_chunk.pop(field, None)
    normalized_feedback = " ".join(str(user_feedback or "").split())
    target_chunk["rerunPromptStored"] = False
    target_chunk["rerunUserFeedbackPresent"] = bool(normalized_feedback)
    target_chunk["rerunUserFeedbackCharCount"] = len(normalized_feedback)
    target_chunk["rerunUserFeedbackSha256"] = (
        _sha256_text(normalized_feedback) if normalized_feedback else ""
    )


def _record_targeted_failed_output(
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
    *,
    round_number: int,
    chunk_id: str,
    validation_attempt: int,
    error: str,
    output_text: str,
) -> None:
    failed_payload = _serialize_failed_output(output_text, error=error)
    event = {
        "event": "validation-retry",
        "round": round_number,
        "chunkId": chunk_id,
        "paragraphIndex": target_chunk.get("paragraphIndex"),
        "chunkIndex": target_chunk.get("chunkIndex"),
        **failed_payload,
        "attempt": validation_attempt,
    }
    _append_compare_validation_event(compare_payload, event)
    failed_attempts = _normalize_failed_attempts(target_chunk.get("failedAttempts"))
    if not failed_attempts:
        failed_attempts = _normalize_failed_attempts(target_chunk.get("rejectedCandidates"))
    failed_attempts.append(
        {
            **failed_payload,
            "attempt": validation_attempt,
        }
    )
    target_chunk.pop("rejectedCandidates", None)
    target_chunk["failedAttempts"] = failed_attempts[-4:]


def _replace_chunk_candidate_selection_event(
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
    selection_event: dict[str, object],
) -> None:
    """Keep the chunk reference and authoritative compare event byte-equal."""

    chunk_id = str(target_chunk.get("chunkId", "") or "")
    events = compare_payload.get("validationEvents")
    if not isinstance(events, list):
        events = []
    compare_payload["validationEvents"] = [
        event
        for event in events
        if not (
            isinstance(event, dict)
            and event.get("event") == "candidate-selection"
            and str(event.get("chunkId", "") or "") == chunk_id
        )
    ]
    compare_payload["validationEvents"].append(selection_event)
    target_chunk["candidateSelection"] = selection_event


def _run_bounded_app_candidate_selection(
    *,
    input_text: str,
    chunk_id: str,
    round_number: int,
    round_dimension: dict[str, object] | None,
    global_style_profile: dict[str, object] | None,
    generate: Callable[[int, str | None], str],
    validate: Callable[[str], str],
    record_hard_failure: Callable[[int, str, str], None],
) -> dict[str, Any]:
    """Use the main-round bounded selector for every app-level rerun path."""

    def evaluate(
        *,
        output_text: str,
        candidate_id: str,
        origin: str,
        attempt: int,
        hard_valid: bool,
        hard_validation_error: str = "",
    ) -> dict[str, object]:
        candidate = _evaluate_rewrite_candidate(
            input_text=input_text,
            output_text=output_text,
            candidate_id=candidate_id,
            origin=origin,
            attempt=attempt,
            hard_valid=hard_valid,
            hard_validation_error=hard_validation_error,
            round_dimension=round_dimension,
            global_style_profile=global_style_profile,
        )
        # Keep app-level injection/regression points aligned with the main
        # deterministic scorer without changing selector semantics.
        if hard_valid and str(output_text or "").strip():
            candidate["stylePenalty"] = round(_score_rewrite_output(input_text, output_text), 4)
        return candidate

    baseline = evaluate(
        output_text=input_text,
        candidate_id="baseline",
        origin="baseline",
        attempt=0,
        hard_valid=True,
    )
    candidates: list[dict[str, object]] = [baseline]
    retry_note: str | None = None
    conditional_retry_count = 0
    errors: list[str] = []

    for attempt in range(1, TARGETED_RERUN_VALIDATION_ATTEMPTS + 1):
        output_for_review = ""
        try:
            output_for_review = str(generate(attempt, retry_note) or "").strip()
            rewritten_output = validate(output_for_review)
            output_for_review = rewritten_output
            candidate = evaluate(
                output_text=rewritten_output,
                candidate_id=f"model-attempt-{attempt}",
                origin="model",
                attempt=attempt,
                hard_valid=True,
            )
            candidates.append(candidate)
            if (
                attempt < TARGETED_RERUN_SELECTION_ATTEMPTS
                and _candidate_requires_conditional_retry(
                    candidate,
                    baseline,
                    round_dimension=round_dimension,
                )
            ):
                conditional_retry_count += 1
                retry_note = _build_candidate_selection_retry_note(
                    candidate,
                    baseline,
                    round_dimension=round_dimension,
                )
                continue
            break
        except StaleRateAuditStrategyPlanError:
            raise
        except Exception as exc:
            safe_error = safe_public_error_message(exc)
            _guard_category, stable_issue_codes = _classify_failed_attempt_diagnostic(safe_error)
            errors.extend(
                code
                for code in stable_issue_codes
                if code not in errors
            )
            candidates.append(
                evaluate(
                    output_text=output_for_review,
                    candidate_id=f"model-attempt-{attempt}",
                    origin="model",
                    attempt=attempt,
                    hard_valid=False,
                    hard_validation_error=safe_error,
                )
            )
            record_hard_failure(attempt, safe_error, output_for_review)
            retry_note = _build_targeted_validation_retry_note(
                safe_error,
                input_text=input_text,
                failed_output=output_for_review,
                final_recovery=(attempt + 1) > TARGETED_RERUN_SELECTION_ATTEMPTS,
            )
            if (
                attempt >= TARGETED_RERUN_SELECTION_ATTEMPTS
                and any(
                    candidate.get("origin") == "model" and candidate.get("hardValid")
                    for candidate in candidates
                )
            ):
                break

    hard_valid_generated = [
        candidate
        for candidate in candidates
        if candidate.get("origin") == "model" and candidate.get("hardValid")
    ]
    if not hard_valid_generated:
        selection_event = _build_candidate_selection_event(
            chunk_id=chunk_id,
            round_number=round_number,
            candidates=candidates,
            selected=baseline,
            reason_codes=["all_model_candidates_failed_hard_validation", "baseline_preserved_but_rerun_failed"],
            conditional_retry_count=conditional_retry_count,
            decision="hard_failure_preserved_baseline",
            run_failed=True,
        )
        return {
            "ok": False,
            "text": input_text,
            "selectionEvent": selection_event,
            "errors": errors,
            "attemptCount": len(candidates) - 1,
            "candidates": candidates,
        }

    selected, reasons = _select_rewrite_candidate(candidates, round_dimension=round_dimension)
    selection_event = _build_candidate_selection_event(
        chunk_id=chunk_id,
        round_number=round_number,
        candidates=candidates,
        selected=selected,
        reason_codes=reasons,
        conditional_retry_count=conditional_retry_count,
    )
    return {
        "ok": True,
        "text": str(selected.get("_text", input_text) or input_text),
        "selectionEvent": selection_event,
        "errors": errors,
        "attemptCount": len(candidates) - 1,
        "candidates": candidates,
    }


def _bump_quality_summary_counter(compare_payload: dict[str, Any], key: str, chunk_id: str) -> None:
    summary = compare_payload.get("qualitySummary")
    if not isinstance(summary, dict):
        summary = {}
        compare_payload["qualitySummary"] = summary
    summary[key] = int(summary.get(key, 0) or 0) + 1
    ids_key = f"{key.removesuffix('Count')}ChunkIds"
    chunk_ids = summary.get(ids_key)
    if not isinstance(chunk_ids, list):
        chunk_ids = []
    if chunk_id not in chunk_ids:
        chunk_ids.append(chunk_id)
    summary[ids_key] = chunk_ids[:24]


def _rerun_compare_chunk_unlocked(output_path: str, chunk_id: str, model_config: dict[str, Any], user_feedback: str = "") -> dict[str, Any]:
    normalized_output_path = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    compare_path = _find_compare_path_for_output(normalized_output_path)
    compare_payload = read_round_compare(str(normalized_output_path))
    compare_payload_before_rerun = copy.deepcopy(compare_payload)
    chunks = compare_payload.get("chunks")
    if not isinstance(chunks, list):
        raise ValueError("Compare data is invalid: chunks missing.")

    target_chunk = None
    for chunk in chunks:
        if isinstance(chunk, dict) and str(chunk.get("chunkId", "")) == chunk_id:
            target_chunk = chunk
            break
    if target_chunk is None:
        raise ValueError(f"Chunk not found in compare data: {chunk_id}")
    if target_chunk.get("rateAuditStrategyReviewRequired") is True:
        raise ValueError(
            "该块是尚未确认的 RateAudit 定点候选；普通重跑已阻断。"
            "请先明确保留原文、确认候选或确认自定义文本，再发起普通重跑。"
        )

    prompt_profile = normalize_prompt_profile(compare_payload.get("promptProfile", model_config.get("promptProfile", DEFAULT_PROMPT_PROFILE)))
    prompt_sequence = normalize_prompt_sequence(prompt_profile, compare_payload.get("promptSequence", model_config.get("promptSequence")))
    round_number = int(compare_payload.get("round", 1) or 1)
    frozen_input_text = str(target_chunk.get("inputText", ""))
    if not frozen_input_text.strip():
        raise ValueError(f"Chunk {chunk_id} has empty input text.")

    preflight_docx_contract = _assert_docx_targeted_rerun_contract(
        normalized_output_path,
        compare_payload,
        target_chunk,
    )
    frozen_manifest_path = (
        _find_manifest_path_for_output(normalized_output_path)
        if preflight_docx_contract is not None
        else None
    )
    frozen_manifest_sha256 = (
        _sha256_file_bytes(frozen_manifest_path)
        if frozen_manifest_path is not None and frozen_manifest_path.exists()
        else ""
    )

    # inputText and the frozen manifest remain the immutable provenance used to
    # prove this chunk's location.  The rewrite baseline, however, must be the
    # text a fresh export would currently use after applying saved review
    # decisions; otherwise a confirmed rewrite/custom edit is silently ignored.
    materialized = _materialize_rate_audit_output(
        normalized_output_path,
        compare_payload=compare_payload,
    )
    raw_effective_chunks = materialized.get("chunks")
    effective_chunk = next(
        (
            item
            for item in raw_effective_chunks
            if isinstance(item, dict) and str(item.get("chunkId", "") or "") == chunk_id
        ),
        None,
    ) if isinstance(raw_effective_chunks, list) else None
    if not isinstance(effective_chunk, dict):
        raise ValueError(f"Chunk {chunk_id} has no fresh review-materialized text.")
    effective_document_texts = {
        str(item.get("chunkId", "") or "").strip(): str(item.get("text", "") or "")
        for item in raw_effective_chunks
        if isinstance(item, dict) and str(item.get("chunkId", "") or "").strip()
    } if isinstance(raw_effective_chunks, list) else {}
    effective_input_text = str(effective_chunk.get("text", "") or "")
    if not effective_input_text.strip():
        raise ValueError(f"Chunk {chunk_id} has empty fresh review-materialized text.")
    saved_review_decisions = (
        copy.deepcopy(materialized.get("decisions"))
        if isinstance(materialized.get("decisions"), dict)
        else {}
    )

    immutable_format_anchors = _resolve_docx_target_model_format_anchors(
        normalized_output_path,
        compare_payload,
        target_chunk,
        effective_input_text,
    )
    protected_chunk = protect_structure_tokens(
        effective_input_text,
        exact_anchors=immutable_format_anchors,
    )
    transform, mode = _build_transform_from_model_config(
        _resolve_round_model_config(model_config, prompt_profile, round_number)
    )
    previous_output_text = str(target_chunk.get("outputText", "")).strip()
    protected_previous_output_text = _mask_format_anchors_in_prompt_context(
        previous_output_text,
        protected_chunk,
        immutable_format_anchors,
    )
    effective_style_texts = [
        str(item.get("text", "") or "")
        for item in raw_effective_chunks
        if isinstance(item, dict) and str(item.get("text", "") or "").strip()
    ] if isinstance(raw_effective_chunks, list) else []
    global_style_profile = build_global_style_profile_from_texts(effective_style_texts or [effective_input_text])
    style_card = build_local_style_card(effective_input_text, global_style_profile)
    model_style_card = (
        _mask_format_anchors_in_prompt_context(
            style_card,
            protected_chunk,
            immutable_format_anchors,
        )
        if style_card
        else None
    )
    rerun_note, strategy_tags, advice_items = _build_rerun_strategy_note(target_chunk, user_feedback=user_feedback)
    issue_cards = _build_rerun_issue_cards(target_chunk)
    if style_card:
        strategy_tags = _unique_strings([*strategy_tags, "global-style-card"])
        advice_items = _unique_strings([*advice_items, "重跑会参考全文高频连接词、模板句和重复开头，避免继续放大全文层面的机械感。"])
    # Every targeted rerun uses the same bounded baseline + at most two model
    # candidates as the main round. Hard validation, readability, factual
    # guards, same-dimension direction and net style gain are evaluated by the
    # shared provider-independent selector.
    round_dimension = resolve_round_dimension(prompt_profile, round_number, prompt_sequence)
    chunk_quality_flags = (
        target_chunk.get("quality", {}).get("advisoryFlags")
        if isinstance(target_chunk.get("quality"), dict) else None
    )
    dimension_closure_enabled = bool(
        round_dimension
        and "dimension_direction_not_effective" in (chunk_quality_flags or [])
    )
    last_candidate_output = {"text": ""}

    def generate_candidate(_attempt: int, retry_note: str | None) -> str:
        last_candidate_output["text"] = ""
        prompt_input = _build_targeted_rerun_prompt_input(
            protected_input_text=protected_chunk.text,
            previous_output_text=protected_previous_output_text,
            prompt_profile=prompt_profile,
            prompt_sequence=prompt_sequence,
            round_number=round_number,
            chunk_id=chunk_id,
            rerun_note=rerun_note,
            issue_cards=issue_cards,
            style_card=model_style_card,
            validation_retry_note=retry_note,
        )
        return transform(protected_chunk.text, prompt_input, round_number, chunk_id)

    def validate_candidate(raw_output: str) -> str:
        protected_output = normalize_chunk_output(protected_chunk.text, raw_output)
        validate_structure_placeholders(protected_output, protected_chunk.tokens, chunk_id)
        rewritten_output = restore_structure_tokens(protected_output, protected_chunk.tokens)
        last_candidate_output["text"] = rewritten_output
        validate_immutable_text_anchors(
            effective_input_text,
            rewritten_output,
            immutable_format_anchors,
            chunk_id,
        )
        validate_chunk_output(effective_input_text, rewritten_output, chunk_id)
        return rewritten_output

    def record_hard_failure(attempt: int, error_text: str, output_text: str) -> None:
        _record_targeted_failed_output(
            compare_payload,
            target_chunk,
            round_number=round_number,
            chunk_id=chunk_id,
            validation_attempt=attempt,
            error=error_text,
            output_text=last_candidate_output["text"] or output_text,
        )

    selection_result = _run_bounded_app_candidate_selection(
        input_text=effective_input_text,
        chunk_id=chunk_id,
        round_number=round_number,
        round_dimension=round_dimension,
        global_style_profile=global_style_profile,
        generate=generate_candidate,
        validate=validate_candidate,
        record_hard_failure=record_hard_failure,
    )
    selection_event = selection_result["selectionEvent"]
    target_chunk["candidateBaselineText"] = effective_input_text.strip()
    _replace_chunk_candidate_selection_event(
        compare_payload,
        target_chunk,
        selection_event,
    )
    _refresh_source_pattern_profile_registry(
        compare_payload,
        global_style_profile=global_style_profile,
    )
    attempts_used = int(selection_result.get("attemptCount", 0) or 0)
    output_errors = list(selection_result.get("errors") or [])
    dimension_converge_directions = [
        dict(direction)
        for candidate in selection_result.get("candidates", [])
        if isinstance(candidate, dict)
        and candidate.get("origin") == "model"
        and isinstance((direction := candidate.get("sameDimensionDirection")), dict)
    ]
    target_chunk["rerunDimensionConverged"] = bool(selection_event.get("publishedRewrite"))
    target_chunk["rerunDimensionConvergeDirections"] = dimension_converge_directions[-3:]

    if not bool(selection_result.get("ok")):
        # Hard failure records diagnostics only. The previous compare candidate
        # and its linked review decision remain byte-for-byte authoritative.
        target_chunk["outputText"] = previous_output_text
        target_chunk["outputCharCount"] = len(previous_output_text)
        target_chunk["rerunAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
        target_chunk["rerunMode"] = mode
        target_chunk["rerunPromptMode"] = "targeted-short-prompt"
        target_chunk["rerunStrategy"] = strategy_tags
        target_chunk["rerunAdvice"] = _unique_strings([
            "定向重跑候选全部未通过硬门禁；上一版保持不变，本次不计为成功改写。",
            *advice_items,
        ])
        target_chunk["rerunAttemptCount"] = attempts_used
        _set_text_free_rerun_prompt_metadata(
            target_chunk,
            user_feedback=user_feedback,
        )
        target_chunk.pop("rerunStatus", None)
        target_chunk.pop("rerunFallbackMode", None)
        target_chunk.pop("rerunFallbackError", None)
        compare_payload["updatedAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
        _write_export_evidence_json_atomically(compare_path, compare_payload)
        failure_detail = "; ".join(output_errors[-3:]) if output_errors else "no hard-valid candidate"
        raise ValueError(
            f"Chunk {chunk_id} failed hard validation after {attempts_used} targeted rerun attempts: {failure_detail}"
        )

    _arbitrate_app_candidate_document_delta(
        compare_payload=compare_payload,
        target_chunk=target_chunk,
        selection_result=selection_result,
        candidate_baseline_text=effective_input_text,
        effective_document_texts=effective_document_texts,
        round_number=round_number,
    )
    selection_event = selection_result["selectionEvent"]
    if selection_event.get("publishedRewrite") is not True:
        original_chunks = compare_payload_before_rerun.get("chunks")
        original_target_chunk = next(
            (
                item
                for item in original_chunks
                if isinstance(item, dict)
                and str(item.get("chunkId", "") or "") == chunk_id
            ),
            None,
        ) if isinstance(original_chunks, list) else None
        if not isinstance(original_target_chunk, dict):
            raise ValueError("Targeted rerun lost its pre-attempt authoritative chunk.")
        # No model candidate beat the review-materialized baseline.  This is a
        # soft no-op: retaining the old output + decision avoids turning a
        # confirmed custom value D into output=D/custom=D, which the release
        # gate correctly treats as self-disguised candidate text.
        return {
            "chunk": original_target_chunk,
            "compare": compare_payload_before_rerun,
            "outputPath": str(normalized_output_path),
            "comparePath": str(compare_path),
            "reviewDecisions": saved_review_decisions,
            "updatedAt": str(compare_payload_before_rerun.get("updatedAt", "") or ""),
            "preservedExisting": True,
            "candidateSelectionAttempt": selection_event,
        }
    target_chunk["rerunDimensionConverged"] = bool(
        selection_event.get("publishedRewrite")
    )
    output_text = str(selection_result.get("text", effective_input_text) or effective_input_text)
    selected_candidate_id = str(selection_event.get("selectedCandidateId", "") or "")
    selected_candidate = next(
        (
            candidate
            for candidate in selection_result.get("candidates", [])
            if isinstance(candidate, dict)
            and str(candidate.get("candidateId", "") or "") == selected_candidate_id
        ),
        {},
    )
    selected_score = float(selected_candidate.get("stylePenalty", 0.0) or 0.0)

    target_chunk["outputText"] = output_text
    target_chunk["outputCharCount"] = len(output_text)
    target_chunk["quality"] = _apply_candidate_selection_quality(
        _build_chunk_quality(
            effective_input_text,
            output_text,
            round_dimension=round_dimension,
        ),
        selection_event,
    )
    target_chunk.pop("rerunStatus", None)
    target_chunk.pop("rerunNonConvergedReason", None)
    target_chunk.pop("rerunFallbackMode", None)
    target_chunk.pop("rerunFallbackError", None)
    target_chunk["rerunMode"] = mode
    target_chunk["rerunPromptMode"] = "targeted-short-prompt"
    target_chunk["rerunStrategy"] = strategy_tags
    target_chunk["rerunAdvice"] = _unique_strings([
        *(
            []
            if selection_event.get("publishedRewrite") is True
            else ["候选没有可测净收益或出现安全/可读性回归，已保留上一版。"]
        ),
        *advice_items,
    ])
    target_chunk["rerunAttemptCount"] = attempts_used
    target_chunk["rerunSelectedScore"] = round(selected_score, 3)
    _set_text_free_rerun_prompt_metadata(
        target_chunk,
        user_feedback=user_feedback,
    )

    # A successful rerun publishes a new candidate, so a saved decision for
    # the old target candidate is no longer valid.  Decisions for every other
    # chunk remain authoritative and must continue to materialize output/body
    # map content.  Failure paths above never reach this mutation.
    remaining_review_decisions = copy.deepcopy(saved_review_decisions)
    if selection_event.get("publishedRewrite") is True:
        remaining_review_decisions.pop(chunk_id, None)
    updated_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    if selection_event.get("publishedRewrite") is True:
        superseded_rate_audit_evidence = {
            key: target_chunk.pop(key)
            for key in list(target_chunk)
            if key.startswith("rateAuditStrategy")
        }
        if superseded_rate_audit_evidence:
            target_chunk["supersededRateAuditStrategyEvidence"] = {
                "supersededAt": updated_at,
                "reason": "legacy_rerun_published_new_candidate",
                **superseded_rate_audit_evidence,
            }
    # Persist the server-authoritative safe default used by this same commit so
    # clients never claim that a new rewrite was adopted when hard-risk flags
    # caused output/body-map materialization to retain source text.
    target_chunk["rerunDefaultDecision"] = _default_export_decision_for_chunk(target_chunk)
    target_chunk["rerunAt"] = updated_at
    compare_payload["updatedAt"] = updated_at
    compare_payload["reviewUpdatedAt"] = updated_at

    # Validate the exact staged compare + retained review decisions before any
    # output, review sidecar, body map, or compare revision is published.
    _assert_document_release_payload(
        compare_payload,
        remaining_review_decisions,
    )

    fresh_docx_contract = _assert_docx_targeted_rerun_contract(
        normalized_output_path,
        compare_payload,
        target_chunk,
    )
    if preflight_docx_contract is not None:
        if fresh_docx_contract is None:
            raise ValueError("DOCX targeted rerun contract disappeared before publication.")
        for key in ("sourceSha256", "scopeDigest", "formatDigest"):
            if str(fresh_docx_contract.get(key, "") or "") != str(preflight_docx_contract.get(key, "") or ""):
                raise ValueError(f"DOCX targeted rerun contract changed before publication: {key}.")
        if (
            frozen_manifest_path is None
            or not frozen_manifest_sha256
            or not frozen_manifest_path.exists()
            or _sha256_file_bytes(frozen_manifest_path) != frozen_manifest_sha256
        ):
            raise ValueError("DOCX targeted rerun manifest changed before publication.")

    _commit_legacy_rerun_stage(
        output_path=normalized_output_path,
        staged_compare=compare_payload,
        review_decisions=remaining_review_decisions,
        updated_at=updated_at,
        expected_docx_contract=preflight_docx_contract,
    )

    return {
        "chunk": target_chunk,
        "compare": compare_payload,
        "outputPath": str(normalized_output_path),
        "comparePath": str(compare_path),
        "reviewDecisions": remaining_review_decisions,
        "updatedAt": updated_at,
    }


def rerun_compare_chunk(output_path: str, chunk_id: str, model_config: dict[str, Any], user_feedback: str = "") -> dict[str, Any]:
    """Serialize legacy single/batch mutations with RateAudit strategy tasks."""

    normalized_output = _resolve_api_path(output_path, allowed_roots=API_OUTPUT_ALLOWED_ROOTS, label="Output path")
    with get_output_rerun_lock(normalized_output):
        return _rerun_compare_chunk_unlocked(
            str(normalized_output),
            chunk_id,
            model_config,
            user_feedback=user_feedback,
        )


def _rate_audit_strategy_prompt_round(prompt_id: str) -> int:
    if prompt_id == "round1":
        return 1
    if prompt_id == "round2":
        return 2
    if prompt_id == "template-repair":
        return 3
    raise ValueError(f"Prompt {prompt_id} is not a registered RateAudit repair prompt.")


def _assert_rate_audit_strategy_model_contract(
    *,
    source_path: Path,
    output_path: Path,
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
    expected_binding: dict[str, Any],
    recommended_prompt_id: str,
) -> str:
    """Fresh-check source/DOCX/prompt evidence immediately before model use."""

    mismatch_codes: list[str] = []
    try:
        current_source_sha256 = _sha256_file_bytes(source_path)
    except OSError:
        current_source_sha256 = ""
    if current_source_sha256 != str(expected_binding.get("sourceSha256", "") or ""):
        mismatch_codes.append("source_sha256_mismatch")

    if source_path.suffix.lower() == ".docx":
        body_map_path = _find_body_map_path_for_output(output_path)
        try:
            current_body_map_sha256 = _sha256_file_bytes(body_map_path) if body_map_path is not None else ""
        except OSError:
            current_body_map_sha256 = ""
        if current_body_map_sha256 != str(expected_binding.get("bodyMapSha256", "") or ""):
            mismatch_codes.append("body_map_sha256_mismatch")
        raw_manifest_path = str(compare_payload.get("manifestPath", "") or "").strip()
        manifest_path = normalize_path(Path(raw_manifest_path)) if raw_manifest_path else _find_manifest_path_for_output(output_path)
        try:
            current_manifest_sha256 = _sha256_file_bytes(manifest_path) if manifest_path is not None else ""
        except OSError:
            current_manifest_sha256 = ""
        if current_manifest_sha256 != str(expected_binding.get("manifestSha256", "") or ""):
            mismatch_codes.append("manifest_sha256_mismatch")
        if mismatch_codes:
            raise StaleRateAuditStrategyPlanError(mismatch_codes)

    try:
        current_contract = _assert_docx_targeted_rerun_contract(output_path, compare_payload, target_chunk)
    except ValueError as exc:
        raise StaleRateAuditStrategyPlanError(["docx_target_contract_mismatch"]) from exc
    if source_path.suffix.lower() == ".docx":
        if current_contract is None:
            mismatch_codes.append("docx_contract_missing")
        else:
            if str(current_contract.get("sourceSha256", "") or "") != str(expected_binding.get("sourceSha256", "") or ""):
                mismatch_codes.append("source_sha256_mismatch")
            if str(current_contract.get("scopeDigest", "") or "") != str(expected_binding.get("scopeDigest", "") or ""):
                mismatch_codes.append("scope_digest_mismatch")
            if str(current_contract.get("formatDigest", "") or "") != str(expected_binding.get("formatDigest", "") or ""):
                mismatch_codes.append("format_digest_mismatch")
            if not bool(current_contract.get("ready")):
                mismatch_codes.append("content_contract_not_ready")

    try:
        prompt_path = resolve_prompt_path(recommended_prompt_id)
        current_prompt_sha256 = _sha256_file_bytes(prompt_path)
        prompt_text = prompt_path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError, ValueError) as exc:
        raise StaleRateAuditStrategyPlanError(["prompt_unavailable"]) from exc
    if current_prompt_sha256 != str(expected_binding.get("promptSha256", "") or ""):
        mismatch_codes.append("prompt_sha256_mismatch")
    if mismatch_codes:
        raise StaleRateAuditStrategyPlanError(mismatch_codes)
    return prompt_text


def _build_rate_audit_strategy_prompt_input(
    *,
    prompt_text: str,
    protected_input_text: str,
    previous_output_text: str,
    chunk_id: str,
    dimension_definition: dict[str, Any],
    prompt_id: str,
    retry_note: str | None,
) -> str:
    round_number = _rate_audit_strategy_prompt_round(prompt_id)
    base_prompt = build_prompt_input(
        prompt_text,
        protected_input_text,
        round_number,
        chunk_id,
        retry_note=retry_note,
    )
    return "\n\n".join(
        section
        for section in (
            "[RATEAUDIT TARGETED STRATEGY]",
            (
                f"[BOUND DIMENSION] {dimension_definition.get('dimensionId', '')} / "
                f"{dimension_definition.get('evaluatorDimensionId', '')} / "
                f"{dimension_definition.get('primaryMetric', '')}"
            ),
            (
                "Repair only this bound dimension. Preserve facts, citations, numbers, units, "
                "technical terms, paragraph role, and every @@FYADR_*@@ placeholder. "
                "Return only the rewritten chunk body."
            ),
            (
                "[PREVIOUS ACCEPTED/DRAFT TEXT - CONTEXT ONLY]\n"
                + _truncate_prompt_block(previous_output_text, 4000)
                if previous_output_text.strip()
                else ""
            ),
            base_prompt,
        )
        if section
    )


def _mark_rate_audit_strategy_failure(
    *,
    compare_payload: dict[str, Any],
    target_chunk: dict[str, Any],
    binding: dict[str, Any],
    dimension_definition: dict[str, Any],
    attempts_used: int,
    directions: list[dict[str, Any]],
    reason: str,
) -> None:
    target_chunk["rerunAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    target_chunk["rerunPromptMode"] = "rate-audit-strategy-prompt"
    target_chunk["rerunStrategy"] = ["rate-audit", str(dimension_definition.get("dimensionId", "") or "")]
    target_chunk["rerunAttemptCount"] = attempts_used
    target_chunk["rerunDimensionConverged"] = False
    target_chunk["rerunDimensionConvergeDirections"] = directions[-3:]
    target_chunk["rerunStatus"] = "non_converged"
    target_chunk["rerunNonConvergedReason"] = reason
    target_chunk["rateAuditStrategyPlanDigest"] = str(binding.get("planDigest", "") or "")
    target_chunk["rateAuditStrategyPromptId"] = str(binding.get("recommendedPromptId", "") or "")
    target_chunk["rateAuditStrategyEvaluatorDimensionId"] = str(dimension_definition.get("evaluatorDimensionId", "") or "")
    target_chunk["rateAuditStrategyInputSource"] = "review_materialized_compare"
    compare_payload["updatedAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _advance_staged_rate_audit_attempt_ledger(
    *,
    compare_payload: dict[str, Any],
    binding: dict[str, Any],
    chunk_id: str,
    effective_input_text: str,
    attempt_count: int,
    status: str,
) -> None:
    """Bind a preserved-baseline attempt to this exact text generation."""

    ledger = compare_payload.get("rateAuditStrategyAttempts")
    if not isinstance(ledger, dict):
        ledger = {}
        compare_payload["rateAuditStrategyAttempts"] = ledger
    dimension_id = str(binding.get("dimensionId", "") or "")
    prompt_id = str(binding.get("recommendedPromptId", "") or "")
    effective_sha256 = hashlib.sha256(effective_input_text.encode("utf-8")).hexdigest()
    ledger_key = f"{dimension_id}:{chunk_id}"
    previous = ledger.get(ledger_key)
    previous_count = 0
    if (
        isinstance(previous, dict)
        and str(previous.get("recommendedPromptId", "") or "") == prompt_id
        and str(previous.get("effectiveTextSha256", "") or "") == effective_sha256
    ):
        previous_count = max(0, int(previous.get("attemptCount", 0) or 0))
    ledger[ledger_key] = {
        "version": 1,
        "dimensionId": dimension_id,
        "chunkId": chunk_id,
        "recommendedPromptId": prompt_id,
        "effectiveTextSha256": effective_sha256,
        "attemptCount": previous_count + max(0, int(attempt_count or 0)),
        "status": str(status or "preserved_baseline"),
        "lastPlanDigest": str(binding.get("planDigest", "") or ""),
        "updatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
    }


def _rerun_rate_audit_strategy_chunk_unlocked(
    *,
    source_path: Path,
    output_path: Path,
    compare_payload: dict[str, Any],
    chunk_id: str,
    effective_input_text: str,
    effective_document_texts: dict[str, str],
    dimension_id: str,
    recommended_prompt_id: str,
    expected_binding: dict[str, Any],
    model_config: dict[str, Any],
    global_style_profile: dict[str, object],
) -> dict[str, Any]:
    """Execute one bound repair without delegating to the legacy rerun path."""

    compare_path = _find_compare_path_for_output(output_path)
    chunks = compare_payload.get("chunks")
    if not isinstance(chunks, list):
        raise StaleRateAuditStrategyPlanError(["compare_invalid"])
    target_chunk = next(
        (
            item
            for item in chunks
            if isinstance(item, dict) and str(item.get("chunkId", "") or "").strip() == chunk_id
        ),
        None,
    )
    if target_chunk is None:
        raise StaleRateAuditStrategyPlanError(["target_chunks_mismatch"])
    compare_payload_before_attempt = copy.deepcopy(compare_payload)
    if not effective_input_text.strip():
        raise StaleRateAuditStrategyPlanError(["effective_target_empty"])

    dimension_definition = get_rate_audit_dimension_definition(dimension_id)
    if (
        not bool(dimension_definition.get("canExecute"))
        or str(dimension_definition.get("targetScope", "") or "") != "chunk"
        or str(dimension_definition.get("repairPromptId", "") or "") != recommended_prompt_id
    ):
        raise StaleRateAuditStrategyPlanError(["dimension_registry_mismatch"])
    evaluator = dict(ROUND_PERTURBATION_DIMENSIONS.get(recommended_prompt_id, {}))
    if (
        not evaluator
        or str(evaluator.get("id", "") or "") != str(dimension_definition.get("evaluatorDimensionId", "") or "")
        or str(evaluator.get("primaryMetric", "") or "") != str(dimension_definition.get("primaryMetric", "") or "")
    ):
        raise StaleRateAuditStrategyPlanError(["dimension_evaluator_mismatch"])

    # This first check deliberately precedes transform construction. A missing
    # or tampered DOCX body-map therefore produces zero model constructors and
    # zero model calls.
    prompt_text = _assert_rate_audit_strategy_model_contract(
        source_path=source_path,
        output_path=output_path,
        compare_payload=compare_payload,
        target_chunk=target_chunk,
        expected_binding=expected_binding,
        recommended_prompt_id=recommended_prompt_id,
    )
    try:
        immutable_format_anchors = _resolve_docx_target_model_format_anchors(
            output_path,
            compare_payload,
            target_chunk,
            effective_input_text,
        )
        protected_chunk = protect_structure_tokens(
            effective_input_text,
            exact_anchors=immutable_format_anchors,
        )
    except ValueError as exc:
        raise StaleRateAuditStrategyPlanError(
            ["docx_format_anchor_binding_invalid"],
            "RateAudit strategy lost its frozen DOCX format-anchor binding.",
        ) from exc
    transform, mode = _build_transform_from_model_config(model_config)
    previous_output_text = str(target_chunk.get("outputText", "") or "")
    protected_previous_output_text = _mask_format_anchors_in_prompt_context(
        previous_output_text,
        protected_chunk,
        immutable_format_anchors,
    )
    previous_output_literal = previous_output_text
    round_number = int(compare_payload.get("round", 0) or 0)
    last_candidate_output = {"text": ""}

    def generate_candidate(_attempt: int, retry_note: str | None) -> str:
        last_candidate_output["text"] = ""
        fresh_prompt_text = _assert_rate_audit_strategy_model_contract(
            source_path=source_path,
            output_path=output_path,
            compare_payload=compare_payload,
            target_chunk=target_chunk,
            expected_binding=expected_binding,
            recommended_prompt_id=recommended_prompt_id,
        )
        prompt_input = _build_rate_audit_strategy_prompt_input(
            prompt_text=fresh_prompt_text,
            protected_input_text=protected_chunk.text,
            previous_output_text=protected_previous_output_text,
            chunk_id=chunk_id,
            dimension_definition=dimension_definition,
            prompt_id=recommended_prompt_id,
            retry_note=retry_note,
        )
        return transform(
            protected_chunk.text,
            prompt_input,
            _rate_audit_strategy_prompt_round(recommended_prompt_id),
            chunk_id,
        )

    def validate_candidate(raw_output: str) -> str:
        protected_output = normalize_chunk_output(protected_chunk.text, raw_output)
        validate_structure_placeholders(protected_output, protected_chunk.tokens, chunk_id)
        rewritten_output = restore_structure_tokens(protected_output, protected_chunk.tokens)
        last_candidate_output["text"] = rewritten_output
        validate_immutable_text_anchors(
            effective_input_text,
            rewritten_output,
            immutable_format_anchors,
            chunk_id,
        )
        validate_chunk_output(effective_input_text, rewritten_output, chunk_id)
        return rewritten_output

    def record_hard_failure(attempt: int, error_text: str, output_text: str) -> None:
        _record_targeted_failed_output(
            compare_payload,
            target_chunk,
            round_number=round_number,
            chunk_id=chunk_id,
            validation_attempt=attempt,
            error=error_text,
            output_text=last_candidate_output["text"] or output_text,
        )

    selection_result = _run_bounded_app_candidate_selection(
        input_text=effective_input_text,
        chunk_id=chunk_id,
        round_number=round_number,
        round_dimension=evaluator,
        global_style_profile=global_style_profile,
        generate=generate_candidate,
        validate=validate_candidate,
        record_hard_failure=record_hard_failure,
    )
    selection_event = selection_result["selectionEvent"]
    target_chunk["candidateBaselineText"] = effective_input_text.strip()
    _replace_chunk_candidate_selection_event(compare_payload, target_chunk, selection_event)
    _refresh_source_pattern_profile_registry(
        compare_payload,
        global_style_profile=global_style_profile,
    )
    attempts_used = int(selection_result.get("attemptCount", 0) or 0)
    directions = [
        dict(direction)
        for candidate in selection_result.get("candidates", [])
        if isinstance(candidate, dict)
        and candidate.get("origin") == "model"
        and isinstance((direction := candidate.get("sameDimensionDirection")), dict)
    ]

    if not bool(selection_result.get("ok")):
        target_chunk["outputText"] = previous_output_literal
        _mark_rate_audit_strategy_failure(
            compare_payload=compare_payload,
            target_chunk=target_chunk,
            binding=expected_binding,
            dimension_definition=dimension_definition,
            attempts_used=attempts_used,
            directions=directions,
            reason="hard_validation_attempt_limit",
        )
        detail = "; ".join(list(selection_result.get("errors") or [])[-3:]) or "no hard-valid candidate"
        raise ValueError(
            f"Chunk {chunk_id} failed the RateAudit hard gate after {attempts_used} attempts; "
            f"previous output was preserved: {detail}"
        )

    _arbitrate_app_candidate_document_delta(
        compare_payload=compare_payload,
        target_chunk=target_chunk,
        selection_result=selection_result,
        candidate_baseline_text=effective_input_text,
        effective_document_texts=effective_document_texts,
        round_number=round_number,
    )
    selection_event = selection_result["selectionEvent"]
    if selection_event.get("publishedRewrite") is not True:
        compare_payload.clear()
        compare_payload.update(compare_payload_before_attempt)
        restored_chunks = compare_payload.get("chunks")
        restored_target = next(
            (
                item
                for item in restored_chunks
                if isinstance(item, dict)
                and str(item.get("chunkId", "") or "") == chunk_id
            ),
            None,
        ) if isinstance(restored_chunks, list) else None
        if not isinstance(restored_target, dict):
            raise ValueError("RateAudit preserved-baseline restore lost its target chunk.")
        _advance_staged_rate_audit_attempt_ledger(
            compare_payload=compare_payload,
            binding=expected_binding,
            chunk_id=chunk_id,
            effective_input_text=effective_input_text,
            attempt_count=attempts_used,
            status="preserved_baseline",
        )
        return {
            "chunk": restored_target,
            "compare": compare_payload,
            "outputPath": str(output_path),
            "comparePath": str(compare_path),
            "effectiveText": effective_input_text,
            "preservedExisting": True,
            "candidateSelectionAttempt": selection_event,
        }
    selected_output = str(selection_result.get("text", effective_input_text) or effective_input_text)
    published_rewrite = selection_event.get("publishedRewrite") is True
    selected_candidate_id = str(selection_event.get("selectedCandidateId", "") or "")
    selected_candidate = next(
        (
            candidate
            for candidate in selection_result.get("candidates", [])
            if isinstance(candidate, dict)
            and str(candidate.get("candidateId", "") or "") == selected_candidate_id
        ),
        {},
    )
    accepted_score = float(selected_candidate.get("stylePenalty", 0.0) or 0.0)

    target_chunk["outputText"] = selected_output
    target_chunk["outputCharCount"] = len(selected_output)
    target_chunk["quality"] = _apply_candidate_selection_quality(
        _build_chunk_quality(
            effective_input_text,
            selected_output,
            round_dimension=evaluator,
        ),
        selection_event,
    )
    target_chunk["rerunAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    target_chunk["rerunMode"] = mode
    target_chunk["rerunPromptMode"] = "rate-audit-strategy-prompt"
    target_chunk["rerunStrategy"] = ["rate-audit", dimension_id]
    target_chunk["rerunAttemptCount"] = attempts_used
    target_chunk["rerunSelectedScore"] = round(accepted_score, 3)
    target_chunk["rerunDimensionConverged"] = published_rewrite
    target_chunk["rerunDimensionConvergeDirections"] = directions[-3:]
    target_chunk["rateAuditStrategyPlanDigest"] = str(expected_binding.get("planDigest", "") or "")
    target_chunk["rateAuditStrategyPromptId"] = recommended_prompt_id
    target_chunk["rateAuditStrategyEvaluatorDimensionId"] = str(dimension_definition.get("evaluatorDimensionId", "") or "")
    target_chunk["rateAuditStrategyInputSource"] = "review_materialized_compare"
    target_chunk["rateAuditStrategyEffectiveInputSha256"] = hashlib.sha256(effective_input_text.encode("utf-8")).hexdigest()
    if published_rewrite:
        target_chunk["rateAuditStrategyReviewRequired"] = True
        target_chunk.pop("rerunStatus", None)
        target_chunk.pop("rerunNonConvergedReason", None)
    else:
        target_chunk.pop("rateAuditStrategyReviewRequired", None)
        target_chunk["rerunStatus"] = "preserved_baseline"
        target_chunk["rerunNonConvergedReason"] = "candidate_selection_preserved_baseline"
        _advance_staged_rate_audit_attempt_ledger(
            compare_payload=compare_payload,
            binding=expected_binding,
            chunk_id=chunk_id,
            effective_input_text=effective_input_text,
            attempt_count=attempts_used,
            status="preserved_baseline",
        )
    target_chunk.pop("rerunFallbackMode", None)
    target_chunk.pop("rerunFallbackError", None)
    target_chunk["rerunDefaultDecision"] = _default_export_decision_for_chunk(target_chunk)

    compare_payload["updatedAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    return {
        "chunk": target_chunk,
        "compare": compare_payload,
        "outputPath": str(output_path),
        "comparePath": str(compare_path),
    }


def _replace_file_bytes_atomically(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.name}.{threading.get_ident()}.{time.monotonic_ns()}.tmp")
    try:
        temporary_path.write_bytes(payload)
        temporary_path.replace(path)
    finally:
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass


def _commit_legacy_rerun_stage(
    *,
    output_path: Path,
    staged_compare: dict[str, Any],
    review_decisions: dict[str, Any],
    updated_at: str,
    expected_docx_contract: dict[str, Any] | None,
) -> None:
    """Publish one legacy rerun without splitting compare/review/text state.

    The compare revision is the commit marker and is written last.  Every file
    is individually replaced atomically; if staging, publication, or the
    post-write consistency check fails, all touched artifacts are restored to
    their exact previous bytes.
    """

    compare_path = _find_compare_path_for_output(output_path)
    decisions_path = _find_review_decisions_path_for_output(output_path)
    reviewed_paragraphs = _build_paragraphs_from_compare_payload(
        output_path,
        staged_compare,
        review_decisions,
    )
    if reviewed_paragraphs is None:
        raise ValueError("Legacy rerun could not materialize the staged compare output.")

    body_map = _load_body_map_for_output(output_path)
    body_map_path = _find_body_map_path_for_output(output_path)
    updated_body_map = None
    body_map_payload: dict[str, Any] | None = None
    if body_map is not None:
        if body_map_path is None:
            raise ValueError("Legacy rerun body-map path is unavailable.")
        reviewed_paragraphs = _ensure_reviewed_paragraph_count_for_body_map(
            output_path,
            staged_compare,
            reviewed_paragraphs,
            len(body_map.units),
            review_decisions,
        )
        updated_body_map = update_docx_body_map_texts(
            body_map,
            reviewed_paragraphs,
            round_number=body_map.round_number,
        )
        source_path = normalize_path(Path(str(updated_body_map.source_path)))
        staged_contract = build_document_edit_contract(
            source_path,
            body_map=updated_body_map,
            candidate_texts=updated_body_map.current_texts(),
            stage="targeted_rerun_staging",
        )
        assert_document_edit_contract_ready(staged_contract, label="DOCX 定点重跑提交")
        if expected_docx_contract is None:
            raise ValueError("DOCX targeted rerun lost its frozen preflight contract.")
        for key in ("sourceSha256", "scopeDigest", "formatDigest"):
            if str(staged_contract.get(key, "") or "") != str(expected_docx_contract.get(key, "") or ""):
                raise ValueError(f"DOCX targeted rerun staged contract changed: {key}.")
        body_map_payload = updated_body_map.to_dict()
        # Older body-map readers ignore unknown metadata.  Keeping the same
        # revision here makes the four-artifact transaction directly auditable
        # without changing the frozen scope signature or body-map version.
        body_map_payload["updatedAt"] = updated_at
    elif expected_docx_contract is not None:
        raise ValueError("DOCX targeted rerun body map disappeared before publication.")

    staged_compare["updatedAt"] = updated_at
    staged_compare["reviewUpdatedAt"] = updated_at
    review_payload = {
        "outputPath": str(output_path),
        "updatedAt": updated_at,
        "compareRevision": updated_at,
        "reviewBaseCompareRevision": updated_at,
        "decisions": review_decisions,
    }
    expected_output = "\n\n".join(reviewed_paragraphs)
    next_payloads: list[tuple[Path, bytes]] = []
    if body_map_payload is not None and body_map_path is not None:
        next_payloads.append(
            (
                body_map_path,
                json.dumps(body_map_payload, ensure_ascii=False, indent=2).encode("utf-8"),
            )
        )
    # Compare is deliberately last: readers that use updatedAt as a revision
    # cannot observe a new revision before its text/review/body-map artifacts.
    next_payloads.extend(
        [
            (output_path, expected_output.encode("utf-8")),
            (
                decisions_path,
                json.dumps(review_payload, ensure_ascii=False, indent=2).encode("utf-8"),
            ),
            (
                compare_path,
                json.dumps(staged_compare, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8"),
            ),
        ]
    )

    originals: dict[Path, bytes | None] = {}
    try:
        for path, _ in next_payloads:
            originals[path] = path.read_bytes() if path.exists() else None
        for path, payload in next_payloads:
            _replace_file_bytes_atomically(path, payload)

        disk_compare = json.loads(compare_path.read_text(encoding="utf-8"))
        disk_review = json.loads(decisions_path.read_text(encoding="utf-8"))
        if disk_compare.get("updatedAt") != updated_at or disk_compare.get("reviewUpdatedAt") != updated_at:
            raise ValueError("Legacy rerun compare revision failed its post-commit check.")
        if disk_review.get("updatedAt") != updated_at or disk_review.get("decisions") != review_decisions:
            raise ValueError("Legacy rerun review decisions failed their post-commit check.")
        if output_path.read_text(encoding="utf-8") != expected_output:
            raise ValueError("Legacy rerun output text failed its post-commit check.")
        if updated_body_map is not None and body_map_path is not None:
            raw_body_map = json.loads(body_map_path.read_text(encoding="utf-8"))
            disk_body_map = load_docx_body_map(body_map_path)
            if raw_body_map.get("updatedAt") != updated_at:
                raise ValueError("Legacy rerun body-map revision failed its post-commit check.")
            if disk_body_map is None or disk_body_map.current_texts() != reviewed_paragraphs:
                raise ValueError("Legacy rerun body-map text failed its post-commit check.")
    except Exception:
        for path, original in reversed(list(originals.items())):
            try:
                if original is None:
                    path.unlink(missing_ok=True)
                else:
                    _replace_file_bytes_atomically(path, original)
            except OSError:
                pass
        raise


def _commit_rate_audit_strategy_stage(
    *,
    source_path: Path,
    output_path: Path,
    staged_compare: dict[str, Any],
    review_decisions: dict[str, Any],
    expected_binding: dict[str, Any],
) -> dict[str, Any]:
    """Publish compare/output/body-map as one rollback-protected commit."""

    compare_path = _find_compare_path_for_output(output_path)
    _assert_document_release_payload(staged_compare, review_decisions)
    reviewed_paragraphs = _build_paragraphs_from_compare_payload(
        output_path,
        staged_compare,
        review_decisions,
    )
    if reviewed_paragraphs is None:
        raise ValueError("RateAudit strategy could not materialize the staged compare output.")

    body_map = _load_body_map_for_output(output_path)
    body_map_path = _find_body_map_path_for_output(output_path)
    updated_body_map = None
    if body_map is not None:
        if body_map_path is None:
            raise ValueError("RateAudit strategy body-map path is unavailable.")
        reviewed_paragraphs = _ensure_reviewed_paragraph_count_for_body_map(
            output_path,
            staged_compare,
            reviewed_paragraphs,
            len(body_map.units),
            review_decisions,
        )
        updated_body_map = update_docx_body_map_texts(
            body_map,
            reviewed_paragraphs,
            round_number=body_map.round_number,
        )
        staged_contract = build_document_edit_contract(
            source_path,
            body_map=updated_body_map,
            candidate_texts=updated_body_map.current_texts(),
            stage="rate_audit_strategy_staging",
        )
        assert_document_edit_contract_ready(staged_contract, label="RateAudit 策略提交")
        staged_mismatches: list[str] = []
        for key, code in (
            ("sourceSha256", "source_sha256_mismatch"),
            ("scopeDigest", "scope_digest_mismatch"),
            ("formatDigest", "format_digest_mismatch"),
        ):
            if str(staged_contract.get(key, "") or "") != str(expected_binding.get(key, "") or ""):
                staged_mismatches.append(code)
        if staged_mismatches:
            raise StaleRateAuditStrategyPlanError(staged_mismatches)

    staged_compare["updatedAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    next_payloads: list[tuple[Path, bytes]] = []
    if updated_body_map is not None and body_map_path is not None:
        next_payloads.append(
            (
                body_map_path,
                json.dumps(updated_body_map.to_dict(), ensure_ascii=False, indent=2).encode("utf-8"),
            )
        )
    next_payloads.extend(
        [
            (output_path, "\n\n".join(reviewed_paragraphs).encode("utf-8")),
            (
                compare_path,
                json.dumps(staged_compare, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8"),
            ),
        ]
    )
    originals: dict[Path, bytes | None] = {}
    try:
        for path, _ in next_payloads:
            originals[path] = path.read_bytes() if path.exists() else None
        for path, payload in next_payloads:
            _replace_file_bytes_atomically(path, payload)
        post_audit = get_document_rate_audit(str(source_path), str(output_path))
        if not isinstance(post_audit, dict):
            raise ValueError("RateAudit strategy post-commit audit is unavailable.")
        return post_audit
    except Exception:
        for path, original in originals.items():
            try:
                if original is None:
                    path.unlink(missing_ok=True)
                else:
                    _replace_file_bytes_atomically(path, original)
            except OSError:
                pass
        raise


def _persist_rate_audit_strategy_attempt_ledger(
    *,
    output_path: Path,
    staged_compare: dict[str, Any],
    failed_chunk_ids: list[str],
    binding: dict[str, Any],
    effective_chunks: dict[str, str],
) -> dict[str, Any]:
    """Persist attempt exhaustion only; never publish staged candidate text."""

    disk_compare = read_round_compare(str(output_path))
    ledger = disk_compare.get("rateAuditStrategyAttempts")
    if not isinstance(ledger, dict):
        ledger = {}
        disk_compare["rateAuditStrategyAttempts"] = ledger
    staged_chunks = staged_compare.get("chunks") if isinstance(staged_compare.get("chunks"), list) else []
    dimension_id = str(binding.get("dimensionId", "") or "")
    prompt_id = str(binding.get("recommendedPromptId", "") or "")
    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    for chunk_id in failed_chunk_ids:
        staged_chunk = next(
            (
                item
                for item in staged_chunks
                if isinstance(item, dict) and str(item.get("chunkId", "") or "") == chunk_id
            ),
            None,
        )
        if not isinstance(staged_chunk, dict) or staged_chunk.get("rerunStatus") != "non_converged":
            continue
        attempt_delta = max(0, int(staged_chunk.get("rerunAttemptCount", 0) or 0))
        effective_sha256 = hashlib.sha256(str(effective_chunks.get(chunk_id, "") or "").encode("utf-8")).hexdigest()
        ledger_key = f"{dimension_id}:{chunk_id}"
        previous = ledger.get(ledger_key)
        previous_count = 0
        if (
            isinstance(previous, dict)
            and str(previous.get("recommendedPromptId", "") or "") == prompt_id
            and str(previous.get("effectiveTextSha256", "") or "") == effective_sha256
        ):
            previous_count = max(0, int(previous.get("attemptCount", 0) or 0))
        ledger[ledger_key] = {
            "version": 1,
            "dimensionId": dimension_id,
            "chunkId": chunk_id,
            "recommendedPromptId": prompt_id,
            "effectiveTextSha256": effective_sha256,
            "attemptCount": previous_count + attempt_delta,
            "status": "non_converged",
            "lastPlanDigest": str(binding.get("planDigest", "") or ""),
            "updatedAt": now,
        }
    disk_compare["updatedAt"] = now
    _write_export_evidence_json_atomically(_find_compare_path_for_output(output_path), disk_compare)
    return disk_compare


def execute_rate_audit_strategy(
    request_payload: dict[str, Any],
    model_config: dict[str, Any],
    *,
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    """Execute a bound strategy under the same per-output lock as legacy reruns."""

    source_path = _resolve_api_path(
        str(request_payload.get("sourcePath", "") or ""),
        allowed_roots=((ROOT_DIR / "origin").resolve(),),
        label="Source file",
    )
    output_path = _resolve_api_path(
        str(request_payload.get("outputPath", "") or ""),
        allowed_roots=API_OUTPUT_ALLOWED_ROOTS,
        label="Output path",
    )
    if not isinstance(model_config, dict):
        raise ValueError("modelConfig is required.")

    def notify(event: dict[str, Any]) -> None:
        if progress_callback is not None:
            progress_callback(event)

    with get_output_rerun_lock(output_path):
        # This is the mandatory second validation after HTTP 202 registration.
        report, binding = validate_rate_audit_strategy_request(request_payload)
        execution_snapshot = read_round_artifact_snapshot(output_path, include_internal=True)
        if str(execution_snapshot.get("artifactSnapshotDigest", "") or "") != str(
            binding.get("artifactSnapshotDigest", "") or ""
        ):
            raise StaleRateAuditStrategyPlanError(["artifact_snapshot_digest_mismatch"])
        if str(execution_snapshot.get("contentRevision", "") or "") != str(
            binding.get("contentRevision", "") or ""
        ):
            raise StaleRateAuditStrategyPlanError(["content_revision_mismatch"])
        materialized = _materialize_rate_audit_output(
            output_path,
            artifact_snapshot=execution_snapshot,
        )
        materialized_sha256 = hashlib.sha256(str(materialized.get("text", "") or "").encode("utf-8")).hexdigest()
        if materialized_sha256 != str(binding.get("effectiveTextSha256", "") or ""):
            raise StaleRateAuditStrategyPlanError(["effective_text_sha256_mismatch"])
        original_compare = materialized.get("comparePayload")
        if not isinstance(original_compare, dict):
            raise StaleRateAuditStrategyPlanError(["compare_invalid"])
        staged_compare = copy.deepcopy(original_compare)
        review_decisions = (
            copy.deepcopy(materialized.get("decisions"))
            if isinstance(materialized.get("decisions"), dict)
            else {}
        )
        effective_chunks = {
            str(item.get("chunkId", "") or "").strip(): str(item.get("text", "") or "")
            for item in materialized.get("chunks", [])
            if isinstance(item, dict) and str(item.get("chunkId", "") or "").strip()
        }
        global_style_profile = build_global_style_profile_from_texts(
            [text for text in effective_chunks.values() if text.strip()]
        )
        target_ids = list(binding.get("targetChunkIds", []) or [])
        dimension_id = str(binding.get("dimensionId", "") or "")
        recommended_prompt_id = str(binding.get("recommendedPromptId", "") or "")
        success_chunk_ids: list[str] = []
        failures: list[dict[str, Any]] = []
        compare_path = str(_find_compare_path_for_output(output_path))

        for index, chunk_id in enumerate(target_ids, start=1):
            if cancel_check is not None and cancel_check():
                break
            notify({"phase": "chunk-start", "index": index, "total": len(target_ids), "chunkId": chunk_id})
            try:
                result = _rerun_rate_audit_strategy_chunk_unlocked(
                    source_path=source_path,
                    output_path=output_path,
                    compare_payload=staged_compare,
                    chunk_id=chunk_id,
                    effective_input_text=effective_chunks.get(chunk_id, ""),
                    effective_document_texts=dict(effective_chunks),
                    dimension_id=dimension_id,
                    recommended_prompt_id=recommended_prompt_id,
                    expected_binding=binding,
                    model_config=model_config,
                    global_style_profile=global_style_profile,
                )
                result_chunk = result.get("chunk")
                if not isinstance(result_chunk, dict):
                    raise ValueError("RateAudit strategy returned no staged chunk.")
                staged_output_text = str(
                    result.get("effectiveText", result_chunk.get("outputText", "")) or ""
                )
                if not staged_output_text.strip():
                    raise ValueError("RateAudit strategy returned an empty staged chunk.")
                # Every later target sees all earlier staged selections.  This
                # closes the batch bypass where several individually-small
                # deltas accumulated under one stale pre-batch profile.
                effective_chunks[chunk_id] = staged_output_text
                global_style_profile = build_global_style_profile_from_texts(
                    [text for text in effective_chunks.values() if text.strip()]
                )
                compare_path = str(result.get("comparePath", compare_path) or compare_path)
                success_chunk_ids.append(chunk_id)
                notify({"phase": "chunk-complete", "index": index, "total": len(target_ids), "chunkId": chunk_id})
            except StaleRateAuditStrategyPlanError:
                raise
            except Exception as exc:
                safe_error = safe_public_error_message(exc)
                failed_chunk = None
                if isinstance(staged_compare.get("chunks"), list):
                    failed_chunk = next(
                        (
                            item
                            for item in staged_compare["chunks"]
                            if isinstance(item, dict) and str(item.get("chunkId", "") or "") == chunk_id
                        ),
                        None,
                    )
                failed_attempts = (
                    _normalize_failed_attempts(failed_chunk.get("failedAttempts"))
                    if isinstance(failed_chunk, dict)
                    else []
                )
                failure = _build_text_free_app_failure(
                    chunk_id=chunk_id,
                    error=safe_error,
                    failed_attempts=failed_attempts,
                )
                if isinstance(failed_chunk, dict):
                    for key in ("rerunStatus", "rerunNonConvergedReason"):
                        if failed_chunk.get(key):
                            failure[key] = failed_chunk.get(key)
                failures.append(failure)
                notify(
                    {
                        "phase": "chunk-failed",
                        "index": index,
                        "total": len(target_ids),
                        "chunkId": chunk_id,
                        "error": failure["error"],
                        "guardCategory": failure["guardCategory"],
                        "issueCodes": failure["issueCodes"],
                        "errorStored": False,
                        "reasoningSuppressed": True,
                        "providerContentStored": False,
                        "failedAttempts": failed_attempts,
                    }
                )
                # The strategy bundle is transactional. Once one target fails,
                # no later target is attempted and no earlier staged candidate
                # is published.
                break

        canceled = bool(cancel_check is not None and cancel_check())
        all_succeeded = bool(
            not canceled
            and not failures
            and len(success_chunk_ids) == len(target_ids)
        )
        if all_succeeded:
            # CAS immediately before publication. Because staged edits have not
            # touched disk, the original plan must still recompute byte-for-byte.
            _, commit_binding = validate_rate_audit_strategy_request(request_payload)
            if str(commit_binding.get("planDigest", "") or "") != str(binding.get("planDigest", "") or ""):
                raise StaleRateAuditStrategyPlanError(["plan_digest_mismatch"])
            post_audit = _commit_rate_audit_strategy_stage(
                source_path=source_path,
                output_path=output_path,
                staged_compare=staged_compare,
                review_decisions=review_decisions,
                expected_binding=binding,
            )
            latest_compare: dict[str, Any] | None = staged_compare
            post_audit_error = ""
        else:
            # Candidates that converged before a later failure were never
            # committed and therefore cannot be reported as successes.
            non_converged_ids = [
                str(item.get("chunkId", "") or "")
                for item in staged_compare.get("chunks", [])
                if isinstance(item, dict)
                and item.get("rerunStatus") == "non_converged"
                and str(item.get("chunkId", "") or "") in {str(failure.get("chunkId", "") or "") for failure in failures}
            ]
            rolled_back_ids = list(success_chunk_ids)
            success_chunk_ids = []
            for rolled_back_id in rolled_back_ids:
                rolled_back_failure = _build_text_free_app_failure(
                    chunk_id=rolled_back_id,
                    error="transaction_rolled_back",
                )
                rolled_back_failure["rerunStatus"] = "transaction_rolled_back"
                failures.append(rolled_back_failure)
            if non_converged_ids:
                # The text transaction stays rolled back, but the server-owned
                # attempt ledger advances so repeated requests cannot bypass
                # the registry's maximum attempt policy.
                validate_rate_audit_strategy_request(request_payload)
                latest_compare = _persist_rate_audit_strategy_attempt_ledger(
                    output_path=output_path,
                    staged_compare=staged_compare,
                    failed_chunk_ids=non_converged_ids,
                    binding=binding,
                    effective_chunks=effective_chunks,
                )
                post_audit_error = ""
                try:
                    post_audit = get_document_rate_audit(str(source_path), str(output_path))
                except Exception as exc:
                    post_audit = None
                    post_audit_error = "rate_audit_refresh_failed"
            else:
                try:
                    latest_compare = read_round_compare(str(output_path))
                except Exception:
                    latest_compare = None
                post_audit = None
                post_audit_error = ""
        completed_count = len(success_chunk_ids) + len(failures)
        resulting_plan = (
            (post_audit or {}).get("strategyPlan")
            if isinstance((post_audit or {}).get("strategyPlan"), dict)
            else report.get("strategyPlan")
            if isinstance(report.get("strategyPlan"), dict)
            else {}
        )
        initial_plan = report.get("strategyPlan") if isinstance(report.get("strategyPlan"), dict) else {}
        blocking_manual_dimensions: list[dict[str, Any]] = []
        seen_manual_dimensions: set[str] = set()
        for raw_item in [
            *((initial_plan or {}).get("blockingManualDimensions", []) or []),
            *((resulting_plan or {}).get("blockingManualDimensions", []) or []),
        ]:
            if not isinstance(raw_item, dict):
                continue
            manual_dimension_id = str(raw_item.get("dimensionId", "") or "").strip()
            if not manual_dimension_id or manual_dimension_id in seen_manual_dimensions:
                continue
            seen_manual_dimensions.add(manual_dimension_id)
            blocking_manual_dimensions.append(dict(raw_item))
        manual_review_required = bool(
            (initial_plan or {}).get("manualReviewRequired")
            or (resulting_plan or {}).get("manualReviewRequired")
        )
        manual_review_still_required = bool(
            (initial_plan or {}).get("manualReviewStillRequired")
            or (resulting_plan or {}).get("manualReviewStillRequired")
            or blocking_manual_dimensions
        )
        return {
            "ok": True,
            "outputPath": str(output_path),
            "comparePath": compare_path,
            "compare": latest_compare,
            "successChunkIds": success_chunk_ids,
            "totalCount": len(target_ids),
            "completedCount": completed_count,
            "successCount": len(success_chunk_ids),
            "failureCount": len(failures),
            "canceled": canceled,
            "failures": failures,
            "strategyBinding": binding,
            "strategyDecision": str((report.get("strategyPlan") or {}).get("decision", "")),
            "resultingStrategyDecision": str((resulting_plan or {}).get("decision", "") or ""),
            "manualReviewRequired": manual_review_required,
            "manualReviewStillRequired": manual_review_still_required,
            "blockingManualDimensions": blocking_manual_dimensions,
            "plateauReached": bool((resulting_plan or {}).get("plateauReached")),
            "plateauReason": str((resulting_plan or {}).get("plateauReason", "") or ""),
            "compareRevisionBefore": str(binding.get("compareRevision", "") or ""),
            "compareRevisionAfter": str(((post_audit or {}).get("strategyBinding") or {}).get("compareRevision", "")),
            "postAudit": post_audit,
            "postAuditError": post_audit_error,
        }


def load_model_config_payload(model_config_json: str | None = None, model_config_file: str | None = None) -> dict[str, Any]:
    if model_config_file:
        config_path = Path(model_config_file).resolve()
        return json.loads(config_path.read_text(encoding="utf-8"))
    if model_config_json:
        return json.loads(model_config_json)
    raise ValueError("Either model_config_json or model_config_file must be provided.")


def cli_main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Desktop app service bridge")
    subparsers = parser.add_subparsers(dest="command", required=True)

    import_parser = subparsers.add_parser("import-document")
    import_parser.add_argument("source_path")

    status_parser = subparsers.add_parser("document-status")
    status_parser.add_argument("source_path")
    status_parser.add_argument("prompt_profile", nargs="?", default=DEFAULT_PROMPT_PROFILE)

    history_parser = subparsers.add_parser("document-history")
    history_parser.add_argument("source_path")

    list_history_parser = subparsers.add_parser("document-history-list")

    delete_history_parser = subparsers.add_parser("delete-document-history")
    delete_history_parser.add_argument("doc_id")
    delete_history_parser.add_argument("--from-round", type=int, default=None)
    delete_history_parser.add_argument("--prompt-profile", default=None)

    run_parser = subparsers.add_parser("run-round")
    run_parser.add_argument("source_path")
    run_parser.add_argument("model_config_json", nargs="?", default=None)
    run_parser.add_argument("--config-file", default=None)
    run_parser.add_argument("--round", type=int, default=None)

    test_parser = subparsers.add_parser("test-connection")
    test_parser.add_argument("model_config_json", nargs="?", default=None)
    test_parser.add_argument("--config-file", default=None)

    export_parser = subparsers.add_parser("export-round")
    export_parser.add_argument("output_path")
    export_parser.add_argument("export_path")
    export_parser.add_argument("target_format", choices=["txt", "docx"])

    preview_parser = subparsers.add_parser("read-output")
    preview_parser.add_argument("output_path")
    preview_parser.add_argument("--max-chars", type=int, default=DEFAULT_PREVIEW_MAX_CHARS)

    compare_parser = subparsers.add_parser("read-compare")
    compare_parser.add_argument("output_path")

    args = parser.parse_args()

    try:
        if args.command == "import-document":
            payload = import_document(args.source_path)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "document-status":
            payload = get_document_status(args.source_path, prompt_profile=args.prompt_profile)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "document-history":
            payload = get_document_history(args.source_path)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "document-history-list":
            payload = list_document_histories()
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "delete-document-history":
            payload = delete_document_history(args.doc_id, args.from_round, args.prompt_profile)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "run-round":
            payload = run_round_for_app(
                args.source_path,
                load_model_config_payload(args.model_config_json, args.config_file),
                args.round,
            )
            emit_result_payload(payload)
        elif args.command == "test-connection":
            payload = test_model_connection(load_model_config_payload(args.model_config_json, args.config_file))
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "export-round":
            payload = export_round_output(args.output_path, args.export_path, args.target_format)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "read-output":
            payload = read_output_text(args.output_path, max_chars=args.max_chars)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        elif args.command == "read-compare":
            payload = read_round_compare(args.output_path)
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        else:
            raise ValueError(f"Unsupported command: {args.command}")
    except Exception as exc:
        if args.command == "run-round":
            emit_error_payload(str(exc))
        raise


if __name__ == "__main__":
    cli_main()
