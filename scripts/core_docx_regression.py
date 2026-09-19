from __future__ import annotations

import io
import json
import os
import struct
import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree

from core_documents import (
    M_NS,
    NS,
    W_NS,
    W_T,
    FormatFidelityError,
    export_docx,
    export_txt,
    import_document,
    load_document,
    save_document,
    update_scope,
)


DOCUMENT_XML = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>第一章 标题</w:t></w:r></w:p>
    <w:p>
      <w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr><w:spacing w:line="360"/></w:pPr>
      <w:r><w:rPr><w:rFonts w:eastAsia="宋体"/></w:rPr><w:t xml:space="preserve">  本文使用</w:t></w:r>
      <w:hyperlink r:id="rIdLink"><w:r><w:rPr><w:b/><w:u w:val="single"/></w:rPr><w:t>Transformer</w:t></w:r></w:hyperlink>
      <w:r><w:rPr><w:sz w:val="24"/></w:rPr><w:t>模型，实验值10，引用见</w:t></w:r>
      <w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>[1]</w:t></w:r>
      <w:r><w:t>，网址</w:t></w:r>
      <w:hyperlink r:id="rIdLink"><w:r><w:rPr><w:u w:val="single"/></w:rPr><w:t xml:space="preserve">https://example.com。  </w:t></w:r></w:hyperlink>
      <w:bookmarkStart w:id="2" w:name="正文锚点"/><w:r><w:t xml:space="preserve"/></w:r><w:bookmarkEnd w:id="2"/>
    </w:p>
    <w:p>
      <w:r><w:t>跨页前正文</w:t><w:lastRenderedPageBreak/><w:t>跨页后正文</w:t></w:r>
    </w:p>
    <w:p>
      <w:pPr><w:sectPr><w:type w:val="continuous"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:pPr>
      <w:r><w:t>分节边界正文</w:t></w:r>
    </w:p>
    <w:tbl><w:tr><w:tc><w:p><w:r><w:t>表格内容不改</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
    <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>TOC \\o "1-3"</w:instrText></w:r><w:r><w:t>目录字段</w:t></w:r></w:p>
    <w:p><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath><w:r><w:t>公式段</w:t></w:r></w:p>
    <w:p><w:r><w:drawing><wp:inline/></w:drawing><w:t>绘图段</w:t></w:r></w:p>
    <w:p><w:ins w:id="1"><w:r><w:t>修订中的文字</w:t></w:r></w:ins></w:p>
    <w:p><w:sdt><w:sdtContent><w:r><w:t>内容控件文字</w:t></w:r></w:sdtContent></w:sdt></w:p>
    <w:sectPr><w:headerReference w:type="default" r:id="rIdHeader"/><w:footerReference w:type="default" r:id="rIdFooter"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr>
  </w:body>
</w:document>'''.encode()


def fixture_docx(document_xml: bytes = DOCUMENT_XML) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
  <Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
  <Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
  <Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
  <Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>
</Types>''',
        )
        archive.writestr(
            "_rels/.rels",
            b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>''',
        )
        archive.writestr("word/document.xml", document_xml)
        archive.writestr(
            "word/_rels/document.xml.rels",
            b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com" TargetMode="External"/>
  <Relationship Id="rIdComments" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>
  <Relationship Id="rIdHeader" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>
  <Relationship Id="rIdFooter" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>
  <Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
  <Relationship Id="rIdNumbering" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>
</Relationships>''',
        )
        archive.writestr("word/header1.xml", f'<w:hdr xmlns:w="{W_NS}"><w:p><w:r><w:t>页眉原文</w:t></w:r></w:p></w:hdr>'.encode())
        archive.writestr("word/footer1.xml", f'<w:ftr xmlns:w="{W_NS}"><w:p><w:r><w:t>页脚原文</w:t></w:r></w:p></w:ftr>'.encode())
        archive.writestr("word/comments.xml", f'<w:comments xmlns:w="{W_NS}"><w:comment w:id="0" w:author="FYADR" w:date="2026-01-01T00:00:00Z"><w:p><w:r><w:t>批注部件原样保留</w:t></w:r></w:p></w:comment></w:comments>'.encode())
        archive.writestr(
            "word/numbering.xml",
            (
                f'<w:numbering xmlns:w="{W_NS}">'
                '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
                '<w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'
                '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>'
            ).encode(),
        )
        archive.writestr(
            "word/styles.xml",
            (
                f'<w:styles xmlns:w="{W_NS}">'
                '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>'
                '<w:style w:type="paragraph" w:styleId="1"><w:name w:val="heading 1"/>'
                '<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'
                '</w:styles>'
            ).encode(),
        )
        archive.comment = b"fyadr-fixture"
    return output.getvalue()


def fixture_with_word_zip_metadata() -> bytes:
    """Model ZIP metadata emitted by Word that zipfile safely normalizes."""

    value = bytearray(fixture_docx())
    end = value.rfind(b"PK\x05\x06")
    if end < 0:
        raise AssertionError("fixture ZIP end record is missing")
    entry_count = struct.unpack_from("<H", value, end + 10)[0]
    cursor = struct.unpack_from("<I", value, end + 16)[0]
    for _ in range(entry_count):
        if value[cursor : cursor + 4] != b"PK\x01\x02":
            raise AssertionError("fixture ZIP central directory is invalid")
        struct.pack_into("<H", value, cursor + 8, 0x0006)
        struct.pack_into("<I", value, cursor + 38, 0)
        local_offset = struct.unpack_from("<I", value, cursor + 42)[0]
        if value[local_offset : local_offset + 4] != b"PK\x03\x04":
            raise AssertionError("fixture ZIP local entry is invalid")
        struct.pack_into("<H", value, local_offset + 6, 0x0006)
        name_length, extra_length, comment_length = struct.unpack_from("<HHH", value, cursor + 28)
        cursor += 46 + name_length + extra_length + comment_length
    return bytes(value)


def fixture_with_digital_signature() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(fixture_docx()), "r") as before, zipfile.ZipFile(
        output, "w"
    ) as after:
        after.comment = before.comment
        for info in before.infolist():
            after.writestr(info, before.read(info.filename))
        after.writestr("_xmlsignatures/sig1.xml", b"<Signature>fixture</Signature>")
    return output.getvalue()


class CoreDocxRegression(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.previous_data = os.environ.get("FYADR_DATA_DIR")
        os.environ["FYADR_DATA_DIR"] = str(Path(self.temporary.name) / "data")

    def tearDown(self) -> None:
        if self.previous_data is None:
            os.environ.pop("FYADR_DATA_DIR", None)
        else:
            os.environ["FYADR_DATA_DIR"] = self.previous_data
        self.temporary.cleanup()

    def test_import_scope_and_in_place_export_preserve_ooxml(self) -> None:
        public = import_document(io.BytesIO(fixture_docx()), "论文.docx")
        self.assertEqual(public["kind"], "docx")
        self.assertEqual(public["safeCount"], 7)
        self.assertEqual(public["excludedCount"], 2)
        heading, body, page_break, section_body, field, formula, drawing, revision, content_control = public["paragraphs"]
        self.assertEqual(heading["styleName"], "heading 1")
        self.assertEqual(heading["outlineLevel"], 0)
        self.assertTrue(heading["safe"])
        self.assertEqual(heading["suggestionReason"], "heading")
        self.assertFalse(heading["selected"])
        self.assertTrue(body["selected"])
        self.assertTrue(page_break["safe"])
        self.assertTrue(page_break["selected"])
        self.assertTrue(section_body["safe"])
        self.assertFalse(field["safe"])
        self.assertTrue(formula["safe"])
        self.assertTrue(drawing["safe"])
        self.assertFalse(revision["safe"])
        self.assertTrue(content_control["safe"])
        self.assertTrue(content_control["selected"])
        protection_map = public["protectionMap"]
        self.assertTrue(protection_map["available"])
        self.assertEqual(protection_map["summary"]["totalUnits"], 10)
        self.assertEqual(protection_map["summary"]["editableUnits"], 6)
        self.assertEqual(protection_map["summary"]["availableUnits"], 1)
        self.assertEqual(protection_map["summary"]["lockedUnits"], 3)
        self.assertEqual(protection_map["summary"]["tableUnits"], 1)
        self.assertTrue(any(item["reason"] == "table_content" for item in protection_map["summary"]["protectionReasons"]))
        serialized_map = json.dumps(protection_map, ensure_ascii=False)
        self.assertNotIn("complex_structure", serialized_map)
        self.assertNotIn("其他复杂结构", serialized_map)

        scoped = update_scope(public["id"], [body["id"]])
        self.assertEqual(scoped["protectionMap"]["summary"]["editableUnits"], 1)
        document = load_document(public["id"])
        source = Path(document["sourcePath"])
        output, audit = export_docx(document, {body["id"]: "改写结果"}, "test-run")
        self.assertTrue(audit["passed"], audit)

        with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            self.assertEqual(before.comment, after.comment)
            for name in before.namelist():
                if name != "word/document.xml":
                    self.assertEqual(before.read(name), after.read(name), name)

            parser = etree.XMLParser(remove_blank_text=False)
            original_root = etree.fromstring(before.read("word/document.xml"), parser)
            exported_root = etree.fromstring(after.read("word/document.xml"), parser)
            original_body = list(original_root.find("w:body", namespaces=NS))
            exported_body = list(exported_root.find("w:body", namespaces=NS))
            body_index = int(body["bodyChildIndex"])
            original_paragraph = original_body[body_index]
            exported_paragraph = exported_body[body_index]
            self.assertEqual("".join(node.text or "" for node in exported_paragraph.iter(W_T)), "  改写结果  ")
            self.assertEqual(len(list(original_paragraph.iter(W_T))), len(list(exported_paragraph.iter(W_T))))
            self.assertEqual(
                etree.tostring(original_paragraph.find("w:pPr", namespaces=NS)),
                etree.tostring(exported_paragraph.find("w:pPr", namespaces=NS)),
            )
            self.assertEqual(
                len(original_paragraph.findall(".//w:hyperlink", namespaces=NS)),
                len(exported_paragraph.findall(".//w:hyperlink", namespaces=NS)),
            )
            self.assertEqual(len(exported_paragraph.findall(".//w:bookmarkStart", namespaces=NS)), 1)
            self.assertEqual(len(exported_paragraph.findall(".//w:bookmarkEnd", namespaces=NS)), 1)
            self.assertEqual(len(list(original_paragraph.iter(W_T))), len(list(exported_paragraph.iter(W_T))))
            self.assertEqual(list(exported_paragraph.iter(W_T))[-1].get("{http://www.w3.org/XML/1998/namespace}space"), "preserve")
            self.assertEqual(
                len(exported_root.findall(".//w:sectPr", namespaces=NS)),
                2,
            )
            for index, (left, right) in enumerate(zip(original_body, exported_body)):
                if index != body_index:
                    self.assertEqual(etree.tostring(left), etree.tostring(right))

    def test_cross_paragraph_toc_field_sets_body_suggestion_after_outer_field_end(self) -> None:
        toc_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>封面上的普通文字</w:t></w:r></w:p>
  <w:p>
    <w:r><w:fldChar w:fldCharType="begin"/></w:r>
    <w:r><w:instrText xml:space="preserve"> TO</w:instrText></w:r>
    <w:r><w:instrText>C \\o "1-3" </w:instrText></w:r>
    <w:r><w:fldChar w:fldCharType="separate"/></w:r>
    <w:r><w:fldChar w:fldCharType="begin"/></w:r>
    <w:r><w:instrText> PAGEREF _Toc1 </w:instrText></w:r>
    <w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>第一章 1</w:t></w:r>
    <w:r><w:fldChar w:fldCharType="end"/></w:r>
  </w:p>
  <w:p>
    <w:r><w:fldChar w:fldCharType="begin"/></w:r>
    <w:r><w:instrText> PAGEREF _Toc2 </w:instrText></w:r>
    <w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>第二章 8</w:t></w:r>
    <w:r><w:fldChar w:fldCharType="end"/></w:r>
  </w:p>
  <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
  <w:p><w:r><w:t>目录后的页数说明</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>第一章 标题</w:t></w:r></w:p>
  <w:p><w:r><w:t>目录之后的正文应进入建议范围。</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(toc_document)), "跨段目录论文.docx")

        cover, toc_start, toc_result, toc_end, page_count, heading, body = public["paragraphs"]
        self.assertEqual(public["suggestionBasis"], "toc_field")
        self.assertEqual(public["suggestionStartBodyChildIndex"], 5)
        self.assertFalse(cover["selected"])
        self.assertEqual(cover["suggestionReason"], "before_body_start")
        self.assertFalse(toc_start["safe"])
        self.assertFalse(toc_result["safe"])
        self.assertFalse(toc_end["safe"])
        self.assertFalse(page_count["selected"])
        self.assertEqual(page_count["suggestionReason"], "before_body_start")
        self.assertFalse(heading["selected"])
        self.assertEqual(heading["suggestionReason"], "heading")
        self.assertTrue(body["selected"])
        self.assertEqual(public["selectedCount"], 1)

    def test_toc_rows_named_like_back_matter_do_not_hide_following_body(self) -> None:
        toc_rows_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>目录</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr><w:r><w:t>参考文献40</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr><w:r><w:t>致谢41</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr><w:r><w:t>附录42</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>第1章 绪论</w:t></w:r></w:p>
  <w:p><w:r><w:t>目录之后的真实正文必须进入改写范围。</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(toc_rows_document)), "目录状态论文.docx")

        toc_heading, references_row, thanks_row, appendix_row, chapter_heading, body = public["paragraphs"]
        for item in (toc_heading, references_row, thanks_row, appendix_row):
            self.assertFalse(item["selected"])
            self.assertEqual(item["suggestionReason"], "toc")
        self.assertFalse(chapter_heading["selected"])
        self.assertEqual(chapter_heading["suggestionReason"], "heading")
        self.assertTrue(body["selected"])
        self.assertEqual(body["suggestionReason"], "body_text")
        self.assertEqual(public["selectedCount"], 1)

    def test_section_then_outline_heading_is_used_when_no_toc_field_exists(self) -> None:
        section_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>封面信息</w:t></w:r></w:p>
  <w:p><w:pPr><w:sectPr/></w:pPr></w:p>
  <w:p><w:r><w:t>分节后的前置说明</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>第一章 标题</w:t></w:r></w:p>
  <w:p><w:r><w:t>真正需要改写的正文。</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(section_document)), "无目录论文.docx")

        cover, section_break, preface, heading, body = public["paragraphs"]
        self.assertEqual(public["suggestionBasis"], "section_heading")
        self.assertEqual(public["suggestionStartBodyChildIndex"], 3)
        self.assertFalse(cover["selected"])
        self.assertFalse(section_break["safe"])
        self.assertFalse(preface["selected"])
        self.assertEqual(preface["suggestionReason"], "before_body_start")
        self.assertFalse(heading["selected"])
        self.assertTrue(body["selected"])

    def test_abstracts_are_body_while_titles_keywords_captions_and_references_stay_original(self) -> None:
        chinese_abstract = "摘要  " + ("中" * 156)
        english_abstract = "Abstract  " + ("a" * 451)
        abstract_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>封面标题</w:t></w:r></w:p>
  <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> TOC \\o "1-3" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>目录结果</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
  <w:p><w:r><w:t>论文中文标题</w:t></w:r></w:p>
  <w:p><w:r><w:t>{chinese_abstract}</w:t></w:r></w:p>
  <w:p><w:r><w:t>关键词：范围；保护</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>1 引言</w:t></w:r></w:p>
  <w:p><w:r><w:t>章节正文应进入改写范围。</w:t></w:r></w:p>
  <w:p><w:r><w:t>图1-1 系统结构图</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>致谢</w:t></w:r></w:p>
  <w:p><w:r><w:t>感谢所有在研究期间提供帮助的人。</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
  <w:p><w:r><w:t>[1] Author. Reference title. 2025.</w:t></w:r></w:p>
  <w:p><w:r><w:t>GENERIC ENGLISH THESIS TITLE</w:t></w:r></w:p>
  <w:p><w:r><w:t>{english_abstract}</w:t></w:r></w:p>
  <w:p><w:r><w:t>KeyWords scope; protection;</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(abstract_document)), "摘要范围论文.docx")

        by_text = {item["displayText"]: item for item in public["paragraphs"]}
        selected = [item for item in public["paragraphs"] if item["selected"]]
        self.assertEqual(public["suggestionBasis"], "abstract_heading")
        self.assertEqual(public["suggestionStartBodyChildIndex"], 3)
        self.assertEqual(len(chinese_abstract), 160)
        self.assertEqual(len(english_abstract), 461)
        self.assertEqual(
            {item["displayText"] for item in selected},
            {
                chinese_abstract,
                "章节正文应进入改写范围。",
                "感谢所有在研究期间提供帮助的人。",
                english_abstract,
            },
        )
        self.assertTrue(by_text[chinese_abstract]["selected"])
        self.assertTrue(by_text[english_abstract]["selected"])
        self.assertEqual(by_text[english_abstract]["suggestionReason"], "body_text")
        for original_text in (
            "封面标题",
            "论文中文标题",
            "关键词：范围；保护",
            "1 引言",
            "图1-1 系统结构图",
            "致谢",
            "参考文献",
            "[1] Author. Reference title. 2025.",
            "GENERIC ENGLISH THESIS TITLE",
            "KeyWords scope; protection;",
        ):
            self.assertFalse(by_text[original_text]["selected"], original_text)

    def test_content_roles_are_suggestions_and_never_lock_writable_text(self) -> None:
        structured_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>正文段落应进入改写范围。</w:t></w:r></w:p>
  <w:p><w:r><w:t>4. 关键技术实现</w:t></w:r></w:p>
  <w:p><w:r><w:t>图3-1展示了系统结构。该结构由三个模块组成。</w:t></w:r></w:p>
  <w:p><w:r><w:t>系统结构如图3-1所示。</w:t></w:r></w:p>
  <w:p><w:r><w:t>表3.1 dm_user表</w:t></w:r></w:p>
  <w:p><w:r><w:t>注：数据来自公开年报</w:t></w:r></w:p>
  <w:p><w:r><w:t>关键词：门禁系统；二维码</w:t></w:r></w:p>
  <w:p><w:r><w:t>(a)</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
  <w:p><w:r><w:t>张三,李四.系统设计与实现.北京,2024</w:t></w:r></w:p>
  <w:p><w:r><w:t>[2] Author. Another reference. 2023.</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>附录</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>附录中的内部标题</w:t></w:r></w:p>
  <w:p><w:r><w:t>附录中的内容不会进入模型。</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(structured_document)), "结构保护论文.docx")

        (
            body,
            numbered_heading,
            figure_prose,
            cross_reference_prose,
            caption,
            note,
            keywords,
            structural_label,
            references_heading,
            reference_one,
            reference_two,
            appendix,
            appendix_heading,
            appendix_body,
        ) = public["paragraphs"]
        self.assertTrue(body["selected"])
        self.assertTrue(numbered_heading["safe"])
        self.assertFalse(numbered_heading["selected"])
        self.assertEqual(numbered_heading["suggestionReason"], "heading")
        self.assertTrue(figure_prose["safe"])
        self.assertTrue(figure_prose["selected"])
        self.assertTrue(cross_reference_prose["safe"])
        self.assertTrue(cross_reference_prose["selected"])
        self.assertTrue(caption["safe"])
        self.assertFalse(caption["selected"])
        self.assertEqual(caption["suggestionReason"], "caption")
        self.assertTrue(note["safe"])
        self.assertFalse(note["selected"])
        self.assertEqual(note["suggestionReason"], "caption_note")
        self.assertTrue(keywords["safe"])
        self.assertFalse(keywords["selected"])
        self.assertEqual(keywords["suggestionReason"], "keywords")
        self.assertTrue(structural_label["safe"])
        self.assertFalse(structural_label["selected"])
        self.assertEqual(structural_label["suggestionReason"], "structural_label")
        for paragraph in (references_heading, reference_one, reference_two):
            self.assertTrue(paragraph["safe"])
            self.assertFalse(paragraph["selected"])
            self.assertEqual(paragraph["suggestionReason"], "references")
        self.assertTrue(appendix["safe"])
        self.assertFalse(appendix["selected"])
        self.assertTrue(appendix_heading["safe"])
        self.assertTrue(appendix_body["safe"])
        self.assertFalse(appendix_body["selected"])

        suggested_labels = {
            item["reason"]: item["label"]
            for item in public["protectionMap"]["summary"]["protectionReasons"]
        }
        self.assertEqual(suggested_labels["caption"], "图名或表名")
        self.assertEqual(suggested_labels["references"], "参考文献")
        self.assertEqual(public["protectionMap"]["summary"]["lockedUnits"], 0)

        # Content-role suggestions must never become a gate. If the automatic
        # boundary guess is wrong, the user can include any writable paragraph.
        corrected = update_scope(
            public["id"],
            [body["id"], caption["id"], reference_one["id"], appendix_body["id"]],
        )
        self.assertEqual(corrected["selectedCount"], 4)
        self.assertEqual(
            {item["id"] for item in corrected["paragraphs"] if item["selected"]},
            {body["id"], caption["id"], reference_one["id"], appendix_body["id"]},
        )

    def test_standalone_equation_numbers_are_not_suggested_as_body(self) -> None:
        equation_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>模型输入为多变量时间序列，计算结果如下。</w:t></w:r></w:p>
  <w:p><w:r><w:t>（3-1）</w:t></w:r></w:p>
  <w:p><w:r><w:t>式（4.2）</w:t></w:r></w:p>
  <w:p><w:r><w:t>（a）</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(equation_document)), "公式编号.docx")
        prose, numbered, prefixed, lettered = public["paragraphs"]

        self.assertTrue(prose["selected"])
        for paragraph in (numbered, prefixed, lettered):
            self.assertTrue(paragraph["safe"])
            self.assertFalse(paragraph["selected"])
            self.assertEqual(paragraph["suggestionReason"], "structural_label")

    def test_upgrade_removes_old_automatic_equation_selection_but_keeps_manual_includes(self) -> None:
        equation_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>正文段落需要继续改写。</w:t></w:r></w:p>
  <w:p><w:r><w:t>（3-1）</w:t></w:r></w:p>
  <w:p><w:r><w:t>表3-1 试验结果</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(equation_document)), "旧公式范围.docx")
        stored = load_document(public["id"])
        body, equation_number, caption = stored["paragraphs"]
        equation_number["selected"] = True
        equation_number["suggestedSelected"] = True
        caption["selected"] = True
        self.assertFalse(caption["suggestedSelected"])
        stored["scopeConfirmed"] = True
        stored["scopeClassifierVersion"] = 10
        save_document(stored)

        upgraded = load_document(public["id"])
        upgraded_body, upgraded_equation, upgraded_caption = upgraded["paragraphs"]
        self.assertTrue(upgraded_body["selected"])
        self.assertFalse(upgraded_equation["selected"])
        self.assertTrue(upgraded_caption["selected"])

    def test_front_declaration_does_not_lock_the_following_body(self) -> None:
        declaration_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>原创性声明</w:t></w:r></w:p>
  <w:p><w:r><w:t>本人确认本文由本人独立完成。</w:t></w:r></w:p>
  <w:p><w:pPr><w:pStyle w:val="1"/></w:pPr><w:r><w:t>第一章 绪论</w:t></w:r></w:p>
  <w:p><w:r><w:t>进入正文后，这个普通段落可以改写。</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(declaration_document)), "声明与正文.docx")
        declaration, declaration_body, heading, body = public["paragraphs"]
        self.assertEqual(declaration["suggestionReason"], "back_matter")
        self.assertEqual(declaration_body["suggestionReason"], "back_matter")
        self.assertEqual(heading["suggestionReason"], "heading")
        self.assertTrue(body["safe"])
        self.assertTrue(body["selected"])

    def test_confirmed_existing_scope_keeps_the_exact_user_choice_after_upgrade(self) -> None:
        structured_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>系统功能包括：</w:t></w:r></w:p>
  <w:p><w:r><w:t>图2.1 系统结构图</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(structured_document)), "旧范围论文.docx")
        stored = load_document(public["id"])
        body = stored["paragraphs"][0]
        body["safe"] = False
        body["selected"] = False
        body["protectionReason"] = "structural_label"
        body["exclusionReason"] = "结构标记"
        caption = stored["paragraphs"][1]
        caption["safe"] = True
        caption["selected"] = True
        caption["protectionReason"] = ""
        caption["exclusionReason"] = ""
        stored["scopeConfirmed"] = True
        stored["scopeClassifierVersion"] = 1
        save_document(stored)

        upgraded = load_document(public["id"])
        upgraded_body = upgraded["paragraphs"][0]
        upgraded_caption = upgraded["paragraphs"][1]
        self.assertTrue(upgraded_body["safe"])
        self.assertFalse(upgraded_body["selected"])
        self.assertTrue(upgraded_caption["safe"])
        self.assertTrue(upgraded_caption["selected"])
        self.assertEqual(upgraded_caption["suggestionReason"], "caption")
        self.assertTrue(upgraded["scopeConfirmed"])

    def test_unconfirmed_existing_scope_uses_fresh_suggestions_after_upgrade(self) -> None:
        structured_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>
  <w:p><w:r><w:t>系统功能包括：</w:t></w:r></w:p>
  <w:p><w:r><w:t>图2.1 系统结构图</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(structured_document)), "未确认旧范围.docx")
        stored = load_document(public["id"])
        body, caption = stored["paragraphs"]
        body["safe"] = False
        body["selected"] = False
        body["protectionReason"] = "structural_label"
        body["exclusionReason"] = "结构标记"
        caption["selected"] = True
        stored["scopeConfirmed"] = False
        stored["scopeClassifierVersion"] = 1
        save_document(stored)

        upgraded = load_document(public["id"])
        upgraded_body, upgraded_caption = upgraded["paragraphs"]
        self.assertTrue(upgraded_body["safe"])
        self.assertTrue(upgraded_body["selected"])
        self.assertTrue(upgraded_caption["safe"])
        self.assertFalse(upgraded_caption["selected"])
        self.assertEqual(upgraded_caption["suggestionReason"], "caption")
        self.assertFalse(upgraded["scopeConfirmed"])

    def test_inline_objects_stay_fixed_while_surrounding_body_text_is_rewritten(self) -> None:
        mixed_document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}" xmlns:m="{M_NS}" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:body>
  <w:p><w:r><w:t>公式前正文</w:t></w:r><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath><w:r><w:t>公式后正文</w:t></w:r></w:p>
  <w:p><w:r><w:t>引用前正文</w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> REF cite1 </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>[1]</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t>引用后正文</w:t></w:r></w:p>
  <w:p><w:r><w:t>图片前正文</w:t></w:r><w:r><w:drawing><wp:inline/></w:drawing></w:r><w:r><w:t>图片后正文</w:t></w:r></w:p>
  <w:p><w:r><w:t>控制符前正文</w:t><w:tab/><w:t>控制符中正文</w:t><w:br/><w:t>控制符后正文</w:t></w:r></w:p>
  <w:p><w:r><w:t>修订前正文</w:t></w:r><w:ins w:id="7"><w:r><w:t>修订对象原样保留</w:t></w:r></w:ins><w:r><w:t>修订后正文</w:t></w:r></w:p>
  <w:p><w:r><w:t>备选对象前正文</w:t></w:r><mc:AlternateContent><mc:Choice Requires="w"><w:r><w:t>备选对象原样保留</w:t></w:r></mc:Choice></mc:AlternateContent><w:r><w:t>备选对象后正文</w:t></w:r></w:p>
  <w:p><w:r><w:t>内容控件前正文</w:t></w:r><w:sdt><w:sdtPr><w:tag w:val="editable"/></w:sdtPr><w:sdtContent><w:r><w:t>内容控件中正文</w:t></w:r></w:sdtContent></w:sdt><w:r><w:t>内容控件后正文</w:t></w:r></w:p>
  <w:sectPr/>
</w:body></w:document>'''.encode()
        public = import_document(io.BytesIO(fixture_docx(mixed_document)), "混合结构论文.docx")
        self.assertEqual(public["safeCount"], 7)
        self.assertTrue(all(item["selected"] for item in public["paragraphs"]))
        field_paragraph = public["paragraphs"][1]
        self.assertEqual(field_paragraph["displayText"], "引用前正文[1]引用后正文")
        self.assertEqual(field_paragraph["text"], "引用前正文引用后正文")
        self.assertEqual(public["paragraphs"][4]["text"], "修订前正文修订后正文")
        self.assertEqual(public["paragraphs"][5]["text"], "备选对象前正文备选对象后正文")
        self.assertEqual(
            public["paragraphs"][6]["text"],
            "内容控件前正文内容控件中正文内容控件后正文",
        )

        replacements = {
            public["paragraphs"][0]["id"]: "改写后的公式前后正文",
            public["paragraphs"][1]["id"]: "改写后的引用前后正文",
            public["paragraphs"][2]["id"]: "改写后的图片前后正文",
            public["paragraphs"][3]["id"]: "改写后的控制符前后正文",
            public["paragraphs"][4]["id"]: "改写后的修订对象前后正文",
            public["paragraphs"][5]["id"]: "改写后的备选对象前后正文",
            public["paragraphs"][6]["id"]: "改写后的内容控件正文",
        }
        document = load_document(public["id"])
        output, audit = export_docx(document, replacements, "inline-anchors")
        self.assertTrue(audit["passed"], audit)

        parser = etree.XMLParser(remove_blank_text=False)
        with zipfile.ZipFile(document["sourcePath"]) as before, zipfile.ZipFile(output) as after:
            before_root = etree.fromstring(before.read("word/document.xml"), parser)
            after_root = etree.fromstring(after.read("word/document.xml"), parser)
        before_body = list(before_root.find("w:body", namespaces=NS))
        after_body = list(after_root.find("w:body", namespaces=NS))
        self.assertEqual(
            etree.tostring(before_body[0].find(f".//{{{M_NS}}}oMath")),
            etree.tostring(after_body[0].find(f".//{{{M_NS}}}oMath")),
        )
        self.assertEqual(
            etree.tostring(before_body[2].find(".//w:drawing", namespaces=NS)),
            etree.tostring(after_body[2].find(".//w:drawing", namespaces=NS)),
        )
        field_values = [node.text or "" for node in after_body[1].iter(W_T)]
        self.assertIn("[1]", field_values)
        self.assertEqual("".join(value for value in field_values if value != "[1]"), replacements[field_paragraph["id"]])
        self.assertEqual(len(after_body[3].findall(".//w:tab", namespaces=NS)), 1)
        self.assertEqual(len(after_body[3].findall(".//w:br", namespaces=NS)), 1)
        self.assertEqual(
            etree.tostring(before_body[4].find(".//w:ins", namespaces=NS)),
            etree.tostring(after_body[4].find(".//w:ins", namespaces=NS)),
        )
        self.assertEqual(
            etree.tostring(before_body[5].find(".//{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")),
            etree.tostring(after_body[5].find(".//{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")),
        )
        revision_values = [node.text or "" for node in after_body[4].iter(W_T)]
        self.assertIn("修订对象原样保留", revision_values)
        self.assertEqual(
            "".join(value for value in revision_values if value != "修订对象原样保留"),
            replacements[public["paragraphs"][4]["id"]],
        )
        alternate_values = [node.text or "" for node in after_body[5].iter(W_T)]
        self.assertIn("备选对象原样保留", alternate_values)
        self.assertEqual(
            "".join(value for value in alternate_values if value != "备选对象原样保留"),
            replacements[public["paragraphs"][5]["id"]],
        )
        self.assertEqual(
            "".join(node.text or "" for node in after_body[6].iter(W_T)),
            replacements[public["paragraphs"][6]["id"]],
        )

    def test_word_zip_compression_flags_do_not_false_block_export(self) -> None:
        public = import_document(io.BytesIO(fixture_with_word_zip_metadata()), "Word压缩元数据.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        output, audit = export_docx(document, {body["id"]: "真实 Word 包仍可正常导出"}, "word-zip-run")
        self.assertTrue(audit["passed"], audit)
        with zipfile.ZipFile(output) as archive:
            self.assertIsNone(archive.testzip())

    def test_pure_objects_have_specific_reasons_without_locking_nearby_text(self) -> None:
        from core_documents import _paragraph_safety

        revision_snippets = (
            '<w:moveFrom w:id="1"><w:r><w:t>移动修订</w:t></w:r></w:moveFrom>',
            '<w:r><w:rPr><w:rPrChange w:id="2"/></w:rPr><w:t>格式修订</w:t></w:r>',
            '<w:r><w:t>批注文字</w:t><w:commentReference w:id="3"/></w:r>',
            '<w:del w:id="5"><w:r><w:delText>只有删除内容</w:delText></w:r></w:del>',
        )
        for snippet in revision_snippets:
            paragraph = etree.fromstring(f'<w:p xmlns:w="{W_NS}">{snippet}</w:p>'.encode())
            safe, _label, reason = _paragraph_safety(paragraph)
            self.assertFalse(safe, snippet)
            self.assertEqual(reason, "revision")

        pure_objects = (
            (
                f'<w:p xmlns:w="{W_NS}" xmlns:m="{M_NS}"><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath></w:p>',
                "formula",
            ),
            (
                f'<w:p xmlns:w="{W_NS}" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><w:r><w:drawing><wp:inline/></w:drawing></w:r></w:p>',
                "graphic_anchor",
            ),
            (
                f'<w:p xmlns:w="{W_NS}"><w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>',
                "generated_field",
            ),
            (
                f'<w:p xmlns:w="{W_NS}"><w:sdt><w:sdtPr><w:checkBox/></w:sdtPr><w:sdtContent/></w:sdt></w:p>',
                "content_control",
            ),
            (
                f'<w:p xmlns:w="{W_NS}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:AlternateContent><mc:Choice Requires="w"><w:r><w:t>备选表示</w:t></w:r></mc:Choice></mc:AlternateContent></w:p>',
                "non_text_object",
            ),
        )
        for xml, expected_reason in pure_objects:
            paragraph = etree.fromstring(xml.encode())
            safe, _label, reason = _paragraph_safety(paragraph)
            self.assertFalse(safe, xml)
            self.assertEqual(reason, expected_reason)
            self.assertNotEqual(reason, "complex_structure")

        range_marker_with_body = etree.fromstring(
            f'<w:p xmlns:w="{W_NS}"><w:customXmlInsRangeStart w:id="4"/><w:r><w:t>标记后的普通正文</w:t></w:r></w:p>'.encode()
        )
        safe, _label, reason = _paragraph_safety(range_marker_with_body)
        self.assertTrue(safe)
        self.assertEqual(reason, "")

    def test_alignment_aware_distribution_keeps_semantic_anchors_in_their_runs(self) -> None:
        public = import_document(io.BytesIO(fixture_docx()), "多格式锚点.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        replacement = "研究表述经过调整，仍使用Transformer模型，实验值10，相关依据见[1]，网址https://example.com。"
        output, audit = export_docx(document, {body["id"]: replacement}, "anchor-run")
        self.assertTrue(audit["passed"], audit)

        with zipfile.ZipFile(output) as archive:
            parser = etree.XMLParser(remove_blank_text=False)
            root = etree.fromstring(archive.read("word/document.xml"), parser)
            paragraph = list(root.find("w:body", namespaces=NS))[int(body["bodyChildIndex"])]
            nodes = list(paragraph.iter(W_T))
            self.assertEqual("".join(node.text or "" for node in nodes), f"  {replacement}  ")
            self.assertEqual(nodes[1].text, "Transformer")
            self.assertEqual(nodes[3].text, "[1]")
            self.assertEqual(nodes[5].text, "https://example.com。  ")
            self.assertEqual(len(paragraph.findall(".//w:hyperlink", namespaces=NS)), 2)
            self.assertIsNotNone(nodes[3].getparent().find("./w:rPr/w:vertAlign", namespaces=NS))

    def test_raw_document_xml_changes_only_selected_text_spans(self) -> None:
        from core_documents import (
            _mask_selected_text_tokens,
            _selected_word_text_node_indexes,
        )

        public = import_document(io.BytesIO(fixture_docx()), "原始字节保真.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        output, audit = export_docx(document, {body["id"]: "  改写后仍保留锚点与空白  "}, "raw-byte")
        self.assertEqual(audit["status"], "passed")
        with zipfile.ZipFile(document["sourcePath"], "r") as before, zipfile.ZipFile(output, "r") as after:
            before_xml = before.read("word/document.xml")
            after_xml = after.read("word/document.xml")
        parser = etree.XMLParser(remove_blank_text=False)
        root = etree.fromstring(before_xml, parser)
        selected_nodes = _selected_word_text_node_indexes(
            root, {int(body["bodyChildIndex"])}
        )
        self.assertEqual(
            _mask_selected_text_tokens(before_xml, selected_nodes),
            _mask_selected_text_tokens(after_xml, selected_nodes),
        )
        self.assertEqual(before_xml.splitlines()[0], after_xml.splitlines()[0])
        self.assertNotIn(b"ns0:", after_xml)

    def test_digital_signature_is_a_warning_and_generated_file_is_kept(self) -> None:
        public = import_document(io.BytesIO(fixture_with_digital_signature()), "签名论文.docx")
        self.assertTrue(public["hasDigitalSignature"])
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        output, audit = export_docx(document, {body["id"]: "签名文件改写结果"}, "signed")
        self.assertTrue(output.exists())
        self.assertFalse(audit["passed"])
        self.assertEqual(audit["status"], "warning")
        signature = next(item for item in audit["checks"] if item["name"] == "数字签名")
        self.assertFalse(signature["passed"])

    def test_docx_txt_fallback_keeps_protected_body_units(self) -> None:
        public = import_document(io.BytesIO(fixture_docx()), "文本回退.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        output, audit = export_txt(document, {body["id"]: "改写后的普通正文"}, "txt-fallback")
        content = output.read_text(encoding="utf-8")
        self.assertTrue(audit["passed"], audit)
        self.assertIn("改写后的普通正文", content)
        self.assertIn("表格内容不改", content)
        self.assertIn("目录字段", content)
        self.assertIn("公式段", content)
        self.assertIn("绘图段", content)

    def test_last_rendered_page_break_is_safe_and_preserved(self) -> None:
        public = import_document(io.BytesIO(fixture_docx()), "跨页论文.docx")
        page_break = public["paragraphs"][2]
        self.assertTrue(page_break["safe"])

        update_scope(public["id"], [page_break["id"]])
        document = load_document(public["id"])
        output, audit = export_docx(document, {page_break["id"]: "改写后的跨页正文"}, "page-break-run")
        self.assertTrue(audit["passed"], audit)

        with zipfile.ZipFile(output) as archive:
            parser = etree.XMLParser(remove_blank_text=False)
            root = etree.fromstring(archive.read("word/document.xml"), parser)
            body = list(root.find("w:body", namespaces=NS))
            paragraph = body[int(page_break["bodyChildIndex"])]
            self.assertEqual("".join(node.text or "" for node in paragraph.iter(W_T)), "改写后的跨页正文")
            self.assertEqual(len(paragraph.findall(".//w:lastRenderedPageBreak", namespaces=NS)), 1)

    def test_changed_source_blocks_docx_export(self) -> None:
        public = import_document(io.BytesIO(fixture_docx()), "paper.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        Path(document["sourcePath"]).write_bytes(b"changed")
        with self.assertRaises(FormatFidelityError):
            export_docx(document, {body["id"]: "result"}, "test-run")

    def test_word_incompatible_control_character_blocks_docx_without_replacing_text(self) -> None:
        public = import_document(io.BytesIO(fixture_docx()), "control-character.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        with self.assertRaisesRegex(FormatFidelityError, "控制字符"):
            export_docx(document, {body["id"]: "正常文字\u000b异常控制符"}, "control-run")

    def test_exported_package_opens_with_python_docx_when_available(self) -> None:
        try:
            from docx import Document as PythonDocxDocument
        except ImportError:
            self.skipTest("python-docx is not installed in this environment")

        public = import_document(io.BytesIO(fixture_docx()), "openable.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        output, audit = export_docx(document, {body["id"]: "可由标准 DOCX 读取器打开的改写结果"}, "open-run")
        self.assertTrue(audit["passed"], audit)
        reopened = PythonDocxDocument(str(output))
        self.assertGreaterEqual(len(reopened.paragraphs), 4)
        self.assertIn("可由标准 DOCX 读取器打开的改写结果", reopened.paragraphs[1].text)

    def test_alternate_word_namespace_prefix_and_empty_text_nodes_are_preserved(self) -> None:
        alternate_xml = DOCUMENT_XML.replace(b"xmlns:w=", b"xmlns:x=").replace(b"w:", b"x:")
        public = import_document(io.BytesIO(fixture_docx(alternate_xml)), "alternate-prefix.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        output, audit = export_docx(document, {body["id"]: "替换后仍保留命名空间和书签"}, "prefix-run")
        self.assertTrue(audit["passed"], audit)

        with zipfile.ZipFile(output) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            paragraph = list(root.find("w:body", namespaces=NS))[int(body["bodyChildIndex"])]
            self.assertEqual(root.prefix, "x")
            self.assertEqual(len(list(paragraph.iter(W_T))), body["textNodeCount"])
            self.assertEqual(len(paragraph.findall(".//w:bookmarkStart", namespaces=NS)), 1)

    def test_audit_rejects_non_text_change_inside_selected_paragraph(self) -> None:
        from core_documents import _audit_docx

        public = import_document(io.BytesIO(fixture_docx()), "tamper-check.docx")
        body = public["paragraphs"][1]
        update_scope(public["id"], [body["id"]])
        document = load_document(public["id"])
        source = Path(document["sourcePath"])
        exported, _audit = export_docx(document, {body["id"]: "正常改写结果"}, "before-tamper")
        tampered = Path(self.temporary.name) / "tampered.docx"

        with zipfile.ZipFile(exported, "r") as before, zipfile.ZipFile(tampered, "w") as after:
            parser = etree.XMLParser(remove_blank_text=False)
            root = etree.fromstring(before.read("word/document.xml"), parser)
            paragraph = list(root.find("w:body", namespaces=NS))[int(body["bodyChildIndex"])]
            spacing = paragraph.find("./w:pPr/w:spacing", namespaces=NS)
            self.assertIsNotNone(spacing)
            spacing.set(f"{{{W_NS}}}line", "240")
            changed_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            after.comment = before.comment
            for info in before.infolist():
                after.writestr(info, changed_xml if info.filename == "word/document.xml" else before.read(info.filename))

        audit = _audit_docx(source, tampered, {int(body["bodyChildIndex"])})
        self.assertFalse(audit["passed"], audit)
        licensed = next(item for item in audit["checks"] if item["name"] == "仅改动所选正文文字")
        self.assertFalse(licensed["passed"])


if __name__ == "__main__":
    unittest.main()
